"""
title: Office Documents
author: tovren
author_url: https://github.com/tovren/open-webui-plugins
version: 0.3.2
license: MIT
required_open_webui_version: 0.11.4
description: Lets the model create Word (.docx), Excel (.xlsx) and PowerPoint (.pptx) files. The model only supplies the content as structured JSON; fixed code builds the file, so no model-written code is ever run. The file is attached to the answer, where it can be previewed and downloaded.

Tools
- create_word_file: title page, optional table of contents, headings, paragraphs, bullet and
  numbered lists (two levels), tables, quotes, page breaks, header text and page numbers; optionally
  with the header, footer and margins of a company template (word_template valve).
- create_excel_file: several sheets with typed columns (text, number, integer, currency, percent,
  date, time), row formulas from a checked whitelist, total rows and native Excel charts.
- create_powerpoint_file: 16:9 slides with the layouts title, section, bullets, two columns,
  table and chart (native PowerPoint charts), plus speaker notes.
Inline **bold**, *italic* and [links](https://...) work in Word and PowerPoint; Excel cells get plain text.

Security & privacy
- The model never writes code: it fills in a fixed JSON schema, the plugin validates it and builds
  the file. No shell, no network access, nothing is loaded from URLs.
- Every formula, whether a column's `formula` field or a formula written into a cell, is checked against a
  whitelist (SUM, IF, ROUND, VLOOKUP ...): only cells of the same workbook, no external references, no DDE, no
  WEBSERVICE/HYPERLINK/INDIRECT.
  A cell value with such a function (e.g. "=HYPERLINK(...)") is written as text, also in the data sheets of
  PowerPoint charts; other formulas that are not allowed go back to the model.
- Links are only created for http, https and mailto. Limits on slides, rows, columns and text
  length keep a runaway model from blocking the server.
- The file is stored in Open WebUI's own file storage under the user's account, so only the user
  (and admins) can open it. Only the libraries bundled with Open WebUI are used.
"""

import asyncio
import base64
import copy
import difflib
import io
import json
import logging
import re
import unicodedata
import zipfile
from collections.abc import Awaitable, Callable, Collection, Iterator
from dataclasses import dataclass, replace
from datetime import date, datetime
from typing import Any, ClassVar, Literal, NamedTuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import CONTENT_TYPE as DOCX_CT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.writer.theme import theme_xml
from openpyxl.worksheet.worksheet import Worksheet
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.xlsx import CategoryWorkbookWriter
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as PPTX_RT
from pptx.oxml.xmlchemy import OxmlElement as PptxOxmlElement
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.text.text import TextFrame, _Paragraph as PptxParagraph
from pptx.util import Emu, Inches, lazyproperty
from pptx.util import Pt as PptxPt
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    PrivateAttr,
    TypeAdapter,
    ValidationError,
    ValidationInfo,
    field_validator,
    model_validator,
)

log = logging.getLogger(__name__)

EventCallback = Callable[[dict[str, Any]], Awaitable[Any]]

# ------------------------------------------------------------------------------------------ constants --

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
DEFAULT_ACCENT = "1F4E79"
MAX_FILENAME_CHARS = 60
SAFE_LINK_PREFIXES = ("http://", "https://", "mailto:")
SUB_ITEM_PREFIX = "- "  # list items starting with this are indented one level

# Limits: an error tells the model what to shorten, so it can call the tool again.
MAX_TEXT_CHARS = 200_000  # all text of one file together
MAX_WORD_BLOCKS = 500
MAX_WORD_TABLE_ROWS = 500
MAX_WORD_TABLE_COLUMNS = 12
MAX_SHEETS = 20
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 50
MAX_CELLS = 250_000  # all sheets together; empty cells count too
MAX_CELL_CHARS = 32_000  # Excel's limit is 32,767
MAX_FORMULA_CHARS = 500
MAX_SLIDES = 60
MAX_SLIDE_ITEMS = 12
MAX_SLIDE_TABLE_ROWS = 16
MAX_SLIDE_TABLE_COLUMNS = 8
MAX_CHART_CATEGORIES = 50
MAX_CHART_SERIES = 10
MAX_SHEET_CHARTS = 4
MAX_ERRORS_SHOWN = 8

MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
MAX_DECIMALS = 6
CURRENCY_FORMATS = {
    "EUR": '#,##0.00 "€"',
    "USD": '"$"#,##0.00',
    "GBP": '"£"#,##0.00',
    "CHF": '"CHF" #,##0.00',
}

TRANSLATIONS: dict[str, dict[str, str]] = {
    "en": {
        "status_word": "Creating Word document …",
        "status_excel": "Creating Excel workbook …",
        "status_powerpoint": "Creating PowerPoint presentation …",
        "status_done": "{name} created",
        "status_failed": "The file could not be created",
        "default_title": "Document",
        "total": "Total",
        "page_footer": "Page {page} of {pages}",
        "toc_title": "Contents",
        "toc_placeholder": "Right-click and choose “Update field” to show the table of contents.",
        "sheet": "Sheet",
        "date_format": "YYYY-MM-DD",
        "paper": "letter",
        "chart": "Chart",
        "notes": "Notes",
    },
    "de": {
        "status_word": "Word-Dokument wird erstellt …",
        "status_excel": "Excel-Arbeitsmappe wird erstellt …",
        "status_powerpoint": "PowerPoint-Präsentation wird erstellt …",
        "status_done": "{name} erstellt",
        "status_failed": "Die Datei konnte nicht erstellt werden",
        "default_title": "Dokument",
        "total": "Summe",
        "page_footer": "Seite {page} von {pages}",
        "toc_title": "Inhalt",
        "toc_placeholder": "Rechtsklick → „Felder aktualisieren“ zeigt das Inhaltsverzeichnis an.",
        "sheet": "Tabelle",
        "date_format": "DD.MM.YYYY",
        "paper": "a4",
        "chart": "Diagramm",
        "notes": "Notizen",
    },
}

# ----------------------------------------------------------------------------------------------- design --
# One accent colour (the accent_color valve) gives the colours of all three formats: titles and header rows in the
# accent, a dark shade for the title slide, light tints for banded rows, total rows and cards, neutral greys for
# text and lines.

WHITE, BLACK = "FFFFFF", "000000"
TEXT_COLOR = "262626"
MUTED_COLOR = "6B6B6B"  # subtitles, header and footer lines, axis labels
RULE_COLOR = "D9D9D9"  # thin lines


def mix(color: str, other: str, share: float) -> str:
    """`color` with `share` (0 to 1) of `other` mixed in, as RRGGBB."""
    return "".join(
        f"{round(a + (b - a) * share):02X}" for a, b in zip(bytes.fromhex(color), bytes.fromhex(other), strict=True)
    )


@dataclass(frozen=True)
class Design:
    """Colours and font of the files: the accent colour (accent_color valve) gives all colours, the font (font
    valve) is used for all text."""

    accent: str = DEFAULT_ACCENT
    font: str = BODY_FONT

    @property
    def dark(self) -> str:
        return mix(self.accent, BLACK, 0.35)

    @property
    def light(self) -> str:  # bars and lines on the dark title slide
        return mix(self.accent, WHITE, 0.6)

    @property
    def tint(self) -> str:  # total rows
        return mix(self.accent, WHITE, 0.85)

    @property
    def band(self) -> str:  # every second table row, cards
        return mix(self.accent, WHITE, 0.93)

    def series(self, count: int) -> list[str]:
        """Chart colours that stay apart in print. Two series (e.g. last year and this year): the second one,
        usually the newer, in the accent."""
        if count == 2:
            return [mix(self.accent, WHITE, 0.55), self.accent]
        colors = [self.accent, mix(self.accent, WHITE, 0.5), "8C8C8C", self.dark, mix(self.accent, WHITE, 0.75),
                  "BFBFBF", mix(self.accent, BLACK, 0.6), "595959"]  # fmt: skip
        return [colors[i % len(colors)] for i in range(count)]


DEFAULT_DESIGN = Design()
FONT_NAME_RE = re.compile(r"[^\W_][\w .-]{0,40}")  # 'Calibri', 'Segoe UI', 'Source Sans 3'
THEME_FONTS_RE = r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")'  # headings and body text of a theme


# Messages for the model (always English; the model answers the user in their language).
RESULT_ATTACHED = (
    "Done: the file '{name}' was created and is attached to your answer, where the user can preview and "
    "download it. Do not add a link and do not repeat the file content. Briefly tell the user what the "
    "file contains."
)
RESULT_DOWNLOADED = (
    "Done: the file '{name}' was created and downloaded in the user's browser. Do not add a link and do "
    "not repeat the file content. Briefly tell the user what the file contains."
)
RESULT_INVALID = (
    "The file was NOT created because the arguments are invalid:\n{errors}\n"
    "Fix these points and call {tool} again with the complete, corrected content."
)
RESULT_FAILED = "The file could not be created because of an internal error. Tell the user to try again later."

# ------------------------------------------------------------------------------------ input schema --
# The schemas below are what the model sees (Open WebUI turns the type hints and descriptions into
# JSON schema). Limits are checked in code instead of maxItems/maxLength: large limits in a schema
# make the grammar of llama.cpp and vLLM huge and slow.


class Schema(BaseModel):
    model_config = ConfigDict(extra="ignore", coerce_numbers_to_str=True, allow_inf_nan=False)
    _index: int = PrivateAttr(default=-1)  # position in the list the model sent, for error messages
    IGNORED_KEYS: ClassVar[frozenset[str]] = frozenset()  # unknown keys that may hold text without losing content

    @model_validator(mode="before")
    @classmethod
    def _null_means_default(cls, data: Any) -> Any:
        """Models often send null for fields they do not use."""
        return {k: v for k, v in data.items() if v is not None} if isinstance(data, dict) else data

    @model_validator(mode="before")
    @classmethod
    def _no_content_in_unknown_keys(cls, data: Any) -> Any:
        """Unknown keys are ignored, so text under another name ('content', 'formulas', a sheet's 'total') would be
        lost without a word. Such a key is sent back instead, with the fields there are."""
        if isinstance(data, dict):
            unknown = [k for k, v in data.items() if k not in cls.model_fields and k not in cls.IGNORED_KEYS]
            key = next((k for k in unknown if has_content(data[k])), None)
            if key is not None:
                fields = ", ".join(cls.model_fields)
                raise ValueError(f"'{key}' is not a field here, its content would be lost (fields: {fields})")
        return data


def has_content(value: Any) -> bool:
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, list | tuple):  # numbers count here: a table of numbers is content, a single "width": 12 not
        return any(
            has_content(item) or (isinstance(item, int | float) and not isinstance(item, bool)) for item in value
        )
    if isinstance(value, dict):
        return any(has_content(item) for item in value.values())
    if isinstance(value, BaseModel):
        return has_content(value.model_dump())
    return False


class WordBlock(Schema):
    type: Literal[
        "heading", "paragraph", "bullets", "numbered", "table", "quote", "page_break", "table_of_contents"
    ] = Field(description="Kind of block.")
    text: str = Field(
        default="",
        description="Text of a heading, paragraph or quote. Inline **bold**, *italic* and [label](https://...) "
        "links are allowed.",
    )
    level: int = Field(default=1, description="Heading level: 1, 2 or 3.")
    items: list[str] = Field(
        default=[],
        description="Items of a bullets or numbered list. Start an item with '- ' to indent it one level.",
    )
    rows: list[list[str]] = Field(
        default=[], description="Table rows as lists of cell texts. The first row is the header row."
    )

    @field_validator("items")
    @classmethod
    def _top_level(cls, items: list[str], info: ValidationInfo) -> list[str]:
        return top_level(items, numbered=info.data.get("type") == "numbered")

    @field_validator("rows", mode="before")
    @classmethod
    def _null_cells(cls, rows: Any) -> Any:
        return null_cells(rows)

    @model_validator(mode="after")
    def _content_in_the_field_of_the_type(self) -> "WordBlock":
        """A paragraph's 'items' or a table's 'text' is not shown; without this the block would be dropped as empty."""
        shown = WORD_CONTENT_FIELDS.get(self.type)
        other = next((f for f in ("text", "items", "rows") if f != shown and has_content(getattr(self, f))), None)
        if other and shown:
            raise ValueError(
                f"a {self.type} block shows only '{shown}'; put the '{other}' into a block of its own type"
            )
        if other:
            raise ValueError(f"a {self.type} block shows no '{other}'; put it into a block of its own type")
        return self


WORD_CONTENT_FIELDS = {
    "heading": "text", "paragraph": "text", "quote": "text", "bullets": "items", "numbered": "items", "table": "rows",
}  # fmt: skip


class ExcelColumn(Schema):
    header: str = Field(description="Column header.")
    type: Literal["text", "number", "integer", "currency", "percent", "date", "time"] = Field(
        default="text",
        description="Value type. percent: send 12.5 for 12.5 % (not 0.125); the cell then holds 0.125, as in "
        "Excel, so formulas use it as it is: =B{row}*(1-C{row}), without /100, and a percent formula column computes "
        "the fraction (=B{row}/C{row}, without *100). date: YYYY-MM-DD. time: hh:mm; the cell holds the fraction of "
        "a day, so a time column with =C{row}-B{row} shows 8:30, and hours as a number (8.5) need a number column "
        "with =(C{row}-B{row})*24.",
    )
    formula: str = Field(
        default="",
        description="Optional formula filled into every data row, e.g. '=B{row}*C{row}' ({row} is the current "
        "row; data starts in row 2). English function names (SUM, AVERAGE, MIN, MAX, COUNT, IF, ROUND, "
        "IFERROR ...), commas between arguments. Cells of another sheet: 'Rates'!$B$3 (one fixed cell) or "
        "'Rates'!B{row}. Put null in the rows for this column; a row that needs another formula (a sum row "
        "below the items: '=SUM(D2:D6)', '=D7*0.19') gets it in its cell instead.",
    )
    total: Literal["none", "sum", "average", "min", "max", "count"] = Field(
        default="none", description="Adds a total row below the data with this function for the column."
    )


class ExcelChart(Schema):
    type: Literal["column", "bar", "line", "pie"] = Field(default="column", description="Chart type.")
    title: str = Field(default="", description="Chart title.")
    category_column: str = Field(
        default="",
        description="Header of the column with the labels (x axis or pie slices). Not used with series_rows.",
    )
    value_columns: list[str] = Field(
        default=[],
        description="Headers of the numeric columns to plot (pie: only one). With series_rows: the columns that form "
        "the x axis, next to each other; default: the number columns after the first column.",
    )
    series_rows: list[str] = Field(
        default=[],
        description="Only when each data series is a row, e.g. months as columns: the labels in the first column of "
        'the rows to plot, e.g. ["Income", "Expenses"].',
    )
    unit: str = Field(default="", description="Unit of the values, e.g. 'EUR' or 'hours'; shown at the value axis.")


class ExcelSheet(Schema):
    name: str = Field(default="", description="Sheet name, at most 31 characters.")
    columns: list[ExcelColumn] = Field(description="Columns from left to right.")
    rows: list[list[str | float | None]] = Field(
        default=[],
        description="Data rows. Each row has one value per column, in column order. Use JSON numbers for "
        "numbers and null for empty cells and formula columns.",
    )
    chart: ExcelChart | list[ExcelChart] | None = Field(
        default=None, description=f"Optional chart next to the table; a list for up to {MAX_SHEET_CHARTS} charts."
    )

    @model_validator(mode="before")
    @classmethod
    def _rows_as_objects(cls, data: Any) -> Any:
        """Rows written as objects with the column headers as keys ({"Artikel": "Stuhl", "Preis (€)": 240}) become
        lists in column order; a key that is no header is sent back, its value would have no column. Columns and
        rows written into the chart are taken back first, in the same validator: pydantic runs the 'before'
        validators from the last to the first, so the order of two would decide."""
        data = take_out_of_chart(data, ("columns", "rows"))
        rows, columns = (data.get("rows"), data.get("columns")) if isinstance(data, dict) else (None, None)
        if not isinstance(rows, list) or not isinstance(columns, list) or not any(isinstance(r, dict) for r in rows):
            return data
        headers = [str(c.get("header") or "").strip() if isinstance(c, dict) else "" for c in columns]
        if len(set(headers)) < len(headers):
            return data  # two columns with the same header: which one is meant is unclear
        converted = []
        for r, row in enumerate(rows):
            if isinstance(row, dict):
                values = {str(key).strip(): value for key, value in row.items()}
                unknown = next((key for key in values if key not in headers), None)
                if unknown is not None:
                    raise ValueError(
                        f"rows[{r}] is an object whose key '{unknown}' is not a column header ({', '.join(headers)}); "
                        "send each row as a list of values in column order"
                    )
                row = [values.get(header) for header in headers]
            converted.append(row)
        return {**data, "rows": converted}

    @field_validator("chart", mode="before")
    @classmethod
    def _empty_chart_is_no_chart(cls, chart: Any) -> Any:
        if isinstance(chart, list):
            return [c for c in chart if c is not None and not (isinstance(c, dict) and not has_content(c))] or None
        return None if isinstance(chart, dict) and not has_content(chart) else chart

    @property
    def charts(self) -> list[ExcelChart]:
        return self.chart if isinstance(self.chart, list) else [self.chart] if self.chart is not None else []


class ChartSeries(Schema):
    name: str = Field(description="Series name, shown in the legend.")
    values: list[float] = Field(description="One number per category.")


class SlideChart(Schema):
    IGNORED_KEYS = frozenset({"title"})  # the slide's title is shown above the chart
    type: Literal["column", "bar", "line", "pie"] = Field(default="column", description="Chart type.")
    categories: list[str] = Field(description="Category labels (x axis or pie slices).")
    series: list[ChartSeries] = Field(description="Data series (pie: only one).")
    unit: str = Field(default="", description="Unit of the values, e.g. 'hours per week'; shown at the value axis.")


class Slide(Schema):
    layout: Literal["title", "section", "bullets", "two_columns", "table", "chart"] = Field(
        description="title: opening slide with title and subtitle. section: divider between parts. "
        "bullets: title with a list. two_columns: two lists side by side. table: title with a table. "
        "chart: title with a chart."
    )
    title: str = Field(description="Slide title.")
    subtitle: str = Field(default="", description="Subtitle, only for title and section slides.")
    bullets: list[str] = Field(
        default=[],
        description="Items for a bullets slide, short phrases rather than sentences. Start an item with '- ' to "
        "indent it.",
    )
    left: list[str] = Field(default=[], description="Items of the left column (two_columns).")
    right: list[str] = Field(default=[], description="Items of the right column (two_columns).")
    left_heading: str = Field(default="", description="Optional heading above the left column.")
    right_heading: str = Field(default="", description="Optional heading above the right column.")
    table: list[list[str]] = Field(default=[], description="Table rows for a table slide; the first row is the header.")
    chart: SlideChart | None = Field(default=None, description="Chart data for a chart slide.")
    notes: str = Field(default="", description="Optional speaker notes.")

    @field_validator("bullets", "left", "right")
    @classmethod
    def _top_level(cls, items: list[str]) -> list[str]:
        return top_level(items)

    @field_validator("table", mode="before")
    @classmethod
    def _null_cells(cls, rows: Any) -> Any:
        return null_cells(rows)

    @model_validator(mode="before")
    @classmethod
    def _fields_inside_the_chart(cls, data: Any) -> Any:
        return take_out_of_chart(data, SLIDE_CONTENT_FIELDS)

    @field_validator("chart", mode="before")
    @classmethod
    def _empty_chart_is_no_chart(cls, chart: Any) -> Any:
        empty = isinstance(chart, dict) and not has_content([chart.get("categories"), chart.get("series")])
        return None if empty else chart


def take_out_of_chart(data: Any, fields: Collection[str]) -> Any:
    """Fields of the slide or sheet written into its chart: the chart's '}' was forgotten, and loads_lenient adds
    it only after them ('"chart": {...}, "notes": "n"' became '"chart": {..., "notes": "n"}'). Unknown keys are
    ignored, so they would be lost; they are taken back unless the slide or sheet has a value of its own."""
    chart = data.get("chart") if isinstance(data, dict) else None
    if not isinstance(chart, dict) or not (moved := [k for k in fields if k in chart and data.get(k) is None]):
        return data
    return {**data, **{k: chart[k] for k in moved}, "chart": {k: v for k, v in chart.items() if k not in moved}}


BULLET_MARKS = ("• ", "* ", "– ", "· ")  # written by the model in front of the items; the list adds its own
ITEM_NUMBER_RE = re.compile(r"(\d{1,3})[.)]\s+(?=\S)")


def top_level(items: list[str], numbered: bool = False) -> list[str]:
    """Models used to Markdown start every item with '- '; without any top-level item these are no sub items. A
    bullet ('• ', '* ', '– ') in front of every top-level item, or in a numbered list the numbers 1., 2., 3. …, would
    be shown next to the list's own ('• • Punkt', '1. 1. Erstens'), so they are removed."""
    texts = [item.strip() for item in items if item.strip()]
    if texts and all(text.startswith(SUB_ITEM_PREFIX) for text in texts):
        return [item.strip()[len(SUB_ITEM_PREFIX) :].strip() for item in items]
    main = [text for text in texts if not text.startswith(SUB_ITEM_PREFIX)]
    mark = next((mark for mark in BULLET_MARKS if main and all(text.startswith(mark) for text in main)), None)
    numbers = [match and int(match[1]) for match in map(ITEM_NUMBER_RE.match, main)] if numbered else []
    if mark is None and not (main and numbers == list(range(1, len(main) + 1))):
        return items
    result = []
    for item in items:
        if not item.strip() or item.strip().startswith(SUB_ITEM_PREFIX):
            result.append(item)
        else:
            result.append(item.strip()[len(mark) :].strip() if mark else ITEM_NUMBER_RE.sub("", item.strip(), count=1))
    return result


def null_cells(rows: Any) -> Any:
    """null in a table cell is an empty cell."""
    if not isinstance(rows, list):
        return rows
    return [["" if cell is None else cell for cell in row] if isinstance(row, list) else row for row in rows]


WORD_BLOCKS = TypeAdapter(list[WordBlock])
EXCEL_SHEETS = TypeAdapter(list[ExcelSheet])
SLIDES = TypeAdapter(list[Slide])
SLIDE_CONTENT_FIELDS = ("subtitle", "bullets", "left", "right", "left_heading", "right_heading", "table", "notes")


class InvalidInput(ValueError):
    """Raised with one message per problem; the messages are returned to the model."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("\n".join(errors))
        self.errors = errors


LOCATION_NOISE_RE = re.compile(r"(str|float|int|bool|list\[.*\]|function-.*|[A-Z]\w*)")
UNION_MEMBER_RE = re.compile(r"[A-Z]\w*|list\[[A-Z]\w*\]")  # 'ExcelChart', 'list[ExcelChart]': one chart or a list
SHAPE_ERRORS = ("model_type", "list_type")


def _location(parts: tuple[Any, ...]) -> str:
    text = ""
    for part in parts:
        if isinstance(part, str) and LOCATION_NOISE_RE.fullmatch(part):
            continue  # union members pydantic adds to the path, e.g. "rows[0][1].str"
        text += f"[{part}]" if isinstance(part, int) else f".{part}" if text else str(part)
    return text


def _union_errors(errors: list[Any]) -> list[Any]:
    """For a field that takes one object or a list of them (a sheet's chart), pydantic reports both: that the value
    is no list and what is wrong in the object. Only the errors of the shape the model sent are kept."""
    kept: list[Any] = []
    for error in errors:
        loc = error["loc"]
        at = next((i for i, part in enumerate(loc) if isinstance(part, str) and UNION_MEMBER_RE.fullmatch(part)), None)
        if at is not None and error["type"] in SHAPE_ERRORS and len(loc) == at + 1:
            rivals = [other for other in errors if other is not error and other["loc"][:at] == loc[:at]]
            if any(len(other["loc"]) > at + 1 or other["type"] not in SHAPE_ERRORS for other in rivals):
                continue  # the other shape has the real problem
            if any(k["loc"][:at] == loc[:at] for k in kept):
                continue
            error = {**error, "msg": "Input should be an object (or a list of objects)"}
        kept.append(error)
    return kept


def parse_argument(adapter: TypeAdapter, value: Any, name: str, extras: dict[str, Any] | None = None) -> Any:
    """Validate a tool argument.

    Some models send nested arguments as a JSON string (see loads_lenient); `extras` receives other tool
    parameters that the model wrote into the same JSON text. Every validated item remembers its position in
    the list (`_index`), so error messages still name the item the model sent after empty blocks were left out.
    """
    if isinstance(value, str):
        value = _loads_argument(value, name, extras, adapter)
    if isinstance(value, dict) and isinstance(value.get(name), list) and len(value) > 1:
        # {"sheets": [...], "title": "T"}: all parameters in one object
        others = {k: v for k, v in value.items() if k != name}
        lost = next((k for k, v in others.items() if k not in TOOL_PARAMETERS and has_content(v)), None)
        if lost is not None:
            raise InvalidInput([f"{name}: '{lost}' is not a parameter of this tool, its content would be lost"])
        if extras is not None:
            extras.update({k: v for k, v in others.items() if k in TOOL_PARAMETERS})
        value = value[name]
    if isinstance(value, dict) and len(value) == 1 and isinstance(next(iter(value.values())), list):
        value = next(iter(value.values()))  # {"slides": [...]} instead of [...]
    value = _unescape_newlines(value)
    if isinstance(value, list):  # broken tool-call syntax: content got lost, so let the model send it again
        broken = [f"{name}[{i}]" for i, item in enumerate(value) if _contains(item, BROKEN_QUOTE) or _garbled(item)]
        if broken:
            raise InvalidInput(
                [f"{where}: the arguments are garbled here (a key or the token {BROKEN_QUOTE} ended up inside a text,"
                 " and text got lost)" for where in broken[:MAX_ERRORS_SHOWN]]
                + [f"Write every key and string with plain double quotes and send the complete {name} again."]
            )  # fmt: skip
    try:
        result = adapter.validate_python(value)
    except ValidationError as e:
        raise InvalidInput(
            [f"{_location((name, *err['loc']))}: {err['msg'].removeprefix('Value error, ')}"
             for err in _union_errors(e.errors())]
        ) from e  # fmt: skip
    for i, item in enumerate(result):
        item._index = i
    return result


def _loads_argument(text: str, name: str, extras: dict[str, Any] | None, adapter: TypeAdapter) -> Any:
    def accept(value: Any) -> bool:  # a repaired text must give valid arguments, or the JSON error is shown
        try:
            parse_argument(adapter, value, name, {})
        except InvalidInput:
            return False
        return True

    try:
        value = loads_lenient(text, extras, accept)
        if isinstance(value, str) and value.strip()[:1] in ("[", "{"):
            value = loads_lenient(value, extras, accept)  # JSON text encoded twice
    except json.JSONDecodeError as e:
        raise InvalidInput([json_error_message(name, e)]) from e
    except RecursionError as e:
        raise InvalidInput([f"{name}: the JSON is nested too deeply"]) from e
    except ValueError as e:  # NaN/Infinity, numbers with thousands of digits
        raise InvalidInput([f"{name}: the JSON could not be read: {str(e)[:200]}"]) from e
    return value


ERROR_MARK = " <<<HERE>>> "


def json_error_message(name: str, error: json.JSONDecodeError) -> str:
    """Show the text on both sides of the place and name the likely slip. Models repeat the same mistake when
    they only see the text before it (an extra ']' right after the place went unnoticed four times in a row)."""
    text, pos = error.doc, error.pos
    start, end = max(0, pos - 60), pos + 40
    before = ("..." if start else "") + text[start:pos].replace("\n", " ").rstrip()
    after = text[pos:end].replace("\n", " ").lstrip() + ("..." if end < len(text) else "")
    return (
        f"{name}: the JSON is invalid at character {pos} ({error.msg}): {before}{ERROR_MARK}{after} - "
        + json_error_hint(text, pos, error.msg)
    )


KEY_AHEAD_RE = re.compile(r'\s*"[^"\\]*"\s*:')


def json_error_hint(text: str, pos: int, msg: str) -> str:
    """The likely slip at the place of a JSON error, as an instruction for the model."""
    char = text[pos : pos + 1]
    previous = text[:pos].rstrip()[-1:]
    if not text.strip():
        return "the arguments are empty: send the content"
    if msg.startswith("Unterminated string"):
        return 'the arguments end inside a text: they were cut off (output limit), or a text ends with \\" instead of "'
    if msg.startswith("Invalid \\escape"):
        return "a backslash inside a text must be written as \\\\"
    if msg.startswith("Invalid control character"):
        return "a line break or tab inside a text must be written as \\n or \\t"
    if pos >= len(text.rstrip()):
        return "the arguments end in the middle (cut off, e.g. at the output limit): send them again completely"
    if msg == "Expecting value" and char in "„“”‘’'":
        return 'a text must start and end with a straight double quote "'
    if msg == "Expecting value" and text.startswith(("True", "False", "None"), pos):
        return "write true, false and null in lower case"
    if msg == "Expecting value" and previous == ":" and char in ",}]":
        return "a value is missing after the colon"
    if char == ":" and open_brackets(text, pos)[-1:] == "]":
        return "a key stands inside a list: the list before it is not closed, add ] before the key"
    if char == ":":  # '{"type": "paragraph": "Text"}'
        return 'a value is followed by ":", so a key is missing: write , "key": value (e.g. , "text": "...")'
    if char in ("]", "}"):
        expected = open_brackets(text, pos)[-1:]
        if expected != char:
            if not expected:
                return f"the {char} at the mark closes nothing that is open: delete it"
            kind = "list" if expected == "]" else "object"
            return (
                f"the {char} at the mark does not fit: the open {kind} must be closed with {expected} first, "
                f"or the {char} is one too many"
            )
        if previous == ",":
            return f"delete the comma before the {char}"
    if msg == "Extra data" and text[pos:].lstrip().startswith("```"):
        return "delete the ``` after the JSON"
    if msg == "Extra data":
        return "more text follows the end of the value: put all items into one list, or delete the ] closing it early"
    if char == "[" and msg.startswith("Expecting property name"):
        return 'a list stands where a key belongs: write "key": [...], or join the lists (table rows: [[...], [...]])'
    if previous == '"' and msg == "Expecting ',' delimiter" and KEY_AHEAD_RE.match(text, pos):
        return "a comma is missing before the mark"
    if msg == "Expecting ':' delimiter" and char in ",}" and open_brackets(text, pos)[-1:] == "}":
        return 'a key without a value, or a list written with { }: an object needs "key": value, a list uses [ ]'
    if previous == '"' and msg in ("Expecting ',' delimiter", "Expecting ':' delimiter"):
        return (
            'the text before the mark ends with a " that closes it too early: write a " inside a text as \\" '
            "or use other quotation marks („…“ or '…')"
        )
    return 'check the quotes and brackets at the mark; a " inside a text must be written as \\"'


def open_brackets(text: str, end: int) -> str:
    """The closing brackets that are still owed at position `end`, innermost last ('}]}' for '{"a": [{')."""
    stack: list[str] = []
    in_string = escaped = False
    for char in text[:end]:
        if in_string:
            escaped, in_string = (False, True) if escaped else (char == "\\", char != '"')
        elif char == '"':
            in_string = True
        elif char in "[{":
            stack.append("]" if char == "[" else "}")
        elif char in "]}" and stack and stack[-1] == char:
            stack.pop()
    return "".join(stack)


HTML_BREAK_RE = re.compile(r"<br\s*/?>", re.IGNORECASE)
ESCAPED_UNICODE_RE = re.compile(r"\\u([0-9a-fA-F]{4})")
# A backslash that starts something other than \\uXXXX or \\n (than \\n). Searched for instead of matching the whole
# text: such a pattern with \\n in two places took quadratic time ('\\n' * 30000 + '\\x': seconds, holding the GIL).
OTHER_THAN_UNICODE_RE = re.compile(r"\\(?!u[0-9a-fA-F]{4}|n)")
OTHER_THAN_NEWLINE_RE = re.compile(r"\\(?!n)")


def _unescape_newlines(value: Any) -> Any:
    """Line breaks written in other ways: '<br>' (from Markdown tables) and 'Line 1\\nLine 2' (escaped twice).
    An escaped one is only replaced in texts whose backslashes are all followed by n, so a path like
    'C:\\neu\\Ordner' stays as it is."""
    if isinstance(value, str):
        value = HTML_BREAK_RE.sub("\n", value)
        if "\\u" in value and not OTHER_THAN_UNICODE_RE.search(value):  # '\\u2013' escaped twice
            value = ESCAPED_UNICODE_RE.sub(lambda m: chr(int(m[1], 16)), value)
        escaped = "\\n" in value and "\n" not in value and not OTHER_THAN_NEWLINE_RE.search(value)
        return value.replace("\\n", "\n") if escaped else value
    if isinstance(value, dict):
        return {k: _unescape_newlines(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_unescape_newlines(item) for item in value]
    return value


GARBLED_KEY_RE = re.compile(r'",[a-z_]+:|\}\s*,\s*\{\s*[a-z_]+:')  # 'paragraph",text:', '},{text:' in a text
PLAIN_KEY_RE = re.compile(r'[^"{}\[\]:\\<>|]+')  # '"name' or '{type' as a key: the syntax broke there


def _garbled(value: Any) -> bool:
    if isinstance(value, str):
        return bool(GARBLED_KEY_RE.search(value))
    if isinstance(value, dict):
        return any(not PLAIN_KEY_RE.fullmatch(str(k)) or _garbled(v) for k, v in value.items())
    if isinstance(value, list):
        return any(_garbled(item) for item in value)
    return False


def _contains(value: Any, token: str) -> bool:
    if isinstance(value, str):
        return token in value
    if isinstance(value, dict):
        return any(_contains(k, token) or _contains(v, token) for k, v in value.items())
    if isinstance(value, list):
        return any(_contains(item, token) for item in value)
    return False


TABLE_LINE_RE = re.compile(r"^\s*\|.*\|\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
LIST_LINE_RE = re.compile(r"^(?P<indent>\s*)(?:(?P<bullet>[-*•])|(?P<number>\d{1,3})[.)])\s+(?P<text>\S.*)$")
HEADING_LINE_RE = re.compile(r"^\s{0,3}(?P<level>#{1,6})\s+(?P<text>\S.*?)(?:\s+#+)?\s*$")
RULE_LINE_RE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")  # ---, ***, ___


def split_paragraph(text: str) -> list[tuple[str, list[str]]]:
    """Split paragraph text into ("text", lines), ("bullets", items), ("numbered", items), ("table", lines) and
    ("heading1" to "heading3", [text]) parts.

    Models often write Markdown into a paragraph: lists (indented list lines become '- ' sub items), tables,
    headings ('## Abschnitt') and rules ('---', left out). Numbered lines only form a list when they count up
    from 1, so lines that merely start with a number ('3. Oktober ist Feiertag.') stay text.
    """
    parts: list[tuple[str, list[str]]] = []
    for line in text.strip().split("\n"):
        if RULE_LINE_RE.match(line):
            parts.append(("break", []))
            continue
        if heading := HEADING_LINE_RE.match(line):
            parts.append((f"heading{min(len(heading['level']), MAX_HEADING_LEVEL)}", [heading["text"]]))
            continue
        if TABLE_LINE_RE.match(line):  # a Markdown table row
            if parts and parts[-1][0] == "table":
                parts[-1][1].append(line)
            else:
                parts.append(("table", [line]))
            continue
        match = LIST_LINE_RE.match(line)
        indented = bool(match) and len(match["indent"].expandtabs(4)) >= 2
        if match and match["number"] and not indented:
            expected = len([i for i in parts[-1][1] if not i.startswith(SUB_ITEM_PREFIX)]) + 1 if (
                parts and parts[-1][0] == "numbered"
            ) else 1  # fmt: skip
            if int(match["number"]) != expected:
                match = None
        if not line.strip():
            kind, item = "break", ""
        elif match:
            kind = "numbered" if match["number"] else "bullets"
            item = (SUB_ITEM_PREFIX if indented else "") + match["text"].strip()
        else:
            kind, item = "text", line.strip()
        if kind == "break":
            parts.append(("break", []))
        elif parts and (
            parts[-1][0] == kind or (parts[-1][0] in ("bullets", "numbered") and item.startswith(SUB_ITEM_PREFIX))
        ):
            parts[-1][1].append(item)
        else:
            parts.append((kind, [item]))
    result = []
    for kind, lines in parts:
        if kind == "table" and not (len(lines) >= 2 and any(TABLE_SEPARATOR_RE.match(line) for line in lines)):
            kind = "text"  # lines with | but no table
        if kind != "break":
            result.append((kind, lines))
    return result


def markdown_table_rows(lines: list[str]) -> list[list[str]]:
    return [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in lines
        if not TABLE_SEPARATOR_RE.match(line)
    ]


GERMAN_QUOTE_RE = re.compile(r'„([^"“”„\n]{0,79}[^"“”„\n\\])"')  # „Text" -> „Text“ (not „Text\")
MAX_JSON_REPAIRS = 20  # repairs of one text
MAX_REPAIR_TRIES = 60  # decoding attempts of one text: a slip can have several repairs, only one fits
# Other tool parameters a model may write into the same JSON text after the value ('[...], "title": "T"').
TOOL_PARAMETERS = {"title", "subtitle", "header_text", "table_of_contents", "landscape"}
TRAILING_PARAMETER_RE = re.compile(r',\s*"(\w+)"\s*:')
LITERAL_BEFORE_RE = re.compile(r":\s*(?:-?\d[\d.eE+-]*|true|false|null)\s*$")  # '"level": 1' before a stray '"'
UNQUOTED_KEY_RE = re.compile(r'[A-Za-z_]\w*"\s*:')  # 'level": 2': the opening quote of a key is missing


def _reject_constant(name: str) -> Any:
    raise ValueError(f"{name} is not a valid number, use null for empty values")


def _no_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    """An object whose key comes twice with different values: after a wrong repair, the keys of the next object
    land in this one and replace its content. The same value twice ("type": "table" written twice) loses nothing."""
    seen: dict[str, Any] = {}
    for key, value in pairs:
        if key in seen and seen[key] != value:
            raise ValueError(f"the key '{key}' twice")
        seen[key] = value
    return seen


JSON_DECODER = json.JSONDecoder(parse_constant=_reject_constant)
REPAIR_DECODER = json.JSONDecoder(parse_constant=_reject_constant, object_pairs_hook=_no_duplicate_keys)


def loads_lenient(text: str, extras: dict[str, Any] | None = None, accept: Callable[[Any], bool] | None = None) -> Any:
    """json.loads that repairs the mistakes models make in long arguments sent as text.

    - German quotes closed with a straight quote („Hoch") end the JSON string early; they become „Hoch“.
    - A forgotten '}' shows up where the next object starts in place of a key ('…"}, {"layout"…');
      the missing brace is added before that comma.
    - Arrays listed without their outer brackets ('"table": ["A", "B"], ["1", "2"]') get them back.
    - A bracket that does not fit where it stands: two closing brackets in the wrong order ('null]}]' for
      'null]]}') are swapped, one too many is deleted, forgotten ones are added before it; an object whose '{' is
      missing in a list ('}, "name": …') gets it.
    - A stray quote after a number ('"level": 1"}') is deleted, a missing opening quote of a key ('level": 2')
      added.
    - Other forgotten closing brackets are added where a different bracket closes, or at the end if the
      text ends with a complete object or array.
    - Other tool parameters written into the same text after the value ('[...], "title": "T"') are
      put into `extras`.
    The text is only changed after strict parsing failed, and only brackets and quotes are changed. Where a slip
    allows several repairs, they are tried in turn. A repair is only accepted if nothing of the text is lost: no
    text left over after the value, no text cut off in the middle of a value, no key twice in an object (the keys
    of another object), and the value passes `accept` (the tool's schema); otherwise the original error is raised
    so the model sends the arguments again.
    """
    try:
        return JSON_DECODER.decode(text)
    except json.JSONDecodeError as first_error:
        error = first_error
    stripped = text.strip()
    for candidate in dict.fromkeys((GERMAN_QUOTE_RE.sub("„\\1“", stripped), stripped)):
        repaired = _repair_json(candidate, accept)
        if repaired is not None:
            value, parameters = repaired
            if extras is not None:
                extras.update(parameters)
            return value
    raise error


def _repair_json(text: str, accept: Callable[[Any], bool] | None = None) -> tuple[Any, dict[str, Any]] | None:
    """Return (value, tool parameters written after it), or None if the text cannot be repaired without losing
    content. Searches the repairs depth first: each failed decoding offers the repairs at its place."""
    tries = 0

    def search(text: str, depth: int, bracket_fix_used: bool) -> tuple[Any, dict[str, Any]] | None:
        nonlocal tries
        tries += 1
        try:
            value, end = REPAIR_DECODER.raw_decode(text)
        except json.JSONDecodeError as e:
            if depth >= MAX_JSON_REPAIRS:
                return None
            for fixed, uses_bracket_fix in _json_repairs(text, e, bracket_fix_used):
                if tries >= MAX_REPAIR_TRIES:
                    return None
                if (found := search(fixed, depth + 1, bracket_fix_used or uses_bracket_fix)) is not None:
                    return found
            return None
        except ValueError:  # a key twice, NaN
            return None
        parameters = _trailing_parameters(text[end:])
        if parameters is None or (accept is not None and not accept(value)):
            return None
        return value, parameters

    return search(text, 0, False)


def _json_repairs(text: str, e: json.JSONDecodeError, bracket_fix_used: bool) -> Iterator[tuple[str, bool]]:
    """The repairs that fit the place of a decoding error, likeliest first: (repaired text, whether it is the
    bracket fix over the whole text, which is tried once)."""
    pos, char = e.pos, text[e.pos : e.pos + 1]
    comma = text.rfind(",", 0, pos)
    key_expected = e.msg == "Expecting property name enclosed in double quotes"
    if key_expected and char == "{" and comma >= 0:
        yield text[:comma] + "}" + text[comma:], False
    if key_expected and char == "[" and (wrapped := wrap_array_run(text, pos)):
        yield wrapped, False
    if key_expected and UNQUOTED_KEY_RE.match(text, pos):
        yield text[:pos] + '"' + text[pos:], False
    if char == '"' and LITERAL_BEFORE_RE.search(text[max(0, pos - 40) : pos]) and re.match(r'"\s*[,}\]]', text[pos:]):
        yield text[:pos] + text[pos + 1 :], False
    owed = open_brackets(text, pos)
    if char and char in "]}" and owed[-1:] != char:
        following = re.match(r"\s*([\]}])", text[pos + 1 :])
        if owed and following and following[1] == owed[-1] and owed[-2:-1] == char:  # '}]' for ']}'
            other = pos + 1 + following.start(1)
            yield text[:pos] + owed[-1] + text[pos + 1 : other] + char + text[other + 1 :], False
        if owed:  # '["E"}': the wrong bracket
            yield text[:pos] + owed[-1] + text[pos + 1 :], False
        if char in owed:  # it closes an outer one: the inner ones were forgotten
            inner = owed[owed.rindex(char) + 1 :]
            yield text[:pos] + inner[::-1] + text[pos:], False
        yield text[:pos] + text[pos + 1 :], False  # one too many
    if char == ":" and owed[-1:] == "]":  # '[…]}, "name": …' in a list: an object lost its '{'
        key = text.rfind('"', 0, text.rfind('"', 0, pos))
        before = text[:key].rstrip()
        if key > 0 and before.endswith(",") and before[:-1].rstrip().endswith("}"):
            start = _object_start(text, len(before[:-1].rstrip()) - 1)
            first = re.match(r'\{\s*("(?:[^"\\]|\\.)*")', text[start:]) if start is not None else None
            if (
                first and first[1] == text[key : text.index('"', key + 1) + 1]
            ):  # the same first key as the object before
                yield text[:key] + "{" + text[key:], False
    if not bracket_fix_used:
        fixed, appended = _close_brackets(text)
        if not (appended and not text.rstrip().endswith(("}", "]"))):  # not cut off in the middle of a value
            yield fixed, True


def _trailing_parameters(rest: str) -> dict[str, Any] | None:
    """Tool parameters written after the value; {} if nothing is left over, None if content would be lost."""
    rest = rest.strip()
    if not rest.strip("}] \t\r\n"):
        return {}
    match = TRAILING_PARAMETER_RE.match(rest)
    if not match or match[1] not in TOOL_PARAMETERS:
        return None
    body = "{" + rest[match.start(1) - 1 :]
    for candidate in (body, body + "}"):
        try:
            parameters, end = JSON_DECODER.raw_decode(candidate)
        except (json.JSONDecodeError, ValueError):
            continue
        if (
            isinstance(parameters, dict)
            and set(parameters) <= TOOL_PARAMETERS
            and not candidate[end:].strip("}] \t\r\n")
        ):
            return parameters
        return None  # other keys or more content after the parameters: it would be lost
    return None


ARRAY_VALUE_RE = re.compile(r'"(?:rows|table)"\s*:\s*\[')  # only table rows: wrapped list items would be lists in lists


def wrap_array_run(text: str, pos: int) -> str | None:
    """'"rows": ["a"], ["b"], ["c"]}' -> '"rows": [["a"], ["b"], ["c"]]}'; None if the text does not look like that.
    Only for table rows ("rows", "table"): list items written that way would become lists inside the list."""
    starts = [match.end() - 1 for match in ARRAY_VALUE_RE.finditer(text, 0, pos)]
    if not starts:
        return None
    decoder = json.JSONDecoder()
    start, end = starts[-1], starts[-1]
    while True:
        try:
            _, end = decoder.raw_decode(text, end)
        except json.JSONDecodeError:
            return None
        following = re.match(r"\s*,\s*(?=\[)", text[end:])
        if not following:
            break
        end += following.end()
    if end <= pos:  # the arrays must reach past the place where the error was found
        return None
    return text[:start] + "[" + text[start:end] + "]" + text[end:]


def _object_start(text: str, close: int) -> int | None:
    """The position of the '{' that the '}' at `close` closes."""
    starts: list[int] = []
    in_string = escaped = False
    for i, char in enumerate(text[: close + 1]):
        if in_string:
            escaped, in_string = (False, True) if escaped else (char == "\\", char != '"')
        elif char == '"':
            in_string = True
        elif char in "[{":
            starts.append(i)
        elif char in "]}" and starts:
            start = starts.pop()
            if i == close:
                return start if char == "}" else None
    return None


def _close_brackets(text: str) -> tuple[str, int]:
    """Add closing brackets the text forgot (before a bracket that closes an outer one, and at the end); returns the
    text and the number of brackets added at the end."""
    out: list[str] = []
    stack: list[str] = []
    in_string = escaped = False
    for char in text:
        if in_string:
            escaped, in_string = (False, True) if escaped else (char == "\\", char != '"')
        elif char == '"':
            in_string = True
        elif char in "[{":
            stack.append("]" if char == "[" else "}")
        elif char in "]}":
            while stack and stack[-1] != char:
                out.append(stack.pop())
            if stack:
                stack.pop()
        out.append(char)
    return "".join(out) + "".join(reversed(stack)), len(stack)


class Checker:
    """Collects limit violations so the model gets all of them at once."""

    def __init__(self) -> None:
        self.errors: list[str] = []
        self.chars = 0

    def require(self, condition: bool, message: str) -> bool:
        if not condition:
            self.errors.append(message)
        return condition

    def text(self, *texts: str) -> None:
        self.chars += sum(len(t) for t in texts)

    def raise_if_failed(self) -> None:
        self.require(self.chars <= MAX_TEXT_CHARS, f"the content is too long (more than {MAX_TEXT_CHARS} characters)")
        if self.errors:
            raise InvalidInput(self.errors)


# ----------------------------------------------------------------------------------- inline markdown --

INLINE_RE = re.compile(
    r"\\(?P<escaped>[\\`*_{}\[\]()#+\-.!|~<>])"
    r"|`(?P<code>[^`]+)`"
    r"|!?\[(?P<label>[^\]]*)\]\((?P<url>(?:[^()\s]|\([^()\s]*\))+)(?:\s+\"[^\"]*\")?\)"
    r"|<(?P<autolink>(?:https?://|mailto:)[^>\s]+)>"
    r"|\*\*\*(?P<bold_italic>.+?)\*\*\*"
    # Bold may contain *italic*, also at its end ('**a *b***'). The alternatives must not overlap, or text
    # like '***a*a*a…' takes exponential time (and the regex holds the GIL, freezing all of Open WebUI).
    r"|\*\*(?P<bold>(?:[^*]|\*[^*\s][^*]*\*(?=\*\*)|\*(?!\*))+?)\*\*"
    r"|(?<!\w)__(?P<bold_u>.{1,500}?)__(?!\w)"  # bounded: unclosed openers would each scan the rest of the text
    r"|~~(?P<strike>.+?)~~"
    r"|\*(?P<italic>[^*\s](?:[^*]*[^*\s])?)\*"
    r"|(?<!\w)_(?P<italic_u>[^_\s](?:[^_]*[^_\s])?)_(?!\w)",
    re.DOTALL,
)
EMPHASIS_GROUPS: dict[str, dict[str, bool]] = {
    "bold_italic": {"bold": True, "italic": True},
    "bold": {"bold": True},
    "bold_u": {"bold": True},
    "strike": {"strike": True},
    "italic": {"italic": True},
    "italic_u": {"italic": True},
}
BROKEN_QUOTE = '<|"|>'  # Gemma's string delimiter; leaks into arguments when its tool-call syntax breaks
XML_ILLEGAL_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]+')


@dataclass(frozen=True)
class Run:
    """A piece of inline text with uniform formatting."""

    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False
    url: str | None = None


PLAIN = Run("")


def clean_text(text: str) -> str:
    """Remove characters that are invalid in XML (docx/xlsx/pptx) and leaked model tokens."""
    return XML_ILLEGAL_RE.sub("", text.replace(BROKEN_QUOTE, '"'))


def parse_inline(text: str, style: Run = PLAIN) -> list[Run]:
    """Split inline Markdown into formatted runs (emphasis may be nested)."""
    text = clean_text(text)
    runs: list[Run] = []
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            runs.append(replace(style, text=text[position : match.start()]))
        runs.extend(_runs_for_match(match, style))
        position = match.end()
    if position < len(text):
        runs.append(replace(style, text=text[position:]))
    return runs


def _runs_for_match(match: re.Match[str], style: Run) -> list[Run]:
    kind = match.lastgroup or ""
    if kind == "escaped":
        return [replace(style, text=match["escaped"])]
    if kind == "code":
        return [replace(style, text=match["code"], code=True)]
    if kind == "url":
        label = to_plain(match["label"]) or match["url"]
        url = match["url"] if is_safe_url(match["url"]) else None
        return [replace(style, text=label, url=url)]
    if kind == "autolink":
        return [replace(style, text=match["autolink"], url=match["autolink"])]
    return parse_inline(match[kind], replace(style, **EMPHASIS_GROUPS[kind]))


def to_plain(text: str) -> str:
    return "".join(run.text for run in parse_inline(text)).strip()


def excel_text(text: str) -> str:
    """Text for an Excel cell: Excel cells have no inline formatting, so Markdown markers are removed."""
    text = clean_text(text)
    return to_plain(text) if any(marker in text for marker in "*_`[~<\\") else text


MAILTO_FIELDS = {"subject", "body", "cc", "bcc"}  # others, e.g. "attach", are acted on by some mail clients
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def is_safe_url(url: str) -> bool:
    url = url.strip()
    if not url.lower().startswith(SAFE_LINK_PREFIXES):
        return False
    if url.lower().startswith("mailto:") and "?" in url:
        fields = [part.split("=", 1)[0].lower() for part in url.split("?", 1)[1].split("&") if part]
        return all(field in MAILTO_FIELDS for field in fields)
    return True


def safe_filename(text: str, fallback: str) -> str:
    """Turn text into a short file name that is valid on Windows, macOS and Linux.

    Invisible format characters are removed: a right-to-left override would make 'x‮fdp.docx' look
    like 'xcod.pdf'.
    """
    text = "".join(char for char in to_plain(text) if unicodedata.category(char) != "Cf")
    name = " ".join(UNSAFE_FILENAME_RE.sub(" ", text).split())[:MAX_FILENAME_CHARS].strip(" .")
    if name.split(".")[0].upper() in WINDOWS_RESERVED_NAMES:
        name += "_"
    return name or fallback


def tidy_title(text: str) -> str:
    """Plain-text title; a file-name-like title ('Report_Q3_2026') gets its spaces back."""
    title = re.sub(r"\.(docx?|xlsx?|pptx?)$", "", " ".join(to_plain(text).split()), flags=re.IGNORECASE)
    if "_" in title and " " not in title:
        title = " ".join(title.replace("_", " ").split())
    return title


def similar_text(a: str, b: str) -> bool:
    """One text contains the other, or they differ only slightly (a title repeated as a heading)."""
    ka, kb = (re.sub(r"\W+", "", to_plain(t)).casefold() for t in (a, b))
    if not ka or not kb:
        return False
    return (min(len(ka), len(kb)) >= 12 and (ka in kb or kb in ka)) or difflib.SequenceMatcher(
        None, ka, kb
    ).ratio() >= 0.9


def same_text(a: str, b: str) -> bool:
    def key(text: str) -> str:
        return re.sub(r"\W+", "", to_plain(text)).casefold()

    return key(a) == key(b)


def split_level(item: str) -> tuple[int, str]:
    """'- text' is a level-1 (indented) item, everything else level 0."""
    stripped = item.strip()
    if stripped.startswith(SUB_ITEM_PREFIX):
        return 1, stripped[len(SUB_ITEM_PREFIX) :].strip()
    return 0, stripped


# ------------------------------------------------------------------------------------------------ Word --

ORDERED_LIST_STYLES = ("List Number", "List Number 2")
BULLET_LIST_STYLES = ("List Bullet", "List Bullet 2")
MAX_HEADING_LEVEL = 3
LONG_DOCUMENT_CHARS = 4000  # roughly two pages
W_NAMESPACE = qn("w:x")[:-1]
SETTINGS_AFTER_UPDATE_FIELDS = {
    *(qn(f"w:{name}") for name in (
        "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "attachedSchema", "themeFontLang",
        "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
        "readModeInkLockDown", "smartTagType", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
    )),
}  # fmt: skip
MAX_TEMPLATE_BYTES = 20_000_000
# With a template, its styles for these are replaced: they get the design. The content uses Body Text, not the
# template's Normal style, which its header and footer build on.
TEMPLATE_STYLES = ("Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Quote", "TOC Heading",
                   "List Bullet", "List Bullet 2", "List Number", "List Number 2")  # fmt: skip
# Settings of a template that would follow every document: a mail merge (Word asks to run its SQL command), write
# protection, and the path of the template it was made from (Word looks for it, e.g. on a network share).
TEMPLATE_SETTINGS_REMOVED = ("w:attachedTemplate", "w:mailMerge", "w:documentProtection", "w:writeProtection")
DOTX_MAIN = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"


def read_word_template(path: str) -> bytes | None:
    """The Word file of the word_template valve, a .dotx turned into a .docx; None if there is none or it cannot be
    read (logged: the documents are made without it)."""
    path = path.strip()
    if not path:
        return None
    try:
        with open(path, "rb") as file:
            data = file.read(MAX_TEMPLATE_BYTES + 1)
        if len(data) > MAX_TEMPLATE_BYTES:
            raise ValueError(f"the file is larger than {MAX_TEMPLATE_BYTES // 1_000_000} MB")
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            if sum(item.file_size for item in package.infolist()) > 5 * MAX_TEMPLATE_BYTES:
                raise ValueError("the file is too large unpacked")
            types = package.read("[Content_Types].xml")
            if DOTX_MAIN not in types:
                return data
            buffer = io.BytesIO()  # python-docx opens documents only
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as document:
                for item in package.infolist():
                    content = package.read(item)
                    if item.filename == "[Content_Types].xml":
                        content = content.replace(DOTX_MAIN, DOCX_CT.WML_DOCUMENT_MAIN.encode())
                    document.writestr(item, content)
            return buffer.getvalue()
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as e:
        log.warning("The Word template %r cannot be used, documents are made without it: %s", path, e)
        return None


@dataclass
class WordOptions:
    title: str
    subtitle: str = ""
    header_text: str = ""
    table_of_contents: bool = False
    landscape: bool = False
    page_numbers: bool = True
    design: Design = DEFAULT_DESIGN
    template: bytes | None = None  # see read_word_template


def word_title(title: str, blocks: list[WordBlock], fallback: str = "") -> tuple[str, list[WordBlock]]:
    """Return the document title and the blocks without a first heading that only repeats it.

    Models often leave out the title and start with it as a heading instead, or pass a file name
    ('Minutes_2026_10_12'). Then the first heading becomes the title. Without either, the title is `fallback`;
    an empty title means no title block (a letter), the file name then falls back to the default title.
    """
    title = re.sub(r"\.(docx?)$", "", to_plain(title).strip(), flags=re.IGNORECASE)
    first = blocks[0] if blocks and blocks[0].type == "heading" and blocks[0].text.strip() else None
    file_name_like = "_" in title and " " not in title
    if first and (not title or same_text(title, first.text) or (file_name_like and first.level == 1)):
        return to_plain(first.text), blocks[1:]
    if first and similar_text(title, first.text):  # 'Protokoll der Teambesprechung' under '… vom 12.10.2026'
        return max(tidy_title(title), to_plain(first.text), key=len), blocks[1:]
    return tidy_title(title) or fallback, blocks


def is_empty_block(block: WordBlock) -> bool:
    """Blocks without content; models add empty paragraphs as spacing, they are left out."""
    if block.type in ("heading", "paragraph", "quote"):
        return not block.text.strip()
    if block.type in ("bullets", "numbered"):
        return not any(item.strip() for item in block.items)
    if block.type == "table":
        return not any(str(cell).strip() for row in block.rows for cell in row)
    return False


def check_word(options: WordOptions, blocks: list[WordBlock]) -> None:
    check = Checker()
    check.require(
        any(b.type not in ("page_break", "table_of_contents") for b in blocks) or bool(options.subtitle.strip()),
        "blocks: the document has no content",
    )
    check.require(len(blocks) <= MAX_WORD_BLOCKS, f"blocks: at most {MAX_WORD_BLOCKS} blocks")
    if options.table_of_contents or any(b.type == "table_of_contents" for b in blocks):
        body = sum(
            len(b.text) + sum(map(len, b.items)) + sum(len(c) for r in b.rows for c in r)
            for b in blocks
            if b.type != "heading"
        )
        check.require(
            any(b.type == "heading" for b in blocks) and body >= 100,
            "blocks: the document has a table of contents but no chapters with text; send the complete content "
            "(headings with their paragraphs, lists and tables)",
        )
    check.text(options.title, options.subtitle, options.header_text)
    for i, block in enumerate(blocks):
        where = f"blocks[{block._index if block._index >= 0 else i}] ({block.type})"
        check.text(block.text, *block.items, *(cell for row in block.rows for cell in row))
        tables = [block.rows] if block.type == "table" else []
        if block.type == "paragraph":  # Markdown tables in the text become tables too
            tables = [markdown_table_rows(lines) for kind, lines in split_paragraph(block.text) if kind == "table"]
        for rows in tables:
            table = "" if block.type == "table" else " (the Markdown table in the text)"
            columns = table_columns(rows)
            check.require(
                columns <= MAX_WORD_TABLE_COLUMNS, f"{where}: at most {MAX_WORD_TABLE_COLUMNS} columns{table}"
            )
            check.require(len(rows) <= MAX_WORD_TABLE_ROWS, f"{where}: at most {MAX_WORD_TABLE_ROWS} rows{table}")
    check.raise_if_failed()


def _set_cell_shading(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_field(paragraph: DocxParagraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, NUMPAGES, TOC ...). Word fills in the value when it lays out the page."""
    for kind, text in (
        ("begin", None),
        ("instr", instruction),
        ("separate", None),
        ("text", placeholder),
        ("end", None),
    ):
        run = paragraph.add_run()
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = f" {text} "
        elif kind == "text":
            run.text = text or ""
            continue
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


# Word wants the children of w:pPr and w:tblPr in a fixed order; new ones go before these.
PPR_AFTER_BORDERS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
TBLPR_AFTER_CELL_MARGINS = ("w:tblLook", "w:tblCaption", "w:tblDescription")
THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
PAPER_SIZES = {"a4": (Cm(21), Cm(29.7)), "letter": (Cm(21.59), Cm(27.94))}


def _set_borders(parent: Any, tag: str, successors: tuple[str, ...], sides: list[tuple[str, int, str]]) -> None:
    """Replace the w:pBdr of a paragraph (or style) or the w:tblBorders of a table. Sides in schema order (top,
    left, bottom, right, insideH, insideV), size in eighths of a point, 0 for no line."""
    old = parent.find(qn(tag))
    if old is not None:
        parent.remove(old)
    borders = OxmlElement(tag)
    for side, size, color in sides:
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single" if size else "nil")
        if size:
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:space"), "4" if tag == "w:pBdr" else "0")
            edge.set(qn("w:color"), color)
        borders.append(edge)
    parent.insert_element_before(borders, *successors)


def _set_cell_margins(table: Any, top: int, bottom: int) -> None:
    """Cell margins in twips; left and right stay at Word's 108 (CELL_PADDING)."""
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", 108), ("bottom", bottom), ("right", 108)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    table._tbl.tblPr.insert_element_before(margins, *TBLPR_AFTER_CELL_MARGINS)


def _set_style(style: Any, size: float, color: str, bold: bool = False, italic: bool = False,
               before: float | None = None, after: float | None = None, font: str = BODY_FONT) -> None:  # fmt: skip
    """Font, colour and spacing of a built-in style; its theme font and theme colour are removed, they would win."""
    style.font.name = font
    for attr in THEME_FONT_ATTRS:
        style.element.rPr.rFonts.attrib.pop(qn(attr), None)
    style.font.size, style.font.bold, style.font.italic = Pt(size), bold, italic
    style.font.color.rgb = RGBColor.from_string(color)
    for attr in ("w:themeShade", "w:themeTint"):
        style.element.rPr.color.attrib.pop(qn(attr), None)
    if before is not None:
        style.paragraph_format.space_before = Pt(before)
    if after is not None:
        style.paragraph_format.space_after = Pt(after)


CHAR_WIDTH = Pt(5.8)  # average character of 11 pt Calibri, bold header text included
CELL_PADDING = Pt(12)  # left and right cell margin
MAX_WORD_CHARS = 20  # longer words may wrap
MAX_TABLE_COLUMN_CHARS = 60


def table_columns(rows: list[list[str]]) -> int:
    """Number of columns: the header's, or more if a row has more non-empty cells (trailing empty cells don't count)."""
    used = [max((i + 1 for i, cell in enumerate(row) if str(cell).strip()), default=0) for row in rows]
    return max(len(rows[0]), *used, 1)


def number_columns(rows: list[list[str]], columns: int) -> list[bool]:
    """Columns whose cells below the header are all numbers ('4.800 €', '12 %', '-3,5'; empty cells and placeholders
    such as '-' aside); Word and slide tables align them right."""
    result = []
    for c in range(columns):
        texts = [to_plain(str(row[c])) for row in rows[1:] if c < len(row) and str(row[c]).strip()]
        texts = [text for text in texts if text.lower() not in PLACEHOLDERS]
        result.append(bool(texts) and all(parse_number_text(text) is not None for text in texts))
    return result


def column_widths(rows: list[list[str]], columns: int, total: int, font_size: float = 11) -> list[int]:
    """Split the text width: every column fits its longest word, the rest goes to the columns with more text."""
    minimum, wanted = [], []
    char_width = CHAR_WIDTH * font_size / 11
    for column in range(columns):
        texts = [to_plain(row[column]) for row in rows if column < len(row)]
        longest_word = max((len(word) for text in texts for word in text.split()), default=1)
        longest_text = max((len(text) for text in texts), default=1)
        minimum.append(min(longest_word, MAX_WORD_CHARS) * char_width + CELL_PADDING)
        wanted.append(max(min(longest_text, MAX_TABLE_COLUMN_CHARS) * char_width + CELL_PADDING, minimum[-1]))
    if sum(wanted) <= total:  # everything fits on one line: widen all columns evenly
        widths = [w * total / sum(wanted) for w in wanted]
    elif sum(minimum) >= total:  # not even the words fit: shrink proportionally
        widths = [m * total / sum(minimum) for m in minimum]
    else:  # words fit: share the rest by how much more each column wants
        extra = total - sum(minimum)
        widths = [m + extra * (w - m) / (sum(wanted) - sum(minimum)) for m, w in zip(minimum, wanted, strict=True)]
    result = [int(w) for w in widths]
    result[-1] += total - sum(result)
    return result


class DocxWriter:
    """Render validated blocks into a Word document."""

    def __init__(self, options: WordOptions, texts: dict[str, str]) -> None:
        self.options = options
        self.texts = texts
        self.design = options.design
        self.template = False
        self.restyled = {"Header", "Footer"}  # with a template only if it has none: its own stay as they are
        if options.template:
            try:
                self.doc, self.template = Document(io.BytesIO(options.template)), True
                self._take_template()
                self._styles()
                return
            except Exception:
                log.warning("The Word template cannot be used, the document is made without it", exc_info=True)
                self.template, self.restyled = False, {"Header", "Footer"}
        self.doc = Document()
        self._styles()

    @property
    def body_style(self) -> str | None:
        return "Body Text" if self.template else None  # None: Normal

    def _take_template(self) -> None:
        """Keep the template's header, footer, page size and margins; remove its content and give the content
        styles the design (TEMPLATE_STYLES)."""
        body = self.doc.element.body
        last = body.get_or_add_sectPr()  # the last section's settings, with its header and footer
        for tag in ("w:headerReference", "w:footerReference"):  # a section without its own takes the previous one's
            for earlier in reversed(body.xpath("./w:p/w:pPr/w:sectPr")):
                for reference in earlier.findall(qn(tag)):
                    if not any(own.get(qn("w:type")) == reference.get(qn("w:type")) for own in last.findall(qn(tag))):
                        last.insert(0, copy.deepcopy(reference))
        for child in list(body):
            if child is not last:
                body.remove(child)
        default = Document()
        styles = self.doc.styles.element
        own = {(style.name_val or "").casefold() for style in styles.findall(qn("w:style"))}
        self.restyled = {name for name in ("Header", "Footer") if name.casefold() not in own}
        # The template's default character style, which Word names after its language
        default_font = next(
            (style.styleId for style in styles.findall(qn("w:style"))
             if style.get(qn("w:type")) == "character" and style.get(qn("w:default")) in ("1", "true")),
            None,
        )  # fmt: skip
        renamed = {"Normal": "BodyText", "DefaultParagraphFont": default_font}
        for name in (*TEMPLATE_STYLES, *sorted(self.restyled)):
            style = default.styles[name].element
            link = style.find(qn("w:link"))
            linked = default.styles.element.get_by_id(link.get(qn("w:val"))) if link is not None else None
            for new in (style, linked) if linked is not None else (style,):
                for old in styles.findall(qn("w:style")):
                    if (old.name_val or "").casefold() == new.name_val.casefold() or old.styleId == new.styleId:
                        old_link = old.find(qn("w:link"))  # e.g. "Überschrift 1 Zchn" of "Überschrift 1"
                        old_linked = styles.get_by_id(old_link.get(qn("w:val"))) if old_link is not None else None
                        if old_linked is not None:
                            styles.remove(old_linked)
                        if old.getparent() is not None:
                            styles.remove(old)
                new = copy.deepcopy(new)
                for tag in ("w:basedOn", "w:next"):
                    reference = new.find(qn(tag))
                    if reference is None or reference.get(qn("w:val")) not in renamed:
                        continue
                    if new.styleId == "BodyText" or not renamed[reference.get(qn("w:val"))]:
                        new.remove(reference)  # nothing of the template's Normal
                    else:
                        reference.set(qn("w:val"), renamed[reference.get(qn("w:val"))])
                styles.append(new)
        self._take_numbering(default)
        try:
            settings_part = self.doc.part.part_related_by(RELATIONSHIP_TYPE.SETTINGS)
        except KeyError:
            settings_part = None  # a file without settings
        for tag in TEMPLATE_SETTINGS_REMOVED if settings_part is not None else ():
            for element in settings_part.element.findall(qn(tag)):
                settings_part.element.remove(element)
                for node in element.iter():
                    if node.get(qn("r:id")):
                        settings_part.drop_rel(node.get(qn("r:id")))
        core = self.doc.core_properties  # the template's author, title ... are not the document's
        core.author = core.last_modified_by = core.subject = core.keywords = core.comments = core.category = ""
        core.title = to_plain(self.options.title)

    def _take_numbering(self, default: Any) -> None:
        """The list styles now point to the default template's numbering: copy it in, with new numbers."""
        source = default.part.numbering_part.element
        try:
            numbering = self.doc.part.part_related_by(RELATIONSHIP_TYPE.NUMBERING).element
        except KeyError:
            package, numbering = self.doc.part.package, parse_xml(f"<w:numbering {nsdecls('w')}/>")
            partname = package.next_partname("/word/numbering%d.xml")
            self.doc.part.relate_to(
                NumberingPart(partname, DOCX_CT.WML_NUMBERING, numbering, package), RELATIONSHIP_TYPE.NUMBERING
            )
        abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        for abstract_id, name in enumerate(TEMPLATE_STYLES[-4:], start=max(abstract_ids, default=-1) + 1):
            num_id = self.doc.styles[name].element.pPr.numPr.numId
            used = source.num_having_numId(num_id.val).abstractNumId.val
            abstract = copy.deepcopy(source.xpath(f"w:abstractNum[@w:abstractNumId='{used}']")[0])
            abstract.set(qn("w:abstractNumId"), str(abstract_id))
            for nsid in abstract.findall(qn("w:nsid")):  # Word's list identity; the template may use the same
                abstract.remove(nsid)
            numbering.insert_element_before(abstract, "w:num", "w:numIdMacAtCleanup")
            num_id.val = numbering.add_num(abstract_id).numId

    def _styles(self) -> None:
        """A calm, modern look: dark grey body text with some air, headings in the accent colour, no theme fonts."""
        styles, accent, font = self.doc.styles, self.design.accent, self.design.font
        body = styles[self.body_style or "Normal"]
        _set_style(body, 10.5, TEXT_COLOR, after=6, font=font)
        body.paragraph_format.line_spacing = 1.15
        _set_style(styles["Title"], 26, accent, bold=True, after=2, font=font)
        title_borders = styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
        if title_borders is not None:
            styles["Title"].element.pPr.remove(title_borders)  # the title block draws its own line
        _set_style(styles["Subtitle"], 13, MUTED_COLOR, after=2, font=font)
        _set_style(styles["Heading 1"], 16, accent, bold=True, before=20, after=6, font=font)
        _set_style(styles["Heading 2"], 13, TEXT_COLOR, bold=True, before=14, after=4, font=font)
        _set_style(styles["Heading 3"], 11, accent, bold=True, before=10, after=2, font=font)
        if "TOC Heading" in self._style_names():
            _set_style(styles["TOC Heading"], 16, accent, bold=True, before=12, after=6, font=font)
        for name in sorted(self.restyled):  # small grey text; their tab stops are for Letter paper and would win
            _set_style(styles[name], 8.5, MUTED_COLOR, after=0, font=font)
            tabs = styles[name].element.pPr.find(qn("w:tabs")) if styles[name].element.pPr is not None else None
            if tabs is not None:
                styles[name].element.pPr.remove(tabs)
        quote = styles["Quote"]
        _set_style(quote, 10.5, MUTED_COLOR, italic=True, before=6, after=10, font=font)
        quote.paragraph_format.left_indent = Cm(0.4)
        _set_borders(quote.element.get_or_add_pPr(), "w:pBdr", PPR_AFTER_BORDERS, [("left", 18, accent)])
        self._for_previews()

    def _for_previews(self) -> None:
        """Word takes fonts and colours from the paragraph styles. Browser previews (docx-preview in Open WebUI)
        also read the document defaults, the theme fonts and the character styles linked to the paragraph
        styles, and they have no Symbol font for Word's bullets (boxes instead of dots). All get the design."""
        font, styles = self.design.font, self.doc.styles
        for fonts in styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts") if not self.template else ():
            for attr in THEME_FONT_ATTRS:
                fonts.attrib.pop(qn(attr), None)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(attr), font)
        for part in self.doc.part.package.iter_parts() if not self.template else ():  # the template's header and
            if str(part.partname).startswith("/word/theme/"):  # footer may use its theme fonts
                part._blob = re.sub(THEME_FONTS_RE.encode(), lambda m: m[1] + font.encode() + m[2], part.blob)
        for style in styles:
            link = style.element.find(qn("w:link"))
            if style.type != WD_STYLE_TYPE.PARAGRAPH or link is None or style.element.rPr is None:
                continue
            linked = styles.element.get_by_id(link.get(qn("w:val")))
            if linked is not None:  # e.g. "Heading 1 Char": same look as its paragraph style
                for old in linked.findall(qn("w:rPr")):
                    linked.remove(old)
                linked.append(copy.deepcopy(style.element.rPr))
        for level in self.doc.part.numbering_part.element.iter(qn("w:lvl")):
            kind = level.find(qn("w:numFmt"))
            if kind is None or kind.get(qn("w:val")) != "bullet":
                continue
            indent = level.find(f"{qn('w:pPr')}/{qn('w:ind')}")
            sub_item = indent is not None and int(indent.get(qn("w:left"), "0")) > 720
            level.find(qn("w:lvlText")).set(qn("w:val"), "\u2013" if sub_item else "\u2022")
            for old in level.findall(qn("w:rPr")):
                level.remove(old)
            props = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            color = OxmlElement("w:color")
            color.set(qn("w:val"), self.design.accent)
            props.append(fonts)
            props.append(color)
            level.append(props)

    def render(self, blocks: list[WordBlock]) -> bytes:
        self.long_document = (
            sum(len(b.text) + sum(map(len, b.items)) for b in blocks) + 40 * len(blocks) > LONG_DOCUMENT_CHARS
        )
        self._page_setup()
        last = self.doc.add_heading(to_plain(self.options.title), level=0) if self.options.title.strip() else None
        if self.options.subtitle.strip():
            last = self.doc.add_paragraph(style="Subtitle")
            self._add_runs(last, parse_inline(self.options.subtitle))
        if last is not None:  # without a title and subtitle (a letter) there is no title block
            _set_borders(last._p.get_or_add_pPr(), "w:pBdr", PPR_AFTER_BORDERS, [("bottom", 12, self.design.accent)])
            last.paragraph_format.space_after = Pt(18)
        if self.options.table_of_contents:
            self._table_of_contents()
        for block in blocks:
            match block.type:
                case "heading":
                    self.doc.add_heading(to_plain(block.text), level=min(max(block.level, 1), MAX_HEADING_LEVEL))
                case "paragraph":
                    self._paragraph(block.text)
                case "quote":
                    self._add_runs(self.doc.add_paragraph(style="Quote"), parse_inline(block.text.strip()))
                case "bullets" | "numbered":
                    self._list(block.items, ordered=block.type == "numbered")
                case "table":
                    self._table(block.rows)
                case "page_break":
                    self.doc.add_page_break()
                case "table_of_contents":
                    self._table_of_contents()
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()

    def _page_setup(self) -> None:
        """Paper, margins, header and footer; a template's own stay, its header and footer are not added to."""
        section = self.doc.sections[0]
        width, height = PAPER_SIZES[self.texts.get("paper", "letter")]
        page = {"page_width": width, "page_height": height, "left_margin": Cm(2.5), "right_margin": Cm(2.5),
                "top_margin": Cm(2.5), "bottom_margin": Cm(2.2), "header_distance": Cm(1.25),
                "footer_distance": Cm(1.25)}  # fmt: skip
        for name, value in page.items():
            if not self.template or getattr(section, name) is None:  # what a template lacks
                setattr(section, name, value)
        if self.options.landscape and section.page_width < section.page_height:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        if self.options.header_text.strip() and section.header.is_linked_to_previous:
            header = section.header.paragraphs[0]
            header.text = clean_text(self.options.header_text.strip())
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if self.options.page_numbers and section.footer.is_linked_to_previous:
            # title on the left, "Page X of Y" on the right, a thin line above
            footer = section.footer.paragraphs[0]
            style = footer.style
            while style is not None:  # a template's Footer style may have tab stops of its own (centre ...)
                for tab in style.paragraph_format.tab_stops:
                    footer.paragraph_format.tab_stops.add_tab_stop(tab.position, WD_TAB_ALIGNMENT.CLEAR)
                style = style.base_style
            width = section.page_width - section.left_margin - section.right_margin
            footer.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            _set_borders(footer._p.get_or_add_pPr(), "w:pBdr", PPR_AFTER_BORDERS, [("top", 4, RULE_COLOR)])
            footer.add_run(clean_text(to_plain(self.options.title))[:90] + "\t")
            before, _, rest = self.texts["page_footer"].partition("{page}")
            between, _, after = rest.partition("{pages}")
            footer.add_run(before)
            _add_field(footer, "PAGE", "1")
            footer.add_run(between)
            _add_field(footer, "NUMPAGES", "1")
            footer.add_run(after)

    def _table_of_contents(self) -> None:
        self.doc.add_paragraph(
            self.texts["toc_title"], style="TOC Heading" if "TOC Heading" in self._style_names() else None
        )
        _add_field(
            self.doc.add_paragraph(style=self.body_style), 'TOC \\o "1-3" \\h \\z \\u', self.texts["toc_placeholder"]
        )
        # Ask Word to update the fields when the file is opened, so the table of contents gets filled in.
        # The settings elements have a fixed order; updateFields comes before compat, rsids etc.
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            update = OxmlElement("w:updateFields")
            update.set(qn("w:val"), "true")
            follower = next(
                (
                    child
                    for child in settings
                    if isinstance(child.tag, str)
                    and (child.tag in SETTINGS_AFTER_UPDATE_FIELDS or not child.tag.startswith(W_NAMESPACE))
                ),
                None,
            )
            if follower is not None:
                follower.addprevious(update)
            else:
                settings.append(update)
        if self.long_document:  # a short document would get a nearly empty first page
            self.doc.add_page_break()

    def _paragraph(self, text: str) -> None:
        """Paragraphs separated by blank lines; Markdown list lines inside the text become real lists."""
        for kind, lines in split_paragraph(text):
            if kind == "text":
                self._add_runs(self.doc.add_paragraph(style=self.body_style), parse_inline("\n".join(lines)))
            elif kind == "table":
                self._table(markdown_table_rows(lines))
            elif kind.startswith("heading"):
                self.doc.add_heading(to_plain(lines[0]), level=int(kind[-1]))
            else:
                self._list(lines, ordered=kind == "numbered")

    def _style_names(self) -> set[str]:
        return {style.name for style in self.doc.styles}

    def _list(self, items: list[str], ordered: bool) -> None:
        styles = ORDERED_LIST_STYLES if ordered else BULLET_LIST_STYLES
        num_ids: list[int | None] = [None, None]  # per level; each list and sub list restarts at 1
        for item in items:
            level, text = split_level(item)
            if not text:
                continue
            paragraph = self.doc.add_paragraph(style=styles[level])
            paragraph.paragraph_format.space_after = Pt(3)
            if ordered:
                if level == 0:
                    num_ids[1] = None  # the next sub list starts again at 1
                if num_ids[level] is None:
                    num_ids[level] = self._new_numbering(paragraph.style)
                paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_ids[level]
            self._add_runs(paragraph, parse_inline(text))

    def _new_numbering(self, style: Any) -> int:
        """Create a numbering instance of the style's list that starts at 1 (python-docx has no API for it)."""
        numbering = self.doc.part.numbering_part.element
        style_num_id = style.element.pPr.numPr.numId.val
        num = numbering.add_num(numbering.num_having_numId(style_num_id).abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def _table(self, rows: list[list[str]]) -> None:
        """Header row in the accent colour, every second row lightly tinted, thin lines between the rows only."""
        columns = table_columns(rows)
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.autofit = False
        lines = [("top", 0, ""), ("left", 0, ""), ("bottom", 8, self.design.accent), ("right", 0, ""),
                 ("insideH", 4, RULE_COLOR), ("insideV", 0, "")]  # fmt: skip
        _set_borders(table._tbl.tblPr, "w:tblBorders", TBLPR_AFTER_BORDERS, lines)
        _set_cell_margins(table, top=50, bottom=50)
        section = self.doc.sections[-1]
        widths = column_widths(rows, columns, section.page_width - section.left_margin - section.right_margin)
        numbers = number_columns(rows, columns)
        cells = table._cells  # once: row.cells and column.cells rebuild the whole grid on every call
        for c, (column, width) in enumerate(zip(table.columns, widths, strict=True)):
            column.width = width  # read by LibreOffice
            for cell in cells[c::columns]:
                cell.width = width  # read by Word
        for row_index, (doc_row, row) in enumerate(zip(table.rows, rows, strict=True)):
            if row_index == 0:  # repeat the header row on every page
                header_flag = OxmlElement("w:tblHeader")
                header_flag.set(qn("w:val"), "true")
                doc_row._tr.get_or_add_trPr().append(header_flag)
            fill = self.design.accent if row_index == 0 else self.design.band if row_index % 2 == 0 else None
            for column, cell in enumerate(cells[row_index * columns : (row_index + 1) * columns]):
                text = row[column] if column < len(row) else ""
                paragraph = cell.paragraphs[0]
                if self.template:
                    paragraph.style = self.body_style
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if numbers[column]:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._add_runs(
                    paragraph, parse_inline(text, Run("", bold=row_index == 0)), WHITE if row_index == 0 else None
                )
                if fill:
                    _set_cell_shading(cell, fill)
        self.doc.add_paragraph(style=self.body_style).paragraph_format.space_after = Pt(4)

    def _add_runs(self, paragraph: DocxParagraph, runs: list[Run], color: str | None = None) -> None:
        for run in runs:
            docx_run = paragraph.add_run(run.text)
            docx_run.bold = run.bold or None
            docx_run.italic = run.italic or None
            if color:
                docx_run.font.color.rgb = RGBColor.from_string(color)
            if run.strike:
                docx_run.font.strike = True
            if run.code:
                docx_run.font.name = CODE_FONT
            if run.url:
                self._make_link(paragraph, docx_run, run.url, color or self.design.accent)

    @staticmethod
    def _make_link(paragraph: DocxParagraph, docx_run: Any, url: str, color: str) -> None:
        docx_run.font.underline = True
        docx_run.font.color.rgb = RGBColor.from_string(color)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        docx_run._r.addprevious(link)
        link.append(docx_run._r)


# The file as Markdown, for the "Content" tab of Open WebUI's file view (data.content, rendered as Markdown). The
# upload skips Open WebUI's own text extraction (no processing, no embeddings), so the tool writes it itself.
MAX_CONTENT_CHARS = 100_000


def _markdown_cell(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ").strip()


def markdown_table(rows: list[list[Any]]) -> str:
    width = max((len(row) for row in rows), default=0)
    lines = [
        "| " + " | ".join(_markdown_cell(row[c]) if c < len(row) else "" for c in range(width)) + " |" for row in rows
    ]
    if lines:
        lines.insert(1, "|" + " --- |" * width)
    return "\n".join(lines)


def markdown_items(items: list[str], ordered: bool = False) -> str:
    lines = []
    for item in items:
        level, text = split_level(item)
        if text:
            lines.append("   " * level + ("1. " if ordered else "- ") + text)
    return "\n".join(lines)


def word_markdown(options: WordOptions, blocks: list[WordBlock]) -> str:
    parts = [f"# {to_plain(options.title)}" if options.title.strip() else ""]
    parts.append(f"*{options.subtitle.strip()}*" if options.subtitle.strip() else "")
    for block in blocks:
        if block.type == "heading":
            parts.append("#" * (min(max(block.level, 1), MAX_HEADING_LEVEL) + 1) + " " + to_plain(block.text))
        elif block.type == "paragraph":
            parts.append(block.text.strip())
        elif block.type == "quote":
            parts.append("\n".join("> " + line for line in block.text.strip().splitlines()))
        elif block.type in ("bullets", "numbered"):
            parts.append(markdown_items(block.items, ordered=block.type == "numbered"))
        elif block.type == "table":
            parts.append(markdown_table(block.rows))
    return "\n\n".join(part for part in parts if part.strip())


def build_docx(options: WordOptions, blocks: list[WordBlock], texts: dict[str, str]) -> tuple[bytes, str]:
    """The file and its text as Markdown."""
    blocks = [block for block in blocks if not is_empty_block(block)]
    while blocks and blocks[-1].type == "page_break":  # would only add an empty page
        blocks = blocks[:-1]
    if any(block.type == "table_of_contents" for block in blocks):
        options = replace(options, table_of_contents=False)  # placed where the model put it instead
    check_word(options, blocks)
    return DocxWriter(options, texts).render(blocks), word_markdown(options, blocks)


# ----------------------------------------------------------------------------------------------- Excel --

GERMAN_FUNCTIONS = {
    "SUMME": "SUM", "MITTELWERT": "AVERAGE", "ANZAHL": "COUNT", "ANZAHL2": "COUNTA", "WENN": "IF", "WENNS": "IFS",
    "WENNFEHLER": "IFERROR", "UND": "AND", "ODER": "OR", "NICHT": "NOT", "RUNDEN": "ROUND", "AUFRUNDEN": "ROUNDUP",
    "ABRUNDEN": "ROUNDDOWN", "ZÄHLENWENN": "COUNTIF", "ZÄHLENWENNS": "COUNTIFS", "SUMMEWENN": "SUMIF",
    "SUMMEWENNS": "SUMIFS", "MITTELWERTWENN": "AVERAGEIF", "MITTELWERTWENNS": "AVERAGEIFS", "MAXWENNS": "MAXIFS",
    "MINWENNS": "MINIFS", "PRODUKT": "PRODUCT", "POTENZ": "POWER", "WURZEL": "SQRT", "REST": "MOD",
    "GANZZAHL": "INT", "BETRAG": "ABS", "SVERWEIS": "VLOOKUP", "XVERWEIS": "XLOOKUP", "VERGLEICH": "MATCH",
    "DATUM": "DATE", "HEUTE": "TODAY", "JAHR": "YEAR", "MONAT": "MONTH", "TAG": "DAY", "EDATUM": "EDATE",
    "MONATSENDE": "EOMONTH", "NETTOARBEITSTAGE": "NETWORKDAYS", "WOCHENTAG": "WEEKDAY", "ZEIT": "TIME",
    "STUNDE": "HOUR", "LINKS": "LEFT", "RECHTS": "RIGHT", "LÄNGE": "LEN", "VERKETTEN": "CONCATENATE",
    "TEXTKETTE": "CONCAT", "TEXTVERKETTEN": "TEXTJOIN", "ISTLEER": "ISBLANK", "ISTZAHL": "ISNUMBER",
    "ISTTEXT": "ISTEXT",
    # not allowed, named in English in the message
    "ZEITWERT": "TIMEVALUE", "DATWERT": "DATEVALUE", "SUMMENPRODUKT": "SUMPRODUCT", "INDIREKT": "INDIRECT",
    "BEREICH.VERSCHIEBEN": "OFFSET", "ZELLE": "CELL", "WEBDIENST": "WEBSERVICE", "XMLFILTERN": "FILTERXML",
}  # fmt: skip
# The whitelist: (fewest arguments, most arguments, how the function is written); a wrong count gets the last back.
FUNCTION_ARGUMENTS: dict[str, tuple[int, int, str]] = {
    "SUM": (1, 255, "SUM(B2:B9)"),
    "AVERAGE": (1, 255, "AVERAGE(B2:B9)"),
    "MIN": (1, 255, "MIN(B2:B9)"),
    "MAX": (1, 255, "MAX(B2:B9)"),
    "MEDIAN": (1, 255, "MEDIAN(B2:B9)"),
    "COUNT": (1, 255, "COUNT(B2:B9)"),
    "COUNTA": (1, 255, "COUNTA(B2:B9)"),
    "COUNTIF": (2, 2, 'COUNTIF(range, criterion), e.g. COUNTIF($A$2:$A$9, "Ja")'),
    "COUNTIFS": (2, 254, "COUNTIFS(range1, criterion1, range2, criterion2, ...)"),
    "SUMIF": (2, 3, "SUMIF(range, criterion, sum_range), e.g. SUMIF($A$2:$A$9, A{row}, $C$2:$C$9)"),
    "SUMIFS": (3, 255, "SUMIFS(sum_range, range1, criterion1, range2, criterion2, ...)"),
    "AVERAGEIF": (2, 3, "AVERAGEIF(range, criterion, average_range)"),
    "AVERAGEIFS": (3, 255, "AVERAGEIFS(average_range, range1, criterion1, ...)"),
    "MAXIFS": (3, 255, "MAXIFS(max_range, range1, criterion1, ...)"),
    "MINIFS": (3, 255, "MINIFS(min_range, range1, criterion1, ...)"),
    "IF": (2, 3, 'IF(condition, value_if_true, value_if_false), e.g. IF(B{row}>100, "hoch", "niedrig")'),
    "IFS": (2, 254, "IFS(condition1, value1, condition2, value2, ...)"),
    "IFERROR": (2, 2, "IFERROR(value, value_if_error), e.g. IFERROR(B{row}/C{row}, 0)"),
    "AND": (1, 255, "AND(condition1, condition2, ...)"),
    "OR": (1, 255, "OR(condition1, condition2, ...)"),
    "NOT": (1, 1, "NOT(condition)"),
    "ROUND": (2, 2, "ROUND(value, digits), e.g. ROUND(B{row}*0.19, 2)"),
    "ROUNDUP": (2, 2, "ROUNDUP(value, digits)"),
    "ROUNDDOWN": (2, 2, "ROUNDDOWN(value, digits)"),
    "ABS": (1, 1, "ABS(value)"),
    "PRODUCT": (1, 255, "PRODUCT(B{row}, C{row})"),
    "POWER": (2, 2, "POWER(number, power)"),
    "SQRT": (1, 1, "SQRT(number)"),
    "MOD": (2, 2, "MOD(number, divisor)"),
    "INT": (1, 1, "INT(number)"),
    "VLOOKUP": (3, 4, "VLOOKUP(value, $A$2:$C$9, column_number, FALSE)"),
    "XLOOKUP": (3, 6, "XLOOKUP(value, $A$2:$A$9, $C$2:$C$9)"),
    "INDEX": (2, 3, "INDEX($C$2:$C$9, row_number)"),
    "MATCH": (2, 3, "MATCH(value, $A$2:$A$9, 0)"),
    "DATE": (3, 3, "DATE(year, month, day)"),
    "TODAY": (0, 0, "TODAY()"),
    "YEAR": (1, 1, "YEAR(date)"),
    "MONTH": (1, 1, "MONTH(date)"),
    "DAY": (1, 1, "DAY(date)"),
    "EDATE": (2, 2, "EDATE(start_date, months)"),
    "EOMONTH": (2, 2, "EOMONTH(start_date, months)"),
    "NETWORKDAYS": (2, 3, "NETWORKDAYS(start_date, end_date, holidays)"),
    "WEEKDAY": (1, 2, "WEEKDAY(date, 2)"),
    "TIME": (3, 3, "TIME(hours, minutes, seconds)"),
    "HOUR": (1, 1, "HOUR(time)"),
    "MINUTE": (1, 1, "MINUTE(time)"),
    "LEFT": (1, 2, "LEFT(text, count)"),
    "RIGHT": (1, 2, "RIGHT(text, count)"),
    "LEN": (1, 1, "LEN(text)"),
    "CONCATENATE": (1, 255, 'CONCATENATE(A{row}, " ", B{row})'),
    "CONCAT": (1, 254, 'CONCAT(A{row}, " ", B{row})'),
    "TEXTJOIN": (3, 252, 'TEXTJOIN(", ", TRUE, A{row}, B{row})'),
    "ISBLANK": (1, 1, "ISBLANK(B{row})"),
    "ISNUMBER": (1, 1, "ISNUMBER(B{row})"),
    "ISTEXT": (1, 1, "ISTEXT(B{row})"),
}
PAIRED_ARGUMENTS = {"SUMIFS": 1, "AVERAGEIFS": 1, "MAXIFS": 1, "MINIFS": 1, "COUNTIFS": 0, "IFS": 0}  # before the pairs
FORMULA_FUNCTIONS = set(FUNCTION_ARGUMENTS)
NEW_FUNCTIONS = {"XLOOKUP", "IFS", "MAXIFS", "MINIFS", "CONCAT", "TEXTJOIN"}  # Excel 2019 and later: _xlfn. in the file
# Functions that load data from outside or reach cells by computed addresses; a cell value with one stays text.
BLOCKED_FUNCTIONS = {
    "HYPERLINK", "WEBSERVICE", "FILTERXML", "INDIRECT", "OFFSET", "CELL", "INFO", "CALL", "REGISTER", "REGISTER.ID",
    "EXEC", "RTD", "DDE", "IMAGE", "IMPORTXML", "IMPORTDATA", "IMPORTRANGE", "ENCODEURL", "EVALUATE",
}  # fmt: skip
FUNCTION_ALTERNATIVES = {  # for functions models often try
    "TIMEVALUE": "give the column with the times the type time, its cells hold them as numbers: =C{row}-B{row}",
    "DATEVALUE": "give the column with the dates the type date, its cells hold them as numbers: =C{row}-B{row}",
    "TEXT": "the column type sets how a value is shown (date, time, currency, percent)",
    "SUMPRODUCT": "add a column with the product of each row (=B{row}*C{row}) and SUM it",
    "DATEDIF": "subtract the dates: =C{row}-B{row} gives the days between them",
    "SWITCH": "use IFS or IF",
}
FORMULA_CONSTANTS = {"TRUE": "TRUE", "FALSE": "FALSE", "WAHR": "TRUE", "FALSCH": "FALSE"}
FORMULA_TOKEN_RE = re.compile(  # [0-9], not \d: \d also matches digits of other scripts
    r"(?P<space>\s+)"
    r'|(?P<string>"(?:[^"]|"")*"|[„“”][^"„“”\n]*[“”])'  # "" is a " in a text; German quotes („ja“) are repaired
    r"|(?P<sheet>(?:'(?P<quoted>(?:[^']|'')+)'|(?P<bare>[^\W\d][\w.]*))!)"  # 'Rates 2026'!, Rates!
    r"|(?P<function>(?:_xlfn\.)?[A-Za-zÄÖÜäöü][A-Za-z0-9ÄÖÜäöü.]*)\s*\("
    r"|(?P<columns>\$?(?P<first>[A-Za-z]{1,3}):\$?(?P<last>[A-Za-z]{1,3}))(?![A-Za-z0-9_(])"  # whole columns: B:B
    r"|(?P<cell>\$?(?P<column>[A-Za-z]{1,3})(?P<row_absolute>\$?)(?P<row>[1-9][0-9]{0,6}|\{row\}))(?![A-Za-z0-9_])"
    r"|(?P<shifted>\$?[A-Za-z]{1,3}\$?\{row\s*[-+]\s*[0-9]+\})"  # B{row-1}
    r"|(?P<name>[A-Za-z]+)(?![A-Za-z0-9_(])"
    r"|(?P<number>[0-9]+(?:\.[0-9]+)?)"
    r"|(?P<operator><=|>=|<>|[-+*/^&=<>(),;:%])"
)
DECIMAL_COMMA_RE = re.compile(r",([0-9]+)(?![0-9A-Za-z(])")
MAX_EXCEL_COLUMN, MAX_EXCEL_ROW = 16_384, 1_048_576  # XFD, the last column; the last row
MAX_FORMULA_DEPTH = 30  # parentheses inside each other
OPERAND_KINDS = ("cell", "number", "string", "name", "columns")
FIRST_DATA_ROW = "2"
OPERATOR_LEVELS = (("=", "<>", "<", ">", "<=", ">="), ("&",), ("+", "-"), ("*", "/"), ("^",))  # lowest first
ARITHMETIC = {"+", "-", "*", "/", "^"}
SHIFTED_ROW_HINT = (
    "{row-1} (another row than the current one) is not possible in a column formula; write such formulas into the "
    "cells instead, e.g. a running balance in column D: '=B2-C2' in the first row, '=D2+B3-C3' in the second, "
    "'=D3+B4-C4' in the third"
)


def _column_number(letters: str) -> int:
    number = 0
    for letter in letters.upper():
        number = number * 26 + ord(letter) - ord("A") + 1
    return number


def sheet_lookup(sheets: list[ExcelSheet], names: list[str]) -> dict[str, str]:
    """The written sheet names by the names a formula may use for them: the written one and the one the model sent
    (before invalid characters were removed), both without case, like Excel (and sheet_names)."""
    lookup = {name.lower(): name for name in names}
    for sheet, name in zip(sheets, names, strict=True):
        lookup.setdefault(" ".join(to_plain(sheet.name).split()).lower(), name)
    return lookup


def _sheet_prefix(name: str) -> str:
    return "'" + name.replace("'", "''") + "'!"


class UnsafeFormula(ValueError):
    """A formula that reaches outside of the workbook or computes cell addresses (other workbooks, DDE, HYPERLINK,
    INDIRECT ...)."""


class FormulaToken(NamedTuple):
    kind: str  # string, function, cell, name, number, operator
    text: str  # in Excel's file syntax: English names, ',' between arguments, a cell with its sheet ('Rates'!$B$3)
    sheet: str = ""  # cells: the written name of another sheet, "" for the formula's own sheet
    column: str = ""  # cells: the column letters in upper case
    row: str = ""  # cells: the row number or "{row}"
    absolute: bool = False  # cells: the row is fixed ($3)


def _outside_texts(text: str) -> str:
    """The formula without its texts, split into tokens as parse_formula splits it: '"[neu]"', '„[neu]“' and
    '"a""[b]"' are texts, '[Book]S!A1' and 'cmd|x' are not."""
    parts, position = [], 1
    while position < len(text):
        match = FORMULA_TOKEN_RE.match(text, position)
        if match is None:
            parts.append(text[position])
            position += 1
        else:
            parts.append("" if match.lastgroup == "string" else match.group())
            position = match.end()
    return "".join(parts)


def parse_formula(
    formula: str,
    sheets: dict[str, str] | None = None,
    last_rows: dict[str, int] | None = None,
    cell: tuple[str, int, int] | None = None,
) -> list[FormulaToken]:
    """Check a formula against the whitelist and return its tokens in Excel's file syntax (without spaces).

    Only cell references of this workbook (other sheets by the names in `sheets`, see sheet_lookup), numbers,
    quoted text, operators and whitelisted functions are allowed. Other workbooks ([...]), DDE (|) and the
    BLOCKED_FUNCTIONS raise UnsafeFormula, everything else that is not allowed ValueError, with a message for the
    model. With ';' as separator (German Excel), commas in numbers are decimal commas; outside of functions a comma
    between digits is always one. Whole columns (B:B) become the data rows, from `last_rows` (the last data row by
    written sheet name, "" for the formula's own sheet). In a formula of one cell (`cell`: the written name of its
    sheet, its column number and row), whole columns with the cell itself end in the row above it: a total below
    the data ('=SUM(B:B)' in B4 adds up $B$2:$B$3).
    """
    text = XML_ILLEGAL_RE.sub("", formula).strip()
    if not text.startswith("="):
        text = "=" + text  # models sometimes leave out the "="
    if len(text) > MAX_FORMULA_CHARS:
        raise ValueError(f"a formula may have at most {MAX_FORMULA_CHARS} characters")
    if text.count('"') % 2:
        raise ValueError('a text in the formula is not closed: a text needs a " at both ends ("" is an empty text)')
    outside_texts = _outside_texts(text)
    if "[" in outside_texts or "|" in outside_texts:  # a sheet name cannot have [ ]
        raise UnsafeFormula("references to other workbooks ([...]) and DDE (|) are not allowed")
    german = ";" in re.sub(r"'(?:[^']|'')*'", "", outside_texts)  # outside of texts and sheet names
    hint = " (cells of another sheet: 'Sheet name'!B3)" if "!" in text else ""
    tokens: list[FormulaToken] = []
    sheet: str | None = None  # the sheet written before the next cell
    depth = 0
    position = 1
    previous = ""  # kind of the previous token, "space" included
    before_space = ""  # kind of the last token before a space
    while position < len(text):
        match = FORMULA_TOKEN_RE.match(text, position)
        if not match:
            raise ValueError(f"'{text[position : position + 10]}' is not allowed in a formula{hint}")
        kind, token = match.lastgroup or "", match.group()
        end = match.end()
        if (
            kind != "space"
            and previous == "space"
            and (before_space in OPERAND_KINDS or before_space == ")")
            and (kind in (*OPERAND_KINDS, "function", "sheet") or token == "(")
        ):
            near = text[position : position + 10]
            raise ValueError(f"a space between two values is not allowed near '{near}'; use an operator")
        if kind == "space":
            before_space = previous if previous != "space" else before_space
        elif kind == "sheet":
            name = match["quoted"].replace("''", "'") if match["quoted"] is not None else match["bare"]
            sheet = (sheets or {}).get(" ".join(name.split()).lower(), "")
            if not sheet:
                known = ", ".join(f"'{known_name}'" for known_name in dict.fromkeys((sheets or {}).values()))
                raise ValueError(
                    f"there is no sheet '{name}' in this workbook" + (f" (sheets: {known})" if known else "")
                )
            following = FORMULA_TOKEN_RE.match(text, end)
            if (tokens and tokens[-1].text == ":") or not following or following.lastgroup not in ("cell", "columns"):
                raise ValueError(f"write a cell of another sheet as '{sheet}'!B3 and a range as '{sheet}'!B2:B9")
        elif kind == "function":
            written = match["function"]
            name = GERMAN_FUNCTIONS.get(written.upper().removeprefix("_XLFN."), written.upper().removeprefix("_XLFN."))
            if name in BLOCKED_FUNCTIONS:
                raise UnsafeFormula(f"the function {written} is not allowed")
            if name not in FORMULA_FUNCTIONS:
                alternative = f": {FUNCTION_ALTERNATIVES[name]}" if name in FUNCTION_ALTERNATIVES else ""
                allowed = ", ".join(sorted(FORMULA_FUNCTIONS))
                raise ValueError(f"the function {written} is not allowed{alternative} (allowed: {allowed})")
            tokens.append(FormulaToken("function", ("_xlfn." if name in NEW_FUNCTIONS else "") + name + "("))
            depth += 1
        elif kind == "cell":
            column, row = match["column"].upper(), match["row"]
            if _column_number(column) > MAX_EXCEL_COLUMN or (row != "{row}" and int(row) > MAX_EXCEL_ROW):
                raise ValueError(f"'{token}' is outside of the sheet")
            owner = sheet  # the end of a range is on the sheet of its start: 'Rates'!B2:B9
            if owner is None:
                range_end = len(tokens) >= 2 and tokens[-1].text == ":" and tokens[-2].kind == "cell"
                owner = tokens[-2].sheet if range_end else ""
            written = (_sheet_prefix(sheet) if sheet else "") + token.upper().replace("{ROW}", "{row}")
            tokens.append(FormulaToken("cell", written, owner, column, row, bool(match["row_absolute"])))
            sheet = None
        elif kind == "columns":
            owner, first, last = sheet or "", match["first"].upper(), match["last"].upper()
            if last_rows is None or owner not in last_rows:
                raise ValueError(f"whole columns ({token}) are not allowed: write the rows, e.g. ${first}$2:${last}$50")
            rows = str(max(last_rows[owner], 2))  # the data rows: not the header, not the total row
            columns = sorted((_column_number(first), _column_number(last)))
            if cell and owner in ("", cell[0]) and columns[0] <= cell[1] <= columns[1] and 2 <= cell[2] <= int(rows):
                here = f"{get_column_letter(cell[1])}{cell[2]}"
                if cell[2] == 2 or cell[2] < int(rows):  # a total above or between the data: which rows is unclear
                    above, below = f"${first}$2:${last}${cell[2] - 1}", f"${first}$3:${last}${max(int(rows), 3)}"
                    example = above if cell[2] > 2 else below
                    raise ValueError(
                        f"{token} in {here} includes {here} itself (a circular reference): write the rows to "
                        f"compute with, e.g. {example}"
                    )
                rows = str(cell[2] - 1)  # a total below the data: the rows above it
            tokens += [
                FormulaToken("cell", f"{_sheet_prefix(sheet) if sheet else ''}${first}$2", owner, first, "2", True),
                FormulaToken("operator", ":"),
                FormulaToken("cell", f"${last}${rows}", owner, last, rows, True),
            ]
            sheet = None
        elif kind == "shifted":
            raise ValueError(SHIFTED_ROW_HINT)
        elif kind == "name":
            if token.upper() not in FORMULA_CONSTANTS:
                raise ValueError(f"'{token}' is not allowed in a formula (only cells like B{{row}} or B2){hint}")
            tokens.append(FormulaToken("name", FORMULA_CONSTANTS[token.upper()]))
        elif kind == "number":
            fraction = DECIMAL_COMMA_RE.match(text, end) if "." not in token and (german or depth == 0) else None
            if fraction:  # 0,19 -> 0.19
                token, end = f"{token}.{fraction[1]}", fraction.end()
            tokens.append(FormulaToken("number", token))
        elif kind == "string":
            tokens.append(FormulaToken("string", '"' + token[1:-1] + '"'))  # „ja“ -> "ja"
        elif kind == "operator":
            if token == ";":
                token = ","
            elif token == "(":
                depth += 1
            elif token == ")":
                depth -= 1
                if depth < 0:
                    raise ValueError("the formula has a ')' without '('")
            tokens.append(FormulaToken("operator", token))
        if depth > MAX_FORMULA_DEPTH:
            raise ValueError(f"the formula is nested too deeply (at most {MAX_FORMULA_DEPTH} levels of parentheses)")
        previous = kind if kind != "operator" or token != ")" else ")"
        position = end
    if depth != 0:
        raise ValueError("the parentheses in the formula do not match")
    return tokens


class FormulaNode(NamedTuple):
    """A part of a parsed formula (see formula_tree)."""

    kind: str  # number, string, bool, cell, range, unary, binary, percent, call, empty (an argument left out)
    text: str  # the operator, the function's name, or the token as written into the file
    args: tuple["FormulaNode", ...] = ()
    cell: FormulaToken | None = None  # cell: its token
    index: int = -1  # cell: its position in the token list


def formula_tree(tokens: list[FormulaToken]) -> FormulaNode:
    """Parse the tokens of parse_formula in Excel's order of operations; raises ValueError for formulas Excel cannot
    read (=B2*, =B2*/2, ==B2) and for a wrong number of arguments (=ROUND(B2), =SUM())."""
    return _FormulaParser(tokens).parse()


def walk(node: FormulaNode) -> Iterator[FormulaNode]:
    yield node
    for arg in node.args:
        yield from walk(arg)


def number_value(node: FormulaNode) -> float | None:
    return float(node.text) if node.kind == "number" else None


class _FormulaParser:
    """Recursive descent; unary minus binds closer than ^, as in Excel."""

    def __init__(self, tokens: list[FormulaToken]) -> None:
        self.tokens, self.i = tokens, 0

    def peek(self) -> str:
        return self.tokens[self.i].text if self.i < len(self.tokens) and self.tokens[self.i].kind == "operator" else ""

    def parse(self) -> FormulaNode:
        if not self.tokens:
            raise ValueError("the formula is empty")
        node = self.binary(0)
        if self.i < len(self.tokens):
            raise ValueError(self.unexpected())
        return node

    def unexpected(self) -> str:
        token = self.tokens[self.i].text
        if token == ",":
            return "',' separates the arguments of a function; outside of one it is not allowed"
        return f"'{token}' follows '{self.tokens[self.i - 1].text}' without an operator between them"

    def binary(self, level: int) -> FormulaNode:
        if level == len(OPERATOR_LEVELS):
            return self.unary()
        node = self.binary(level + 1)
        while self.peek() in OPERATOR_LEVELS[level]:
            operator = self.peek()
            self.i += 1
            node = FormulaNode("binary", operator, (node, self.binary(level + 1)))
        return node

    def unary(self) -> FormulaNode:
        if self.peek() in ("+", "-"):
            operator = self.peek()
            self.i += 1
            return FormulaNode("unary", operator, (self.unary(),))
        node = self.primary()
        while self.peek() == "%":
            self.i += 1
            node = FormulaNode("percent", "%", (node,))
        return node

    def primary(self) -> FormulaNode:
        if self.i >= len(self.tokens):
            raise ValueError(f"the formula ends with '{self.tokens[-1].text}': a value is missing after it")
        token = self.tokens[self.i]
        self.i += 1
        if token.kind in ("number", "string"):
            return FormulaNode(token.kind, token.text)
        if token.kind == "name":
            return FormulaNode("bool", token.text)
        if token.kind == "function":
            return self.call(token.text[:-1].removeprefix("_xlfn."))
        if token.kind == "cell":
            node = FormulaNode("cell", token.text, cell=token, index=self.i - 1)
            if self.peek() != ":":
                return node
            self.i += 1
            if self.i >= len(self.tokens) or self.tokens[self.i].kind != "cell":
                raise ValueError("':' joins two cells to a range, e.g. B2:B9")
            end = self.tokens[self.i]
            self.i += 1
            return FormulaNode("range", ":", (node, FormulaNode("cell", end.text, cell=end, index=self.i - 1)))
        if token.text == "(":
            node = self.binary(0)
            if self.peek() != ")":
                raise ValueError(self.unexpected())
            self.i += 1
            return node
        before = self.tokens[self.i - 2].text if self.i >= 2 else "="
        if token.text == "=" and self.i == 1:
            raise ValueError("the formula starts with '==': write a single '='")
        if token.text in (",", ")") and (before == "," or before.endswith("(")):
            raise ValueError(f"an argument is missing before '{token.text}'")
        if token.text == ":":
            raise ValueError("':' joins two cells to a range, e.g. B2:B9")
        raise ValueError(f"'{token.text}' follows '{before}' without a value between them")

    def call(self, name: str) -> FormulaNode:
        args: list[FormulaNode] = []
        if self.peek() == ")":
            self.i += 1
        else:
            while True:  # an argument after the first may be left out, as in Excel: IF(B2,,1), VLOOKUP(…,2,)
                empty = args and self.peek() in (",", ")")
                args.append(FormulaNode("empty", "") if empty else self.binary(0))
                if self.peek() not in (",", ")"):
                    raise ValueError(self.unexpected())
                self.i += 1
                if self.tokens[self.i - 1].text == ")":
                    break
        fewest, most, usage = FUNCTION_ARGUMENTS[name]
        before_pairs = PAIRED_ARGUMENTS.get(name)
        if not fewest <= len(args) <= most or (before_pairs is not None and (len(args) - before_pairs) % 2):
            count = f"{len(args)} argument{'' if len(args) == 1 else 's'}" if args else "no arguments"
            raise ValueError(f"{name} has {count} here; write it as {usage}")
        return FormulaNode("call", name, tuple(args))


def normalize_formula(
    formula: str,
    infer_row: bool = True,
    sheets: dict[str, str] | None = None,
    last_rows: dict[str, int] | None = None,
    cell: tuple[str, int, int] | None = None,
) -> str:
    """Check a formula (parse_formula, formula_tree) and return it in Excel's file syntax (English names, commas).
    A column formula (infer_row) gets its rows made clear, see column_formula_rows; `cell` is the place of a formula
    in one cell, see parse_formula. Raises ValueError with a message for the model."""
    tokens = parse_formula(formula, sheets, last_rows, cell)
    tree = formula_tree(tokens)
    if infer_row:
        tokens = column_formula_rows(tokens, tree)
    return "=" + "".join(token.text for token in tokens)


def _absolute(token: FormulaToken) -> FormulaToken:
    prefix = token.text[: token.text.rfind("!") + 1]
    return token._replace(text=f"{prefix}${token.column}${token.row}", absolute=True)


def column_formula_rows(tokens: list[FormulaToken], tree: FormulaNode) -> list[FormulaToken]:
    """The cells of a column formula, which is filled into every data row, with their rows made clear.

    A range over several rows without {row} ('SUM(B2:B9)', 'SUM($B$2:B9)') is the same range in every row and gets
    written as one ($B$2:$B$9); so does the start of 'B2:B{row}'. If every other cell of its own sheet is in the
    first data row ('=B2*C2', '=B2/SUM(B2:B4)'), these get {row}, so every row computes its own values. Any other cell
    without {row} or $ in one row ('=B3*C3', '=B1*C1', '=B{row}*C2') is sent back: models count the rows wrong (row 1
    is the header), or mean a fixed cell; cells of several rows ('SUM(D2*F2,D3*F3)') are fixed cells. A cell of
    another sheet has to say which: 'Rates'!$B$3 or 'Rates'!B{row}.
    """
    tokens = list(tokens)
    ends: dict[int, int] = {}
    for node in walk(tree):
        if node.kind == "range":
            a, b = (arg.index for arg in node.args)
            ends[a], ends[b] = b, a

    def open_row(token: FormulaToken) -> bool:
        return not token.absolute and token.row != "{row}"

    fixed, open_rows = set(), []
    for i, token in enumerate(tokens):
        if token.kind != "cell" or not open_row(token):
            continue
        other = tokens[ends[i]] if i in ends else None
        if other is not None and other.row != token.row:
            fixed.add(i)  # B2:B9, $B$2:B9, or the start of B2:B{row}
        else:
            open_rows.append(i)
    for i in fixed:
        tokens[i] = _absolute(tokens[i])
    own = [i for i in open_rows if not tokens[i].sheet]
    current_row = any(token.kind == "cell" and token.row == "{row}" and not token.sheet for token in tokens)
    if own and not current_row and all(tokens[i].row == FIRST_DATA_ROW for i in own):
        for i in own:
            tokens[i] = tokens[i]._replace(text=tokens[i].text[: -len(FIRST_DATA_ROW)] + "{row}", row="{row}")
        own = []
    other_sheet = [i for i in open_rows if tokens[i].sheet]
    if other_sheet:
        sheet, column, row = tokens[other_sheet[0]].sheet, tokens[other_sheet[0]].column, tokens[other_sheet[0]].row
        raise ValueError(
            f"'{sheet}'!{column}{row} in a column formula: write '{sheet}'!${column}${row} for this one cell in "
            f"every row, or '{sheet}'!{column}{{row}} for the cell in the same row"
        )
    if len({tokens[i].row for i in own}) == 1:  # cells of several rows ('SUM(D2*F2,D3*F3)') are meant as they are
        cells = ", ".join(dict.fromkeys(tokens[i].text for i in own))
        first = tokens[own[0]]
        header = " (row 1 is the header row)" if any(tokens[i].row == "1" for i in own) else ""
        raise ValueError(
            f"{cells} in a column formula: the formula is filled into every row, so a cell without {{row}} is the "
            f"same cell in all of them{header}. The data starts in row 2: write {first.column}{{row}} for the cell in "
            f"the same row (e.g. =B{{row}}*C{{row}}), ${first.column}${first.row} for one fixed cell, and $B$2:$B$9 "
            "for a fixed range"
        )
    return tokens


NUMBER_RE = re.compile(  # thousands separators: . , ' ’ space, no-break space, narrow no-break space
    r"(?P<prefix>€|\$|£|EUR|USD|GBP|CHF)?\s*"
    r"(?P<sign>[-+]?)"
    r"(?P<int>\d{1,3}(?:(?P<group>[.,' \u00a0\u202f\u2019])\d{3}(?:(?P=group)\d{3})*)?|\d+)"
    r"(?:(?!(?P=group))(?P<decimal>[.,])(?P<fraction>\d+))?"
    r"\s*(?P<suffix>%|€|\$|£|EUR|USD|GBP|CHF)?"
)
CURRENCY_CODES = {"€": "EUR", "EUR": "EUR", "$": "USD", "USD": "USD", "£": "GBP", "GBP": "GBP", "CHF": "CHF"}
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M",
                "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M")  # fmt: skip
DATE_LIKE_RE = re.compile(r"[0-9]{1,4}([./-])[0-9]{1,2}(?:\1[0-9]{1,4})?(?:[ T].*)?|[0-9]{1,2}[./][0-9]{4}")
MONTH_ONLY_RE = re.compile(r"[0-9]{4}-[0-9]{1,2}|[0-9]{1,2}[./-][0-9]{4}")
TIME_RE = re.compile(
    r"(?P<hours>[0-9]{1,3}):(?P<minutes>[0-5][0-9])(?::(?P<seconds>[0-5][0-9]))?(?:\s*(?:Uhr|h))?", re.IGNORECASE
)
TIME_LIKE_RE = re.compile(r"[0-9]{1,2}(?:[.,][0-9]{1,2})?\s*(?:Uhr|h)|[0-9]{1,2}[.,][0-9]{2}", re.IGNORECASE)  # 8 Uhr


def parse_number_text(text: str, decimal_comma: bool | None = None) -> tuple[float | int, bool] | None:
    """'1.234,56 €' -> (1234.56, False), '12,5 %' -> (0.125, True). None if the text is not a number.

    A dot or comma followed by exactly three digits is read as a thousands separator, except after a
    leading 0 ('0,125') and, if `decimal_comma` says which one is the decimal mark, for that mark when
    it appears only once ('1,799' is 1.799 in German, '1.500' is 1.5 in English).
    """
    match = NUMBER_RE.fullmatch(text.strip().replace("\u2212", "-"))
    if not match or (match["prefix"] and match["suffix"]):
        return None
    if match["group"] in (",", ".") and not match["fraction"]:
        head, *groups = match["int"].split(match["group"])
        decimal_mark = None if decimal_comma is None else "," if decimal_comma else "."
        if len(groups) == 1 and (head == "0" or match["group"] == decimal_mark):
            number = float(f"{match['sign']}{head}.{groups[0]}")
            return (number / 100, True) if match["suffix"] == "%" else (number, False)
    digits, fraction = re.sub(r"\D", "", match["int"]), match["fraction"] or ""
    number: float | int = float(f"{match['sign']}{digits}.{fraction}") if fraction else int(match["sign"] + digits)
    if match["suffix"] == "%":
        return number / 100, True
    return number, False


def value_currency(value: Any) -> str | None:
    """The currency written with an amount ('$1,200.00' -> 'USD'); None without one."""
    match = NUMBER_RE.fullmatch(value.strip().replace("\u2212", "-")) if isinstance(value, str) else None
    return CURRENCY_CODES.get(match["prefix"] or match["suffix"] or "") if match else None


def parse_date(text: str) -> date | datetime | None:
    for date_format in DATE_FORMATS:
        try:
            parsed = datetime.strptime(text.strip(), date_format)
        except ValueError:
            continue
        return parsed if "%H" in date_format else parsed.date()
    return None


def parse_time(text: Any) -> float | None:
    """'08:30' -> 0.354 (the fraction of a day, as Excel stores times), '25:30' -> 1.0625 (a duration); None if the
    text is no time."""
    match = TIME_RE.fullmatch(text.strip()) if isinstance(text, str) else None
    if not match:
        return None
    return (int(match["hours"]) * 3600 + int(match["minutes"]) * 60 + int(match["seconds"] or 0)) / 86400


def date_problem(value: Any) -> str | None:
    """Why a value of a date column cannot be written as a date; None for dates, and for words that may stay
    text (e.g. 'offen')."""
    if isinstance(value, int | float) and not isinstance(value, bool):
        return f"{value:g} is a number, not a date: write dates as YYYY-MM-DD"
    text = str(value).strip()
    if not text or parse_date(text):
        return None
    if parse_time(text) is not None:
        return f"'{text}' is a time: write dates as YYYY-MM-DD, and times into a column of type time"
    if not DATE_LIKE_RE.fullmatch(text):
        return None
    if "/" in text:
        return f"'{text}' could be day/month or month/day: write dates as YYYY-MM-DD (or DD.MM.YYYY)"
    if MONTH_ONLY_RE.fullmatch(text):
        return f"'{text}' has no day: write a whole date (e.g. 2026-11-01), or give the column the type text for months"
    return f"'{text}' is not a valid date: write dates as YYYY-MM-DD, e.g. 2026-11-02"


def time_problem(value: Any) -> str | None:
    """Why a value of a time column cannot be written as a time; None for times, and for words that may stay text."""
    if isinstance(value, int | float) and not isinstance(value, bool):
        return f"{value:g} is a number, not a time: write times as hh:mm, e.g. 08:30"
    text = str(value).strip()
    if not text or parse_time(text) is not None:
        return None
    if parse_date(text):
        return f"'{text}' is a date: a time column takes times only (hh:mm); a date with a time goes into a date column"
    if TIME_LIKE_RE.fullmatch(text):
        return f"'{text}': write times as hh:mm, e.g. 08:30"
    return None


NULL_TEXTS = {"null", "none", "nan"}


def convert_value(value: str | float | None, column_type: str, decimal_comma: bool | None = None) -> tuple[Any, bool]:
    """Return (cell value, is_typed). Values that do not fit the column type stay text (is_typed False).

    JSON numbers stay numbers in text columns too, so sums and charts over them work. A time is the fraction of a
    day, as in Excel.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if column_type == "text":
        if isinstance(value, float):
            return (int(value) if value.is_integer() else value), True
        return excel_text(value), True
    if column_type in ("date", "time"):
        parsed = (parse_date if column_type == "date" else parse_time)(value) if isinstance(value, str) else None
        return (parsed, True) if parsed is not None else (excel_text(str(value)), False)
    if isinstance(value, str):
        parsed_number = parse_number_text(value, decimal_comma)
        if parsed_number is None:
            return excel_text(value), False
        number, is_fraction = parsed_number
    else:
        number, is_fraction = value, False
    if column_type == "percent" and not is_fraction:
        number = number / 100
    if column_type == "integer" and float(number).is_integer():
        number = int(number)
    elif isinstance(number, float) and number.is_integer() and column_type != "percent":
        number = int(number)
    return number, True


def computable(value: Any) -> bool:
    """Whether Excel can compute with a cell value: empty, a number or date, or text that it reads as one ('12',
    '08:00')."""
    if value is None or isinstance(value, int | float | date):
        return True
    text = str(value)
    return parse_number_text(text) is not None or parse_date(text) is not None or parse_time(text) is not None


def _decimals(values: list[Any]) -> int:
    decimals = 0
    for value in values:
        if isinstance(value, float):
            text = f"{value:.{MAX_DECIMALS}f}".rstrip("0")
            decimals = max(decimals, len(text.split(".")[1]) if "." in text else 0)
    return decimals


ROW_LINE = Border(bottom=Side(style="thin", color=RULE_COLOR))
# Excel cells have no padding: an indent keeps text and numbers of neighbouring columns apart.
ALIGN = {
    "text": Alignment(horizontal="left", vertical="center", indent=1),
    "number": Alignment(horizontal="right", vertical="center", indent=1),
    "date": Alignment(horizontal="center", vertical="center"),
    "lines": Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1),
}
WIDE_SHEET_COLUMNS = 6  # wider sheets, and sheets with a chart, are printed in landscape
NARROW_FONTS = {"calibri", "carlito"}  # most other fonts are wider: columns get more room, header rows grow
SHEET_NAME_INVALID_RE = re.compile(r"[\[\]:*?/\\]")
TOTAL_FUNCTIONS = {"sum": "SUM", "average": "AVERAGE", "min": "MIN", "max": "MAX", "count": "COUNTA"}
RESERVED_SHEET_NAMES = {"history"}  # reserved by Excel
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")  # text starting like this may turn into a formula when edited
# A cell value that is meant as a formula: '=' and a function, a cell or a sheet ('=> see below' is text)
FORMULA_LIKE_RE = re.compile(
    r"=\s*=?\s*(?:[A-Za-zÄÖÜäöü_][\w.]*\s*\(|\$?[A-Za-z]{1,3}\$?(?:[0-9]|\{row)|'|[^\W\d][\w.]*!|\()"
)
NUMBER_TYPES = ("number", "integer", "currency", "percent")
TIME_TYPES = ("date", "time")
PLACEHOLDERS = {"-", "\u2013", "\u2014", "/", "n/a", "na", "k.a.", "k. a.", "tbd", "offen", "?"}  # in number columns
UNIT_TEXT_RE = re.compile(  # '3 h', '500 Stk', 'ca. 500'
    r"(?P<about>(?:ca|circa|etwa|rund|approx|about|~)\.?\s*)?(?P<number>[-+]?[0-9][0-9.,'\u2019 \u00a0]*?)\s*"
    r"(?P<unit>[^\W\d_][^\d]{0,15})?",
    re.IGNORECASE,
)


def sheet_names(sheets: list[ExcelSheet], fallback: str = "Sheet") -> list[str]:
    """Make valid, unique sheet names (31 characters, no []:*?/\\)."""
    names: list[str] = []
    for n, sheet in enumerate(sheets, start=1):
        base = " ".join(SHEET_NAME_INVALID_RE.sub(" ", to_plain(sheet.name)).split()).strip("'")[:31].strip("' ")
        if base.lower() in RESERVED_SHEET_NAMES:
            base += "_"
        base = base or f"{fallback}{n}"
        name, k = base, 2
        while name.lower() in (existing.lower() for existing in names):
            suffix = f" ({k})"
            name, k = base[: 31 - len(suffix)] + suffix, k + 1
        names.append(name)
    return names


def align_rows(sheet: ExcelSheet) -> list[list[str | float | None]]:
    """Rows that left out the formula columns get None inserted at those positions; short rows are filled up."""
    formula_positions = [i for i, column in enumerate(sheet.columns) if column.formula.strip()]
    short = len(sheet.columns) - len(formula_positions)
    aligned = []
    for row in sheet.rows:
        if formula_positions and len(row) == short:
            row = list(row)
            for position in formula_positions:
                row.insert(position, None)
        aligned.append(list(row) + [None] * (len(sheet.columns) - len(row)))
    return aligned


def prepare_sheet(sheet: ExcelSheet, decimal_comma: bool | None = None) -> ExcelSheet:
    """The sheet as the checks and the writer use it: rows aligned to the columns (align_rows), the types of
    columns sent without one inferred (infer_column_types), 'null' as text in typed columns an empty cell."""
    sheet = sheet.model_copy(update={"rows": align_rows(sheet)})
    return blank_null_texts(infer_column_types(sheet, decimal_comma))


def infer_column_types(sheet: ExcelSheet, decimal_comma: bool | None = None) -> ExcelSheet:
    """Columns sent without a type whose values are all numbers become number, currency or percent columns, so the
    numbers are not stored as text (which would make sums 0 and charts empty); all dates ('2026-11-02') or times
    ('08:00') make date or time columns."""
    columns = []
    for c, column in enumerate(sheet.columns):
        if "type" not in column.model_fields_set and not column.formula.strip():
            values = [
                row[c]
                for row in sheet.rows
                if c < len(row) and row[c] is not None and str(row[c]).strip().lower() not in {"", *NULL_TEXTS}
            ]
            kind = _number_kind(values, decimal_comma) or _time_kind(values)
            if kind is not None:
                column = column.model_copy(update={"type": kind})
        columns.append(column)
    return sheet.model_copy(update={"columns": columns})


def _number_kind(values: list[Any], decimal_comma: bool | None) -> str | None:
    kinds = set()
    for value in values:
        if isinstance(value, float):
            kinds.add("number")
            continue
        if isinstance(value, str) and value.strip().startswith("="):
            continue  # a formula in a cell says nothing about the type
        text = str(value).strip().replace("\u2212", "-")
        match = NUMBER_RE.fullmatch(text)
        digits = re.sub(r"\D", "", text)
        if not match or parse_number_text(text, decimal_comma) is None or len(digits) > 15:
            return None
        if re.match(r"[-+]?0[0-9]", text):
            return None  # leading zero: a code such as a postal code, not a number
        kinds.add(
            "percent" if match["suffix"] == "%" else "currency" if match["prefix"] or match["suffix"] else "number"
        )
    special = kinds - {"number"}
    if not kinds or len(special) > 1:
        return None
    return special.pop() if special else "number"


def _time_kind(values: list[Any]) -> str | None:
    texts = [value.strip() for value in values if isinstance(value, str) and not value.strip().startswith("=")]
    if not texts or len(texts) < len([v for v in values if not str(v).strip().startswith("=")]):
        return None
    if all(parse_date(text) for text in texts):
        return "date"
    if all(parse_time(text) is not None for text in texts):
        return "time"
    return None


def blank_null_texts(sheet: ExcelSheet) -> ExcelSheet:
    """'null' written as text is an empty cell in number and date columns (in a text column it may be a word)."""
    types = [column.type for column in sheet.columns]
    rows = [
        [None if c < len(types) and types[c] != "text" and str(v).strip().lower() in NULL_TEXTS else v
         for c, v in enumerate(row)]
        for row in sheet.rows
    ]  # fmt: skip
    return sheet.model_copy(update={"rows": rows})


def has_totals(sheet: ExcelSheet) -> bool:
    return bool(sheet.rows) and any(column.total != "none" for column in sheet.columns)


# What a column or formula holds, for the checks and the number format of formula columns without a type: the
# column types, plus 'hours' (times * 24), 'bool', 'text' and 'unknown'.
NUMBER_KINDS = {"number", "integer", "currency", "percent", "hours", "bool"}
COUNTING_FUNCTIONS = {"COUNT", "COUNTA", "COUNTIF", "COUNTIFS", "LEN", "NETWORKDAYS", "WEEKDAY", "YEAR", "MONTH",
                      "DAY", "MATCH", "HOUR", "MINUTE"}  # fmt: skip
TEXT_FUNCTIONS = {"LEFT", "RIGHT", "CONCATENATE", "CONCAT", "TEXTJOIN"}
CONDITION_FUNCTIONS = {"AND", "OR", "NOT", "ISBLANK", "ISNUMBER", "ISTEXT"}


def common_kind(kinds: list[str]) -> str:
    """The kind of SUM, MIN, IF … over values of these kinds."""
    unique = list(dict.fromkeys(kinds))
    if len(unique) == 1:
        return unique[0]
    if unique and set(unique) <= NUMBER_KINDS:
        return next((kind for kind in ("currency", "hours") if kind in unique), "number")
    return "unknown"


def operation_kind(operator: str, a: str, b: str, left: FormulaNode, right: FormulaNode) -> str:
    """The kind of `a operator b`."""
    if operator in OPERATOR_LEVELS[0]:
        return "bool"
    if operator == "&":
        return "text"
    if "unknown" in (a, b) or "text" in (a, b):
        return "unknown"
    numbers = a in NUMBER_KINDS and b in NUMBER_KINDS
    if operator in ("+", "-"):
        if numbers:
            return common_kind([a, b])
        if a == b == "date":
            return "integer" if operator == "-" else "unknown"  # the days between two dates
        if a == "date" or (b == "date" and operator == "+"):
            return "date"  # a deadline: date + days
        return "time" if "time" in (a, b) and {a, b} <= NUMBER_KINDS | {"time"} else "unknown"
    if operator == "*":
        if "time" in (a, b):
            other, node = (b, right) if a == "time" else (a, left)
            return "hours" if number_value(node) == 24 else "time" if other in NUMBER_KINDS else "unknown"
        if not numbers:
            return "unknown"
        if "currency" in (a, b):
            return "unknown" if a == b else "currency"
        return a if a == b and a in ("percent", "integer") else "number"
    if operator == "/":
        if a == "time":
            return "number" if b == "time" else "time" if b in NUMBER_KINDS else "unknown"
        if a == "hours" and number_value(right) == 24:
            return "time"
        if not numbers:
            return "unknown"
        return ("number" if b == "currency" else "currency") if a == "currency" else "number"
    return "number" if numbers else "unknown"  # ^


def text_kind(values: list[Any]) -> str:
    """What a text column holds for formulas: Excel computes with numbers, dates and times written as text too."""
    values = [v for v in values if v is not None and str(v).strip()]
    if not values:
        return "unknown"
    if all(isinstance(v, int | float) or parse_number_text(str(v)) is not None for v in values):
        return "number"
    if all(isinstance(v, str) and parse_date(v) for v in values):
        return "date"
    if all(isinstance(v, str) and parse_time(v) is not None for v in values):
        return "time"
    return "text"


class FormulaUse(NamedTuple):
    """A cell that a formula computes with (+ - * / ^)."""

    place: str  # of the formula, for messages
    sheet: int  # the formula's sheet
    token: FormulaToken
    target: tuple[int, int]  # the cell's sheet and column
    rows: list[int]  # the data rows it reads (from 0; -1 is the header row)


class Book:
    """The workbook the checks and the writer work on: the prepared sheets (prepare_sheet), their written names, and
    the checked formulas of the columns and of single cells."""

    def __init__(
        self, sheets: list[ExcelSheet], names: list[str], decimal_comma: bool | None = None, currency: str = "EUR"
    ) -> None:
        self.sheets, self.names, self.decimal_comma, self.currency = sheets, names, decimal_comma, currency
        self.lookup = sheet_lookup(sheets, names)
        self.columns: dict[tuple[int, int], str] = {}  # (sheet, column): the column formula, with {row}
        self.cells: dict[tuple[int, int, int], str] = {}  # (sheet, data row from 0, column): a formula in one cell
        self._parsed: dict[str, tuple[list[FormulaToken], FormulaNode]] = {}
        self._kinds: dict[tuple[int, int], str] = {}
        self._inputs: dict[tuple[int, int, int], bool] = {}

    def where(self, s: int) -> str:
        sheet = self.sheets[s]
        return f"sheets[{sheet._index if sheet._index >= 0 else s}]"

    def last_rows(self, s: int) -> dict[str, int]:
        """The last data row of every sheet by written name, "" for sheet s (for whole columns such as B:B)."""
        rows = {name: len(sheet.rows) + 1 for name, sheet in zip(self.names, self.sheets, strict=True)}
        return {**rows, "": len(self.sheets[s].rows) + 1}

    def columns_hint(self, s: int) -> str:
        """'columns of 'S': A=Artikel, B=Menge, …', for messages about formulas: models miscount the letters."""
        columns = self.sheets[s].columns
        listed = [f"{get_column_letter(c + 1)}={to_plain(column.header.strip())[:25]}" for c, column in
                  enumerate(columns[:26])]  # fmt: skip
        return f"columns of '{self.names[s]}': " + ", ".join(listed) + (", …" if len(columns) > 26 else "")

    def add_cell_formula(self, s: int, r: int, c: int, value: str) -> str | None:
        """Check a formula written into a cell and keep it; returns a message if it is sent back. In a column of type
        text it stays text, and so does a formula that reaches outside (HYPERLINK, ...) unless the column takes
        numbers. A formula that is merely not allowed or broken is sent back, so the model can rewrite it."""
        column = self.sheets[s].columns[c]
        if column.type == "text" and "type" in column.model_fields_set:
            return None
        try:
            place = (self.names[s], c + 1, r + 2)
            self.cells[(s, r, c)] = normalize_formula(value, False, self.lookup, self.last_rows(s), place)
        except UnsafeFormula as e:
            if column.type in NUMBER_TYPES:
                return f"the formula '{value.strip()[:60]}' is not allowed: {e}"
        except ValueError as e:
            if column.type in NUMBER_TYPES or FORMULA_LIKE_RE.match(value.strip()):
                return f"the formula '{value.strip()[:60]}' in '{column.header.strip()}' is not allowed: {e}"
        return None

    def has_inputs(self, s: int, r: int, c: int) -> bool:
        """Whether data row r (from 0) has a value in a column that the formula of column c uses in its own row;
        True for a formula that uses no such column. Formula columns that use each other count as filled."""
        key = (s, r, c)
        if key not in self._inputs:
            self._inputs[key] = True  # while it is worked out: a cycle counts as filled
            inputs = {self.target(t, s)[1] for t in self.parsed(self.columns[(s, c)])[0]
                      if t.kind == "cell" and t.row == "{row}" and self.target(t, s)[0] == s} - {c}  # fmt: skip
            self._inputs[key] = not inputs or any(_cell_filled(self, s, r + 2, u) for u in inputs)
        return self._inputs[key]

    def copies_column_formula(self, s: int, r: int, c: int) -> bool:
        """Whether the formula of a cell in a formula column is the column formula written out for a row: the same
        but for the row numbers ('=B2*C2', also '=B3*C3' in another row)."""
        column = [t._replace(text="", row="") if t.kind == "cell" else t for t in self.parsed(self.columns[(s, c)])[0]]
        cell = [t._replace(text="", row="") if t.kind == "cell" else t for t in self.parsed(self.cells[(s, r, c)])[0]]
        return [(t.kind, t.text.upper(), t.sheet, t.column) for t in column] == [
            (t.kind, t.text.upper(), t.sheet, t.column) for t in cell
        ]

    def trailing_sum_rows(self, s: int) -> int:
        """How many rows at the end are sum rows (or empty rows between them): charts and the filter end above."""
        rows, count = self.sheets[s].rows, 0
        for r in range(len(rows) - 1, 0, -1):
            if self.sum_row(s, r):
                count = len(rows) - r
            elif any(v is not None and str(v).strip() for v in rows[r]):
                break
        return count

    def sum_row(self, s: int, r: int) -> bool:
        """A row with a formula of its own in a formula column (net, VAT, gross below the items)."""
        return any((s, r, c) in self.cells for (t, c) in self.columns if t == s)

    def parsed(self, formula: str) -> tuple[list[FormulaToken], FormulaNode]:
        """The tokens and tree of a checked formula."""
        if formula not in self._parsed:
            tokens = parse_formula(formula, self.lookup)
            self._parsed[formula] = tokens, formula_tree(tokens)
        return self._parsed[formula]

    def formulas(self) -> list[tuple[int, int, int | None, str, str]]:
        """Every formula as (sheet, column, data row or None for a column formula, place for messages, formula)."""
        result = [(s, c, None, f"{self.where(s)}.columns[{c}].formula", f) for (s, c), f in self.columns.items()]
        return result + [(s, c, r, f"{self.where(s)}.rows[{r}]", f) for (s, r, c), f in self.cells.items()]

    def target(self, token: FormulaToken, s: int) -> tuple[int, int]:
        """(sheet, column) of a cell in a formula of sheet s."""
        return (self.names.index(token.sheet) if token.sheet else s), _column_number(token.column) - 1

    def cell_name(self, token: FormulaToken) -> str:
        return token.text if "!" in token.text or not token.sheet else _sheet_prefix(token.sheet) + token.text

    def value(self, s: int, r: int, c: int) -> Any:
        """What is written into data row r (from 0) of column c: a number, date, text or None."""
        rows = self.sheets[s].rows
        raw = rows[r][c] if 0 <= r < len(rows) and c < len(rows[r]) else None
        return convert_value(raw, self.sheets[s].columns[c].type, self.decimal_comma)[0]

    def used_columns(self) -> set[tuple[int, int]]:
        """(sheet, column) of every column a formula uses, the columns of ranges included."""
        used = set()
        for s, _, _, _, formula in self.formulas():
            for node in walk(self.parsed(formula)[1]):
                if node.kind == "cell" and node.cell is not None:
                    used.add(self.target(node.cell, s))
                elif node.kind == "range":
                    (t, a), (_, b) = (self.target(arg.cell, s) for arg in node.args if arg.cell is not None)
                    used.update((t, c) for c in range(min(a, b), max(a, b) + 1))
        return used

    def arithmetic_uses(self) -> list[FormulaUse]:
        """The cells that formulas compute with: operands of + - * / ^, also -B2 and B2%."""
        uses = []
        for s, _, r, place, formula in self.formulas():
            for node in walk(self.parsed(formula)[1]):
                if node.kind in ("unary", "percent") or (node.kind == "binary" and node.text in ARITHMETIC):
                    for arg in node.args:
                        if arg.kind == "cell" and arg.cell is not None:
                            if arg.cell.row == "{row}":
                                rows = list(range(len(self.sheets[s].rows))) if r is None else [r]
                            else:
                                rows = [int(arg.cell.row) - 2]
                            uses.append(FormulaUse(place, s, arg.cell, self.target(arg.cell, s), rows))
        return uses

    def kind(self, s: int, c: int) -> str:
        """What the cells of a column hold: its type; for a formula column without a type what the formula gives;
        for a text column what its values are."""
        key = (s, c)
        if key not in self._kinds:
            self._kinds[key] = "unknown"  # a formula that uses its own column (sent back elsewhere)
            column = self.sheets[s].columns[c]
            if column.type != "text":
                kind = column.type
            elif key in self.columns:
                kind = self.expression_kind(self.parsed(self.columns[key])[1], s)
            else:
                kind = text_kind([self.value(s, r, c) for r in range(len(self.sheets[s].rows))])
            self._kinds[key] = kind
        return self._kinds[key]

    def expression_kind(self, node: FormulaNode, s: int) -> str:
        """What a formula (or part of one) gives, from the kinds of the cells it uses."""
        if node.kind == "empty":
            return "unknown"
        if node.kind == "number":
            return "number" if "." in node.text else "integer"
        if node.kind in ("string", "bool"):
            return "text" if node.kind == "string" else "bool"
        if node.kind in ("cell", "range"):
            cell = node.cell if node.kind == "cell" else node.args[0].cell
            assert cell is not None
            t, c = self.target(cell, s)
            if cell.row == "1" or c >= len(self.sheets[t].columns):
                return "text" if cell.row == "1" else "unknown"
            return self.kind(t, c)
        if node.kind == "unary":
            return self.expression_kind(node.args[0], s)
        if node.kind == "percent":
            return "number"
        if node.kind == "binary":
            left, right = node.args
            a, b = self.expression_kind(left, s), self.expression_kind(right, s)
            return operation_kind(node.text, a, b, left, right)
        return self._call_kind(node, s)

    def _call_kind(self, node: FormulaNode, s: int) -> str:
        name, args = node.text, node.args

        def kinds(nodes: Collection[FormulaNode]) -> list[str]:  # "" in IF(…, "", B*C) says nothing
            return [self.expression_kind(n, s) for n in nodes if n.kind not in ("string", "empty")] or ["text"]

        if name in ("SUM", "MIN", "MAX", "AVERAGE", "MEDIAN", "PRODUCT"):
            return common_kind(kinds(args)) if name != "PRODUCT" or len(args) == 1 else "number"
        if name in ("ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "MOD", "INT"):
            kind = self.expression_kind(args[0], s)
            return "integer" if name == "INT" and kind == "number" else kind
        if name in ("SUMIF", "AVERAGEIF"):
            return self.expression_kind(args[2] if len(args) > 2 else args[0], s)
        if name in ("SUMIFS", "AVERAGEIFS", "MAXIFS", "MINIFS", "INDEX"):
            return self.expression_kind(args[0], s)
        if name == "XLOOKUP":
            return self.expression_kind(args[2], s)
        if name in ("IF", "IFERROR", "IFS"):
            return common_kind(kinds(args[1:] if name == "IF" else args if name == "IFERROR" else args[1::2]))
        if name in COUNTING_FUNCTIONS:
            return "integer"
        if name in ("DATE", "TODAY", "EDATE", "EOMONTH"):
            return "date"
        if name == "TIME":
            return "time"
        if name in TEXT_FUNCTIONS:
            return "text"
        if name in CONDITION_FUNCTIONS:
            return "bool"
        if name == "VLOOKUP" and args[1].kind == "range" and (index := number_value(args[2])) is not None:
            first = args[1].args[0].cell
            assert first is not None
            t, c = self.target(first, s)
            column = c + int(index) - 1
            return self.kind(t, column) if 0 <= column < len(self.sheets[t].columns) else "unknown"
        return "number" if name in ("POWER", "SQRT") else "unknown"

    def evaluate(self, node: FormulaNode, s: int, r: int) -> float | None:
        """The value of a formula of + - * / ^ with numbers and number cells of its own sheet in data row r; None if
        it uses anything else (functions, text, other sheets, formula columns)."""
        if node.kind == "number":
            return float(node.text)
        if node.kind == "cell" and node.cell is not None:
            t, c = self.target(node.cell, s)
            row = r if node.cell.row == "{row}" else int(node.cell.row) - 2
            if t != s or c >= len(self.sheets[s].columns) or not 0 <= row < len(self.sheets[s].rows):
                return None
            if (s, c) in self.columns or (s, row, c) in self.cells:
                return None
            if self.sheets[s].columns[c].type not in NUMBER_TYPES:
                return None
            value = self.value(s, row, c)
            return 0.0 if value is None else float(value) if isinstance(value, int | float) else None
        values = [self.evaluate(arg, s, r) for arg in node.args]
        if not values or None in values or node.kind not in ("unary", "percent", "binary"):
            return None
        if node.kind == "unary":
            return -values[0] if node.text == "-" else values[0]
        if node.kind == "percent":
            return values[0] / 100
        a, b = values
        try:
            match node.text:
                case "+":
                    return a + b
                case "-":
                    return a - b
                case "*":
                    return a * b
                case "/":
                    return a / b if b else None
                case "^":
                    return float(a**b) if abs(b) <= 100 else None
        except (OverflowError, ZeroDivisionError, TypeError):
            return None
        return None

    def currencies(self, s: int, c: int) -> list[str]:
        """The currencies written with the values of a column ('$1,200.00'), in order of appearance."""
        return list(dict.fromkeys(code for row in self.sheets[s].rows if (code := value_currency(row[c]))))

    def currency_of(self, s: int, c: int, seen: frozenset[tuple[int, int]] = frozenset()) -> str:
        """The currency of a column: the one written with its values, for a formula column the one of the first
        currency column it uses, otherwise the currency valve."""
        if (s, c) in self.columns and (s, c) not in seen:
            for token in self.parsed(self.columns[(s, c)])[0]:
                if token.kind == "cell":
                    t, u = self.target(token, s)
                    if u < len(self.sheets[t].columns) and (t, u) != (s, c) and self.kind(t, u) == "currency":
                        return self.currency_of(t, u, seen | {(s, c)})
            return self.currency
        found = self.currencies(s, c)
        return found[0] if len(found) == 1 else self.currency

    def number_format(self, s: int, c: int, values: list[Any], date_format: str) -> str | None:
        """The number format of a column from its type; a formula column without a type gets the format of what its
        formula gives (a deadline =B{row}+14 a date, quantity * price the currency)."""
        column = self.sheets[s].columns[c]
        formula = (s, c) in self.columns
        kind = self.kind(s, c) if formula and column.type == "text" else column.type
        match kind:
            case "integer":
                return "#,##0"
            case "number" | "hours":
                decimals = _decimals(values) if not formula else 2
                return "#,##0" + ("." + "0" * decimals if decimals else "")
            case "currency":
                return CURRENCY_FORMATS[self.currency_of(s, c)]
            case "percent":
                fractional = any(isinstance(v, float) and not round(v * 100, 9).is_integer() for v in values)
                return "0.0%" if formula or fractional else "0%"
            case "date":
                timed = any(isinstance(v, datetime) and (v.hour, v.minute, v.second) != (0, 0, 0) for v in values)
                return date_format + (" hh:mm" if timed else "")
            case "time":  # a duration may be 24 hours or more: [h]
                if formula or any(isinstance(v, float) and v >= 1 for v in values):
                    return "[h]:mm"
                return "hh:mm:ss" if any(isinstance(v, float) and round(v * 86400) % 60 for v in values) else "hh:mm"
        return None


def check_excel(
    sheets: list[ExcelSheet], decimal_comma: bool | None = None, names: list[str] | None = None, currency: str = "EUR"
) -> Book:
    """Validate the prepared sheets (prepare_sheet) and return the workbook with its checked formulas. `names` are
    the written sheet names (see sheet_names), which formulas use for cells of other sheets."""
    check = Checker()
    book = Book(sheets, names or sheet_names(sheets), decimal_comma, currency)
    check.require(bool(sheets), "sheets: the workbook has no sheets")
    check.require(len(sheets) <= MAX_SHEETS, f"sheets: at most {MAX_SHEETS} sheets")
    cells = sum(len(sheet.rows) * len(sheet.columns) for sheet in sheets)
    check.require(cells <= MAX_CELLS, f"sheets: at most {MAX_CELLS} cells in all sheets together ({cells} sent)")
    complete, charted = [], set()
    for s, sheet in enumerate(sheets):
        where = book.where(s)
        headers = [column.header.strip() for column in sheet.columns]
        check.require(bool(sheet.columns), f"{where}.columns: the sheet has no columns")
        check.require(len(sheet.columns) <= MAX_SHEET_COLUMNS, f"{where}.columns: at most {MAX_SHEET_COLUMNS} columns")
        check.require(len(sheet.rows) <= MAX_SHEET_ROWS, f"{where}.rows: at most {MAX_SHEET_ROWS} rows")
        check.text(sheet.name, *headers)
        hint = book.columns_hint(s) if sheet.columns else ""
        for c, column in enumerate(sheet.columns):
            if column.formula.strip():
                try:
                    book.columns[(s, c)] = normalize_formula(column.formula, True, book.lookup, book.last_rows(s))
                except ValueError as e:
                    check.require(False, f"{where}.columns[{c}].formula: {e} ({hint})")
        fits = all(len(row) <= len(sheet.columns) for row in sheet.rows)
        for r, row in enumerate(sheet.rows):
            check.require(
                len(row) <= len(sheet.columns),
                f"{where}.rows[{r}]: {len(row)} values but only {len(sheet.columns)} columns",
            )
            for c, value in enumerate(row):
                if not isinstance(value, str):
                    continue
                check.text(value)
                check.require(
                    len(value) <= MAX_CELL_CHARS, f"{where}.rows[{r}]: a cell has more than {MAX_CELL_CHARS} characters"
                )
                formula = fits and value.strip().startswith("=")  # in a formula column it replaces the formula
                if formula and (problem := book.add_cell_formula(s, r, c, value)):
                    check.errors.append(f"{where}.rows[{r}]: {problem} ({hint})")
                elif (
                    formula
                    and (s, r, c) in book.cells
                    and (s, c) in book.columns
                    and book.copies_column_formula(s, r, c)
                ):
                    del book.cells[(s, r, c)]  # '=B2*C2' (or '=B3*C3') in every row: the column formula applies
        check.require(len(sheet.charts) <= MAX_SHEET_CHARTS, f"{where}.chart: at most {MAX_SHEET_CHARTS} charts")
        for k, chart in enumerate(sheet.charts):
            place = f"{where}.chart" + (f"[{k}]" if isinstance(sheet.chart, list) else "")
            if check.require(bool(sheet.rows), f"{place}: the sheet has no rows to plot"):
                try:
                    layout = chart_layout(sheet, chart)
                except ValueError as e:
                    check.errors.extend(f"{place}.{error}" for error in e.args[0])
                    continue
                first, last = layout.categories[:2]
                plotted = range(first, last + 1) if layout.by_rows else (index for _, index in layout.series)
                charted.update((s, column - 1) for column in plotted)
        if len(check.errors) > MAX_ERRORS_SHOWN:
            break
        if sheet.rows and fits:
            complete.append(s)
    if sheets and not any(sheet.rows for sheet in sheets):  # a test call or a copy without its data
        check.errors.append(
            "sheets: no sheet has rows, so the file would be empty. Send the data in 'rows'; for a template to fill "
            "in, send its rows with empty cells (null)"
        )
    used = book.used_columns()
    for s in complete:
        check.errors.extend(_column_problems(book, s, used, charted))
    if not check.errors:  # all sheets are complete: cells of other sheets can be looked up
        for problems in (
            _reference_problems,
            _row_formula_problems,
            _percent_problems,
            _arithmetic_problems,
            _kind_problems,
            _value_problems,
        ):
            check.errors.extend(problems(book))
    check.raise_if_failed()
    return book


def number_advice(text: str, header: str) -> str:
    """How to send a text of a number column as a number: '3 h' -> the unit into the header."""
    match = UNIT_TEXT_RE.fullmatch(text.strip())
    if match and match["about"]:
        number, about = match["number"].strip(), match["about"].strip()
        return f"send the number ({number}); a note such as '{about}' goes into a text column"
    if match and match["unit"]:
        unit = match["unit"].strip()
        return f"send the number ({match['number'].strip()}) and put the unit into the header, e.g. '{header} ({unit})'"
    return "send a number, or null for an empty cell"


def _column_problems(book: Book, s: int, used: set[tuple[int, int]], charted: set[tuple[int, int]]) -> list[str]:
    """Columns whose content cannot be what the model meant; these usually come from rows whose values are
    shifted against the columns, or from values the model forgot or wrote as text."""
    sheet, where, rows = book.sheets[s], book.where(s), book.sheets[s].rows
    formulas = {c: f for (t, c), f in book.columns.items() if t == s}
    problems: list[str] = []
    filled = [any(row[c] is not None and str(row[c]).strip() for row in rows) for c in range(len(sheet.columns))]
    for c, column in enumerate(sheet.columns):
        header = column.header.strip()
        name = f"{where}.columns[{c}] ('{header}')"
        if c in formulas:
            continue
        if column.total != "none" and not filled[c]:
            problems.append(
                f"{name}: the total needs values in this column: fill them in, give the column a formula "
                "(e.g. '=B{row}*C{row}'), or put the total on the column with the numbers"
            )
        uses = [what for what, given in (("the total", column.total != "none"), ("the formulas", (s, c) in used),
                                         ("the chart", (s, c) in charted)) if given]  # fmt: skip
        values = [(r, row[c]) for r, row in enumerate(rows)
                  if row[c] is not None and str(row[c]).strip() and (s, r, c) not in book.cells]  # fmt: skip
        texts = [(r, v) for r, v in values if not convert_value(v, column.type, book.decimal_comma)[1]]
        words = [(r, v) for r, v in texts if str(v).strip().lower() not in PLACEHOLDERS]
        if column.type in NUMBER_TYPES:
            if len(texts) >= 2 and len(texts) * 2 > len(values):
                example = str(texts[0][1])
                advice = (
                    "If they are numbers with a unit: " + number_advice(example, header)
                    if UNIT_TEXT_RE.fullmatch(example.strip())
                    else "Each row needs one value per column, in column order - is a column missing? Or give the "
                    "column type 'text'."
                )
                count = f"{len(texts)} of {len(values)}"
                problems.append(f"{name}: {count} values are not numbers (e.g. '{example}'). {advice}")
            elif column.type == "percent" and (fraction := _fractions_only([v for _, v in values], book)) is not None:
                problems.append(
                    f"{name}: all values are below 1 (e.g. {fraction}), but a percent column takes percent numbers: "
                    f"5 means 5 %, so {fraction} would show as {fraction} %. Send 5 for 5 % (or '5 %'); a value that "
                    f"really is below 1 % needs the sign, e.g. '0.5 %'."
                )
            elif words and uses:
                r, value = words[0]
                more = f" (and {len(words) - 1} more)" if len(words) > 1 else ""
                problems.append(
                    f"{where}.rows[{r}]: '{value}' in '{header}' is not a number{more}, so {' and '.join(uses)} "
                    f"cannot count it: {number_advice(str(value), header)}"
                )
            if column.type == "currency" and len(found := book.currencies(s, c)) > 1:
                problems.append(
                    f"{name}: the values have different currencies ({', '.join(found)}); a column has one currency: "
                    "put each currency into a column of its own, or convert the amounts"
                )
        elif column.type in TIME_TYPES:
            problem = date_problem if column.type == "date" else time_problem
            found = next(((r, text) for r, v in values if (text := problem(v))), None)
            if found:
                problems.append(f"{where}.rows[{found[0]}]: {found[1]} (column '{header}')")
            elif words and uses:
                r, value = words[0]
                problems.append(
                    f"{where}.rows[{r}]: '{value}' in '{header}' is not a {column.type}, so {' and '.join(uses)} "
                    "cannot use it: leave the cell empty (null) and put such notes into a text column"
                )
    hint = book.columns_hint(s)
    for c, formula in formulas.items():
        tokens = book.parsed(formula)[0]
        own_sheet = [
            (i, token) for i, token in enumerate(tokens)
            if token.kind == "cell" and book.target(token, s)[0] == s
        ]  # fmt: skip
        letter = get_column_letter(c + 1)
        if any(token.column == letter for _, token in own_sheet):
            problems.append(
                f"{where}.columns[{c}] ('{sheet.columns[c].header.strip()}'): the formula uses its own column "
                f"{letter}; it may only use the other columns (e.g. quantity * price). A sum row below the items "
                f"gets its own formula in its row instead, e.g. '=SUM({letter}2:{letter}6)' ({hint})"
            )
            continue
        for i, token in own_sheet:
            used_column = _column_number(token.column) - 1
            in_range = (i and tokens[i - 1].text == ":") or (i + 1 < len(tokens) and tokens[i + 1].text == ":")
            if used_column >= len(sheet.columns) and not in_range:
                problems.append(
                    f"{where}.columns[{c}] ('{sheet.columns[c].header.strip()}'): the formula uses {token.text}, but "
                    f"the sheet has only the columns A to {get_column_letter(len(sheet.columns))} ({hint})"
                )
                break
        for letters in dict.fromkeys(token.column for _, token in own_sheet):
            used_column = _column_number(letters) - 1
            if used_column < len(sheet.columns) and used_column not in formulas and not filled[used_column]:
                problems.append(
                    f"{where}.columns[{used_column}] ('{sheet.columns[used_column].header.strip()}') is empty, but the "
                    f"formula of '{sheet.columns[c].header.strip()}' uses it: fill in its values ({hint})"
                )
    if problem := _total_row_problem(book, s, formulas):
        problems.append(problem)
    return problems


def _fractions_only(values: list[Any], book: Book) -> float | None:
    """The first value if every value of a percent column but 0 is a plain number between -1 and 1: the model sent
    fractions (0.05 for 5 %), which the column would show as 0.05 %. None if any value has a '%' or is 1 or more."""
    numbers = []
    for value in values:
        if isinstance(value, bool) or str(value).strip().lower() in PLACEHOLDERS:
            continue
        if isinstance(value, int | float):
            numbers.append(value)
        elif isinstance(value, str) and (parsed := parse_number_text(value, book.decimal_comma)) and not parsed[1]:
            numbers.append(parsed[0])
        else:
            return None
    numbers = [n for n in numbers if n != 0]
    return numbers[0] if numbers and all(abs(n) < 1 for n in numbers) else None


def factors(node: FormulaNode) -> list[FormulaNode]:
    """The cells a part of a formula multiplies: B{row} and C{row} in -B{row}*C{row}/D{row}, none in (1-C{row})."""
    if node.kind == "cell":
        return [node]
    if node.kind == "binary" and node.text in ("*", "/"):
        return factors(node.args[0]) + (factors(node.args[1]) if node.text == "*" else [])
    return factors(node.args[0]) if node.kind == "unary" else []


def _percent_problems(book: Book) -> list[str]:
    """Formulas that treat a percent cell as a percent number: a percent cell holds the fraction (5 % is 0.05), so
    '/100', '*0.01' or '%' on it divides a second time, and '*100' in a percent column shows 100 times the share
    (2500 %)."""
    problems: list[str] = []
    for s, c, _, place, formula in book.formulas():
        nodes = list(walk(book.parsed(formula)[1]))
        column = book.sheets[s].columns[c]
        if column.type == "percent" and any(
            node.kind == "binary" and node.text == "*" and 100 in map(number_value, node.args) for node in nodes
        ):
            problems.append(
                f"{place}: '{column.header.strip()}' is a percent column, whose cells hold the fraction "
                f"(0.25 shows as 25 %): drop the *100, or 25 % shows as 2500 %"
            )
        for node in nodes:
            if node.kind == "binary" and node.text == "/" and number_value(node.args[1]) == 100:
                cells, written = factors(node.args[0]), "/100"
            elif node.kind == "binary" and node.text == "*" and 0.01 in map(number_value, node.args):
                cells, written = [cell for arg in node.args for cell in factors(arg)], "*0.01"
            elif node.kind == "percent":
                cells, written = factors(node.args[0]), "%"
            else:
                continue
            for cell in cells:
                assert cell.cell is not None
                t, target = book.target(cell.cell, s)
                if target < len(book.sheets[t].columns) and book.sheets[t].columns[target].type == "percent":
                    name = book.cell_name(cell.cell)
                    problems.append(
                        f"{place}: {name} is a percent cell ('{book.sheets[t].columns[target].header.strip()}'), which "
                        f"already holds the fraction (5 % is 0.05): write {name} instead of {name}{written}"
                    )
    return problems


def _reference_problems(book: Book) -> list[str]:
    """Single cells of other sheets that are empty: models count the rows there wrong (row 1 is the header), and
    an empty cell silently counts as 0. Ranges may reach beyond the data, e.g. in SUMIF."""
    sheets = book.sheets

    def filled(t: int, row: int, column: int) -> bool:
        target = sheets[t]
        if not 1 <= column <= len(target.columns):
            return False
        if row == 1:
            return bool(target.columns[column - 1].header.strip())
        if row == len(target.rows) + 2 and has_totals(target):  # the total row, with its label in column A
            return target.columns[column - 1].total != "none" or column == 1
        if not 2 <= row <= len(target.rows) + 1:
            return False
        value = target.rows[row - 2][column - 1]
        return (t, column - 1) in book.columns or (value is not None and str(value).strip() != "")

    problems: list[str] = []
    for s, _, _, place, formula in book.formulas():
        tokens, tree = book.parsed(formula)
        in_ranges = {arg.index for node in walk(tree) if node.kind == "range" for arg in node.args}
        for i, token in enumerate(tokens):
            if token.kind != "cell" or i in in_ranges or not token.sheet or book.target(token, s)[0] == s:
                continue
            t, column = book.names.index(token.sheet), _column_number(token.column)
            target, cell = sheets[t], f"'{token.sheet}'!{token.column}{token.row}"
            if token.row == "{row}":
                if not any(filled(t, row, column) for row in range(2, len(target.rows) + 2)):
                    problems.append(
                        f"{place}: {cell} is empty: '{token.sheet}' has no values in column {token.column} (it has "
                        f"the columns A to {get_column_letter(len(target.columns))})"
                    )
                elif len(target.rows) < len(sheets[s].rows):
                    problems.append(
                        f"{place}: {cell} is empty in the last rows: '{token.sheet}' has {len(target.rows)} "
                        f"data rows, this sheet {len(sheets[s].rows)}"
                    )
            elif not filled(t, int(token.row), column):
                total = f", total row {len(target.rows) + 2}" if has_totals(target) else ""
                problems.append(
                    f"{place}: {cell} is an empty cell. On '{token.sheet}', row 1 is the header and the data "
                    f"is in rows 2 to {len(target.rows) + 1}{total}, columns A to "
                    f"{get_column_letter(len(target.columns))}"
                )
    return problems


def _row_formula_problems(book: Book) -> list[str]:
    """Formulas in single cells (sum rows below the items: Netto, MwSt, Brutto, subtotals) that reach their own
    cell (a circular reference), that a column total counts a second time, or that use an empty cell of their own
    sheet: models count the rows wrong (row 1 is the header), and an empty cell silently counts as 0."""
    problems: list[str] = []
    for s, c, r, place, formula in book.formulas():
        if r is None:
            continue
        sheet, header = book.sheets[s], book.sheets[s].columns[c].header.strip()
        tokens, tree = book.parsed(formula)
        own_row, letter = r + 2, get_column_letter(c + 1)
        in_ranges = {arg.index for node in walk(tree) if node.kind == "range" for arg in node.args}
        cells, spans = [], []  # (column, row) of single cells; (first column, last column, first row, last row)
        for node in walk(tree):
            if node.kind == "range" and all(arg.cell is not None and not arg.cell.sheet for arg in node.args):
                (a, ra), (b, rb) = ((_column_number(arg.cell.column) - 1, _row_of(arg.cell, own_row))
                                    for arg in node.args if arg.cell is not None)  # fmt: skip
                spans.append((min(a, b), max(a, b), min(ra, rb), max(ra, rb)))
        for i, token in enumerate(tokens):
            if token.kind == "cell" and not token.sheet and i not in in_ranges:
                cells.append((_column_number(token.column) - 1, _row_of(token, own_row)))
        shown = formula.replace("{row}", str(own_row))
        if (c, own_row) in cells or any(a <= c <= b and ra <= own_row <= rb for a, b, ra, rb in spans):
            problems.append(
                f"{place}: the formula {shown} in '{header}' uses its own cell {letter}{own_row} (a circular "
                f"reference); a sum of the rows above ends in the row above, e.g. =SUM({letter}2:{letter}{own_row - 1})"
            )
            continue
        if (
            sheet.columns[c].total != "none"
            and tree.kind == "call"
            and tree.text in ("SUM", "SUMIF", "SUMIFS")
            and any(a <= c <= b for a, b, _, _ in spans)
        ):  # a subtotal; a discount row (=-ROUND(SUM(…)*0.05,2)) is meant to be in the total
            problems.append(
                f"{place}: the formula {shown} computes from '{header}' itself, and the column's total "
                f'("{sheet.columns[c].total}") adds up every row, this one too, so it counts twice: either drop '
                f'"total" on \'{header}\' and write all sum rows yourself, or drop the sum rows and keep "total"'
            )
            continue
        empty = list(dict.fromkeys((column, row) for column, row in cells if _miscounted(book, s, row, column)))
        if empty:
            names = ", ".join(f"{get_column_letter(column + 1)}{row}" for column, row in empty)
            problems.append(
                f"{place}: the formula {shown} uses {names}, which "
                f"{'is' if len(empty) == 1 else 'are'} empty, so it computes with 0. Row 1 is the header, the data is "
                f"in rows 2 to {len(sheet.rows) + 1} ({_row_labels(book, s)}); is the row number right?"
            )
    return problems


def _miscounted(book: Book, s: int, row: int, c: int) -> bool:
    """A cell that a formula of the same sheet can only mean by a wrong row count: outside the table, in a row
    with nothing in it, or a formula column's cell that has nothing to compute in its row. An empty input cell
    in a filled row may be left for the user (a template, an IF that checks it)."""
    sheet = book.sheets[s]
    if _cell_filled(book, s, row, c):
        return False
    r = row - 2
    if not 0 <= r < len(sheet.rows) or (s, c) in book.columns:
        return True
    return not any(v is not None and str(v).strip() for v in sheet.rows[r]) and not book.sum_row(s, r)


def _row_of(token: FormulaToken, own_row: int) -> int:
    return own_row if token.row == "{row}" else int(token.row)


def _cell_filled(book: Book, s: int, row: int, c: int) -> bool:
    """Whether a cell of sheet s (row as in Excel, column from 0) holds something: the header, a value, a formula
    of the cell, the column formula with values to compute with, or the total row the tool adds."""
    sheet = book.sheets[s]
    if not 0 <= c < len(sheet.columns):
        return False
    if row == 1:
        return bool(sheet.columns[c].header.strip())
    if row == len(sheet.rows) + 2 and has_totals(sheet):
        return sheet.columns[c].total != "none" or c == 0
    r = row - 2
    if not 0 <= r < len(sheet.rows):
        return False
    if (s, r, c) in book.cells:
        return True
    if (s, c) in book.columns:  # with no values in the columns it uses, the formula gives 0 (or an error)
        return book.has_inputs(s, r, c)  # the writer leaves it out of a sum row only without inputs
    value = sheet.rows[r][c]
    return value is not None and str(value).strip() != ""


def _row_labels(book: Book, s: int) -> str:
    """'2=Laptop, 3=Monitor, 8=Netto, …': the rows by their text column with the most labels (the first of
    those), for messages about row numbers."""
    sheet = book.sheets[s]
    texts = [c for c, column in enumerate(sheet.columns) if column.type == "text" and (s, c) not in book.columns]
    count = {c: sum(isinstance(row[c], str) and bool(row[c].strip()) for row in sheet.rows) for c in texts}
    c = max(texts, key=lambda c: (count[c], -c)) if texts else 0
    labels = [f"{r + 2}={to_plain(str(row[c]).strip())[:20]}" for r, row in enumerate(sheet.rows[:30])
              if row[c] is not None and str(row[c]).strip()]  # fmt: skip
    more = ", …" if len(sheet.rows) > 30 else ""
    return "rows: " + ", ".join(labels) + more if labels else "the rows have no labels"


def _arithmetic_problems(book: Book) -> list[str]:
    """Formulas that compute (+ - * / ^) with text, which Excel shows as #VALUE!: a text column (usually the column
    letters were miscounted: '=A{row}*B{row}' with the article names in A), the header row, or single texts in the
    rows (a word or '-' where a number belongs). Comparisons, SUM and the like may use text."""
    problems: list[str] = []
    reported = set()
    for use in book.arithmetic_uses():
        t, column = use.target
        target = book.sheets[t]
        key = (use.place, use.target)
        if column >= len(target.columns) or use.target in book.columns or key in reported:
            continue
        header, name = target.columns[column].header.strip(), book.cell_name(use.token)
        if -1 in use.rows:
            reported.add(key)
            problems.append(
                f"{use.place}: {name} is in the header row ('{header}'), and computing with its text gives #VALUE!; "
                f"the data starts in row 2 ({book.columns_hint(t)})"
            )
            continue
        bad = [(r, value) for r in use.rows if 0 <= r < len(target.rows) and (t, r, column) not in book.cells
               and not computable(value := book.value(t, r, column))]  # fmt: skip
        if not bad:
            continue
        reported.add(key)
        values = [book.value(t, r, column) for r in range(len(target.rows))]
        if not any(value is not None and computable(value) for value in values):
            problems.append(
                f"{use.place}: the formula computes with {name}, but '{header}' holds text (e.g. '{bad[0][1]}'), which "
                f"gives #VALUE!; is the column letter right? ({book.columns_hint(t)})"
            )
        else:
            r, value = bad[0]
            more = f" (and {len(bad) - 1} more rows)" if len(bad) > 1 else ""
            problems.append(
                f"{use.place}: the formula computes with '{header}' ({name}), but {book.where(t)}.rows[{r}] has "
                f"'{value}' there{more}, which gives #VALUE!: send a number, or null for an empty cell"
            )
    return problems


def _kind_problems(book: Book) -> list[str]:
    """Column formulas whose result the column type would show wrong: times subtracted in a number column (a
    fraction of a day: 0.35 for 8:30), times * 24 in a time column (204:00), a date in a number column."""
    problems: list[str] = []
    for (s, c), formula in book.columns.items():
        column = book.sheets[s].columns[c]
        kind = book.expression_kind(book.parsed(formula)[1], s)
        place = f"{book.where(s)}.columns[{c}] ('{column.header.strip()}')"
        if kind == "time" and column.type in NUMBER_TYPES:
            problems.append(
                f"{place}: the formula computes with times, which Excel holds as fractions of a day (8 hours are "
                f"0.33), so this {column.type} column would show 0.35 for 8:30. For hours as a number multiply by 24 "
                f"(=({formula[1:]})*24), or give the column the type time to show 8:30"
            )
        elif kind == "hours" and column.type == "time":
            problems.append(
                f"{place}: a time column shows the fraction of a day as hh:mm, but the formula multiplies the times "
                "by 24, so 8:30 would show as 204:00: drop the *24, or give the column the type number for hours as "
                "a number (8.5)"
            )
        elif kind == "date" and column.type in NUMBER_TYPES:
            problems.append(
                f"{place}: the formula gives a date, which this {column.type} column would show as a number "
                "(e.g. 46328): give the column the type date"
            )
    return problems


def _value_problems(book: Book) -> list[str]:
    """Values sent in a formula column that differ from what its formula computes: the formula fills the column,
    so either it uses the wrong columns (miscounted letters, a percent trap) or the values are wrong (e.g. sums of
    rows below the data, which the formula overwrites with 0). Only formulas of + - * / with numbers and cells of
    the same sheet are computed (see Book.evaluate); 0 and the same value in every row are placeholders."""
    problems: list[str] = []
    for (s, c), formula in book.columns.items():
        sheet, column = book.sheets[s], book.sheets[s].columns[c]
        tree = book.parsed(formula)[1]
        differ, sent_values = [], set()
        for r, row in enumerate(sheet.rows):
            sent = row[c]
            if sent is None or isinstance(sent, bool) or (isinstance(sent, str) and sent.strip()[:1] in ("", "=")):
                continue
            sent_values.add(str(sent))
            computed = book.evaluate(tree, s, r)
            number, typed = convert_value(sent, column.type if column.type in NUMBER_TYPES else "number",
                                          book.decimal_comma)  # fmt: skip
            if computed is None or not typed or not isinstance(number, int | float) or number == 0:
                continue  # 0 for "the formula fills it in"
            plain = sent if isinstance(sent, int | float) else number  # a percent sent as 0.25 for 25 %
            if not any(abs(x - computed) <= max(0.015, 0.01 * abs(computed)) for x in (number, plain)):
                differ.append((r, sent, computed))
        if differ and len(sent_values) > 1:  # the same value in every row is a placeholder too
            r, sent, computed = differ[0]
            shown = f"{round(computed * 100, 4):g} %" if column.type == "percent" else f"{round(computed, 4):g}"
            more = f" ({len(differ)} of {len(sheet.rows)} rows differ)" if len(differ) > 1 else ""
            header = column.header.strip()
            problems.append(
                f"{book.where(s)}.rows[{r}]: the formula {formula} of '{header}' gives {shown} in this row, but the "
                f"row sends {json.dumps(sent, ensure_ascii=False)}{more}. The formula fills the column, the values "
                f"sent for it are not used: check which columns it uses ({book.columns_hint(s)}), and send null for "
                f"'{header}'; a sum row (net, VAT, gross) needs its own formula there instead, e.g. '=SUM(...)'"
            )
    return problems


# 'Summe', 'Netto-Summe', 'Summe Ausgaben', 'Gesamt (EUR):' …; the whole label: 'Total Quality Management' is none
TOTAL_ROW_LABEL_RE = re.compile(
    r"(?:(?:netto|brutto|zwischen|gesamt)[- ]?)?"
    r"(?:summe(?:\s+[^\W\d_]+)?|gesamt(?:summe|betrag|kosten|preis)?(?:\s+(?:netto|brutto))?|insgesamt|total|sum)"
    r"\s*(?:\([^()]*\))?\s*:?",
    re.IGNORECASE,
)


def _total_row_problem(book: Book, s: int, formulas: dict[int, str]) -> str | None:
    """A total row written into the rows: with column formulas it gets them too (quantity * price of the total
    row), next to a column's "total" it is counted twice, and without its own values it adds up nothing. The
    column's "total" builds a correct one. A row with its own sum formulas, whose other formula columns have
    nothing to compute with, is the model's own total (net, VAT, gross)."""
    sheet, where, rows = book.sheets[s], book.where(s), book.sheets[s].rows
    label = rows[-1][0] if len(rows) >= 2 and rows[-1] else None
    if not isinstance(label, str) or not TOTAL_ROW_LABEL_RE.fullmatch(to_plain(label).strip()):
        return None
    numbers = bool(formulas) or any(c.total != "none" or c.type in NUMBER_TYPES for c in sheet.columns[1:])
    others = [v for c, v in enumerate(rows[-1]) if c > 0 and c not in formulas and v is not None and str(v).strip()]
    totals = any(c.total != "none" for c in sheet.columns)
    last = len(rows) - 1
    own = all((s, last, c) in book.cells or not book.has_inputs(s, last, c) for c in formulas) and (
        bool(others) if not formulas else book.sum_row(s, last)
    )
    if not numbers or (not totals and own):
        return None  # nothing to add up, or the model's own sums ('=SUM(…)' cells are written as formulas)
    return (
        f"{where}.rows[{len(rows) - 1}] ('{label.strip()}') is a total row: remove it and set \"total\" "
        '(e.g. "sum") on the columns to add up; the tool adds the total row with the right formulas. Sums that '
        '"total" cannot build (net, VAT, gross, subtotals) need their own formula in every column that has one, '
        "e.g. '=SUM(D2:D6)' in the row, and no \"total\" on those columns"
    )


@dataclass
class ChartLayout:
    """1-based columns/rows: series are (title, column or row), categories are a column or a row range."""

    by_rows: bool
    series: list[tuple[str, int]]
    categories: tuple[int, int, int, int]  # min_col, max_col, min_row, max_row


def chart_layout(sheet: ExcelSheet, chart: ExcelChart) -> ChartLayout:
    """Work out what to plot (the rows of a prepared sheet, see prepare_sheet); raises ValueError with a list of
    messages for the model."""
    headers = [column.header.strip() for column in sheet.columns]
    last_row = len(sheet.rows) + 1
    listed = f"({', '.join(headers)})"
    errors: list[str] = []
    if chart.series_rows:
        labels = [to_plain(str(row[0] if row and row[0] is not None else "")) for row in sheet.rows]
        keys = [label.casefold() for label in labels]
        series = []
        for name in chart.series_rows:
            if to_plain(name).casefold() in keys:
                row = keys.index(to_plain(name).casefold())
                series.append((labels[row], row + 2))
            else:
                errors.append(f"series_rows: '{name}' is not in the first column ({', '.join(labels[:15])})")
        if chart.value_columns:
            missing = [name for name in chart.value_columns if name.strip() not in headers]
            errors += [f"value_columns: '{name}' is not a column header {listed}" for name in missing]
            columns = sorted(headers.index(name.strip()) + 1 for name in chart.value_columns if name not in missing)
        else:
            columns = []
            for c, column in enumerate(sheet.columns[1:], start=2):
                if column.type not in NUMBER_TYPES or column.formula.strip():
                    break
                columns.append(c)
        if columns and columns != list(range(columns[0], columns[-1] + 1)):
            errors.append("value_columns: with series_rows the columns must be next to each other")
        elif not columns and not errors:
            errors.append("value_columns: no number columns found after the first column; list them in value_columns")
        if errors:
            raise ValueError(errors)
        return ChartLayout(True, series, (columns[0], columns[-1], 1, 1))
    if not chart.category_column.strip() and not chart.value_columns:  # only type/title: plot the number columns
        numbers = [h for c, h in enumerate(headers) if c > 0 and sheet.columns[c].type in NUMBER_TYPES]
        chart = chart.model_copy(update={"category_column": headers[0] if headers else "", "value_columns": numbers})
    if chart.category_column.strip() not in headers:
        errors.append(f"category_column: '{chart.category_column}' is not a column header {listed}")
    if not chart.value_columns:
        errors.append("value_columns: is empty (or use series_rows when the data series are rows)")
    errors += [
        f"value_columns: '{name}' is not a column header {listed} - if these are row labels, use series_rows"
        for name in chart.value_columns
        if name.strip() not in headers
    ]
    for name in chart.value_columns:
        if name.strip() in headers:
            c = headers.index(name.strip())
            column = sheet.columns[c]
            has_numbers = any(
                isinstance(row[c], float)
                or (
                    isinstance(row[c], str)
                    and (parse_number_text(row[c]) is not None or row[c].strip().startswith("="))
                )
                for row in sheet.rows
            )
            if not column.formula.strip() and not has_numbers:
                errors.append(f"value_columns: '{name}' has no numbers to plot")
    if errors:
        raise ValueError(errors)
    category = headers.index(chart.category_column.strip()) + 1
    series = [(name.strip(), headers.index(name.strip()) + 1) for name in chart.value_columns]
    return ChartLayout(False, series, (category, category, 2, last_row))


CHART_ROWS = 22  # charts below each other next to the table, one every 22 rows


class XlsxWriter:
    def __init__(self, currency: str, texts: dict[str, str], decimal_comma: bool | None = None,
                 design: Design = DEFAULT_DESIGN, title: str = "") -> None:  # fmt: skip
        self.currency = currency
        self.texts = texts
        self.decimal_comma = decimal_comma
        self.design = design
        self.title = title
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        # Charts and cells the user adds later take the theme's fonts.
        self.workbook.loaded_theme = re.sub(THEME_FONTS_RE, lambda m: m[1] + design.font + m[2], theme_xml)
        self.body_font = Font(name=design.font, size=11, color=TEXT_COLOR)
        self.bold_font = Font(name=design.font, size=11, bold=True, color=TEXT_COLOR)
        self.header_font = Font(name=design.font, size=11, bold=True, color=WHITE)
        self.chart_titles: dict[str, list[str]] = {}  # per sheet, for the text version
        self.narrow_font = design.font.casefold() in NARROW_FONTS  # the column widths are measured for Calibri

    def render(self, book: Book) -> bytes:
        self.book = book
        for s, name in enumerate(book.names):
            self._sheet(self.workbook.create_sheet(title=name), s)
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    def _sheet(self, ws: Worksheet, s: int) -> None:
        """Header row in the accent colour with white text, every second row lightly tinted, thin lines between the
        rows, no grid lines; printed on one page width with the header row on every page."""
        sheet = self.book.sheets[s]
        header_fill = PatternFill("solid", fgColor=self.design.accent)
        rows = sheet.rows
        last_row = len(rows) + 1
        widths = [len(column.header.strip()) for column in sheet.columns]
        for c, column in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=1, column=c, value=excel_text(column.header.strip()) or None)
            cell.data_type = "s"
            cell.quotePrefix = str(cell.value or "").startswith(FORMULA_START)
            cell.font = self.header_font
            cell.fill = header_fill
            kind = (
                "date"
                if column.type in TIME_TYPES
                else "number"
                if column.type in NUMBER_TYPES or column.formula
                else "text"
            )
            cell.alignment = Alignment(
                horizontal=ALIGN[kind].horizontal, vertical="center", wrap_text=True, indent=ALIGN[kind].indent
            )
        if self.narrow_font and all("\n" not in column.header for column in sheet.columns):
            ws.row_dimensions[1].height = 22  # otherwise Excel fits the height to the wrapped header
        columns_values: list[list[Any]] = [[] for _ in sheet.columns]
        for r, row in enumerate(rows, start=2):
            for c, (column, raw) in enumerate(zip(sheet.columns, row, strict=True), start=1):
                cell = ws.cell(row=r, column=c)
                cell.font, cell.border = self.body_font, ROW_LINE
                formula = self.book.cells.get((s, r - 2, c - 1)) or self.book.columns.get((s, c - 1))
                if (
                    formula
                    and (s, r - 2, c - 1) not in self.book.cells
                    and self.book.sum_row(s, r - 2)
                    and not self.book.has_inputs(s, r - 2, c - 1)
                ):
                    formula = ""  # a sum row: the line item formula would show 0 there
                if formula:  # checked against the whitelist
                    cell.value = formula.replace("{row}", str(r))
                    cell.alignment = ALIGN["date" if column.type in TIME_TYPES else "number"]
                    widths[c - 1] = max(widths[c - 1], 12)
                    continue
                value, typed = convert_value(raw, column.type, self.decimal_comma)
                cell.value = value
                cell.alignment = ALIGN[
                    "date"
                    if isinstance(value, date) or (column.type == "time" and typed)
                    else "number"
                    if isinstance(value, int | float)
                    else "text"
                ]
                if isinstance(value, str):
                    cell.data_type = "s"  # other text never becomes a live formula ("=HYPERLINK(...)" stays text)
                    cell.quotePrefix = value.startswith(FORMULA_START)  # ... not even when the user edits the cell
                    if "\n" in value:
                        cell.alignment = ALIGN["lines"]
                    widths[c - 1] = max(widths[c - 1], *(len(line) for line in value.split("\n")))
                elif value is not None:
                    columns_values[c - 1].append(value)
                    wide = column.type in ("currency", *TIME_TYPES)
                    widths[c - 1] = max(widths[c - 1], 12 if wide else len(str(value)) + 2)
        date_format = self.texts["date_format"]
        format_rows = last_row + (1 if has_totals(sheet) else 0)
        formats = []
        for c, column in enumerate(sheet.columns, start=1):
            fmt = self.book.number_format(s, c - 1, columns_values[c - 1], date_format)
            formats.append(fmt)
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, max_row=format_rows, min_col=c, max_col=c):
                    cell.number_format = fmt
            width = (widths[c - 1] + 4) * (1 if self.narrow_font else 1.15)  # + 4: the indent on both sides
            ws.column_dimensions[get_column_letter(c)].width = min(max(width, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        ws.freeze_panes = "A2"
        data_last = last_row - self.book.trailing_sum_rows(s)  # sorting must not move the sums
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{data_last}"
        ws.sheet_view.showGridLines = False
        if len(rows) >= 2:  # every second row tinted by a rule, not by cell fills: it stays in place when sorted
            band = PatternFill("solid", start_color=self.design.band, end_color=self.design.band)
            ws.conditional_formatting.add(
                f"A2:{get_column_letter(len(sheet.columns))}{last_row}",
                FormulaRule(formula=["MOD(ROW(),2)=1"], fill=band),
            )
        self._totals(ws, sheet, last_row, formats)
        for k, chart in enumerate(sheet.charts):
            self._chart(ws, sheet, chart, data_last, 2 + k * CHART_ROWS)
        self._print_setup(ws, landscape=len(sheet.columns) > WIDE_SHEET_COLUMNS or bool(sheet.charts))

    def markdown(self) -> str:
        """The written sheets as Markdown tables, numbers shown like in the cells, formulas as formulas."""
        parts = []
        for ws in self.workbook.worksheets:
            rows = [[self._cell_text(cell) for cell in row] for row in ws.iter_rows()]
            parts += [f"## {ws.title}", markdown_table([row for row in rows if any(row)])]
            parts += [f"*{self.texts['chart']}: {title}*" for title in self.chart_titles.get(ws.title, [])]
        return "\n\n".join(part for part in parts if part)

    def _cell_text(self, cell: Any) -> str:
        """A cell's value as it is shown: decimals, %, currency, date and time from its number format."""
        value, fmt = cell.value, str(cell.number_format)
        if value is None or isinstance(value, str):
            return value or ""
        if isinstance(value, datetime | date):
            shown = value.strftime("%d.%m.%Y" if self.texts["date_format"].startswith("DD") else "%Y-%m-%d")
            return shown + (value.strftime(" %H:%M") if isinstance(value, datetime) and "hh:mm" in fmt else "")
        if ":mm" in fmt and isinstance(value, int | float):  # a time: the fraction of a day
            minutes = round(value * 1440)
            hours = minutes // 60 if fmt.startswith("[h]") else minutes // 60 % 24
            return f"{hours:02d}:{minutes % 60:02d}" if not fmt.startswith("[h]") else f"{hours}:{minutes % 60:02d}"
        percent = "%" in fmt
        value = value * 100 if percent else value
        decimals = re.search(r"\.(0+)", fmt)
        if fmt == "General":
            text = f"{value:,}"
        else:
            text = f"{value:,.{len(decimals[1]) if decimals else 0}f}"
        if self.decimal_comma:
            text = text.replace(",", "\0").replace(".", ",").replace("\0", ".")
        symbol = next((s for s in ("€", "$", "£", "CHF") if s in fmt), "")
        text += " %" if percent else ""
        return symbol + text if symbol in ("$", "£") else f"{text} {symbol}".strip()

    def _print_setup(self, ws: Worksheet, landscape: bool) -> None:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4 if self.texts.get("paper") == "a4" else ws.PAPERSIZE_LETTER
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0  # one page wide, as many pages tall as needed
        ws.print_title_rows = "1:1"
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)
        for part, text in ((ws.oddHeader.left, self.title), (ws.oddHeader.right, ws.title)):
            part.text, part.size, part.color = text.replace("&", "&&")[:100], 9, MUTED_COLOR
            part.font = f"{self.design.font},Regular"
        page = self.texts["page_footer"].replace("{page}", "&P").replace("{pages}", "&N")
        footer = ws.oddFooter.right
        footer.text, footer.size, footer.color, footer.font = page, 9, MUTED_COLOR, f"{self.design.font},Regular"

    def _totals(self, ws: Worksheet, sheet: ExcelSheet, last_row: int, formats: list[str | None]) -> None:
        if not has_totals(sheet):
            return
        row = last_row + 1
        lines = Border(top=Side(style="medium", color=self.design.accent), bottom=Side(style="thin", color=RULE_COLOR))
        fill = PatternFill("solid", fgColor=self.design.tint)
        for c, column in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=row, column=c)
            if column.total != "none":
                letter = get_column_letter(c)
                cell.value = f"={TOTAL_FUNCTIONS[column.total]}({letter}2:{letter}{last_row})"
                if column.total == "count":
                    cell.number_format = "0"  # a count, not an amount or a percentage
                elif column.total == "sum" and formats[c - 1] in ("hh:mm", "hh:mm:ss"):
                    cell.number_format = "[h]:mm"  # the hours of all rows, not the time of day they add up to
            elif c == 1:  # no room for the label when the first column has a total itself
                cell.value = self.texts["total"]
                cell.data_type = "s"
            cell.font, cell.border, cell.fill = self.bold_font, lines, fill
            cell.alignment = ALIGN["number" if column.total != "none" else "text"]

    def _chart(self, ws: Worksheet, sheet: ExcelSheet, chart: ExcelChart, last_row: int, top: int) -> None:
        layout = chart_layout(sheet, chart)
        series = layout.series[:1] if chart.type == "pie" else layout.series
        if chart.type == "pie":
            excel_chart: Any = PieChart()
        elif chart.type == "line":
            excel_chart = LineChart()
        else:
            excel_chart = BarChart()
            excel_chart.type = "bar" if chart.type == "bar" else "col"
        min_col, max_col, min_row, max_row = layout.categories
        if not layout.by_rows:
            max_row = min(max_row, last_row)  # not the sum rows below the data
        for title, index in series:
            if layout.by_rows:
                values = Reference(ws, min_col=min_col, max_col=max_col, min_row=index, max_row=index)
            else:
                values = Reference(ws, min_col=index, min_row=2, max_row=last_row)
            excel_chart.series.append(Series(values, title=excel_text(title)))
        excel_chart.set_categories(Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row, max_row=max_row))
        if chart.title.strip():
            excel_chart.title = excel_text(chart.title.strip())
        if len(series) == 1 and chart.type != "pie":
            excel_chart.legend = None  # a single series needs no legend
        elif excel_chart.legend is not None:
            excel_chart.legend.position = "b"
        categories = (max_col - min_col + 1) if layout.by_rows else (max_row - min_row + 1)
        self._chart_style(excel_chart, chart.type, categories)
        if chart.unit.strip() and chart.type != "pie":
            excel_chart.y_axis.title = excel_text(chart.unit.strip())[:60]
        excel_chart.width, excel_chart.height = 18, 10
        ws.add_chart(excel_chart, f"{get_column_letter(len(sheet.columns) + 2)}{top}")
        names = ", ".join(title for title, _ in series)
        self.chart_titles.setdefault(ws.title, []).append(f"{to_plain(chart.title.strip()) or names} ({chart.type})")

    def _chart_style(self, excel_chart: Any, kind: str, categories: int) -> None:
        """Colours of the design, light grid lines only on the value axis, values on the bars when there are few."""
        if kind == "pie":
            (series,) = excel_chart.series
            for i, color in enumerate(self.design.series(categories)):
                point = DataPoint(idx=i)
                point.graphicalProperties.solidFill = color
                point.graphicalProperties.line.solidFill = WHITE
                series.dPt.append(point)
            # outside the slices: dark text on the dark ones would not be readable
            excel_chart.dataLabels = DataLabelList(
                showPercent=True,
                showVal=False,
                showCatName=False,
                showSerName=False,
                showLegendKey=False,
                dLblPos="outEnd",
            )
            return
        for series, color in zip(excel_chart.series, self.design.series(len(excel_chart.series)), strict=True):
            if kind == "line":
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 28575  # 2.25 pt
                series.marker.symbol, series.marker.size = "circle", 6
                series.marker.graphicalProperties = GraphicalProperties(
                    solidFill=color, ln=LineProperties(solidFill=color)
                )
                series.smooth = False
            else:
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color
        if kind != "line":
            excel_chart.gapWidth = 60
            if len(excel_chart.series) > 1:
                excel_chart.overlap = -10
        excel_chart.x_axis.delete = excel_chart.y_axis.delete = False  # openpyxl 3.1 leaves it open; Excel hides them
        excel_chart.y_axis.majorGridlines = ChartLines(
            spPr=GraphicalProperties(ln=LineProperties(solidFill=RULE_COLOR))
        )
        excel_chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        excel_chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=RULE_COLOR))
        if len(excel_chart.series) == 1 and categories <= 8:
            excel_chart.dataLabels = DataLabelList(showVal=True, showPercent=False, showCatName=False,
                                                   showSerName=False, showLegendKey=False)  # fmt: skip


def build_xlsx(
    sheets: list[ExcelSheet], currency: str, texts: dict[str, str], decimal_comma: bool | None = None,
    design: Design = DEFAULT_DESIGN, title: str = "",
) -> tuple[bytes, str]:  # fmt: skip
    """The file and its text as Markdown."""
    sheets = [prepare_sheet(sheet, decimal_comma) for sheet in sheets]
    book = check_excel(sheets, decimal_comma, sheet_names(sheets, texts["sheet"]), currency)
    writer = XlsxWriter(currency, texts, decimal_comma, design, title)
    return writer.render(book), writer.markdown()


# ------------------------------------------------------------------------------------------ PowerPoint --

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
TITLE_TOP, TITLE_HEIGHT = Inches(0.45), Inches(1.0)
BAR_TOP = Inches(1.5)
BODY_TOP = Inches(1.8)
BODY_HEIGHT = SLIDE_HEIGHT - BODY_TOP - Inches(0.6)
COLUMN_GAP = Inches(0.5)
CARD_PADDING = Inches(0.25)  # two_columns: the columns are light cards with the text inside
COLUMN_TEXT_WIDTH = (CONTENT_WIDTH - COLUMN_GAP) // 2 - 2 * CARD_PADDING
COLUMN_TEXT_HEIGHT = BODY_HEIGHT - 2 * CARD_PADDING
FOOTER_TOP = SLIDE_HEIGHT - Inches(0.45)
SLIDE_TEXT_COLOR = "3A3A3A"
# Layout indices of python-pptx's default template.
LAYOUT_TITLE, LAYOUT_CONTENT, LAYOUT_SECTION, LAYOUT_TWO_CONTENT, LAYOUT_TITLE_ONLY = 0, 1, 2, 3, 5
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


SHOWN_ON = {  # the layouts that show a field; on other slides its content would be lost
    "subtitle": ("title", "section"), "bullets": ("bullets",), "left": ("two_columns",), "right": ("two_columns",),
    "left_heading": ("two_columns",), "right_heading": ("two_columns",), "table": ("table",), "chart": ("chart",),
}  # fmt: skip


def check_slides(title: str, slides: list[Slide]) -> None:
    check = Checker()
    check.text(title)
    check.require(bool(slides), "slides: the presentation has no slides")
    check.require(len(slides) <= MAX_SLIDES, f"slides: at most {MAX_SLIDES} slides")
    for i, slide in enumerate(slides):
        where = f"slides[{slide._index if slide._index >= 0 else i}] ({slide.layout})"
        check.text(slide.title, slide.subtitle, slide.notes, *slide.bullets, *slide.left, *slide.right)
        check.require(
            bool(slide.title.strip() or (slide.layout == "title" and title.strip())), f"{where}: 'title' is empty"
        )
        for field_name in ("bullets", "left", "right"):
            items = getattr(slide, field_name)
            check.require(
                len(items) <= MAX_SLIDE_ITEMS,
                f"{where}.{field_name}: at most {MAX_SLIDE_ITEMS} items per slide, split the content over more slides",
            )
        for field_name, layouts in SHOWN_ON.items():
            check.require(
                slide.layout in layouts or not has_content(getattr(slide, field_name)),
                f"{where}: '{field_name}' is not shown on a {slide.layout} slide; use layout {' or '.join(layouts)}"
                " for it, or leave it out",
            )
        too_much = f"{where}: too much text for one slide; split it over two slides or move details into the notes"
        if slide.layout == "bullets":
            if check.require(any(b.strip() for b in slide.bullets), f"{where}: 'bullets' is empty"):
                check.require(
                    len(slide.bullets) > MAX_SLIDE_ITEMS or body_font_size(slide.bullets) is not None, too_much
                )
        elif slide.layout == "two_columns":
            if check.require(
                any(b.strip() for b in slide.left + slide.right), f"{where}: 'left' and 'right' are empty"
            ):
                lines = 1 if slide.left_heading.strip() or slide.right_heading.strip() else 0
                check.require(
                    max(len(slide.left), len(slide.right)) > MAX_SLIDE_ITEMS
                    or all(
                        body_font_size(items, 22, COLUMN_TEXT_WIDTH, lines, COLUMN_TEXT_HEIGHT) is not None
                        for items in (slide.left, slide.right)
                    ),
                    too_much,
                )
        elif slide.layout == "table":
            if check.require(bool(slide.table), f"{where}: 'table' is empty") and check.require(
                len(slide.table) > 1,
                f"{where}.table: only a header row; put the header and all data rows into one table",
            ):
                check.text(*(cell for row in slide.table for cell in row))
                check.require(
                    len(slide.table) <= MAX_SLIDE_TABLE_ROWS,
                    f"{where}.table: at most {MAX_SLIDE_TABLE_ROWS} rows per slide, split the table over more slides",
                )
                check.require(
                    max(len(row) for row in slide.table) <= MAX_SLIDE_TABLE_COLUMNS,
                    f"{where}.table: at most {MAX_SLIDE_TABLE_COLUMNS} columns",
                )
                check.require(
                    len(slide.table) > MAX_SLIDE_TABLE_ROWS
                    or table_font_size(slide.table, table_columns(slide.table)) is not None,
                    f"{where}.table: the cells hold too much text to fit on one slide; shorten them or split the table",
                )
        elif slide.layout == "chart":
            chart = slide.chart
            if check.require(chart is not None, f"{where}: 'chart' is missing"):
                assert chart is not None
                categories = len(chart.categories)
                check.require(
                    0 < categories <= MAX_CHART_CATEGORIES,
                    f"{where}.chart.categories: 1 to {MAX_CHART_CATEGORIES} labels",
                )
                check.require(
                    0 < len(chart.series) <= MAX_CHART_SERIES, f"{where}.chart.series: 1 to {MAX_CHART_SERIES} series"
                )
                for k, series in enumerate(chart.series):
                    check.require(
                        len(series.values) == categories,
                        f"{where}.chart.series[{k}]: {len(series.values)} values for {categories} categories",
                    )
    check.raise_if_failed()


MIN_BODY_FONT_SIZE = 12
EMU_PER_POINT = 12700


class TextLabelWorkbookWriter(CategoryWorkbookWriter):
    """Writes the chart's embedded workbook with category and series names as plain text.

    python-pptx uses xlsxwriter's write(), which turns '=...' into a formula and URLs into links; Excel
    would evaluate them when the user clicks "Edit Data" in PowerPoint.
    """

    def _write_cat_column(self, worksheet: Any, col: int, level: Any, num_format: Any) -> None:
        worksheet.set_column(col, col, 10)
        for offset, name in level:
            worksheet.write_string(offset + 1, col, str(name), num_format)

    def _write_series(self, workbook: Any, worksheet: Any) -> None:
        col_offset = self._chart_data.categories.depth
        for index, series in enumerate(self._chart_data):
            num_format = workbook.add_format({"num_format": series.number_format})
            worksheet.write_string(0, index + col_offset, str(series.name))
            for row, value in enumerate(series.values, start=1):
                worksheet.write_number(row, index + col_offset, value, num_format)


class TextLabelChartData(CategoryChartData):
    @lazyproperty
    def _workbook_writer(self) -> CategoryWorkbookWriter:
        return TextLabelWorkbookWriter(self)


def body_font_size(
    items: list[str], base: int = 24, width: int = CONTENT_WIDTH, extra_lines: int = 0, height: int = BODY_HEIGHT
) -> int | None:
    """Largest font size (from `base` down) at which the items fit the text box, or None if they do not fit
    even at the smallest size; python-pptx cannot shrink text to fit. Estimates half an em per character,
    1.3 line height and 6 pt before each item; sub items are set smaller (see sub_item_size)."""
    for size in range(base, MIN_BODY_FONT_SIZE - 1, -2):
        lines = extra_lines
        for item in items:
            level, text = split_level(item)
            item_size = sub_item_size(size) if level else size
            chars_per_line = max(int(width / EMU_PER_POINT / (item_size * 0.5)), 1)
            lines += max(1, -(-len(to_plain(text)) // chars_per_line)) * item_size / size
        if lines * size * 1.3 + len(items) * 6 <= height / EMU_PER_POINT:
            return size
    return None


def sub_item_size(size: int) -> int:
    return size - 4 if size >= 18 else size - 2


def table_font_size(rows: list[list[str]], columns: int) -> tuple[int, int] | None:
    """(font size, row height in EMU) at which the table fits below the title, or None. Long cells wrap,
    so each row is as high as its cell with the most lines."""
    column_width = CONTENT_WIDTH / max(columns, 1) / EMU_PER_POINT - 10  # minus the cell margins
    longest_words = [
        max((len(word) for row in rows if c < len(row) for word in to_plain(row[c]).split()), default=1)
        for c in range(columns)
    ]
    for size in (18, 16, 14, 12, 11, 10):
        words_fit = sum(min(n, MAX_WORD_CHARS) * CHAR_WIDTH * size / 11 + CELL_PADDING for n in longest_words)
        if words_fit > CONTENT_WIDTH and size > 10:
            continue  # long words would be broken in the middle
        chars_per_line = max(int(column_width / (size * 0.5)), 1)
        line_counts = [
            max((max(1, -(-len(to_plain(cell)) // chars_per_line)) for cell in row), default=1) for row in rows
        ]
        row_height = int(Inches(0.1) + PptxPt(size * 1.25))
        height = sum(int(Inches(0.1) + PptxPt(size * 1.25 * lines)) for lines in line_counts)
        if height <= BODY_HEIGHT:
            return size, row_height
    return None


def _no_bullet(paragraph: PptxParagraph) -> None:
    paragraph._p.get_or_add_pPr().append(PptxOxmlElement("a:buNone"))


class PptxWriter:
    """Title and section slides on a dark or accent background, content slides white with the title in the accent
    colour, a thin footer line with the presentation title and the slide number."""

    def __init__(self, title: str, design: Design = DEFAULT_DESIGN) -> None:
        self.title = title
        self.design = design
        self.accent = PptxRGBColor.from_string(design.accent)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
        # All text uses the theme's fonts, also the charts and the slides the user adds later.
        master = self.prs.slide_masters[0].part
        theme = next(rel.target_part for rel in master.rels.values() if rel.reltype == PPTX_RT.THEME)
        theme._blob = re.sub(THEME_FONTS_RE.encode(), lambda m: m[1] + design.font.encode() + m[2], theme.blob)
        self.number = 0

    def render(self, slides: list[Slide]) -> bytes:
        for self.number, slide in enumerate(slides, start=1):
            getattr(self, f"_slide_{slide.layout}")(slide)
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()

    # -- slide types --

    def _slide_title(self, data: Slide) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUT_TITLE])
        self._background(slide, self.design.dark)
        title, subtitle = slide.placeholders[0], slide.placeholders[1]
        text = data.title.strip() or self.title
        self._place(title, MARGIN, Inches(1.5), CONTENT_WIDTH, Inches(2.6))
        self._text(title, [text], size=40 if len(text) <= 50 else 32, bold=True, color=WHITE, anchor=MSO_ANCHOR.BOTTOM)
        self._bar(slide, Inches(4.3), MARGIN, Inches(1.4), self.design.light)
        self._place(subtitle, MARGIN, Inches(4.55), CONTENT_WIDTH, Inches(1.5))
        lines = [data.subtitle] if data.subtitle.strip() else []
        self._text(subtitle, lines, size=20, color=mix(self.design.accent, WHITE, 0.8), anchor=MSO_ANCHOR.TOP)
        self._finish(slide, data)

    def _slide_section(self, data: Slide) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUT_SECTION])
        self._background(slide, self.design.accent)
        title, subtitle = slide.placeholders[0], slide.placeholders[1]
        self._place(title, MARGIN, Inches(2.2), CONTENT_WIDTH, Inches(1.7))
        self._text(title, [data.title], size=36, bold=True, color=WHITE, anchor=MSO_ANCHOR.BOTTOM)
        self._bar(slide, Inches(4.05), MARGIN, Inches(1.4), mix(self.design.accent, WHITE, 0.7))
        self._place(subtitle, MARGIN, Inches(4.3), CONTENT_WIDTH, Inches(1.2))
        lines = [data.subtitle] if data.subtitle.strip() else []
        self._text(subtitle, lines, size=20, color=mix(self.design.accent, WHITE, 0.85), anchor=MSO_ANCHOR.TOP)
        self._finish(slide, data)

    def _slide_bullets(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_CONTENT, data)
        body = slide.placeholders[1]
        self._place(body, MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_HEIGHT)
        self._items(body.text_frame, data.bullets, body_font_size(data.bullets) or MIN_BODY_FONT_SIZE)
        self._finish(slide, data)

    def _slide_two_columns(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TWO_CONTENT, data)
        width = (CONTENT_WIDTH - COLUMN_GAP) // 2
        heading_lines = 1 if data.left_heading.strip() or data.right_heading.strip() else 0
        size = min(
            body_font_size(items, 22, COLUMN_TEXT_WIDTH, heading_lines, COLUMN_TEXT_HEIGHT) or MIN_BODY_FONT_SIZE
            for items in (data.left, data.right)
        )
        for placeholder, left, heading, items in (
            (slide.placeholders[1], MARGIN, data.left_heading, data.left),
            (slide.placeholders[2], MARGIN + width + COLUMN_GAP, data.right_heading, data.right),
        ):
            if heading.strip() or any(item.strip() for item in items):
                self._card(slide, left, BODY_TOP, width, BODY_HEIGHT)
            self._place(
                placeholder, left + CARD_PADDING, BODY_TOP + CARD_PADDING, COLUMN_TEXT_WIDTH, COLUMN_TEXT_HEIGHT
            )
            self._items(placeholder.text_frame, items, size, heading=heading)
        self._finish(slide, data)

    def _slide_table(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TITLE_ONLY, data)
        rows, columns = len(data.table), table_columns(data.table)
        size, row_height = table_font_size(data.table, columns) or (10, int(Inches(0.1) + PptxPt(12.5)))
        shape = slide.shapes.add_table(rows, columns, MARGIN, BODY_TOP, CONTENT_WIDTH, row_height * rows)
        widths = column_widths(data.table, columns, CONTENT_WIDTH, size)
        numbers = number_columns(data.table, columns)
        for column, width in zip(shape.table.columns, widths, strict=True):
            column.width = Emu(width)  # by content, so long words are not broken
        for r, row in enumerate(data.table):
            fill = self.design.accent if r == 0 else self.design.band if r % 2 == 0 else WHITE
            for c in range(columns):
                cell = shape.table.cell(r, c)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = PptxRGBColor.from_string(fill)
                frame = cell.text_frame
                frame.word_wrap = True
                if numbers[c]:
                    frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                runs = parse_inline(row[c] if c < len(row) else "")
                self._runs(
                    frame.paragraphs[0], runs, size, bold=r == 0 or None, color=WHITE if r == 0 else SLIDE_TEXT_COLOR
                )
        self._finish(slide, data)

    def _slide_chart(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TITLE_ONLY, data)
        assert data.chart is not None
        chart_data = TextLabelChartData()
        chart_data.categories = [clean_text(to_plain(c)) for c in data.chart.categories]
        series_list = data.chart.series[:1] if data.chart.type == "pie" else data.chart.series
        for series in series_list:
            chart_data.add_series(clean_text(to_plain(series.name)), series.values)
        frame = slide.shapes.add_chart(
            CHART_TYPES[data.chart.type], MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_HEIGHT, chart_data
        )
        chart = frame.chart
        chart.has_title = False  # the slide title says it
        chart.font.size = PptxPt(14)
        chart.font.color.rgb = PptxRGBColor.from_string(SLIDE_TEXT_COLOR)
        if data.chart.unit.strip() and data.chart.type != "pie":
            axis = chart.value_axis
            axis.has_title = True
            axis.axis_title.text_frame.text = clean_text(to_plain(data.chart.unit))[:60]
        chart.has_legend = data.chart.type == "pie" or len(series_list) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if data.chart.type == "pie":
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_value = False
            plot.data_labels.number_format = "0%"
            plot.data_labels.number_format_is_linked = False
        self._chart_style(chart, data.chart.type, len(series_list), len(data.chart.categories))
        self._finish(slide, data)

    def _chart_style(self, chart: Any, kind: str, series_count: int, categories: int) -> None:
        """Colours of the design, light grid lines on the value axis only, values at the bars when there are few."""
        plot = chart.plots[0]
        if kind == "pie":
            for point, color in zip(plot.series[0].points, self.design.series(categories), strict=False):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = PptxRGBColor.from_string(color)
                point.format.line.color.rgb = PptxRGBColor.from_string(WHITE)
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END  # dark text on the dark slices: not readable
            return
        for series, color in zip(plot.series, self.design.series(series_count), strict=False):
            rgb = PptxRGBColor.from_string(color)
            if kind == "line":
                series.smooth = False
                series.format.line.color.rgb = rgb
                series.format.line.width = PptxPt(2.75)
                series.marker.style, series.marker.size = XL_MARKER_STYLE.CIRCLE, 8
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = rgb
                series.marker.format.line.color.rgb = rgb
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = rgb
                series.format.line.fill.background()
        if kind != "line":
            plot.gap_width = 60
            if series_count > 1:
                plot.overlap = -10
        value_axis, category_axis = chart.value_axis, chart.category_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = PptxRGBColor.from_string(RULE_COLOR)
        value_axis.format.line.fill.background()
        value_axis.major_tick_mark = category_axis.major_tick_mark = XL_TICK_MARK.NONE
        value_axis.tick_labels.font.size = PptxPt(12)
        value_axis.tick_labels.font.color.rgb = PptxRGBColor.from_string(MUTED_COLOR)
        category_axis.format.line.color.rgb = PptxRGBColor.from_string(RULE_COLOR)
        if series_count == 1 and categories <= 8:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.font.size = PptxPt(12)  # no numFmt: without formatCode PowerPoint repairs the file
            labels.position = XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END

    # -- helpers --

    def _content_slide(self, layout: int, data: Slide) -> Any:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        title = slide.placeholders[0]
        self._place(title, MARGIN, TITLE_TOP, CONTENT_WIDTH, TITLE_HEIGHT)
        self._text(
            title,
            [data.title],
            size=28 if len(data.title) <= 60 else 22,
            bold=True,
            color=self.accent,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        self._bar(slide, BAR_TOP, MARGIN, Inches(0.9))
        self._footer(slide)
        return slide

    def _footer(self, slide: Any) -> None:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, FOOTER_TOP - Inches(0.06), CONTENT_WIDTH, Emu(9525))
        line.fill.solid()
        line.fill.fore_color.rgb = PptxRGBColor.from_string(RULE_COLOR)
        line.line.fill.background()
        number_width = Inches(1.0)
        for text, left, width, align in (
            (clean_text(to_plain(self.title))[:120], MARGIN, CONTENT_WIDTH - number_width, PP_ALIGN.LEFT),
            (str(self.number), MARGIN + CONTENT_WIDTH - number_width, number_width, PP_ALIGN.RIGHT),
        ):
            if not text.strip():
                continue
            box = slide.shapes.add_textbox(left, FOOTER_TOP, width, Inches(0.3))
            frame = box.text_frame
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            frame.word_wrap = True
            frame.paragraphs[0].alignment = align
            self._runs(frame.paragraphs[0], [Run(text)], 10, color=MUTED_COLOR)

    def _finish(self, slide: Any, data: Slide) -> None:
        for placeholder in list(slide.placeholders):  # remove placeholders left empty ("Click to add text")
            if not placeholder.has_text_frame or not placeholder.text_frame.text.strip():
                placeholder._element.getparent().remove(placeholder._element)
        if data.notes.strip():
            slide.notes_slide.notes_text_frame.text = clean_text(data.notes.strip())

    @staticmethod
    def _place(shape: SlidePlaceholder, left: int, top: int, width: int, height: int) -> None:
        shape.left, shape.top, shape.width, shape.height = Emu(left), Emu(top), Emu(width), Emu(height)

    @staticmethod
    def _background(slide: Any, color: str) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptxRGBColor.from_string(color)

    def _bar(self, slide: Any, top: int, left: int, width: int, color: str | None = None) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PptxRGBColor.from_string(color) if color else self.accent
        bar.line.fill.background()

    def _card(self, slide: Any, left: int, top: int, width: int, height: int) -> None:
        """A light rounded box behind the text, moved to the back so the text stays on top."""
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.adjustments[0] = 0.04
        card.fill.solid()
        card.fill.fore_color.rgb = PptxRGBColor.from_string(self.design.band)
        card.line.fill.background()
        card.shadow.inherit = False
        tree = slide.shapes._spTree
        tree.remove(card._element)
        tree.insert(2, card._element)  # after the group's own properties

    def _text(self, shape: SlidePlaceholder, lines: list[str], size: int, bold: bool | None = None,
              color: Any = None, anchor: Any = None, align: Any = PP_ALIGN.LEFT) -> None:  # fmt: skip
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        if anchor is not None:
            frame.vertical_anchor = anchor
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.alignment = align
            self._runs(paragraph, parse_inline(line.strip()), size, bold=bold, color=color)

    def _items(self, frame: TextFrame, items: list[str], size: int, heading: str = "") -> None:
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        first = True
        if heading.strip():
            paragraph = frame.paragraphs[0]
            _no_bullet(paragraph)
            self._runs(paragraph, parse_inline(heading.strip()), size + 2, bold=True, color=self.accent)
            first = False
        for item in items:
            level, text = split_level(item)
            if not text:
                continue
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.level = level
            self._runs(paragraph, parse_inline(text), sub_item_size(size) if level else size, color=SLIDE_TEXT_COLOR)
            paragraph.space_before = PptxPt(6 if level == 0 else 2)
            self._bullet_color(paragraph)

    def _bullet_color(self, paragraph: PptxParagraph) -> None:
        color = PptxOxmlElement("a:buClr")
        rgb = PptxOxmlElement("a:srgbClr")
        rgb.set("val", self.design.accent)
        color.append(rgb)
        paragraph._p.get_or_add_pPr().insert_element_before(
            color, "a:buSzTx", "a:buSzPct", "a:buSzPts", "a:buFontTx", "a:buFont", "a:buNone", "a:buAutoNum",
            "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
        )  # fmt: skip

    @staticmethod
    def _runs(
        paragraph: PptxParagraph, runs: list[Run], size: int, bold: bool | None = None, color: Any = None
    ) -> None:
        for run in runs:
            pptx_run = paragraph.add_run()
            pptx_run.text = run.text
            pptx_run.font._rPr.set("cap", "none")
            font = pptx_run.font
            font.size = PptxPt(size)
            font.bold = True if (bold or run.bold) else None
            font.italic = True if run.italic else None
            if run.code:
                font.name = CODE_FONT
            if color is not None:
                font.color.rgb = color if isinstance(color, PptxRGBColor) else PptxRGBColor.from_string(color)
            if run.url:
                pptx_run.hyperlink.address = run.url


def slides_markdown(title: str, slides: list[Slide], texts: dict[str, str]) -> str:
    parts = [f"# {to_plain(title)}"] if title.strip() else []
    for number, slide in enumerate(slides, start=1):
        parts += [f"## {number}. {to_plain(slide.title.strip() or title)}", slide.subtitle.strip()]
        parts.append(markdown_items(slide.bullets))
        for heading, items in ((slide.left_heading, slide.left), (slide.right_heading, slide.right)):
            parts += [f"**{heading.strip()}**" if heading.strip() else "", markdown_items(items)]
        parts.append(markdown_table(slide.table))
        if slide.chart is not None:
            chart = slide.chart
            unit = f", {chart.unit.strip()}" if chart.unit.strip() else ""
            series = chart.series[:1] if chart.type == "pie" else chart.series
            rows = [["", *chart.categories], *([s.name, *(f"{v:g}" for v in s.values)] for s in series)]
            parts += [f"*{texts['chart']} ({chart.type}{unit})*", markdown_table(rows)]
        if slide.notes.strip():
            parts.append(f"> {texts['notes']}: " + " ".join(slide.notes.split()))
    return "\n\n".join(part for part in parts if part.strip())


def build_pptx(
    title: str, slides: list[Slide], design: Design = DEFAULT_DESIGN, texts: dict[str, str] | None = None
) -> tuple[bytes, str]:
    """The file and its text as Markdown."""
    check_slides(title, slides)
    return PptxWriter(title, design).render(slides), slides_markdown(title, slides, texts or TRANSLATIONS["en"])


# ------------------------------------------------------------------------------------------- delivery --


def download_js(data: bytes, filename: str, mime: str) -> str:
    """Browser snippet that downloads the file; all dynamic values are JSON-encoded."""
    return f"""
const bytes = Uint8Array.from(atob({json.dumps(base64.b64encode(data).decode())}), (c) => c.charCodeAt(0));
const url = URL.createObjectURL(new Blob([bytes], {{ type: {json.dumps(mime)} }}));
const link = document.createElement('a');
link.href = url;
link.download = {json.dumps(filename)};
document.body.appendChild(link);
link.click();
link.remove();
setTimeout(() => URL.revokeObjectURL(url), 10000);
return true;
"""


@dataclass
class Context:
    """What Open WebUI passes to a tool call."""

    user: dict[str, Any] | None
    request: Any
    chat_id: str | None
    message_id: str | None
    emitter: EventCallback | None
    caller: EventCallback | None

    async def status(self, description: str, done: bool = False) -> None:
        if self.emitter is not None:
            await self.emitter({"type": "status", "data": {"description": description, "done": done}})


async def attach_file(data: bytes, filename: str, mime: str, ctx: Context, text: str = "") -> None:
    """Store the file in Open WebUI's file storage for the user and attach it to the current answer.

    Uses the same internal functions as Open WebUI's built-in image generation (0.11.4).
    """
    from open_webui.models.chats import Chats
    from open_webui.models.files import Files
    from open_webui.models.users import UserModel
    from open_webui.routers.files import upload_file_handler
    from open_webui.utils.chat_id import is_saved_chat_id
    from starlette.datastructures import Headers, UploadFile

    if not ctx.user or ctx.request is None or ctx.emitter is None:
        raise RuntimeError("no user, request or event emitter context")
    user = UserModel(**ctx.user)
    upload = UploadFile(file=io.BytesIO(data), filename=filename, headers=Headers({"content-type": mime}))
    file_item = await upload_file_handler(
        ctx.request,
        file=upload,
        metadata={"chat_id": ctx.chat_id, "message_id": ctx.message_id},
        process=False,
        user=user,
    )
    file_id = file_item.id
    if text.strip():  # shown in the "Content" tab of the file view
        try:
            await Files.update_file_data_by_id(file_id, {"content": text[:MAX_CONTENT_CHARS]})
        except Exception:
            log.exception("Storing the text of %s failed", filename)
    try:
        url = ctx.request.app.url_path_for("get_file_content_by_id", id=file_id)
    except Exception:
        url = f"/api/v1/files/{file_id}/content"
    entry = {"type": "file", "id": file_id, "url": str(url), "name": filename, "size": len(data), "content_type": mime}
    files = [entry]
    if ctx.message_id and is_saved_chat_id(ctx.chat_id):
        try:  # the file is stored now; if linking it to the chat fails, the event below still attaches it
            await Chats.insert_chat_files(
                chat_id=ctx.chat_id, message_id=ctx.message_id, file_ids=[file_id], user_id=user.id
            )
            files = await Chats.add_message_files_by_id_and_message_id(ctx.chat_id, ctx.message_id, [entry]) or files
        except Exception:
            log.exception("Linking %s to the chat failed", filename)
    await ctx.emitter({"type": "chat:message:files", "data": {"files": files}})


DOWNLOAD_TIMEOUT = 60  # seconds; Open WebUI waits forever if the user has switched to another chat


async def download_file(data: bytes, filename: str, mime: str, ctx: Context) -> None:
    """Hand the file to the browser. Open WebUI does not raise when the browser is gone: it returns an error
    dict, or never answers, so anything but the snippet's `true` counts as failure."""
    if ctx.caller is None:
        raise RuntimeError("__event_call__ is not available, cannot download in the browser")
    result = await asyncio.wait_for(
        ctx.caller({"type": "execute", "data": {"code": download_js(data, filename, mime)}}), DOWNLOAD_TIMEOUT
    )
    if result is not True:
        raise RuntimeError(f"the browser did not confirm the download: {str(result)[:200]}")


# ------------------------------------------------------------------------------------------------ Tools --


class Tools:
    class Valves(BaseModel):
        language: Literal["en", "de"] = Field(
            default="en",
            description="Language of status messages and of texts inside the files (total row, page footer).",
        )
        delivery: Literal["attach", "download"] = Field(
            default="attach",
            description="attach: store the file in Open WebUI and attach it to the answer (preview, stays in the chat). "
            "download: send it straight to the browser as a download (nothing stored, gone after reload). "
            "If attaching fails, the plugin falls back to download.",
        )
        currency: Literal["EUR", "USD", "GBP", "CHF"] = Field(
            default="EUR", description="Currency of Excel columns of type 'currency'."
        )
        accent_color: str = Field(
            default=DEFAULT_ACCENT,
            description="Accent colour of all files (titles, header rows, charts, title slide) as a hex value without "
            "'#', e.g. 1F4E79. Lighter and darker tones are derived from it.",
        )
        font: str = Field(
            default=BODY_FONT,
            description="Font of all files, e.g. Calibri, Arial or Aptos. It has to be installed where the files are "
            "opened; otherwise Office shows a similar one.",
        )
        page_numbers: bool = Field(default=True, description="Add page numbers to the footer of Word documents.")
        word_template: str = Field(
            default="",
            description="Optional Word file (.docx or .dotx) on the Open WebUI server, e.g. a letterhead: Word documents "
            "get its header, footer, page size and margins, e.g. /app/backend/data/templates/letterhead.dotx.",
        )

    def __init__(self) -> None:
        self.valves = self.Valves()

    async def create_word_file(
        self,
        blocks: list[WordBlock],
        title: str = "",
        subtitle: str = "",
        header_text: str = "",
        table_of_contents: bool = False,
        landscape: bool = False,
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create a Word document (.docx) and attach it to the answer. Use it when the user wants a Word file or a
        document to download, e.g. a report, letter, minutes or concept. Write the complete text yourself.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param blocks: The content in reading order, e.g. [{"type": "heading", "text": "Introduction", "level": 1}, {"type": "paragraph", "text": "..."}, {"type": "bullets", "items": ["...", "- sub item"]}, {"type": "table", "rows": [["Header A", "Header B"], ["a", "b"]]}].
        :param title: Readable document title with spaces, e.g. "Minutes of the team meeting on 12 October 2026". It is shown at the top and used as the file name, so do not repeat it as the first heading.
        :param subtitle: Optional subtitle below the title.
        :param header_text: Optional text in the page header, e.g. company or project name.
        :param table_of_contents: Add a table of contents (useful for long documents with several headings).
        :param landscape: Use landscape pages (for wide tables).
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]

        def build() -> tuple[str, bytes, str]:
            extras: dict[str, Any] = {}  # parameters the model wrote into the blocks text
            parsed = parse_argument(WORD_BLOCKS, blocks, "blocks", extras)
            parsed = [block for block in parsed if not is_empty_block(block)]
            given = {k: _unescape_newlines(str(v or extras.get(k) or "")) for k, v in
                     (("title", title), ("subtitle", subtitle), ("header_text", header_text))}  # fmt: skip
            display_title, parsed = word_title(given["title"], parsed)
            options = WordOptions(
                title=display_title,
                subtitle=given["subtitle"],
                header_text=given["header_text"],
                table_of_contents=bool(table_of_contents or extras.get("table_of_contents") is True),
                landscape=bool(landscape or extras.get("landscape") is True),
                page_numbers=self.valves.page_numbers,
                design=self._design(),
                template=read_word_template(self.valves.word_template),
            )
            return display_title, *build_docx(options, parsed, texts)

        return await self._create("create_word_file", "docx", DOCX_MIME, build, texts["status_word"], ctx)

    async def create_excel_file(
        self,
        sheets: list[ExcelSheet],
        title: str = "",
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create an Excel workbook (.xlsx) and attach it to the answer. Use it when the user wants an Excel file or a
        spreadsheet, e.g. a list, budget, plan or calculation. Numbers are real numbers; formulas, total rows and
        charts are possible.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param sheets: One or more sheets, e.g. [{"name": "Budget", "columns": [{"header": "Item"}, {"header": "Amount", "type": "currency", "total": "sum"}], "rows": [["Rent", 1200], ["Travel", 350.5]]}].
        :param title: Readable title of the workbook with spaces, used as the file name.
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]

        def build() -> tuple[str, bytes, str]:
            extras: dict[str, Any] = {}
            parsed = parse_argument(EXCEL_SHEETS, sheets, "sheets", extras)
            given = str(title or extras.get("title") or "")
            display_title = tidy_title(given) or (tidy_title(parsed[0].name) if parsed else "")
            decimal_comma = self.valves.language == "de"
            design = self._design()
            return display_title, *build_xlsx(parsed, self.valves.currency, texts, decimal_comma, design, display_title)

        return await self._create("create_excel_file", "xlsx", XLSX_MIME, build, texts["status_excel"], ctx)

    async def create_powerpoint_file(
        self,
        slides: list[Slide],
        title: str = "",
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create a PowerPoint presentation (.pptx, 16:9) and attach it to the answer. Use it when the user wants
        slides or a presentation. Start with a 'title' slide, keep bullets short (at most 6 per slide), and put
        details into the speaker notes.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param slides: The slides in order, e.g. [{"layout": "title", "title": "...", "subtitle": "..."}, {"layout": "bullets", "title": "...", "bullets": ["...", "- sub point"], "notes": "..."}, {"layout": "chart", "title": "...", "chart": {"type": "column", "categories": ["Q1", "Q2"], "series": [{"name": "Sales", "values": [10, 12]}]}}].
        :param title: Readable title of the presentation with spaces, used as the file name.
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]
        design = self._design()

        def build() -> tuple[str, bytes, str]:
            extras: dict[str, Any] = {}
            parsed = parse_argument(SLIDES, slides, "slides", extras)
            display_title = tidy_title(str(title or extras.get("title") or "")) or next(
                (tidy_title(s.title) for s in parsed if s.title.strip()), ""
            )
            return display_title, *build_pptx(display_title, parsed, design, texts)

        return await self._create("create_powerpoint_file", "pptx", PPTX_MIME, build, texts["status_powerpoint"], ctx)

    def _design(self) -> Design:
        accent = self.valves.accent_color.strip().lstrip("#")
        font = " ".join(self.valves.font.split())
        return Design(
            accent.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", accent) else DEFAULT_ACCENT,
            font if FONT_NAME_RE.fullmatch(font) else BODY_FONT,
        )

    async def _create(
        self,
        tool: str,
        extension: str,
        mime: str,
        build: Callable[[], tuple[str, bytes, str]],
        status: str,
        ctx: Context,
    ) -> str:
        texts = TRANSLATIONS[self.valves.language]
        await ctx.status(status)
        try:
            title, data, text = await asyncio.to_thread(build)
        except RecursionError:
            await ctx.status(texts["status_failed"], done=True)
            return RESULT_INVALID.format(errors="- the arguments are nested too deeply", tool=tool)
        except InvalidInput as e:
            await ctx.status(texts["status_failed"], done=True)
            errors = e.errors[:MAX_ERRORS_SHOWN] + (["..."] if len(e.errors) > MAX_ERRORS_SHOWN else [])
            return RESULT_INVALID.format(errors="\n".join(f"- {error}" for error in errors), tool=tool)
        except Exception:
            log.exception("Creating a .%s file failed", extension)
            await ctx.status(texts["status_failed"], done=True)
            return RESULT_FAILED
        filename = f"{safe_filename(title, texts['default_title'])}.{extension}"
        result = await self._deliver(data, filename, mime, ctx, text)
        await ctx.status(texts["status_done"].format(name=filename) if result else texts["status_failed"], done=True)
        return result.format(name=filename) if result else RESULT_FAILED

    async def _deliver(self, data: bytes, filename: str, mime: str, ctx: Context, text: str = "") -> str:
        """Attach or download the file; returns the result template, or '' if both ways failed."""
        if self.valves.delivery == "attach":
            try:
                await attach_file(data, filename, mime, ctx, text)
                return RESULT_ATTACHED
            except Exception:
                log.exception("Attaching %s failed, falling back to a browser download", filename)
        try:
            await download_file(data, filename, mime, ctx)
            return RESULT_DOWNLOADED
        except Exception:
            log.exception("Downloading %s failed", filename)
            return ""
