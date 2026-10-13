"""
title: Office Documents
author: <AUTHOR NAME>
author_url: <URL>
version: 0.1.0
license: MIT
required_open_webui_version: 0.11.4
description: Lets the model create Word (.docx), Excel (.xlsx) and PowerPoint (.pptx) files. The model only supplies the content as structured JSON; fixed code builds the file, so no model-written code is ever run. The file is attached to the answer, where it can be previewed and downloaded.

Tools
- create_word_file: title page, optional table of contents, headings, paragraphs, bullet and
  numbered lists (two levels), tables, quotes, page breaks, header text and page numbers.
- create_excel_file: several sheets with typed columns (text, number, integer, currency, percent,
  date), row formulas from a checked whitelist, total rows and native Excel charts.
- create_powerpoint_file: 16:9 slides with the layouts title, section, bullets, two columns,
  table and chart (native PowerPoint charts), plus speaker notes.
Inline **bold**, *italic* and [links](https://...) work in Word and PowerPoint; Excel cells get plain text.

Security & privacy
- The model never writes code: it fills in a fixed JSON schema, the plugin validates it and builds
  the file. No shell, no network access, nothing is loaded from URLs.
- Every formula, whether a column's `formula` field or a formula written into a cell, is checked against a
  whitelist (SUM, IF, ROUND ...): no external references, no other sheets, no DDE, no WEBSERVICE/HYPERLINK.
  Anything else (e.g. "=HYPERLINK(...)") is written as text, also in the data sheets of PowerPoint charts.
- Links are only created for http, https and mailto. Limits on slides, rows, columns and text
  length keep a runaway model from blocking the server.
- The file is stored in Open WebUI's own file storage under the user's account, so only the user
  (and admins) can open it. Only the libraries bundled with Open WebUI are used.
"""

import asyncio
import base64
import difflib
import io
import json
import logging
import re
import unicodedata
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, replace
from datetime import date, datetime
from typing import Any, Literal

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.xlsx import CategoryWorkbookWriter
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement as PptxOxmlElement
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.text.text import TextFrame, _Paragraph as PptxParagraph
from pptx.util import Emu, Inches, lazyproperty
from pptx.util import Pt as PptxPt
from pydantic import BaseModel, ConfigDict, Field, PrivateAttr, TypeAdapter, ValidationError, model_validator

log = logging.getLogger(__name__)

EventCallback = Callable[[dict[str, Any]], Awaitable[Any]]

# ------------------------------------------------------------------------------------------ constants --

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

BODY_FONT = "Calibri"
BODY_FONT_SIZE = Pt(11)
CODE_FONT = "Consolas"
HEADER_FILL = "DDE4EE"
LINK_COLOR = "0563C1"
MAX_FILENAME_CHARS = 60
SAFE_LINK_PREFIXES = ("http://", "https://", "mailto:")
SUB_ITEM_PREFIX = "- "  # list items starting with this are indented one level

# Limits: an error tells the model what to shorten, so it can call the tool again.
MAX_TEXT_CHARS = 200_000  # all text of one file together
MAX_WORD_BLOCKS = 500
MAX_WORD_TABLE_ROWS = 500
MAX_WORD_TABLE_COLUMNS = 12
MAX_SHEETS = 20
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 50
MAX_CELLS = 250_000  # all sheets together; empty cells count too
MAX_CELL_CHARS = 32_000  # Excel's limit is 32,767
MAX_FORMULA_CHARS = 500
MAX_SLIDES = 60
MAX_SLIDE_ITEMS = 12
MAX_SLIDE_TABLE_ROWS = 16
MAX_SLIDE_TABLE_COLUMNS = 8
MAX_CHART_CATEGORIES = 50
MAX_CHART_SERIES = 10
MAX_ERRORS_SHOWN = 8

MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
MAX_DECIMALS = 6
CURRENCY_FORMATS = {
    "EUR": '#,##0.00 "€"',
    "USD": '"$"#,##0.00',
    "GBP": '"£"#,##0.00',
    "CHF": '"CHF" #,##0.00',
}

TRANSLATIONS: dict[str, dict[str, str]] = {
    "en": {
        "status_word": "Creating Word document …",
        "status_excel": "Creating Excel workbook …",
        "status_powerpoint": "Creating PowerPoint presentation …",
        "status_done": "{name} created",
        "status_failed": "The file could not be created",
        "default_title": "Document",
        "total": "Total",
        "page_footer": "Page {page} of {pages}",
        "toc_title": "Contents",
        "toc_placeholder": "Right-click and choose “Update field” to show the table of contents.",
        "sheet": "Sheet",
        "date_format": "YYYY-MM-DD",
    },
    "de": {
        "status_word": "Word-Dokument wird erstellt …",
        "status_excel": "Excel-Arbeitsmappe wird erstellt …",
        "status_powerpoint": "PowerPoint-Präsentation wird erstellt …",
        "status_done": "{name} erstellt",
        "status_failed": "Die Datei konnte nicht erstellt werden",
        "default_title": "Dokument",
        "total": "Summe",
        "page_footer": "Seite {page} von {pages}",
        "toc_title": "Inhalt",
        "toc_placeholder": "Rechtsklick → „Felder aktualisieren“ zeigt das Inhaltsverzeichnis an.",
        "sheet": "Tabelle",
        "date_format": "DD.MM.YYYY",
    },
}

# Messages for the model (always English; the model answers the user in their language).
RESULT_ATTACHED = (
    "Done: the file '{name}' was created and is attached to your answer, where the user can preview and "
    "download it. Do not add a link and do not repeat the file content. Briefly tell the user what the "
    "file contains."
)
RESULT_DOWNLOADED = (
    "Done: the file '{name}' was created and downloaded in the user's browser. Do not add a link and do "
    "not repeat the file content. Briefly tell the user what the file contains."
)
RESULT_INVALID = (
    "The file was NOT created because the arguments are invalid:\n{errors}\n"
    "Fix these points and call {tool} again with the complete, corrected content."
)
RESULT_FAILED = "The file could not be created because of an internal error. Tell the user to try again later."

# ------------------------------------------------------------------------------------ input schema --
# The schemas below are what the model sees (Open WebUI turns the type hints and descriptions into
# JSON schema). Limits are checked in code instead of maxItems/maxLength: large limits in a schema
# make the grammar of llama.cpp and vLLM huge and slow.


class Schema(BaseModel):
    model_config = ConfigDict(extra="ignore", coerce_numbers_to_str=True, allow_inf_nan=False)
    _index: int = PrivateAttr(default=-1)  # position in the list the model sent, for error messages

    @model_validator(mode="before")
    @classmethod
    def _null_means_default(cls, data: Any) -> Any:
        """Models often send null for fields they do not use."""
        return {k: v for k, v in data.items() if v is not None} if isinstance(data, dict) else data


class WordBlock(Schema):
    type: Literal[
        "heading", "paragraph", "bullets", "numbered", "table", "quote", "page_break", "table_of_contents"
    ] = Field(description="Kind of block.")
    text: str = Field(
        default="",
        description="Text of a heading, paragraph or quote. Inline **bold**, *italic* and [label](https://...) "
        "links are allowed.",
    )
    level: int = Field(default=1, description="Heading level: 1, 2 or 3.")
    items: list[str] = Field(
        default=[],
        description="Items of a bullets or numbered list. Start an item with '- ' to indent it one level.",
    )
    rows: list[list[str]] = Field(
        default=[], description="Table rows as lists of cell texts. The first row is the header row."
    )


class ExcelColumn(Schema):
    header: str = Field(description="Column header.")
    type: Literal["text", "number", "integer", "currency", "percent", "date"] = Field(
        default="text",
        description="Value type. percent: 12.5 means 12.5 %. date: YYYY-MM-DD.",
    )
    formula: str = Field(
        default="",
        description="Optional formula filled into every data row, e.g. '=B{row}*C{row}' ({row} is the current "
        "row; data starts in row 2). English function names (SUM, AVERAGE, MIN, MAX, COUNT, IF, ROUND, "
        "IFERROR ...), commas between arguments, only cells of this sheet. Put null in the rows for this column.",
    )
    total: Literal["none", "sum", "average", "min", "max", "count"] = Field(
        default="none", description="Adds a total row below the data with this function for the column."
    )


class ExcelChart(Schema):
    type: Literal["column", "bar", "line", "pie"] = Field(default="column", description="Chart type.")
    title: str = Field(default="", description="Chart title.")
    category_column: str = Field(
        default="",
        description="Header of the column with the labels (x axis or pie slices). Not used with series_rows.",
    )
    value_columns: list[str] = Field(
        default=[],
        description="Headers of the numeric columns to plot (pie: only one). With series_rows: the columns that form "
        "the x axis, next to each other; default: the number columns after the first column.",
    )
    series_rows: list[str] = Field(
        default=[],
        description="Only when each data series is a row, e.g. months as columns: the labels in the first column of "
        'the rows to plot, e.g. ["Income", "Expenses"].',
    )


class ExcelSheet(Schema):
    name: str = Field(default="", description="Sheet name, at most 31 characters.")
    columns: list[ExcelColumn] = Field(description="Columns from left to right.")
    rows: list[list[str | float | None]] = Field(
        default=[],
        description="Data rows. Each row has one value per column, in column order. Use JSON numbers for "
        "numbers and null for empty cells and formula columns.",
    )
    chart: ExcelChart | None = Field(default=None, description="Optional chart next to the table.")


class ChartSeries(Schema):
    name: str = Field(description="Series name, shown in the legend.")
    values: list[float] = Field(description="One number per category.")


class SlideChart(Schema):
    type: Literal["column", "bar", "line", "pie"] = Field(default="column", description="Chart type.")
    categories: list[str] = Field(description="Category labels (x axis or pie slices).")
    series: list[ChartSeries] = Field(description="Data series (pie: only one).")
    unit: str = Field(default="", description="Unit of the values, e.g. 'hours per week'; shown at the value axis.")


class Slide(Schema):
    layout: Literal["title", "section", "bullets", "two_columns", "table", "chart"] = Field(
        description="title: opening slide with title and subtitle. section: divider between parts. "
        "bullets: title with a list. two_columns: two lists side by side. table: title with a table. "
        "chart: title with a chart."
    )
    title: str = Field(description="Slide title.")
    subtitle: str = Field(default="", description="Subtitle, only for title and section slides.")
    bullets: list[str] = Field(
        default=[],
        description="Items for a bullets slide, short phrases rather than sentences. Start an item with '- ' to "
        "indent it.",
    )
    left: list[str] = Field(default=[], description="Items of the left column (two_columns).")
    right: list[str] = Field(default=[], description="Items of the right column (two_columns).")
    left_heading: str = Field(default="", description="Optional heading above the left column.")
    right_heading: str = Field(default="", description="Optional heading above the right column.")
    table: list[list[str]] = Field(default=[], description="Table rows for a table slide; the first row is the header.")
    chart: SlideChart | None = Field(default=None, description="Chart data for a chart slide.")
    notes: str = Field(default="", description="Optional speaker notes.")


WORD_BLOCK_TYPES = ("heading", "paragraph", "bullets", "numbered", "table", "quote", "page_break", "table_of_contents")
WORD_TYPE_ALIASES = {
    "list": "bullets", "bullet": "bullets", "bullet_list": "bullets", "bulleted_list": "bullets",
    "unordered_list": "bullets", "ordered_list": "numbered", "numbered_list": "numbered", "number_list": "numbered",
    "text": "paragraph", "para": "paragraph", "subheading": "heading", "title": "heading", "blockquote": "quote",
    "pagebreak": "page_break", "page-break": "page_break", "toc": "table_of_contents",
}  # fmt: skip
WORD_BLOCKS = TypeAdapter(list[WordBlock])
EXCEL_SHEETS = TypeAdapter(list[ExcelSheet])
SLIDES = TypeAdapter(list[Slide])
SLIDE_ONLY_FIELDS = set(Slide.model_fields) - set(SlideChart.model_fields)


class InvalidInput(ValueError):
    """Raised with one message per problem; the messages are returned to the model."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("\n".join(errors))
        self.errors = errors


LOCATION_NOISE_RE = re.compile(r"(str|float|int|bool|list\[.*\]|function-.*)")
POSITION_KEY = "__position__"  # added to list items before normalizing, so errors name the item the model sent


def _location(parts: tuple[Any, ...]) -> str:
    text = ""
    for part in parts:
        if isinstance(part, str) and LOCATION_NOISE_RE.fullmatch(part):
            continue  # union members pydantic adds to the path, e.g. "rows[0][1].str"
        text += f"[{part}]" if isinstance(part, int) else f".{part}" if text else str(part)
    return text


def parse_argument(
    adapter: TypeAdapter,
    value: Any,
    name: str,
    normalize: Callable[[Any, int], Any] | None = None,
    normalize_list: Callable[[list[Any]], list[Any]] | None = None,
    extras: dict[str, Any] | None = None,
    rows_key: str | None = None,
) -> Any:
    """Validate a tool argument.

    Some models send nested arguments as a JSON string, and small models often use slightly different
    shapes (e.g. {"bullets": [...]} instead of {"type": "bullets", "items": [...]}). `normalize` repairs
    such items before validation, which saves the model a second attempt. `extras` receives other tool
    parameters that the model wrote into the same JSON text. `rows_key` names the table field whose
    unclosed rows may have swallowed the following items (see take_out_items). Every validated item
    remembers its position in the list the model sent (`_index`), so later error messages point to the
    right item.
    """
    if isinstance(value, str):
        value = _loads_argument(value, name, extras)
    if isinstance(value, dict) and len(value) == 1 and isinstance(next(iter(value.values())), list):
        value = next(iter(value.values()))  # {"slides": [...]} instead of [...]
    value = _unescape_newlines(value)
    if isinstance(value, list):  # broken tool-call syntax: content got lost, so let the model send it again
        broken = [f"{name}[{i}]" for i, item in enumerate(value) if _contains(item, BROKEN_QUOTE) or _garbled(item)]
        if broken:
            raise InvalidInput(
                [f"{where}: the arguments are garbled here (a key or the token {BROKEN_QUOTE} ended up inside a text,"
                 " and text got lost)" for where in broken[:MAX_ERRORS_SHOWN]]
                + [f"Write every key and string with plain double quotes and send the complete {name} again."]
            )  # fmt: skip
        value = [{**item, POSITION_KEY: i} if isinstance(item, dict) else item for i, item in enumerate(value)]
        if rows_key is not None:
            value = take_out_items(value, rows_key)
    if normalize is not None and isinstance(value, list):
        value = [normalize(item, i) if isinstance(item, dict) else item for i, item in enumerate(value)]
    if normalize_list is not None and isinstance(value, list):
        value = normalize_list(value)
    origin = (
        [item.get(POSITION_KEY, i) if isinstance(item, dict) else i for i, item in enumerate(value)]
        if isinstance(value, list)
        else []
    )
    try:
        result = adapter.validate_python(value)
    except ValidationError as e:
        errors = []
        for err in e.errors():
            loc = err["loc"]
            if loc and isinstance(loc[0], int) and loc[0] < len(origin):
                loc = (origin[loc[0]], *loc[1:])
            errors.append(f"{_location((name, *loc))}: {err['msg']}")
        raise InvalidInput(errors) from e
    for item, position in zip(result, origin, strict=False):
        item._index = position
    return result


def _loads_argument(text: str, name: str, extras: dict[str, Any] | None) -> Any:
    try:
        value = loads_lenient(text, extras)
        if isinstance(value, str) and value.strip()[:1] in ("[", "{"):
            value = loads_lenient(value, extras)  # JSON text encoded twice
    except json.JSONDecodeError as e:
        before = e.doc[max(0, e.pos - 80) : e.pos].replace("\n", " ")
        raise InvalidInput(
            [f"{name}: the JSON is invalid at character {e.pos} ({e.msg}), right after: ...{before}"
             f" - check the quotes and brackets there; a \" inside a text must be written as \\\""]
        ) from e  # fmt: skip
    except RecursionError as e:
        raise InvalidInput([f"{name}: the JSON is nested too deeply"]) from e
    except ValueError as e:  # NaN/Infinity, numbers with thousands of digits
        raise InvalidInput([f"{name}: the JSON could not be read: {str(e)[:200]}"]) from e
    return value


HTML_BREAK_RE = re.compile(r"<br\s*/?>", re.IGNORECASE)
ESCAPED_UNICODE_RE = re.compile(r"\\u([0-9a-fA-F]{4})")
ESCAPED_UNICODE_ONLY_RE = re.compile(r"(?:[^\\]|\\u[0-9a-fA-F]{4}|\\n)*", re.DOTALL)
ESCAPED_NEWLINES_RE = re.compile(r"(?:[^\\]|\\n)*\\n(?:[^\\]|\\n)*", re.DOTALL)


def _unescape_newlines(value: Any) -> Any:
    """Line breaks written in other ways: '<br>' (from Markdown tables) and 'Line 1\\nLine 2' (escaped twice).
    An escaped one is only replaced in texts whose backslashes are all followed by n, so a path like
    'C:\\neu\\Ordner' stays as it is."""
    if isinstance(value, str):
        value = HTML_BREAK_RE.sub("\n", value)
        if "\\u" in value and ESCAPED_UNICODE_ONLY_RE.fullmatch(value):  # '\\u2013' escaped twice
            value = ESCAPED_UNICODE_RE.sub(lambda m: chr(int(m[1], 16)), value)
        return value.replace("\\n", "\n") if "\n" not in value and ESCAPED_NEWLINES_RE.fullmatch(value) else value
    if isinstance(value, dict):
        return {k: _unescape_newlines(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_unescape_newlines(item) for item in value]
    return value


ITEM_KEYS = {"type", "layout", "title", "text", "heading", "bullets", "items", "rows", "table", "chart"}


def take_out_items(items: list[Any], rows_key: str) -> list[Any]:
    """'"rows": [[...], [...]], {"type": "paragraph"…}' with the table's ']' forgotten puts the following
    items into the rows: [[...], [...], {"type": "paragraph"…}, …]. They are taken out again and follow
    the table, in their order."""
    result: list[Any] = []
    for item in items:
        rows = item.get(rows_key) if isinstance(item, dict) else None
        first = (
            next((i for i, row in enumerate(rows) if isinstance(row, dict)), None) if isinstance(rows, list) else None
        )
        if first and rows is not None and set(rows[first]) & ITEM_KEYS:
            position = item.get(POSITION_KEY)
            result.append({**item, rows_key: rows[:first]})
            following = [{**f, POSITION_KEY: position} if isinstance(f, dict) else f for f in rows[first:]]
            result.extend(take_out_items(following, rows_key))
        else:
            result.append(item)
    return result


GARBLED_KEY_RE = re.compile(r'",[a-z_]+:|\}\s*,\s*\{\s*[a-z_]+:')  # 'paragraph",text:', '},{text:' in a text
PLAIN_KEY_RE = re.compile(r"[\w -]+")  # '"name' or '{type' as a key: the syntax broke there


def _garbled(value: Any) -> bool:
    if isinstance(value, str):
        return bool(GARBLED_KEY_RE.search(value))
    if isinstance(value, dict):
        return any(not PLAIN_KEY_RE.fullmatch(str(k)) or _garbled(v) for k, v in value.items())
    if isinstance(value, list):
        return any(_garbled(item) for item in value)
    return False


def _contains(value: Any, token: str) -> bool:
    if isinstance(value, str):
        return token in value
    if isinstance(value, dict):
        return any(_contains(k, token) or _contains(v, token) for k, v in value.items())
    if isinstance(value, list):
        return any(_contains(item, token) for item in value)
    return False


TABLE_LINE_RE = re.compile(r"^\s*\|.*\|\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
LIST_LINE_RE = re.compile(r"^(?P<indent>\s*)(?:(?P<bullet>[-*•])|(?P<number>\d{1,3})[.)])\s+(?P<text>\S.*)$")


def split_paragraph(text: str) -> list[tuple[str, list[str]]]:
    """Split paragraph text into ("text", lines), ("bullets", items) and ("numbered", items) parts.

    Models often write Markdown lists into a paragraph; indented list lines become '- ' sub items.
    Numbered lines only form a list when they count up from 1, so lines that merely start with a
    number ('3. Oktober ist Feiertag.') stay text.
    """
    parts: list[tuple[str, list[str]]] = []
    for line in text.strip().split("\n"):
        if TABLE_LINE_RE.match(line):  # a Markdown table row
            if parts and parts[-1][0] == "table":
                parts[-1][1].append(line)
            else:
                parts.append(("table", [line]))
            continue
        match = LIST_LINE_RE.match(line)
        indented = bool(match) and len(match["indent"].expandtabs(4)) >= 2
        if match and match["number"] and not indented:
            expected = len([i for i in parts[-1][1] if not i.startswith(SUB_ITEM_PREFIX)]) + 1 if (
                parts and parts[-1][0] == "numbered"
            ) else 1  # fmt: skip
            if int(match["number"]) != expected:
                match = None
        if not line.strip():
            kind, item = "break", ""
        elif match:
            kind = "numbered" if match["number"] else "bullets"
            item = (SUB_ITEM_PREFIX if indented else "") + match["text"].strip()
        else:
            kind, item = "text", line.strip()
        if kind == "break":
            parts.append(("break", []))
        elif parts and (
            parts[-1][0] == kind or (parts[-1][0] in ("bullets", "numbered") and item.startswith(SUB_ITEM_PREFIX))
        ):
            parts[-1][1].append(item)
        else:
            parts.append((kind, [item]))
    result = []
    for kind, lines in parts:
        if kind == "table" and not (len(lines) >= 2 and any(TABLE_SEPARATOR_RE.match(line) for line in lines)):
            kind = "text"  # lines with | but no table
        if kind != "break":
            result.append((kind, lines))
    return result


def markdown_table_rows(lines: list[str]) -> list[list[str]]:
    return [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in lines
        if not TABLE_SEPARATOR_RE.match(line)
    ]


GERMAN_QUOTE_RE = re.compile(r'„([^"“”„\n]{0,79}[^"“”„\n\\])"')  # „Text" -> „Text“ (not „Text\")
MAX_JSON_REPAIRS = 20
# Other tool parameters a model may write into the same JSON text after the value ('[...], "title": "T"').
TOOL_PARAMETERS = {"title", "subtitle", "header_text", "table_of_contents", "landscape"}
TRAILING_PARAMETER_RE = re.compile(r',\s*"(\w+)"\s*:')


def _reject_constant(name: str) -> Any:
    raise ValueError(f"{name} is not a valid number, use null for empty values")


JSON_DECODER = json.JSONDecoder(parse_constant=_reject_constant)


def loads_lenient(text: str, extras: dict[str, Any] | None = None) -> Any:
    """json.loads that repairs the mistakes models make in long arguments sent as text.

    - German quotes closed with a straight quote („Hoch") end the JSON string early; they become „Hoch“.
    - A forgotten '}' shows up where the next object starts in place of a key ('…"}, {"layout"…');
      the missing brace is added before that comma.
    - Arrays listed without their outer brackets ('"table": ["A", "B"], ["1", "2"]') get them back.
    - Other forgotten closing brackets are added where a different bracket closes, or at the end if the
      text ends with a complete object or array.
    - Other tool parameters written into the same text after the value ('[...], "title": "T"') are
      put into `extras`.
    The text is only changed after strict parsing failed. A repair is only accepted if nothing of the
    text is lost: if text is left over after the value, or the text was cut off in the middle of a value,
    the original error is raised so the model sends the arguments again.
    """
    try:
        return JSON_DECODER.decode(text)
    except json.JSONDecodeError as first_error:
        error = first_error
    stripped = text.strip()
    for candidate in dict.fromkeys((GERMAN_QUOTE_RE.sub("„\\1“", stripped), stripped)):
        repaired = _repair_json(candidate)
        if repaired is None:
            continue
        value, rest = repaired
        parameters = _trailing_parameters(rest)
        if parameters is not None:
            if extras is not None:
                extras.update(parameters)
            return value
    raise error


def _repair_json(text: str) -> tuple[Any, str] | None:
    """Return (value, text left over after it), or None if the text cannot be repaired without losing content."""
    bracket_fix_used = False
    for _ in range(MAX_JSON_REPAIRS):
        try:
            value, end = JSON_DECODER.raw_decode(text)
            return value, text[end:]
        except json.JSONDecodeError as e:
            comma = text.rfind(",", 0, e.pos)
            key_expected = e.msg == "Expecting property name enclosed in double quotes"
            if key_expected and text.startswith("{", e.pos) and comma >= 0:
                text = text[:comma] + "}" + text[comma:]
            elif key_expected and text.startswith("[", e.pos) and (wrapped := wrap_array_run(text, e.pos)):
                text = wrapped
            elif not bracket_fix_used:
                fixed, appended = _close_brackets(text)
                if appended and not text.rstrip().endswith(("}", "]")):
                    return None  # cut off in the middle of a value (e.g. the output limit was reached)
                text, bracket_fix_used = fixed, True
            else:
                return None
    return None


def _trailing_parameters(rest: str) -> dict[str, Any] | None:
    """Tool parameters written after the value; {} if nothing is left over, None if content would be lost."""
    rest = rest.strip()
    if not rest.strip("}] \t\r\n"):
        return {}
    match = TRAILING_PARAMETER_RE.match(rest)
    if not match or match[1] not in TOOL_PARAMETERS:
        return None
    body = "{" + rest[match.start(1) - 1 :]
    for candidate in (body, body + "}"):
        try:
            parameters, _end = JSON_DECODER.raw_decode(candidate)
        except (json.JSONDecodeError, ValueError):
            continue
        return {k: v for k, v in parameters.items() if k in TOOL_PARAMETERS} if isinstance(parameters, dict) else {}
    return {}


ARRAY_VALUE_RE = re.compile(r'"[^"\\]*"\s*:\s*\[')


def wrap_array_run(text: str, pos: int) -> str | None:
    """'"rows": ["a"], ["b"], ["c"]}' -> '"rows": [["a"], ["b"], ["c"]]}'; None if the text does not look like that."""
    starts = [match.end() - 1 for match in ARRAY_VALUE_RE.finditer(text, 0, pos)]
    if not starts:
        return None
    decoder = json.JSONDecoder()
    start, end = starts[-1], starts[-1]
    while True:
        try:
            _, end = decoder.raw_decode(text, end)
        except json.JSONDecodeError:
            return None
        following = re.match(r"\s*,\s*(?=\[)", text[end:])
        if not following:
            break
        end += following.end()
    if end <= pos:  # the arrays must reach past the place where the error was found
        return None
    return text[:start] + "[" + text[start:end] + "]" + text[end:]


def close_brackets(text: str) -> str:
    """Add closing brackets the text forgot: before a bracket that closes an outer one, and at the end."""
    return _close_brackets(text)[0]


def _close_brackets(text: str) -> tuple[str, int]:
    """close_brackets, plus the number of brackets added at the end."""
    out: list[str] = []
    stack: list[str] = []
    in_string = escaped = False
    for char in text:
        if in_string:
            escaped, in_string = (False, True) if escaped else (char == "\\", char != '"')
        elif char == '"':
            in_string = True
        elif char in "[{":
            stack.append("]" if char == "[" else "}")
        elif char in "]}":
            while stack and stack[-1] != char:
                out.append(stack.pop())
            if stack:
                stack.pop()
        out.append(char)
    return "".join(out) + "".join(reversed(stack)), len(stack)


def split_lines(text: str) -> list[str]:
    """A list sent as one text: one item per line, without list markers."""
    return [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s+", "", line).strip() for line in text.splitlines() if line.strip()]


def normalize_word_block(block: dict[str, Any], index: int) -> dict[str, Any]:
    block = dict(block)
    if not any(_filled(block.get(k)) for k in ("text", "items", "rows")):
        nested = next(
            (v[0] for v in block.values() if isinstance(v, list) and len(v) == 1 and isinstance(v[0], dict)), None
        )
        if nested is not None and set(nested) & {"type", "text", "items", "rows"}:
            # a block inside a block: {"type": "paragraph", "bullets": [{"type": "bullets", "items": [...]}]}
            return normalize_word_block({**nested, POSITION_KEY: block.get(POSITION_KEY)}, index)
    if isinstance(block.get("type"), str):
        kind = block["type"].strip().lower()
        block["type"] = WORD_TYPE_ALIASES.get(kind, kind)
    if "type" not in block:
        kind = next((k for k in WORD_BLOCK_TYPES if k in block), None)
        if kind is not None:  # {"heading": "Text"}, {"bullets": [...]}, {"table": {"rows": [...]}}
            value = block.pop(kind)
            block["type"] = kind
            if isinstance(value, dict):
                block = {**value, **block}
            elif kind in ("bullets", "numbered"):
                block.setdefault("items", value)
            elif kind == "table":
                block.setdefault("rows", value)
            elif isinstance(value, str):
                block.setdefault("text", value)
        elif "rows" in block:
            block["type"] = "table"
        elif "items" in block:
            block["type"] = "bullets"
        elif "text" in block:
            block["type"] = "paragraph"
    content_key = WORD_CONTENT_KEYS.get(block.get("type"))
    if content_key and not _filled(block.get(content_key)) and not _filled(block.get("text")):
        other = _other_content(block, {"type", "level", POSITION_KEY, "text", "items", "rows"})
        if other is not None:  # e.g. {"type": "paragraph", "content": "..."}, {"type": "bullets", "list": [...]}
            block[content_key] = other
    if block.get("type") in ("bullets", "numbered") and isinstance(block.get("text"), list) and not block.get("items"):
        block["items"] = block.pop("text")  # {"type": "bullets", "text": [...]}
    elif isinstance(block.get("text"), list):
        block["text"] = "\n".join(_line_text(line) for line in block["text"])
    if isinstance(block.get("items"), str):
        block["items"] = split_lines(block["items"])
    if isinstance(block.get("items"), list):  # [{"text": "..."}]
        block["items"] = [_line_text(item) if isinstance(item, dict) else item for item in block["items"]]
        if block.get("type") in ("heading", "paragraph", "quote") and not _filled(block.get("text")):
            block["text"] = "\n".join(str(item) for item in flatten(block.pop("items")))
    block["items"] = flatten(block.get("items"))
    if isinstance(block["items"], list):
        block["items"] = top_level(split_item_lines(block["items"]))
    if isinstance(block.get("rows"), list):
        block["rows"] = table_rows(block["rows"])
    return block


WORD_CONTENT_KEYS = {
    "heading": "text", "paragraph": "text", "quote": "text", "bullets": "items", "numbered": "items", "table": "rows",
}  # fmt: skip


def _line_text(line: Any) -> str:
    """Text of one entry of a list sent as text; an object is a nested block, not something to print."""
    if isinstance(line, dict):
        items = line.get("items")
        return "\n".join(map(str, items)) if isinstance(items, list) else str(line.get("text") or "")
    return str(line)


def _filled(value: Any) -> bool:
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, list):
        return any(_filled(item) for item in value)
    return value is not None


def _other_content(item: dict[str, Any], known: set[str]) -> Any:
    """The value of the only other key that holds text or a list (the model used another name for the field)."""
    others = [v for k, v in item.items() if k not in known and isinstance(v, str | list) and _filled(v)]
    return others[0] if len(others) == 1 else None


def top_level(items: list[Any]) -> list[Any]:
    """Models used to Markdown start every item with '- '; without any top-level item these are no sub items."""
    texts = [item.strip() for item in items if isinstance(item, str) and item.strip()]
    if texts and all(text.startswith(SUB_ITEM_PREFIX) for text in texts):
        return [item.strip()[len(SUB_ITEM_PREFIX) :].strip() if isinstance(item, str) else item for item in items]
    return items


def split_item_lines(items: list[Any]) -> list[Any]:
    """An item with line breaks ('a\\n• b\\n• c') holds several items."""
    result: list[Any] = []
    for item in items:
        if isinstance(item, str) and "\n" in item.strip():
            result.extend(re.sub(r"^\s*[•▪◦·*]\s*", "", line).rstrip() for line in item.split("\n") if line.strip())
        else:
            result.append(item)
    return result


def flatten(items: Any) -> Any:
    """[["a", "b"], ["c"]] -> ["a", "b", "c"] (items written as several arrays); a list right after an item
    holds its sub items: ["a", ["a1", "a2"]] -> ["a", "- a1", "- a2"]."""
    if not isinstance(items, list):
        return [] if items is None else items
    result: list[Any] = []
    for entry in items:
        if isinstance(entry, list):
            sub_items = bool(result) and not isinstance(items[0], list)
            for item in entry:
                if isinstance(item, list):
                    result.extend(_deep_cells(item))  # [["a"], [["b"], [["c"]]]]
                elif sub_items and isinstance(item, str) and not item.strip().startswith(SUB_ITEM_PREFIX):
                    result.append(SUB_ITEM_PREFIX + item.strip())
                else:
                    result.append(item)
        else:
            result.append(entry)
    return result


def table_rows(rows: list[Any]) -> list[Any]:
    """Table rows in the shapes models send: a row as plain text becomes a row with one cell; a "row" that
    only holds lists holds the next rows ([["1", "a"], [["2", "b"]]]); lists inside a row hold more cells
    (["TOP", ["Beschluss"]]); an object in a cell gives its value ({"formula": "=SUM(D2:D6)"})."""
    result: list[Any] = []
    for row in rows:
        if isinstance(row, list) and row and all(isinstance(cell, list) for cell in row):
            result.extend(table_rows(row))
        elif isinstance(row, list):
            result.append([_cell_value(cell) for cell in _deep_cells(row)])
        else:
            result.append([_cell_value(row)])
    return result


def _deep_cells(row: list[Any]) -> list[Any]:
    return [cell for item in row for cell in (_deep_cells(item) if isinstance(item, list) else [item])]


def _cell_value(cell: Any) -> Any:
    if isinstance(cell, dict):
        return next((cell[k] for k in ("formula", "value", "text") if isinstance(cell.get(k), str | int | float)), "")
    return cell


NULL_TEXTS = {"null", "none", "nan"}
QUOTED_RE = re.compile(r'"([^"]*)"')


def _sheet_cell(value: Any) -> Any:
    """'null' written as text is an empty cell; '"Rent"' (quotes inside the text) is 'Rent'."""
    if isinstance(value, str):
        if value.strip().lower() in NULL_TEXTS:
            return None
        if quoted := QUOTED_RE.fullmatch(value.strip()):
            return quoted[1]
    return value


def normalize_sheet(sheet: dict[str, Any], _: int) -> dict[str, Any]:
    sheet = dict(sheet)
    if "columns" not in sheet and isinstance(sheet.get("headers"), list):
        sheet["columns"] = sheet.pop("headers")
    if isinstance(sheet.get("columns"), list):
        columns = []
        for column in sheet["columns"]:
            if isinstance(column, str):
                column = {"header": column}  # ["Item", "Amount"]
            elif isinstance(column, dict) and "header" not in column:
                column = {**column, "header": str(column.get("name") or column.get("title") or "")}
            columns.append(column)
        sheet["columns"] = columns
    if isinstance(sheet.get("rows"), list):
        sheet["rows"] = [[_sheet_cell(v) for v in row] for row in table_rows(sheet["rows"])]
    chart = sheet.get("chart")
    if isinstance(chart, dict) and isinstance(chart.get("series_rows"), list):
        names = _deep_cells(chart["series_rows"])  # [["Income", "Expenses"]]
        if all(isinstance(name, str) for name in names):
            sheet["chart"] = {**chart, "series_rows": names}
    if isinstance(chart, dict) and not any(_filled(v) for k, v in chart.items() if k != POSITION_KEY):
        sheet.pop("chart")  # an empty chart object means "no chart"
    return sheet


def move_lonely_charts(sheets: list[Any]) -> list[Any]:
    """A sheet with only a chart (the model wanted the chart on its own sheet) gives its chart to the sheet
    that has the columns the chart refers to."""
    result = list(sheets)
    for lonely in [
        s for s in sheets if isinstance(s, dict) and not s.get("columns") and isinstance(s.get("chart"), dict)
    ]:
        chart = lonely["chart"]
        wanted = {
            str(chart.get("category_column") or "").strip(),
            *(str(v).strip() for v in chart.get("value_columns") or []),
        }
        for target in result:
            headers = (
                {str(c.get("header", "")).strip() for c in target.get("columns") or [] if isinstance(c, dict)}
                if isinstance(target, dict)
                else set()
            )
            if (
                isinstance(target, dict)
                and target is not lonely
                and not target.get("chart")
                and wanted - {""} <= headers
                and headers
            ):
                target["chart"] = chart
                result.remove(lonely)
                break
    return result


SLIDE_LAYOUTS = ("title", "section", "bullets", "two_columns", "table", "chart")
SLIDE_LAYOUT_ALIASES = {
    "two_column": "two_columns", "two-columns": "two_columns", "columns": "two_columns", "comparison": "two_columns",
    "list": "bullets", "bullet": "bullets", "title_slide": "title", "divider": "section",
}  # fmt: skip
SLIDE_ITEM_KEYS = ("items", "points", "content", "text", "body", "list")  # other names models use for bullets
CHART_KINDS = ("column", "bar", "line", "pie")


def normalize_slide(slide: dict[str, Any], index: int) -> dict[str, Any]:
    slide = dict(slide)
    if "layout" not in slide and isinstance(slide.get("type"), str):
        slide["layout"] = slide.pop("type")  # "type": "chart" instead of "layout": "chart"
    if isinstance(slide.get("layout"), str):
        layout = slide["layout"].strip().lower()
        slide["layout"] = SLIDE_LAYOUT_ALIASES.get(layout, layout)
    if not slide.get("chart") and ("categories" in slide or "series" in slide):
        chart = {key: slide.pop(key) for key in ("categories", "series") if key in slide}
        kind = next((key for key in ("type", "layout") if slide.get(key) in CHART_KINDS), None)
        if kind is not None:
            chart["type"] = slide.pop(kind)  # {"type": "pie", "categories": ...}
        slide["chart"] = chart
    if slide.get("layout") in ("title", "section") and not slide.get("subtitle"):
        text = next((slide.pop(k) for k in SLIDE_ITEM_KEYS if isinstance(slide.get(k), str) and slide[k].strip()), None)
        if text is not None:
            slide["subtitle"] = text
    target = "table" if slide.get("layout") == "table" else "bullets"
    if not _filled(slide.get(target)):  # e.g. {"layout": "bullets", "items": [...]}
        key = next((k for k in SLIDE_ITEM_KEYS if isinstance(slide.get(k), str | list) and _filled(slide[k])), None)
        if key is not None:
            slide[target] = slide.pop(key)
    left = slide.get("left")
    if (
        isinstance(left, list)
        and len(left) == 2
        and all(isinstance(column, list) for column in left)
        and not _filled(slide.get("right"))
    ):  # '"left": [...], [...]' with the "right" key forgotten
        slide["left"], slide["right"] = left
    for key in ("bullets", "left", "right"):
        if isinstance(slide.get(key), str):
            slide[key] = split_lines(slide[key])
        slide[key] = flatten(slide.get(key))
        if isinstance(slide[key], list):
            slide[key] = top_level(split_item_lines(slide[key]))
    if isinstance(slide.get("table"), dict) and "rows" in slide["table"]:
        slide["table"] = slide["table"]["rows"]
    if isinstance(slide.get("table"), list):
        slide["table"] = table_rows(slide["table"])
    if isinstance(slide.get("chart"), dict):  # e.g. "notes" written inside the chart after a missing '}'
        chart = dict(slide["chart"])
        for key in set(chart) & SLIDE_ONLY_FIELDS:
            slide.setdefault(key, chart.pop(key))
        if slide.get("layout") in CHART_KINDS:  # "type": "pie" next to the chart object
            chart.setdefault("type", slide["layout"])
            slide["layout"] = "chart"
        slide["chart"] = chart
    if not slide.get("title") and slide.get("left_heading") and slide.get("right_heading"):
        slide["title"] = f"{slide['left_heading']} & {slide['right_heading']}"  # comparison slide without title
    content = {
        "chart": bool(slide.get("chart")),
        "table": _filled(slide.get("table")),
        "two_columns": _filled(slide.get("left")) or _filled(slide.get("right")),
        "bullets": _filled(slide.get("bullets")),
    }
    present = next((layout for layout, filled in content.items() if filled), None)
    layout = slide.get("layout")
    if layout is None:
        slide["layout"] = present or ("title" if index == 0 else "section")
    elif present and (
        layout not in SLIDE_LAYOUTS or layout == "section" or (layout in content and not content[layout])
    ):
        slide["layout"] = present  # e.g. a bullets slide that carries a table: show the table
    elif layout in ("bullets", "two_columns") and not present:
        slide["layout"] = "section"  # a list slide without any items: show its title as a divider
    return slide


class Checker:
    """Collects limit violations so the model gets all of them at once."""

    def __init__(self) -> None:
        self.errors: list[str] = []
        self.chars = 0

    def require(self, condition: bool, message: str) -> bool:
        if not condition:
            self.errors.append(message)
        return condition

    def text(self, *texts: str) -> None:
        self.chars += sum(len(t) for t in texts)

    def raise_if_failed(self) -> None:
        self.require(self.chars <= MAX_TEXT_CHARS, f"the content is too long (more than {MAX_TEXT_CHARS} characters)")
        if self.errors:
            raise InvalidInput(self.errors)


# ----------------------------------------------------------------------------------- inline markdown --

INLINE_RE = re.compile(
    r"\\(?P<escaped>[\\`*_{}\[\]()#+\-.!|~<>])"
    r"|`(?P<code>[^`]+)`"
    r"|!?\[(?P<label>[^\]]*)\]\((?P<url>(?:[^()\s]|\([^()\s]*\))+)(?:\s+\"[^\"]*\")?\)"
    r"|<(?P<autolink>(?:https?://|mailto:)[^>\s]+)>"
    r"|\*\*\*(?P<bold_italic>.+?)\*\*\*"
    # Bold may contain *italic*, also at its end ('**a *b***'). The alternatives must not overlap, or text
    # like '***a*a*a…' takes exponential time (and the regex holds the GIL, freezing all of Open WebUI).
    r"|\*\*(?P<bold>(?:[^*]|\*[^*\s][^*]*\*(?=\*\*)|\*(?!\*))+?)\*\*"
    r"|(?<!\w)__(?P<bold_u>.{1,500}?)__(?!\w)"  # bounded: unclosed openers would each scan the rest of the text
    r"|~~(?P<strike>.+?)~~"
    r"|\*(?P<italic>[^*\s](?:[^*]*[^*\s])?)\*"
    r"|(?<!\w)_(?P<italic_u>[^_\s](?:[^_]*[^_\s])?)_(?!\w)",
    re.DOTALL,
)
EMPHASIS_GROUPS: dict[str, dict[str, bool]] = {
    "bold_italic": {"bold": True, "italic": True},
    "bold": {"bold": True},
    "bold_u": {"bold": True},
    "strike": {"strike": True},
    "italic": {"italic": True},
    "italic_u": {"italic": True},
}
BROKEN_QUOTE = '<|"|>'  # Gemma's string delimiter; leaks into arguments when its tool-call syntax breaks
XML_ILLEGAL_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]+')


@dataclass(frozen=True)
class Run:
    """A piece of inline text with uniform formatting."""

    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False
    url: str | None = None


PLAIN = Run("")


def clean_text(text: str) -> str:
    """Remove characters that are invalid in XML (docx/xlsx/pptx) and leaked model tokens."""
    return XML_ILLEGAL_RE.sub("", text.replace(BROKEN_QUOTE, '"'))


def parse_inline(text: str, style: Run = PLAIN) -> list[Run]:
    """Split inline Markdown into formatted runs (emphasis may be nested)."""
    text = clean_text(text)
    runs: list[Run] = []
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            runs.append(replace(style, text=text[position : match.start()]))
        runs.extend(_runs_for_match(match, style))
        position = match.end()
    if position < len(text):
        runs.append(replace(style, text=text[position:]))
    return runs


def _runs_for_match(match: re.Match[str], style: Run) -> list[Run]:
    kind = match.lastgroup or ""
    if kind == "escaped":
        return [replace(style, text=match["escaped"])]
    if kind == "code":
        return [replace(style, text=match["code"], code=True)]
    if kind == "url":
        label = to_plain(match["label"]) or match["url"]
        url = match["url"] if is_safe_url(match["url"]) else None
        return [replace(style, text=label, url=url)]
    if kind == "autolink":
        return [replace(style, text=match["autolink"], url=match["autolink"])]
    return parse_inline(match[kind], replace(style, **EMPHASIS_GROUPS[kind]))


def to_plain(text: str) -> str:
    return "".join(run.text for run in parse_inline(text)).strip()


def excel_text(text: str) -> str:
    """Text for an Excel cell: Excel cells have no inline formatting, so Markdown markers are removed."""
    text = clean_text(text)
    return to_plain(text) if any(marker in text for marker in "*_`[~<\\") else text


MAILTO_FIELDS = {"subject", "body", "cc", "bcc"}  # others, e.g. "attach", are acted on by some mail clients
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def is_safe_url(url: str) -> bool:
    url = url.strip()
    if not url.lower().startswith(SAFE_LINK_PREFIXES):
        return False
    if url.lower().startswith("mailto:") and "?" in url:
        fields = [part.split("=", 1)[0].lower() for part in url.split("?", 1)[1].split("&") if part]
        return all(field in MAILTO_FIELDS for field in fields)
    return True


def safe_filename(text: str, fallback: str) -> str:
    """Turn text into a short file name that is valid on Windows, macOS and Linux.

    Invisible format characters are removed: a right-to-left override would make 'x‮fdp.docx' look
    like 'xcod.pdf'.
    """
    text = "".join(char for char in to_plain(text) if unicodedata.category(char) != "Cf")
    name = " ".join(UNSAFE_FILENAME_RE.sub(" ", text).split())[:MAX_FILENAME_CHARS].strip(" .")
    if name.split(".")[0].upper() in WINDOWS_RESERVED_NAMES:
        name += "_"
    return name or fallback


def tidy_title(text: str) -> str:
    """Plain-text title; a file-name-like title ('Report_Q3_2026') gets its spaces back."""
    title = re.sub(r"\.(docx?|xlsx?|pptx?)$", "", " ".join(to_plain(text).split()), flags=re.IGNORECASE)
    if "_" in title and " " not in title:
        title = " ".join(title.replace("_", " ").split())
    return title


def similar_text(a: str, b: str) -> bool:
    """One text contains the other, or they differ only slightly (a title repeated as a heading)."""
    ka, kb = (re.sub(r"\W+", "", to_plain(t)).casefold() for t in (a, b))
    if not ka or not kb:
        return False
    return (min(len(ka), len(kb)) >= 12 and (ka in kb or kb in ka)) or difflib.SequenceMatcher(
        None, ka, kb
    ).ratio() >= 0.9


def same_text(a: str, b: str) -> bool:
    def key(text: str) -> str:
        return re.sub(r"\W+", "", to_plain(text)).casefold()

    return key(a) == key(b)


def split_level(item: str) -> tuple[int, str]:
    """'- text' is a level-1 (indented) item, everything else level 0."""
    stripped = item.strip()
    if stripped.startswith(SUB_ITEM_PREFIX):
        return 1, stripped[len(SUB_ITEM_PREFIX) :].strip()
    return 0, stripped


# ------------------------------------------------------------------------------------------------ Word --

ORDERED_LIST_STYLES = ("List Number", "List Number 2")
BULLET_LIST_STYLES = ("List Bullet", "List Bullet 2")
MAX_HEADING_LEVEL = 3
LONG_DOCUMENT_CHARS = 4000  # roughly two pages
W_NAMESPACE = qn("w:x")[:-1]
SETTINGS_AFTER_UPDATE_FIELDS = {
    *(qn(f"w:{name}") for name in (
        "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "attachedSchema", "themeFontLang",
        "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
        "readModeInkLockDown", "smartTagType", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
    )),
}  # fmt: skip


@dataclass
class WordOptions:
    title: str
    subtitle: str = ""
    header_text: str = ""
    table_of_contents: bool = False
    landscape: bool = False
    page_numbers: bool = True


def word_title(title: str, blocks: list[WordBlock], fallback: str) -> tuple[str, list[WordBlock]]:
    """Return the document title and the blocks without a first heading that only repeats it.

    Models often leave out the title and start with it as a heading instead, or pass a file name
    ('Minutes_2026_10_12'). Then the first heading becomes the title.
    """
    title = re.sub(r"\.(docx?)$", "", to_plain(title).strip(), flags=re.IGNORECASE)
    first = blocks[0] if blocks and blocks[0].type == "heading" and blocks[0].text.strip() else None
    file_name_like = "_" in title and " " not in title
    if first and (not title or same_text(title, first.text) or (file_name_like and first.level == 1)):
        return to_plain(first.text), blocks[1:]
    if first and similar_text(title, first.text):  # 'Protokoll der Teambesprechung' under '… vom 12.10.2026'
        return max(tidy_title(title), to_plain(first.text), key=len), blocks[1:]
    return tidy_title(title) or fallback, blocks


def is_empty_block(block: WordBlock) -> bool:
    """Blocks without content; models add empty paragraphs as spacing, they are left out."""
    if block.type in ("heading", "paragraph", "quote"):
        return not block.text.strip()
    if block.type in ("bullets", "numbered"):
        return not any(item.strip() for item in block.items)
    if block.type == "table":
        return not any(str(cell).strip() for row in block.rows for cell in row)
    return False


def check_word(options: WordOptions, blocks: list[WordBlock]) -> None:
    check = Checker()
    check.require(
        any(b.type not in ("page_break", "table_of_contents") for b in blocks) or bool(options.subtitle.strip()),
        "blocks: the document has no content",
    )
    check.require(len(blocks) <= MAX_WORD_BLOCKS, f"blocks: at most {MAX_WORD_BLOCKS} blocks")
    if options.table_of_contents or any(b.type == "table_of_contents" for b in blocks):
        body = sum(
            len(b.text) + sum(map(len, b.items)) + sum(len(c) for r in b.rows for c in r)
            for b in blocks
            if b.type != "heading"
        )
        check.require(
            any(b.type == "heading" for b in blocks) and body >= 100,
            "blocks: the document has a table of contents but no chapters with text; send the complete content "
            "(headings with their paragraphs, lists and tables)",
        )
    check.text(options.title, options.subtitle, options.header_text)
    for i, block in enumerate(blocks):
        where = f"blocks[{block._index if block._index >= 0 else i}] ({block.type})"
        check.text(block.text, *block.items, *(cell for row in block.rows for cell in row))
        if block.type == "table":
            columns = table_columns(block.rows)
            check.require(columns <= MAX_WORD_TABLE_COLUMNS, f"{where}: at most {MAX_WORD_TABLE_COLUMNS} columns")
            check.require(len(block.rows) <= MAX_WORD_TABLE_ROWS, f"{where}: at most {MAX_WORD_TABLE_ROWS} rows")
    check.raise_if_failed()


def _set_cell_shading(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_field(paragraph: DocxParagraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, NUMPAGES, TOC ...). Word fills in the value when it lays out the page."""
    for kind, text in (
        ("begin", None),
        ("instr", instruction),
        ("separate", None),
        ("text", placeholder),
        ("end", None),
    ):
        run = paragraph.add_run()
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = f" {text} "
        elif kind == "text":
            run.text = text or ""
            continue
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


CHAR_WIDTH = Pt(5.8)  # average character of 11 pt Calibri, bold header text included
CELL_PADDING = Pt(12)  # left and right cell margin
MAX_WORD_CHARS = 20  # longer words may wrap
MAX_TABLE_COLUMN_CHARS = 60


def table_columns(rows: list[list[str]]) -> int:
    """Number of columns: the header's, or more if a row has more non-empty cells (trailing empty cells don't count)."""
    used = [max((i + 1 for i, cell in enumerate(row) if str(cell).strip()), default=0) for row in rows]
    return max(len(rows[0]), *used, 1)


def column_widths(rows: list[list[str]], columns: int, total: int, font_size: float = 11) -> list[int]:
    """Split the text width: every column fits its longest word, the rest goes to the columns with more text."""
    minimum, wanted = [], []
    char_width = CHAR_WIDTH * font_size / 11
    for column in range(columns):
        texts = [to_plain(row[column]) for row in rows if column < len(row)]
        longest_word = max((len(word) for text in texts for word in text.split()), default=1)
        longest_text = max((len(text) for text in texts), default=1)
        minimum.append(min(longest_word, MAX_WORD_CHARS) * char_width + CELL_PADDING)
        wanted.append(max(min(longest_text, MAX_TABLE_COLUMN_CHARS) * char_width + CELL_PADDING, minimum[-1]))
    if sum(wanted) <= total:  # everything fits on one line: widen all columns evenly
        widths = [w * total / sum(wanted) for w in wanted]
    elif sum(minimum) >= total:  # not even the words fit: shrink proportionally
        widths = [m * total / sum(minimum) for m in minimum]
    else:  # words fit: share the rest by how much more each column wants
        extra = total - sum(minimum)
        widths = [m + extra * (w - m) / (sum(wanted) - sum(minimum)) for m, w in zip(minimum, wanted, strict=True)]
    result = [int(w) for w in widths]
    result[-1] += total - sum(result)
    return result


class DocxWriter:
    """Render validated blocks into a Word document."""

    def __init__(self, options: WordOptions, texts: dict[str, str]) -> None:
        self.options = options
        self.texts = texts
        self.doc = Document()
        normal_font = self.doc.styles["Normal"].font
        normal_font.name = BODY_FONT
        normal_font.size = BODY_FONT_SIZE

    def render(self, blocks: list[WordBlock]) -> bytes:
        self.long_document = (
            sum(len(b.text) + sum(map(len, b.items)) for b in blocks) + 40 * len(blocks) > LONG_DOCUMENT_CHARS
        )
        self._page_setup()
        self.doc.add_heading(to_plain(self.options.title), level=0)
        if self.options.subtitle.strip():
            self._add_runs(self.doc.add_paragraph(style="Subtitle"), parse_inline(self.options.subtitle))
        if self.options.table_of_contents:
            self._table_of_contents()
        for block in blocks:
            match block.type:
                case "heading":
                    self.doc.add_heading(to_plain(block.text), level=min(max(block.level, 1), MAX_HEADING_LEVEL))
                case "paragraph":
                    self._paragraph(block.text)
                case "quote":
                    self._add_runs(self.doc.add_paragraph(style="Quote"), parse_inline(block.text.strip()))
                case "bullets" | "numbered":
                    self._list(block.items, ordered=block.type == "numbered")
                case "table":
                    self._table(block.rows)
                case "page_break":
                    self.doc.add_page_break()
                case "table_of_contents":
                    self._table_of_contents()
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()

    def _page_setup(self) -> None:
        section = self.doc.sections[0]
        if self.options.landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        if self.options.header_text.strip():
            header = section.header.paragraphs[0]
            header.text = clean_text(self.options.header_text.strip())
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if self.options.page_numbers:
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            before, _, rest = self.texts["page_footer"].partition("{page}")
            between, _, after = rest.partition("{pages}")
            footer.add_run(before)
            _add_field(footer, "PAGE", "1")
            footer.add_run(between)
            _add_field(footer, "NUMPAGES", "1")
            footer.add_run(after)

    def _table_of_contents(self) -> None:
        self.doc.add_paragraph(
            self.texts["toc_title"], style="TOC Heading" if "TOC Heading" in self._style_names() else None
        )
        _add_field(self.doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', self.texts["toc_placeholder"])
        # Ask Word to update the fields when the file is opened, so the table of contents gets filled in.
        # The settings elements have a fixed order; updateFields comes before compat, rsids etc.
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            update = OxmlElement("w:updateFields")
            update.set(qn("w:val"), "true")
            follower = next(
                (
                    child
                    for child in settings
                    if isinstance(child.tag, str)
                    and (child.tag in SETTINGS_AFTER_UPDATE_FIELDS or not child.tag.startswith(W_NAMESPACE))
                ),
                None,
            )
            if follower is not None:
                follower.addprevious(update)
            else:
                settings.append(update)
        if self.long_document:  # a short document would get a nearly empty first page
            self.doc.add_page_break()

    def _paragraph(self, text: str) -> None:
        """Paragraphs separated by blank lines; Markdown list lines inside the text become real lists."""
        for kind, lines in split_paragraph(text):
            if kind == "text":
                self._add_runs(self.doc.add_paragraph(), parse_inline("\n".join(lines)))
            elif kind == "table":
                self._table(markdown_table_rows(lines)[:MAX_WORD_TABLE_ROWS])
            else:
                self._list(lines, ordered=kind == "numbered")

    def _style_names(self) -> set[str]:
        return {style.name for style in self.doc.styles}

    def _list(self, items: list[str], ordered: bool) -> None:
        styles = ORDERED_LIST_STYLES if ordered else BULLET_LIST_STYLES
        num_ids: list[int | None] = [None, None]  # per level; each list and sub list restarts at 1
        for item in items:
            level, text = split_level(item)
            if not text:
                continue
            paragraph = self.doc.add_paragraph(style=styles[level])
            paragraph.paragraph_format.space_after = Pt(3)
            if ordered:
                if level == 0:
                    num_ids[1] = None  # the next sub list starts again at 1
                if num_ids[level] is None:
                    num_ids[level] = self._new_numbering(paragraph.style)
                paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_ids[level]
            self._add_runs(paragraph, parse_inline(text))

    def _new_numbering(self, style: Any) -> int:
        """Create a numbering instance of the style's list that starts at 1 (python-docx has no API for it)."""
        numbering = self.doc.part.numbering_part.element
        style_num_id = style.element.pPr.numPr.numId.val
        num = numbering.add_num(numbering.num_having_numId(style_num_id).abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def _table(self, rows: list[list[str]]) -> None:
        columns = table_columns(rows)
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        table.autofit = False
        section = self.doc.sections[-1]
        widths = column_widths(rows, columns, section.page_width - section.left_margin - section.right_margin)
        cells = table._cells  # once: row.cells and column.cells rebuild the whole grid on every call
        for c, (column, width) in enumerate(zip(table.columns, widths, strict=True)):
            column.width = width  # read by LibreOffice
            for cell in cells[c::columns]:
                cell.width = width  # read by Word
        for row_index, (doc_row, row) in enumerate(zip(table.rows, rows, strict=True)):
            if row_index == 0:  # repeat the header row on every page
                header_flag = OxmlElement("w:tblHeader")
                header_flag.set(qn("w:val"), "true")
                doc_row._tr.get_or_add_trPr().append(header_flag)
            for column, cell in enumerate(cells[row_index * columns : (row_index + 1) * columns]):
                text = row[column] if column < len(row) else ""
                self._add_runs(cell.paragraphs[0], parse_inline(text, Run("", bold=row_index == 0)))
                if row_index == 0:
                    _set_cell_shading(cell, HEADER_FILL)
        self.doc.add_paragraph()

    def _add_runs(self, paragraph: DocxParagraph, runs: list[Run]) -> None:
        for run in runs:
            docx_run = paragraph.add_run(run.text)
            docx_run.bold = run.bold or None
            docx_run.italic = run.italic or None
            if run.strike:
                docx_run.font.strike = True
            if run.code:
                docx_run.font.name = CODE_FONT
            if run.url:
                self._make_link(paragraph, docx_run, run.url)

    @staticmethod
    def _make_link(paragraph: DocxParagraph, docx_run: Any, url: str) -> None:
        docx_run.font.underline = True
        docx_run.font.color.rgb = RGBColor.from_string(LINK_COLOR)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        docx_run._r.addprevious(link)
        link.append(docx_run._r)


def build_docx(options: WordOptions, blocks: list[WordBlock], texts: dict[str, str]) -> bytes:
    blocks = [block for block in blocks if not is_empty_block(block)]
    while blocks and blocks[-1].type == "page_break":  # would only add an empty page
        blocks = blocks[:-1]
    if any(block.type == "table_of_contents" for block in blocks):
        options = replace(options, table_of_contents=False)  # placed where the model put it instead
    check_word(options, blocks)
    return DocxWriter(options, texts).render(blocks)


# ----------------------------------------------------------------------------------------------- Excel --

GERMAN_FUNCTIONS = {
    "SUMME": "SUM",
    "MITTELWERT": "AVERAGE",
    "ANZAHL": "COUNT",
    "ANZAHL2": "COUNTA",
    "WENN": "IF",
    "WENNFEHLER": "IFERROR",
    "UND": "AND",
    "ODER": "OR",
    "NICHT": "NOT",
    "RUNDEN": "ROUND",
    "AUFRUNDEN": "ROUNDUP",
    "ABRUNDEN": "ROUNDDOWN",
    "ZÄHLENWENN": "COUNTIF",
    "SUMMEWENN": "SUMIF",
    "MITTELWERTWENN": "AVERAGEIF",
    "PRODUKT": "PRODUCT",
    "POTENZ": "POWER",
    "WURZEL": "SQRT",
    "REST": "MOD",
    "GANZZAHL": "INT",
    "BETRAG": "ABS",
}
FORMULA_FUNCTIONS = {
    "SUM", "AVERAGE", "MIN", "MAX", "MEDIAN", "COUNT", "COUNTA", "COUNTIF", "SUMIF", "AVERAGEIF",
    "IF", "IFERROR", "AND", "OR", "NOT", "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "PRODUCT", "POWER",
    "SQRT", "MOD", "INT",
}  # fmt: skip
FORMULA_CONSTANTS = {"TRUE": "TRUE", "FALSE": "FALSE", "WAHR": "TRUE", "FALSCH": "FALSE"}
FORMULA_TOKEN_RE = re.compile(  # [0-9], not \d: \d also matches digits of other scripts
    r"(?P<space>\s+)"
    r'|(?P<string>"[^"]*")'
    r"|(?P<function>[A-Za-zÄÖÜäöü][A-Za-z0-9ÄÖÜäöü.]*)\s*\("
    r"|(?P<cell>\$?(?P<column>[A-Za-z]{1,3})(?P<row_absolute>\$?)(?P<row>[1-9][0-9]{0,6}|\{row\}))(?![A-Za-z0-9_])"
    r"|(?P<name>[A-Za-z]+)(?![A-Za-z0-9_(])"
    r"|(?P<number>[0-9]+(?:\.[0-9]+)?)"
    r"|(?P<operator><=|>=|<>|[-+*/^&=<>(),;:%])"
)
DECIMAL_COMMA_RE = re.compile(r",([0-9]+)(?![0-9A-Za-z(])")
MAX_EXCEL_COLUMN, MAX_EXCEL_ROW = 16_384, 1_048_576  # XFD, the last column; the last row
OPERAND_KINDS = ("cell", "number", "string", "name")
FIRST_DATA_ROW = "2"


def _column_number(letters: str) -> int:
    number = 0
    for letter in letters.upper():
        number = number * 26 + ord(letter) - ord("A") + 1
    return number


def normalize_formula(formula: str, infer_row: bool = True) -> str:
    """Check a formula against the whitelist and return it in Excel's file syntax (English names, commas).

    Only cell references of the same sheet, numbers, quoted text, operators and whitelisted functions are
    allowed. Sheet references (!), external workbooks ([...]), DDE (|) and every other function are rejected.
    With ';' as separator (German Excel), commas in numbers are decimal commas; outside of functions a comma
    between digits is always one. A formula without {row} whose cells all point to the first data row
    ('=B2*C2') gets {row} there, so every row computes its own values. Raises ValueError with a message for
    the model.
    """
    text = XML_ILLEGAL_RE.sub("", formula).strip()
    if not text.startswith("="):
        text = "=" + text  # models sometimes leave out the "="
    if len(text) > MAX_FORMULA_CHARS:
        raise ValueError(f"a formula may have at most {MAX_FORMULA_CHARS} characters")
    german = ";" in re.sub(r'"[^"]*"', "", text)
    parts: list[str] = []
    cells: list[tuple[int, re.Match[str]]] = []  # (index in parts, match) of every cell reference
    depth = 0
    position = 1
    previous = ""  # kind of the previous token, "space" included
    before_space = ""  # kind of the last token before a space
    while position < len(text):
        match = FORMULA_TOKEN_RE.match(text, position)
        if not match:
            raise ValueError(f"'{text[position : position + 10]}' is not allowed in a formula")
        kind, token = match.lastgroup, match.group()
        end = match.end()
        if (
            kind != "space"
            and previous == "space"
            and (before_space in OPERAND_KINDS or before_space == ")")
            and (kind in (*OPERAND_KINDS, "function") or token == "(")
        ):
            near = text[position : position + 10]
            raise ValueError(f"a space between two values is not allowed near '{near}'; use an operator")
        if kind == "function":
            name = match["function"].upper()
            name = GERMAN_FUNCTIONS.get(name, name)
            if name not in FORMULA_FUNCTIONS:
                raise ValueError(
                    f"the function {match['function']} is not allowed (allowed: {', '.join(sorted(FORMULA_FUNCTIONS))})"
                )
            token = f"{name}("
            depth += 1
        elif kind == "cell":
            if _column_number(match["column"]) > MAX_EXCEL_COLUMN or (
                match["row"] != "{row}" and int(match["row"]) > MAX_EXCEL_ROW
            ):
                raise ValueError(f"'{token}' is outside of the sheet")
            token = token.upper().replace("{ROW}", "{row}")
            cells.append((len(parts), match))
        elif kind == "name":
            if token.upper() not in FORMULA_CONSTANTS:
                raise ValueError(f"'{token}' is not allowed in a formula (only cells like B{{row}} or B2)")
            token = FORMULA_CONSTANTS[token.upper()]
        elif kind == "number":
            fraction = DECIMAL_COMMA_RE.match(text, end) if "." not in token and (german or depth == 0) else None
            if fraction:  # 0,19 -> 0.19
                token, end = f"{token}.{fraction[1]}", fraction.end()
        elif kind == "operator":
            if token == ";":
                token = ","
            elif token == "(":
                depth += 1
            elif token == ")":
                depth -= 1
                if depth < 0:
                    raise ValueError("the formula has a ')' without '('")
        elif kind == "space":
            before_space = previous if previous != "space" else before_space
            token = ""
        parts.append(token)
        previous = kind if kind != "operator" or token != ")" else ")"
        position = end
    if depth != 0:
        raise ValueError("the parentheses in the formula do not match")
    relative_rows = {match["row"] for _, match in cells if not match["row_absolute"]}
    if infer_row and "{row}" not in relative_rows and relative_rows == {FIRST_DATA_ROW}:
        for index, match in cells:
            if not match["row_absolute"]:
                parts[index] = parts[index][: -len(FIRST_DATA_ROW)] + "{row}"
    return "=" + "".join(parts)


NUMBER_RE = re.compile(  # thousands separators: . , ' space, no-break space, narrow no-break space
    r"(?P<prefix>€|\$|£|EUR|USD|GBP|CHF)?\s*"
    r"(?P<sign>[-+]?)"
    r"(?P<int>\d{1,3}(?:(?P<group>[.,' \u00a0\u202f])\d{3}(?:(?P=group)\d{3})*)?|\d+)"
    r"(?:(?!(?P=group))(?P<decimal>[.,])(?P<fraction>\d+))?"
    r"\s*(?P<suffix>%|€|\$|£|EUR|USD|GBP|CHF)?"
)
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M")


def parse_number_text(text: str, decimal_comma: bool | None = None) -> tuple[float | int, bool] | None:
    """'1.234,56 €' -> (1234.56, False), '12,5 %' -> (0.125, True). None if the text is not a number.

    A dot or comma followed by exactly three digits is read as a thousands separator, except after a
    leading 0 ('0,125') and, if `decimal_comma` says which one is the decimal mark, for that mark when
    it appears only once ('1,799' is 1.799 in German, '1.500' is 1.5 in English).
    """
    match = NUMBER_RE.fullmatch(text.strip().replace("\u2212", "-"))
    if not match or (match["prefix"] and match["suffix"]):
        return None
    if match["group"] in (",", ".") and not match["fraction"]:
        head, *groups = match["int"].split(match["group"])
        decimal_mark = None if decimal_comma is None else "," if decimal_comma else "."
        if len(groups) == 1 and (head == "0" or match["group"] == decimal_mark):
            number = float(f"{match['sign']}{head}.{groups[0]}")
            return (number / 100, True) if match["suffix"] == "%" else (number, False)
    digits, fraction = re.sub(r"\D", "", match["int"]), match["fraction"] or ""
    number: float | int = float(f"{match['sign']}{digits}.{fraction}") if fraction else int(match["sign"] + digits)
    if match["suffix"] == "%":
        return number / 100, True
    return number, False


def parse_date(text: str) -> date | datetime | None:
    for date_format in DATE_FORMATS:
        try:
            parsed = datetime.strptime(text.strip(), date_format)
        except ValueError:
            continue
        return parsed if "%H" in date_format else parsed.date()
    return None


def convert_value(value: str | float | None, column_type: str, decimal_comma: bool | None = None) -> tuple[Any, bool]:
    """Return (cell value, is_typed). Values that do not fit the column type stay text (is_typed False).

    JSON numbers stay numbers in text columns too, so sums and charts over them work.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if column_type == "text":
        if isinstance(value, float):
            return (int(value) if value.is_integer() else value), True
        return excel_text(value), True
    if column_type == "date":
        parsed = parse_date(str(value)) if isinstance(value, str) else None
        return (parsed, True) if parsed else (excel_text(str(value)), False)
    if isinstance(value, str):
        parsed_number = parse_number_text(value, decimal_comma)
        if parsed_number is None:
            return excel_text(value), False
        number, is_fraction = parsed_number
    else:
        number, is_fraction = value, False
    if column_type == "percent" and not is_fraction:
        number = number / 100
    if column_type == "integer" and float(number).is_integer():
        number = int(number)
    elif isinstance(number, float) and number.is_integer() and column_type != "percent":
        number = int(number)
    return number, True


def _decimals(values: list[Any]) -> int:
    decimals = 0
    for value in values:
        if isinstance(value, float):
            text = f"{value:.{MAX_DECIMALS}f}".rstrip("0")
            decimals = max(decimals, len(text.split(".")[1]) if "." in text else 0)
    return decimals


def number_format(column: ExcelColumn, values: list[Any], currency: str, date_format: str) -> str | None:
    match column.type:
        case "integer":
            return "#,##0"
        case "number":
            decimals = _decimals(values) if not column.formula else 2
            return "#,##0" + ("." + "0" * decimals if decimals else "")
        case "currency":
            return CURRENCY_FORMATS[currency]
        case "percent":
            fractional = any(isinstance(v, float) and not round(v * 100, 9).is_integer() for v in values)
            return "0.0%" if column.formula or fractional else "0%"
        case "date":
            return date_format
    return None


BOLD_FONT = Font(name=BODY_FONT, size=11, bold=True)
SHEET_NAME_INVALID_RE = re.compile(r"[\[\]:*?/\\]")
TOTAL_FUNCTIONS = {"sum": "SUM", "average": "AVERAGE", "min": "MIN", "max": "MAX", "count": "COUNTA"}
RESERVED_SHEET_NAMES = {"history"}  # reserved by Excel
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")  # text starting like this may turn into a formula when edited


def sheet_names(sheets: list[ExcelSheet], fallback: str = "Sheet") -> list[str]:
    """Make valid, unique sheet names (31 characters, no []:*?/\\)."""
    names: list[str] = []
    for n, sheet in enumerate(sheets, start=1):
        base = " ".join(SHEET_NAME_INVALID_RE.sub(" ", to_plain(sheet.name)).split()).strip("'")[:31].strip("' ")
        if base.lower() in RESERVED_SHEET_NAMES:
            base += "_"
        base = base or f"{fallback}{n}"
        name, k = base, 2
        while name.lower() in (existing.lower() for existing in names):
            suffix = f" ({k})"
            name, k = base[: 31 - len(suffix)] + suffix, k + 1
        names.append(name)
    return names


def align_rows(sheet: ExcelSheet) -> list[list[str | float | None]]:
    """Rows that left out the formula columns get None inserted at those positions."""
    formula_positions = [i for i, column in enumerate(sheet.columns) if column.formula.strip()]
    short = len(sheet.columns) - len(formula_positions)
    aligned = []
    for row in sheet.rows:
        if formula_positions and len(row) == short:
            row = list(row)
            for position in formula_positions:
                row.insert(position, None)
        aligned.append(list(row) + [None] * (len(sheet.columns) - len(row)))
    return aligned


def check_excel(sheets: list[ExcelSheet], decimal_comma: bool | None = None) -> dict[tuple[int, int], str]:
    """Validate the workbook and return the normalized formulas by (sheet, column)."""
    check = Checker()
    formulas: dict[tuple[int, int], str] = {}
    check.require(bool(sheets), "sheets: the workbook has no sheets")
    check.require(len(sheets) <= MAX_SHEETS, f"sheets: at most {MAX_SHEETS} sheets")
    cells = sum(len(sheet.rows) * len(sheet.columns) for sheet in sheets)
    check.require(cells <= MAX_CELLS, f"sheets: at most {MAX_CELLS} cells in all sheets together ({cells} sent)")
    for s, sheet in enumerate(sheets):
        where = f"sheets[{sheet._index if sheet._index >= 0 else s}]"
        headers = [column.header.strip() for column in sheet.columns]
        check.require(bool(sheet.columns), f"{where}.columns: the sheet has no columns")
        check.require(len(sheet.columns) <= MAX_SHEET_COLUMNS, f"{where}.columns: at most {MAX_SHEET_COLUMNS} columns")
        check.require(len(sheet.rows) <= MAX_SHEET_ROWS, f"{where}.rows: at most {MAX_SHEET_ROWS} rows")
        check.text(sheet.name, *headers)
        for c, column in enumerate(sheet.columns):
            if column.formula.strip():
                try:
                    formulas[(s, c)] = normalize_formula(column.formula)
                except ValueError as e:
                    check.require(False, f"{where}.columns[{c}].formula: {e}")
        for r, row in enumerate(sheet.rows):
            check.require(
                len(row) <= len(sheet.columns),
                f"{where}.rows[{r}]: {len(row)} values but only {len(sheet.columns)} columns",
            )
            for value in row:
                if isinstance(value, str):
                    check.text(value)
                    check.require(
                        len(value) <= MAX_CELL_CHARS,
                        f"{where}.rows[{r}]: a cell has more than {MAX_CELL_CHARS} characters",
                    )
        if len(check.errors) > MAX_ERRORS_SHOWN:
            break
        if sheet.rows and all(len(row) <= len(sheet.columns) for row in sheet.rows):
            check.errors.extend(
                _column_problems(sheet, where, {c: f for (i, c), f in formulas.items() if i == s}, decimal_comma)
            )
        if sheet.chart:
            if check.require(bool(sheet.rows), f"{where}.chart: the sheet has no rows to plot"):
                try:
                    chart_layout(sheet, sheet.chart)
                except ValueError as e:
                    check.errors.extend(f"{where}.chart.{error}" for error in e.args[0])
    check.raise_if_failed()
    return formulas


PLACEHOLDERS = {"-", "\u2013", "\u2014", "/", "n/a", "na", "k.a.", "k. a.", "tbd", "offen", "?"}  # in number columns
FORMULA_CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?(?:[0-9]+|\{row\})")


def _column_problems(sheet: ExcelSheet, where: str, formulas: dict[int, str], decimal_comma: bool | None) -> list[str]:
    """Columns whose content cannot be what the model meant; these usually come from rows whose values are
    shifted against the columns, or from values the model forgot."""
    rows = align_rows(sheet)
    problems: list[str] = []
    filled = [any(row[c] is not None and str(row[c]).strip() for row in rows) for c in range(len(sheet.columns))]
    for c, column in enumerate(sheet.columns):
        name = f"{where}.columns[{c}] ('{column.header.strip()}')"
        if c in formulas:
            continue
        if column.total != "none" and not filled[c]:
            problems.append(
                f"{name}: the total needs values in this column: fill them in, give the column a formula "
                "(e.g. '=B{row}*C{row}'), or put the total on the column with the numbers"
            )
        if column.type in NUMBER_TYPES:
            values = [row[c] for row in rows if row[c] is not None and str(row[c]).strip()]
            texts = [
                v for v in values
                if str(v).strip().lower() not in PLACEHOLDERS and not convert_value(v, column.type, decimal_comma)[1]
                and not cell_formula(v, column)
            ]  # fmt: skip
            bad_formula = next((str(v).strip() for v in texts if str(v).strip().startswith("=")), None)
            if bad_formula is not None:
                try:
                    normalize_formula(bad_formula, infer_row=False)
                except ValueError as e:
                    problems.append(f"{name}: the formula '{bad_formula[:60]}' is not allowed: {e}")
            elif len(texts) >= 2 and len(texts) * 2 > len(values):
                problems.append(
                    f"{name}: {len(texts)} of {len(values)} values are not numbers (e.g. '{texts[0]}'). Each row needs "
                    f"one value per column, in column order - is a column missing? Or give the column type 'text'."
                )
    for c, formula in formulas.items():
        if get_column_letter(c + 1) in FORMULA_CELL_RE.findall(re.sub(r'"[^"]*"', "", formula)):
            problems.append(
                f"{where}.columns[{c}] ('{sheet.columns[c].header.strip()}'): the formula uses its own column "
                f"{get_column_letter(c + 1)}; it may only use the other columns (e.g. quantity * price)"
            )
            continue
        for letters in dict.fromkeys(FORMULA_CELL_RE.findall(re.sub(r'"[^"]*"', "", formula))):
            used = _column_number(letters) - 1
            if used < len(sheet.columns) and used not in formulas and not filled[used] and used != c:
                problems.append(
                    f"{where}.columns[{used}] ('{sheet.columns[used].header.strip()}') is empty, but the formula of "
                    f"'{sheet.columns[c].header.strip()}' uses it: fill in its values"
                )
    return problems


ROW_TOTAL_RE = re.compile(
    r"=(SUM|SUMME|AVERAGE|MITTELWERT|MIN|MAX|COUNTA?|ANZAHL2?)\(\$?([A-Z]{1,3})\$?2:\$?([A-Z]{1,3})\$?([0-9]+)\)",
    re.IGNORECASE,
)
ROW_TOTAL_FUNCTIONS = {
    "SUM": "sum", "SUMME": "sum", "AVERAGE": "average", "MITTELWERT": "average", "MIN": "min", "MAX": "max",
    "COUNT": "count", "COUNTA": "count", "ANZAHL": "count", "ANZAHL2": "count",
}  # fmt: skip
TOTAL_LABEL_RE = re.compile(r"\b(summe|gesamt\w*|insgesamt|total|sum)\b", re.IGNORECASE)
ROW_CELL_RE = re.compile(r"(?<![A-Za-z0-9_$])(\$?[A-Za-z]{1,3})([0-9]+)(?![0-9A-Za-z_(])")


def formulas_from_rows(sheet: ExcelSheet) -> ExcelSheet:
    """Formulas the model wrote into the rows instead of the column definition: '=B2-C2', '=B3-C3' …
    become the column formula '=B{row}-C{row}', and a last row of '=SUM(B2:B7)' cells becomes the
    total row. Cell values are never written as formulas, so otherwise the user would see formula text."""
    rows = align_rows(sheet)
    columns = list(sheet.columns)
    if len(rows) >= 2:
        totals: dict[int, str] = {}
        for c, value in enumerate(rows[-1]):
            if isinstance(value, str) and value.strip().startswith("="):
                match = ROW_TOTAL_RE.fullmatch(re.sub(r"\s+", "", value))
                letter = get_column_letter(c + 1)
                if not match or match[2].upper() != letter or match[3].upper() != letter or int(match[4]) != len(rows):
                    totals = {}
                    break
                totals[c] = ROW_TOTAL_FUNCTIONS[match[1].upper()]
        others = [v for c, v in enumerate(rows[-1]) if c not in totals and v is not None and str(v).strip()]
        if totals and all(isinstance(v, str) for v in others):  # only a label such as "Summe" besides the sums
            rows = rows[:-1]
            for c, function in totals.items():
                columns[c] = columns[c].model_copy(update={"total": function})
        elif (
            len(others) == 1
            and isinstance(rows[-1][0], str)
            and len(rows[-1][0]) <= 40
            and TOTAL_LABEL_RE.search(rows[-1][0])
            and not any(column.formula.strip() for column in columns[:1])
        ):  # an empty "Summe" row: the model wanted a total row, the tool adds a computed one
            rows = rows[:-1]
            if not any(column.total != "none" for column in columns):
                factors = {  # quantity and price of '=B{row}*C{row}': their sum means nothing
                    _column_number(letters) - 1
                    for column in columns
                    if re.search(r"[*/]", column.formula)
                    for letters in FORMULA_CELL_RE.findall(column.formula.upper().replace("{ROW}", "{row}"))
                }
                for c, column in enumerate(columns[1:], start=1):
                    numbers = any(isinstance(row[c], float) for row in rows)
                    if c not in factors and (column.type in NUMBER_TYPES or (numbers and not column.formula.strip())):
                        columns[c] = column.model_copy(update={"total": "sum"})
    for c, column in enumerate(columns):
        values = [(r, row[c]) for r, row in enumerate(rows) if row[c] is not None and str(row[c]).strip()]
        if column.formula.strip() or not values:
            continue
        if not all(isinstance(v, str) and v.strip().startswith("=") for _, v in values):
            continue
        patterns = {
            ROW_CELL_RE.sub(lambda m, n=r + 2: m[1] + "{row}" if int(m[2]) == n else m[0], v.strip()) for r, v in values
        }
        if len(patterns) == 1:
            columns[c] = column.model_copy(update={"formula": patterns.pop()})
            rows = [[None if i == c else v for i, v in enumerate(row)] for row in rows]
    return sheet.model_copy(update={"columns": columns, "rows": rows})


def cell_formula(value: Any, column: ExcelColumn) -> str | None:
    """A formula written into a single cell ('=B2-B3'), checked against the same whitelist as column formulas;
    None if the value is no allowed formula (then it is written as text). Columns the model marked as 'text'
    keep such values as text."""
    if not isinstance(value, str) or not value.strip().startswith("="):
        return None
    if column.type == "text" and "type" in column.model_fields_set:
        return None
    try:
        return normalize_formula(value, infer_row=False)
    except ValueError:
        return None


def infer_column_types(sheet: ExcelSheet, decimal_comma: bool | None = None) -> ExcelSheet:
    """Columns sent without a type whose values are all numbers become number, currency or percent columns,
    so the numbers are not stored as text (which would make sums 0 and charts empty)."""
    rows = align_rows(sheet)
    columns = []
    for c, column in enumerate(sheet.columns):
        if "type" not in column.model_fields_set and not column.formula.strip():
            kind = _number_kind([row[c] for row in rows if row[c] is not None and str(row[c]).strip()], decimal_comma)
            if kind is not None:
                column = column.model_copy(update={"type": kind})
        columns.append(column)
    return sheet.model_copy(update={"columns": columns})


def _number_kind(values: list[Any], decimal_comma: bool | None) -> str | None:
    kinds = set()
    for value in values:
        if isinstance(value, float):
            kinds.add("number")
            continue
        if isinstance(value, str) and value.strip().startswith("="):
            continue  # a formula in a cell says nothing about the type
        text = str(value).strip().replace("\u2212", "-")
        match = NUMBER_RE.fullmatch(text)
        digits = re.sub(r"\D", "", text)
        if not match or parse_number_text(text, decimal_comma) is None or len(digits) > 15:
            return None
        if re.match(r"[-+]?0[0-9]", text):
            return None  # leading zero: a code such as a postal code, not a number
        kinds.add(
            "percent" if match["suffix"] == "%" else "currency" if match["prefix"] or match["suffix"] else "number"
        )
    special = kinds - {"number"}
    if not kinds or len(special) > 1:
        return None
    return special.pop() if special else "number"


def has_totals(sheet: ExcelSheet) -> bool:
    return bool(sheet.rows) and any(column.total != "none" for column in sheet.columns)


NUMBER_TYPES = ("number", "integer", "currency", "percent")


@dataclass
class ChartLayout:
    """1-based columns/rows: series are (title, column or row), categories are a column or a row range."""

    by_rows: bool
    series: list[tuple[str, int]]
    categories: tuple[int, int, int, int]  # min_col, max_col, min_row, max_row


def chart_layout(sheet: ExcelSheet, chart: ExcelChart) -> ChartLayout:
    """Work out what to plot; raises ValueError with a list of messages for the model."""
    headers = [column.header.strip() for column in sheet.columns]
    last_row = len(sheet.rows) + 1
    listed = f"({', '.join(headers)})"
    errors: list[str] = []
    if chart.series_rows:
        labels = [to_plain(str(row[0] if row and row[0] is not None else "")) for row in align_rows(sheet)]
        keys = [label.casefold() for label in labels]
        series = []
        for name in chart.series_rows:
            if to_plain(name).casefold() in keys:
                row = keys.index(to_plain(name).casefold())
                series.append((labels[row], row + 2))
            else:
                errors.append(f"series_rows: '{name}' is not in the first column ({', '.join(labels[:15])})")
        if chart.value_columns:
            missing = [name for name in chart.value_columns if name.strip() not in headers]
            errors += [f"value_columns: '{name}' is not a column header {listed}" for name in missing]
            columns = sorted(headers.index(name.strip()) + 1 for name in chart.value_columns if name not in missing)
        else:
            columns = []
            for c, column in enumerate(sheet.columns[1:], start=2):
                if column.type not in NUMBER_TYPES or column.formula.strip():
                    break
                columns.append(c)
        if columns and columns != list(range(columns[0], columns[-1] + 1)):
            errors.append("value_columns: with series_rows the columns must be next to each other")
        elif not columns and not errors:
            errors.append("value_columns: no number columns found after the first column; list them in value_columns")
        if errors:
            raise ValueError(errors)
        return ChartLayout(True, series, (columns[0], columns[-1], 1, 1))
    if not chart.category_column.strip() and not chart.value_columns:  # only type/title: plot the number columns
        numbers = [h for c, h in enumerate(headers) if c > 0 and sheet.columns[c].type in NUMBER_TYPES]
        chart = chart.model_copy(update={"category_column": headers[0] if headers else "", "value_columns": numbers})
    if chart.category_column.strip() not in headers:
        errors.append(f"category_column: '{chart.category_column}' is not a column header {listed}")
    if not chart.value_columns:
        errors.append("value_columns: is empty (or use series_rows when the data series are rows)")
    errors += [
        f"value_columns: '{name}' is not a column header {listed} - if these are row labels, use series_rows"
        for name in chart.value_columns
        if name.strip() not in headers
    ]
    rows = align_rows(sheet)
    for name in chart.value_columns:
        if name.strip() in headers:
            c = headers.index(name.strip())
            column = sheet.columns[c]
            has_numbers = any(
                isinstance(row[c], float)
                or (
                    isinstance(row[c], str)
                    and (parse_number_text(row[c]) is not None or row[c].strip().startswith("="))
                )
                for row in rows
            )
            if not column.formula.strip() and not has_numbers:
                errors.append(f"value_columns: '{name}' has no numbers to plot")
    if errors:
        raise ValueError(errors)
    category = headers.index(chart.category_column.strip()) + 1
    series = [(name.strip(), headers.index(name.strip()) + 1) for name in chart.value_columns]
    return ChartLayout(False, series, (category, category, 2, last_row))


class XlsxWriter:
    def __init__(self, currency: str, texts: dict[str, str], decimal_comma: bool | None = None) -> None:
        self.currency = currency
        self.texts = texts
        self.decimal_comma = decimal_comma
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)

    def render(self, sheets: list[ExcelSheet], formulas: dict[tuple[int, int], str]) -> bytes:
        for s, (sheet, name) in enumerate(zip(sheets, sheet_names(sheets, self.texts["sheet"]), strict=True)):
            self._sheet(
                self.workbook.create_sheet(title=name), sheet, {c: f for (i, c), f in formulas.items() if i == s}
            )
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()

    def _sheet(self, ws: Worksheet, sheet: ExcelSheet, formulas: dict[int, str]) -> None:
        header_fill = PatternFill("solid", fgColor=HEADER_FILL)
        rows = align_rows(sheet)
        last_row = len(rows) + 1
        widths = [len(column.header.strip()) for column in sheet.columns]
        for c, column in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=1, column=c, value=excel_text(column.header.strip()) or None)
            cell.data_type = "s"
            cell.quotePrefix = str(cell.value or "").startswith(FORMULA_START)
            cell.font = BOLD_FONT
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        columns_values: list[list[Any]] = [[] for _ in sheet.columns]
        for r, row in enumerate(rows, start=2):
            for c, (column, raw) in enumerate(zip(sheet.columns, row, strict=True), start=1):
                if (c - 1) in formulas:
                    ws.cell(row=r, column=c, value=formulas[c - 1].replace("{row}", str(r)))
                    widths[c - 1] = max(widths[c - 1], 12)
                    continue
                if formula := cell_formula(raw, column):  # checked against the whitelist
                    ws.cell(row=r, column=c, value=formula.replace("{row}", str(r)))
                    widths[c - 1] = max(widths[c - 1], 12)
                    continue
                value, typed = convert_value(raw, column.type, self.decimal_comma)
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # other text never becomes a live formula ("=HYPERLINK(...)" stays text)
                    cell.quotePrefix = value.startswith(FORMULA_START)  # ... not even when the user edits the cell
                    if "\n" in value:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    widths[c - 1] = max(widths[c - 1], *(len(line) for line in value.split("\n")))
                elif value is not None:
                    columns_values[c - 1].append(value)
                    widths[c - 1] = max(
                        widths[c - 1], 12 if column.type in ("currency", "date") else len(str(value)) + 2
                    )
        date_format = self.texts["date_format"]
        format_rows = last_row + (1 if has_totals(sheet) else 0)
        for c, column in enumerate(sheet.columns, start=1):
            fmt = number_format(column, columns_values[c - 1], self.currency, date_format)
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, max_row=format_rows, min_col=c, max_col=c):
                    cell.number_format = fmt
            ws.column_dimensions[get_column_letter(c)].width = min(
                max(widths[c - 1] + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{last_row}"
        self._totals(ws, sheet, last_row)
        if sheet.chart:
            self._chart(ws, sheet, sheet.chart, last_row)

    def _totals(self, ws: Worksheet, sheet: ExcelSheet, last_row: int) -> None:
        if not has_totals(sheet):
            return
        row = last_row + 1
        top = Border(top=Side(style="thin"))
        for c, column in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=row, column=c)
            if column.total != "none":
                letter = get_column_letter(c)
                cell.value = f"={TOTAL_FUNCTIONS[column.total]}({letter}2:{letter}{last_row})"
                if column.total == "count":
                    cell.number_format = "0"  # a count, not an amount or a percentage
            elif c == 1:  # no room for the label when the first column has a total itself
                cell.value = self.texts["total"]
                cell.data_type = "s"
            cell.font = BOLD_FONT
            cell.border = top

    def _chart(self, ws: Worksheet, sheet: ExcelSheet, chart: ExcelChart, last_row: int) -> None:
        layout = chart_layout(sheet, chart)
        series = layout.series[:1] if chart.type == "pie" else layout.series
        if chart.type == "pie":
            excel_chart: Any = PieChart()
        elif chart.type == "line":
            excel_chart = LineChart()
        else:
            excel_chart = BarChart()
            excel_chart.type = "bar" if chart.type == "bar" else "col"
        min_col, max_col, min_row, max_row = layout.categories
        for title, index in series:
            if layout.by_rows:
                values = Reference(ws, min_col=min_col, max_col=max_col, min_row=index, max_row=index)
            else:
                values = Reference(ws, min_col=index, min_row=2, max_row=last_row)
            excel_chart.series.append(Series(values, title=excel_text(title)))
        excel_chart.set_categories(Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row, max_row=max_row))
        if chart.title.strip():
            excel_chart.title = excel_text(chart.title.strip())
        if len(series) == 1 and chart.type != "pie":
            excel_chart.legend = None  # a single series needs no legend
        excel_chart.width, excel_chart.height = 18, 9
        ws.add_chart(excel_chart, f"{get_column_letter(len(sheet.columns) + 2)}2")


def build_xlsx(
    sheets: list[ExcelSheet], currency: str, texts: dict[str, str], decimal_comma: bool | None = None
) -> bytes:
    sheets = [infer_column_types(formulas_from_rows(sheet), decimal_comma) for sheet in sheets]
    formulas = check_excel(sheets, decimal_comma)
    return XlsxWriter(currency, texts, decimal_comma).render(sheets, formulas)


# ------------------------------------------------------------------------------------------ PowerPoint --

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
TITLE_TOP, TITLE_HEIGHT = Inches(0.45), Inches(1.0)
BAR_TOP = Inches(1.5)
BODY_TOP = Inches(1.8)
BODY_HEIGHT = SLIDE_HEIGHT - BODY_TOP - Inches(0.6)
COLUMN_GAP = Inches(0.5)
TEXT_GRAY = "595959"
# Layout indices of python-pptx's default template.
LAYOUT_TITLE, LAYOUT_CONTENT, LAYOUT_SECTION, LAYOUT_TWO_CONTENT, LAYOUT_TITLE_ONLY = 0, 1, 2, 3, 5
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def check_slides(title: str, slides: list[Slide]) -> None:
    check = Checker()
    check.text(title)
    check.require(bool(slides), "slides: the presentation has no slides")
    check.require(len(slides) <= MAX_SLIDES, f"slides: at most {MAX_SLIDES} slides")
    for i, slide in enumerate(slides):
        where = f"slides[{slide._index if slide._index >= 0 else i}] ({slide.layout})"
        check.text(slide.title, slide.subtitle, slide.notes, *slide.bullets, *slide.left, *slide.right)
        check.require(
            bool(slide.title.strip() or (slide.layout == "title" and title.strip())), f"{where}: 'title' is empty"
        )
        for field_name in ("bullets", "left", "right"):
            items = getattr(slide, field_name)
            check.require(
                len(items) <= MAX_SLIDE_ITEMS,
                f"{where}.{field_name}: at most {MAX_SLIDE_ITEMS} items per slide, split the content over more slides",
            )
        too_much = f"{where}: too much text for one slide; split it over two slides or move details into the notes"
        if slide.layout == "bullets":
            if check.require(any(b.strip() for b in slide.bullets), f"{where}: 'bullets' is empty"):
                check.require(
                    len(slide.bullets) > MAX_SLIDE_ITEMS or body_font_size(slide.bullets) is not None, too_much
                )
        elif slide.layout == "two_columns":
            if check.require(
                any(b.strip() for b in slide.left + slide.right), f"{where}: 'left' and 'right' are empty"
            ):
                width = (CONTENT_WIDTH - COLUMN_GAP) // 2
                lines = 1 if slide.left_heading.strip() or slide.right_heading.strip() else 0
                check.require(
                    max(len(slide.left), len(slide.right)) > MAX_SLIDE_ITEMS
                    or all(body_font_size(items, 22, width, lines) is not None for items in (slide.left, slide.right)),
                    too_much,
                )
        elif slide.layout == "table":
            if check.require(bool(slide.table), f"{where}: 'table' is empty") and check.require(
                len(slide.table) > 1,
                f"{where}.table: only a header row; put the header and all data rows into one table",
            ):
                check.text(*(cell for row in slide.table for cell in row))
                check.require(
                    len(slide.table) <= MAX_SLIDE_TABLE_ROWS,
                    f"{where}.table: at most {MAX_SLIDE_TABLE_ROWS} rows per slide, split the table over more slides",
                )
                check.require(
                    max(len(row) for row in slide.table) <= MAX_SLIDE_TABLE_COLUMNS,
                    f"{where}.table: at most {MAX_SLIDE_TABLE_COLUMNS} columns",
                )
                check.require(
                    len(slide.table) > MAX_SLIDE_TABLE_ROWS
                    or table_font_size(slide.table, table_columns(slide.table)) is not None,
                    f"{where}.table: the cells hold too much text to fit on one slide; shorten them or split the table",
                )
        elif slide.layout == "chart":
            chart = slide.chart
            if check.require(chart is not None, f"{where}: 'chart' is missing"):
                assert chart is not None
                categories = len(chart.categories)
                check.require(
                    0 < categories <= MAX_CHART_CATEGORIES,
                    f"{where}.chart.categories: 1 to {MAX_CHART_CATEGORIES} labels",
                )
                check.require(
                    0 < len(chart.series) <= MAX_CHART_SERIES, f"{where}.chart.series: 1 to {MAX_CHART_SERIES} series"
                )
                for k, series in enumerate(chart.series):
                    check.require(
                        len(series.values) == categories,
                        f"{where}.chart.series[{k}]: {len(series.values)} values for {categories} categories",
                    )
    check.raise_if_failed()


MIN_BODY_FONT_SIZE = 12
EMU_PER_POINT = 12700


class TextLabelWorkbookWriter(CategoryWorkbookWriter):
    """Writes the chart's embedded workbook with category and series names as plain text.

    python-pptx uses xlsxwriter's write(), which turns '=...' into a formula and URLs into links; Excel
    would evaluate them when the user clicks "Edit Data" in PowerPoint.
    """

    def _write_cat_column(self, worksheet: Any, col: int, level: Any, num_format: Any) -> None:
        worksheet.set_column(col, col, 10)
        for offset, name in level:
            worksheet.write_string(offset + 1, col, str(name), num_format)

    def _write_series(self, workbook: Any, worksheet: Any) -> None:
        col_offset = self._chart_data.categories.depth
        for index, series in enumerate(self._chart_data):
            num_format = workbook.add_format({"num_format": series.number_format})
            worksheet.write_string(0, index + col_offset, str(series.name))
            for row, value in enumerate(series.values, start=1):
                worksheet.write_number(row, index + col_offset, value, num_format)


class TextLabelChartData(CategoryChartData):
    @lazyproperty
    def _workbook_writer(self) -> CategoryWorkbookWriter:
        return TextLabelWorkbookWriter(self)


def body_font_size(items: list[str], base: int = 24, width: int = CONTENT_WIDTH, extra_lines: int = 0) -> int | None:
    """Largest font size (from `base` down) at which the items fit the text box, or None if they do not fit
    even at the smallest size; python-pptx cannot shrink text to fit. Estimates half an em per character,
    1.3 line height and 6 pt before each item; sub items are set smaller (see sub_item_size)."""
    for size in range(base, MIN_BODY_FONT_SIZE - 1, -2):
        lines = extra_lines
        for item in items:
            level, text = split_level(item)
            item_size = sub_item_size(size) if level else size
            chars_per_line = max(int(width / EMU_PER_POINT / (item_size * 0.5)), 1)
            lines += max(1, -(-len(to_plain(text)) // chars_per_line)) * item_size / size
        if lines * size * 1.3 + len(items) * 6 <= BODY_HEIGHT / EMU_PER_POINT:
            return size
    return None


def sub_item_size(size: int) -> int:
    return size - 4 if size >= 18 else size - 2


def table_font_size(rows: list[list[str]], columns: int) -> tuple[int, int] | None:
    """(font size, row height in EMU) at which the table fits below the title, or None. Long cells wrap,
    so each row is as high as its cell with the most lines."""
    column_width = CONTENT_WIDTH / max(columns, 1) / EMU_PER_POINT - 10  # minus the cell margins
    longest_words = [
        max((len(word) for row in rows if c < len(row) for word in to_plain(row[c]).split()), default=1)
        for c in range(columns)
    ]
    for size in (18, 16, 14, 12, 11, 10):
        words_fit = sum(min(n, MAX_WORD_CHARS) * CHAR_WIDTH * size / 11 + CELL_PADDING for n in longest_words)
        if words_fit > CONTENT_WIDTH and size > 10:
            continue  # long words would be broken in the middle
        chars_per_line = max(int(column_width / (size * 0.5)), 1)
        line_counts = [
            max((max(1, -(-len(to_plain(cell)) // chars_per_line)) for cell in row), default=1) for row in rows
        ]
        row_height = int(Inches(0.1) + PptxPt(size * 1.25))
        height = sum(int(Inches(0.1) + PptxPt(size * 1.25 * lines)) for lines in line_counts)
        if height <= BODY_HEIGHT:
            return size, row_height
    return None


def _no_bullet(paragraph: PptxParagraph) -> None:
    paragraph._p.get_or_add_pPr().append(PptxOxmlElement("a:buNone"))


class PptxWriter:
    def __init__(self, title: str, accent: str) -> None:
        self.title = title
        self.accent = PptxRGBColor.from_string(accent)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT

    def render(self, slides: list[Slide]) -> bytes:
        for slide in slides:
            getattr(self, f"_slide_{slide.layout}")(slide)
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()

    # -- slide types --

    def _slide_title(self, data: Slide) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUT_TITLE])
        title, subtitle = slide.placeholders[0], slide.placeholders[1]
        self._place(title, MARGIN, Inches(2.0), CONTENT_WIDTH, Inches(1.8))
        self._text(
            title,
            [data.title.strip() or self.title],
            size=40,
            bold=True,
            color=self.accent,
            anchor=MSO_ANCHOR.BOTTOM,
            align=PP_ALIGN.CENTER,
        )
        self._bar(slide, Inches(3.95), MARGIN + CONTENT_WIDTH // 2 - Inches(0.8), Inches(1.6))
        self._place(subtitle, MARGIN, Inches(4.2), CONTENT_WIDTH, Inches(1.4))
        self._text(
            subtitle,
            [data.subtitle] if data.subtitle.strip() else [],
            size=22,
            color=TEXT_GRAY,
            anchor=MSO_ANCHOR.TOP,
            align=PP_ALIGN.CENTER,
        )
        self._finish(slide, data)

    def _slide_section(self, data: Slide) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUT_SECTION])
        title, subtitle = slide.placeholders[0], slide.placeholders[1]
        self._place(title, MARGIN, Inches(2.4), CONTENT_WIDTH, Inches(1.5))
        self._text(title, [data.title], size=36, bold=True, color=self.accent, anchor=MSO_ANCHOR.BOTTOM)
        self._bar(slide, Inches(4.05), MARGIN, Inches(1.6))
        self._place(subtitle, MARGIN, Inches(4.3), CONTENT_WIDTH, Inches(1.2))
        self._text(
            subtitle, [data.subtitle] if data.subtitle.strip() else [], size=20, color=TEXT_GRAY, anchor=MSO_ANCHOR.TOP
        )
        self._finish(slide, data)

    def _slide_bullets(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_CONTENT, data)
        body = slide.placeholders[1]
        self._place(body, MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_HEIGHT)
        self._items(body.text_frame, data.bullets, body_font_size(data.bullets) or MIN_BODY_FONT_SIZE)
        self._finish(slide, data)

    def _slide_two_columns(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TWO_CONTENT, data)
        width = (CONTENT_WIDTH - COLUMN_GAP) // 2
        heading_lines = 1 if data.left_heading.strip() or data.right_heading.strip() else 0
        size = min(
            body_font_size(items, base=22, width=width, extra_lines=heading_lines) or MIN_BODY_FONT_SIZE
            for items in (data.left, data.right)
        )
        for placeholder, left, heading, items in (
            (slide.placeholders[1], MARGIN, data.left_heading, data.left),
            (slide.placeholders[2], MARGIN + width + COLUMN_GAP, data.right_heading, data.right),
        ):
            self._place(placeholder, left, BODY_TOP, width, BODY_HEIGHT)
            self._items(placeholder.text_frame, items, size, heading=heading)
        self._finish(slide, data)

    def _slide_table(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TITLE_ONLY, data)
        rows, columns = len(data.table), table_columns(data.table)
        size, row_height = table_font_size(data.table, columns) or (10, int(Inches(0.1) + PptxPt(12.5)))
        shape = slide.shapes.add_table(rows, columns, MARGIN, BODY_TOP, CONTENT_WIDTH, row_height * rows)
        widths = column_widths(data.table, columns, CONTENT_WIDTH, size)
        for column, width in zip(shape.table.columns, widths, strict=True):
            column.width = Emu(width)  # by content, so long words are not broken
        for r, row in enumerate(data.table):
            for c in range(columns):
                cell = shape.table.cell(r, c)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                frame = cell.text_frame
                frame.word_wrap = True
                self._runs(frame.paragraphs[0], parse_inline(row[c] if c < len(row) else ""), size, bold=r == 0 or None)
        self._finish(slide, data)

    def _slide_chart(self, data: Slide) -> None:
        slide = self._content_slide(LAYOUT_TITLE_ONLY, data)
        assert data.chart is not None
        chart_data = TextLabelChartData()
        chart_data.categories = [clean_text(to_plain(c)) for c in data.chart.categories]
        series_list = data.chart.series[:1] if data.chart.type == "pie" else data.chart.series
        for series in series_list:
            chart_data.add_series(clean_text(to_plain(series.name)), series.values)
        frame = slide.shapes.add_chart(
            CHART_TYPES[data.chart.type], MARGIN, BODY_TOP, CONTENT_WIDTH, BODY_HEIGHT, chart_data
        )
        chart = frame.chart
        chart.has_title = False  # the slide title says it
        chart.font.size = PptxPt(14)
        if data.chart.unit.strip() and data.chart.type != "pie":
            axis = chart.value_axis
            axis.has_title = True
            axis.axis_title.text_frame.text = clean_text(to_plain(data.chart.unit))[:60]
        chart.has_legend = data.chart.type == "pie" or len(series_list) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if data.chart.type == "pie":
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_value = False
            plot.data_labels.number_format = "0%"
            plot.data_labels.number_format_is_linked = False
        self._finish(slide, data)

    # -- helpers --

    def _content_slide(self, layout: int, data: Slide) -> Any:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        title = slide.placeholders[0]
        self._place(title, MARGIN, TITLE_TOP, CONTENT_WIDTH, TITLE_HEIGHT)
        self._text(
            title,
            [data.title],
            size=30 if len(data.title) <= 60 else 24,
            bold=True,
            color=self.accent,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        self._bar(slide, BAR_TOP, MARGIN, Inches(1.2))
        return slide

    def _finish(self, slide: Any, data: Slide) -> None:
        for placeholder in list(slide.placeholders):  # remove placeholders left empty ("Click to add text")
            if not placeholder.has_text_frame or not placeholder.text_frame.text.strip():
                placeholder._element.getparent().remove(placeholder._element)
        if data.notes.strip():
            slide.notes_slide.notes_text_frame.text = clean_text(data.notes.strip())

    @staticmethod
    def _place(shape: SlidePlaceholder, left: int, top: int, width: int, height: int) -> None:
        shape.left, shape.top, shape.width, shape.height = Emu(left), Emu(top), Emu(width), Emu(height)

    def _bar(self, slide: Any, top: int, left: int, width: int) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.accent
        bar.line.fill.background()

    def _text(self, shape: SlidePlaceholder, lines: list[str], size: int, bold: bool | None = None,
              color: Any = None, anchor: Any = None, align: Any = PP_ALIGN.LEFT) -> None:  # fmt: skip
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        if anchor is not None:
            frame.vertical_anchor = anchor
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.alignment = align
            self._runs(paragraph, parse_inline(line.strip()), size, bold=bold, color=color)

    def _items(self, frame: TextFrame, items: list[str], size: int, heading: str = "") -> None:
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        first = True
        if heading.strip():
            paragraph = frame.paragraphs[0]
            _no_bullet(paragraph)
            self._runs(paragraph, parse_inline(heading.strip()), size + 2, bold=True, color=self.accent)
            first = False
        for item in items:
            level, text = split_level(item)
            if not text:
                continue
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.level = level
            self._runs(paragraph, parse_inline(text), sub_item_size(size) if level else size)
            paragraph.space_before = PptxPt(6 if level == 0 else 2)

    @staticmethod
    def _runs(
        paragraph: PptxParagraph, runs: list[Run], size: int, bold: bool | None = None, color: Any = None
    ) -> None:
        for run in runs:
            pptx_run = paragraph.add_run()
            pptx_run.text = run.text
            pptx_run.font._rPr.set("cap", "none")
            font = pptx_run.font
            font.size = PptxPt(size)
            font.bold = True if (bold or run.bold) else None
            font.italic = True if run.italic else None
            if run.code:
                font.name = CODE_FONT
            if color is not None:
                font.color.rgb = color if isinstance(color, PptxRGBColor) else PptxRGBColor.from_string(color)
            if run.url:
                pptx_run.hyperlink.address = run.url


def build_pptx(title: str, slides: list[Slide], accent: str) -> bytes:
    check_slides(title, slides)
    return PptxWriter(title, accent).render(slides)


# ------------------------------------------------------------------------------------------- delivery --


def download_js(data: bytes, filename: str, mime: str) -> str:
    """Browser snippet that downloads the file; all dynamic values are JSON-encoded."""
    return f"""
const bytes = Uint8Array.from(atob({json.dumps(base64.b64encode(data).decode())}), (c) => c.charCodeAt(0));
const url = URL.createObjectURL(new Blob([bytes], {{ type: {json.dumps(mime)} }}));
const link = document.createElement('a');
link.href = url;
link.download = {json.dumps(filename)};
document.body.appendChild(link);
link.click();
link.remove();
setTimeout(() => URL.revokeObjectURL(url), 10000);
return true;
"""


@dataclass
class Context:
    """What Open WebUI passes to a tool call."""

    user: dict[str, Any] | None
    request: Any
    chat_id: str | None
    message_id: str | None
    emitter: EventCallback | None
    caller: EventCallback | None

    async def status(self, description: str, done: bool = False) -> None:
        if self.emitter is not None:
            await self.emitter({"type": "status", "data": {"description": description, "done": done}})


async def attach_file(data: bytes, filename: str, mime: str, ctx: Context) -> None:
    """Store the file in Open WebUI's file storage for the user and attach it to the current answer.

    Uses the same internal functions as Open WebUI's built-in image generation (0.11.4).
    """
    from open_webui.models.chats import Chats
    from open_webui.models.users import UserModel
    from open_webui.routers.files import upload_file_handler
    from open_webui.utils.chat_id import is_saved_chat_id
    from starlette.datastructures import Headers, UploadFile

    if not ctx.user or ctx.request is None or ctx.emitter is None:
        raise RuntimeError("no user, request or event emitter context")
    user = UserModel(**ctx.user)
    upload = UploadFile(file=io.BytesIO(data), filename=filename, headers=Headers({"content-type": mime}))
    file_item = await upload_file_handler(
        ctx.request,
        file=upload,
        metadata={"chat_id": ctx.chat_id, "message_id": ctx.message_id},
        process=False,
        user=user,
    )
    file_id = file_item.id
    try:
        url = ctx.request.app.url_path_for("get_file_content_by_id", id=file_id)
    except Exception:
        url = f"/api/v1/files/{file_id}/content"
    entry = {"type": "file", "id": file_id, "url": str(url), "name": filename, "size": len(data), "content_type": mime}
    files = [entry]
    if ctx.message_id and is_saved_chat_id(ctx.chat_id):
        try:  # the file is stored now; if linking it to the chat fails, the event below still attaches it
            await Chats.insert_chat_files(
                chat_id=ctx.chat_id, message_id=ctx.message_id, file_ids=[file_id], user_id=user.id
            )
            files = await Chats.add_message_files_by_id_and_message_id(ctx.chat_id, ctx.message_id, [entry]) or files
        except Exception:
            log.exception("Linking %s to the chat failed", filename)
    await ctx.emitter({"type": "chat:message:files", "data": {"files": files}})


DOWNLOAD_TIMEOUT = 60  # seconds; Open WebUI waits forever if the user has switched to another chat


async def download_file(data: bytes, filename: str, mime: str, ctx: Context) -> None:
    """Hand the file to the browser. Open WebUI does not raise when the browser is gone: it returns an error
    dict, or never answers, so anything but the snippet's `true` counts as failure."""
    if ctx.caller is None:
        raise RuntimeError("__event_call__ is not available, cannot download in the browser")
    result = await asyncio.wait_for(
        ctx.caller({"type": "execute", "data": {"code": download_js(data, filename, mime)}}), DOWNLOAD_TIMEOUT
    )
    if result is not True:
        raise RuntimeError(f"the browser did not confirm the download: {str(result)[:200]}")


# ------------------------------------------------------------------------------------------------ Tools --


class Tools:
    class Valves(BaseModel):
        language: Literal["en", "de"] = Field(
            default="en",
            description="Language of status messages and of texts inside the files (total row, page footer).",
        )
        delivery: Literal["attach", "download"] = Field(
            default="attach",
            description="attach: store the file in Open WebUI and attach it to the answer (preview, stays in the chat). "
            "download: send it straight to the browser as a download (nothing stored, gone after reload). "
            "If attaching fails, the plugin falls back to download.",
        )
        currency: Literal["EUR", "USD", "GBP", "CHF"] = Field(
            default="EUR", description="Currency of Excel columns of type 'currency'."
        )
        accent_color: str = Field(
            default="1F4E79", description="Accent colour of PowerPoint titles as a hex value without '#', e.g. 1F4E79."
        )
        page_numbers: bool = Field(default=True, description="Add page numbers to the footer of Word documents.")

    def __init__(self) -> None:
        self.valves = self.Valves()

    async def create_word_file(
        self,
        blocks: list[WordBlock],
        title: str = "",
        subtitle: str = "",
        header_text: str = "",
        table_of_contents: bool = False,
        landscape: bool = False,
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create a Word document (.docx) and attach it to the answer. Use it when the user wants a Word file or a
        document to download, e.g. a report, letter, minutes or concept. Write the complete text yourself.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param blocks: The content in reading order, e.g. [{"type": "heading", "text": "Introduction", "level": 1}, {"type": "paragraph", "text": "..."}, {"type": "bullets", "items": ["...", "- sub item"]}, {"type": "table", "rows": [["Header A", "Header B"], ["a", "b"]]}].
        :param title: Readable document title with spaces, e.g. "Minutes of the team meeting on 12 October 2026". It is shown at the top and used as the file name, so do not repeat it as the first heading.
        :param subtitle: Optional subtitle below the title.
        :param header_text: Optional text in the page header, e.g. company or project name.
        :param table_of_contents: Add a table of contents (useful for long documents with several headings).
        :param landscape: Use landscape pages (for wide tables).
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]

        def build() -> tuple[str, bytes]:
            extras: dict[str, Any] = {}  # parameters the model wrote into the blocks text
            parsed = parse_argument(WORD_BLOCKS, blocks, "blocks", normalize_word_block, extras=extras, rows_key="rows")
            parsed = [block for block in parsed if not is_empty_block(block)]
            given = {k: _unescape_newlines(str(v or extras.get(k) or "")) for k, v in
                     (("title", title), ("subtitle", subtitle), ("header_text", header_text))}  # fmt: skip
            display_title, parsed = word_title(given["title"], parsed, texts["default_title"])
            options = WordOptions(
                title=display_title,
                subtitle=given["subtitle"],
                header_text=given["header_text"],
                table_of_contents=bool(table_of_contents or extras.get("table_of_contents") is True),
                landscape=bool(landscape or extras.get("landscape") is True),
                page_numbers=self.valves.page_numbers,
            )
            return display_title, build_docx(options, parsed, texts)

        return await self._create("create_word_file", "docx", DOCX_MIME, build, texts["status_word"], ctx)

    async def create_excel_file(
        self,
        sheets: list[ExcelSheet],
        title: str = "",
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create an Excel workbook (.xlsx) and attach it to the answer. Use it when the user wants an Excel file or a
        spreadsheet, e.g. a list, budget, plan or calculation. Numbers are real numbers; formulas, total rows and
        charts are possible.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param sheets: One or more sheets, e.g. [{"name": "Budget", "columns": [{"header": "Item"}, {"header": "Amount", "type": "currency", "total": "sum"}], "rows": [["Rent", 1200], ["Travel", 350.5]]}].
        :param title: Readable title of the workbook with spaces, used as the file name.
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]

        def build() -> tuple[str, bytes]:
            extras: dict[str, Any] = {}
            parsed = parse_argument(EXCEL_SHEETS, sheets, "sheets", normalize_sheet, move_lonely_charts, extras)
            given = str(title or extras.get("title") or "")
            display_title = tidy_title(given) or (tidy_title(parsed[0].name) if parsed else "")
            decimal_comma = self.valves.language == "de"
            return display_title, build_xlsx(parsed, self.valves.currency, texts, decimal_comma)

        return await self._create("create_excel_file", "xlsx", XLSX_MIME, build, texts["status_excel"], ctx)

    async def create_powerpoint_file(
        self,
        slides: list[Slide],
        title: str = "",
        __user__: dict | None = None,
        __request__: Any = None,
        __chat_id__: str | None = None,
        __message_id__: str | None = None,
        __event_emitter__: Any = None,
        __event_call__: Any = None,
    ) -> str:
        """
        Create a PowerPoint presentation (.pptx, 16:9) and attach it to the answer. Use it when the user wants
        slides or a presentation. Start with a 'title' slide, keep bullets short (at most 6 per slide), and put
        details into the speaker notes.
        The file only exists once this tool has answered with 'Done'; never say it was created otherwise.
        :param slides: The slides in order, e.g. [{"layout": "title", "title": "...", "subtitle": "..."}, {"layout": "bullets", "title": "...", "bullets": ["...", "- sub point"], "notes": "..."}, {"layout": "chart", "title": "...", "chart": {"type": "column", "categories": ["Q1", "Q2"], "series": [{"name": "Sales", "values": [10, 12]}]}}].
        :param title: Readable title of the presentation with spaces, used as the file name.
        """
        ctx = Context(__user__, __request__, __chat_id__, __message_id__, __event_emitter__, __event_call__)
        texts = TRANSLATIONS[self.valves.language]
        accent = self.valves.accent_color.strip().lstrip("#")
        if not re.fullmatch(r"[0-9A-Fa-f]{6}", accent):
            accent = Tools.Valves().accent_color

        def build() -> tuple[str, bytes]:
            extras: dict[str, Any] = {}
            parsed = parse_argument(SLIDES, slides, "slides", normalize_slide, extras=extras, rows_key="table")
            display_title = tidy_title(str(title or extras.get("title") or "")) or next(
                (tidy_title(s.title) for s in parsed if s.title.strip()), ""
            )
            return display_title, build_pptx(display_title, parsed, accent)

        return await self._create("create_powerpoint_file", "pptx", PPTX_MIME, build, texts["status_powerpoint"], ctx)

    async def _create(
        self, tool: str, extension: str, mime: str, build: Callable[[], tuple[str, bytes]], status: str, ctx: Context
    ) -> str:
        texts = TRANSLATIONS[self.valves.language]
        await ctx.status(status)
        try:
            title, data = await asyncio.to_thread(build)
        except RecursionError:
            await ctx.status(texts["status_failed"], done=True)
            return RESULT_INVALID.format(errors="- the arguments are nested too deeply", tool=tool)
        except InvalidInput as e:
            await ctx.status(texts["status_failed"], done=True)
            errors = e.errors[:MAX_ERRORS_SHOWN] + (["..."] if len(e.errors) > MAX_ERRORS_SHOWN else [])
            return RESULT_INVALID.format(errors="\n".join(f"- {error}" for error in errors), tool=tool)
        except Exception:
            log.exception("Creating a .%s file failed", extension)
            await ctx.status(texts["status_failed"], done=True)
            return RESULT_FAILED
        filename = f"{safe_filename(title, texts['default_title'])}.{extension}"
        result = await self._deliver(data, filename, mime, ctx)
        await ctx.status(texts["status_done"].format(name=filename) if result else texts["status_failed"], done=True)
        return result.format(name=filename) if result else RESULT_FAILED

    async def _deliver(self, data: bytes, filename: str, mime: str, ctx: Context) -> str:
        """Attach or download the file; returns the result template, or '' if both ways failed."""
        if self.valves.delivery == "attach":
            try:
                await attach_file(data, filename, mime, ctx)
                return RESULT_ATTACHED
            except Exception:
                log.exception("Attaching %s failed, falling back to a browser download", filename)
        try:
            await download_file(data, filename, mime, ctx)
            return RESULT_DOWNLOADED
        except Exception:
            log.exception("Downloading %s failed", filename)
            return ""
