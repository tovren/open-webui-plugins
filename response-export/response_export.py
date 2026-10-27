"""
title: Response Export
author: tovren
author_url: https://github.com/tovren/open-webui-plugins
version: 1.3.2
license: MIT
required_open_webui_version: 0.12.4
description: Adds four buttons below every assistant message - export as Word (.docx), export its tables as Excel (.xlsx), download as PDF, and open as an e-mail draft. Files are built in memory and handed straight to the browser; nothing is stored on the server.

Buttons
- Word (.docx) with headings, lists (one nesting level), tables, bold/italic/code, clickable links,
  code blocks, quotes and horizontal rules, in the design of Office Documents (accent_color and
  font valves) and optionally on a company template such as a letterhead (word_template valve).
- Excel (.xlsx) with one sheet per Markdown table (named after a heading directly above it), a
  styled and frozen header row, auto filter, fitted column widths, clickable links, real dates and
  real numbers (also German and Swiss formats such as "1.234,56 EUR", "12,5 %" or "CHF 1'234.50");
  phone numbers such as "+49 711 123456" stay text.
- PDF (.pdf) with the same content and design as the Word file (accent_color valve), page numbers
  and an optional small footer. The plugin writes the PDF itself with the standard PDF fonts
  (Helvetica, Courier), so no PDF library is needed.
- E-mail via a mailto link; answers too long for a mailto link are downloaded as an .eml draft
  (marked "X-Unsent", so Outlook opens it as a new editable message).
Reasoning and tool-call blocks are removed before exporting. If the answer cites sources ([1]), Word,
PDF and e-mail end with the list of its sources.

Security & privacy
- No network access, no shell, nothing is written to disk or stored on the server. Only the
  standard library plus python-docx and openpyxl (both bundled with Open WebUI) are used.
- The browser receives the generated file (base64) or a mailto link, wrapped in a fixed JavaScript
  snippet run via the "execute" event. All dynamic values are JSON-encoded.
- With language 'auto' or an empty timezone, a fixed snippet first reads the interface language
  and time zone of the browser; nothing else is read from the page.
- The page's Content-Security-Policy must allow 'unsafe-eval' for the "execute" event to work. If
  the browser does not run the snippet (or does not answer within 30 seconds), an error is shown.
- The word_template valve reads one Word file from the server (at most 20 MB); its text is removed,
  only header, footer, page size and margins are kept.
- Raw HTML in the response is always escaped, links are limited to http, https, mailto and tel
  in the e-mail and to http, https and mailto in Word, PDF and Excel, and images are never loaded
  (web images become links, embedded ones their description). Spreadsheet cells are written as
  text, numbers or dates, never as formulas.
"""

import asyncio
import base64
import copy
import html
import io
import json
import logging
import re
import unicodedata
import zipfile
import zlib
from collections.abc import Awaitable, Callable, Iterator
from dataclasses import dataclass, field, replace
from datetime import date, datetime, tzinfo
from email.message import EmailMessage
from typing import TYPE_CHECKING, Any, Literal
from urllib.parse import quote, unquote
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import CONTENT_TYPE as DOCX_CT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

if TYPE_CHECKING:
    from docx.styles.style import ParagraphStyle

log = logging.getLogger(__name__)

EventCallback = Callable[[dict[str, Any]], Awaitable[Any]]

# ------------------------------------------------------------------------------------------ constants --

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EML_MIME = "message/rfc822"
PDF_MIME = "application/pdf"

BODY_FONT = "Calibri"  # e-mail, and the default of the font valve (Word)
BODY_FONT_SIZE = Pt(11)
CODE_FONT = "Consolas"
CODE_FONT_SIZE = Pt(9.5)
MAX_DOCX_HEADING_LEVEL = 4
ORDERED_LIST_STYLES = ("List Number", "List Number 2")
BULLET_LIST_STYLES = ("List Bullet", "List Bullet 2")
TABLE_ALIGNMENTS = {  # (starts with ":", ends with ":") -> alignment
    (True, True): "center",
    (False, True): "right",
    (True, False): "left",
}
DOCX_ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

HEADER_FILL = "DDE4EE"
MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
MAX_EXACT_DIGITS = 15  # Excel keeps 15 significant digits; longer numbers stay text
EURO_FORMAT = '#,##0.00 "€"'
DOLLAR_FORMAT = '"$"#,##0.00'
POUND_FORMAT = '"£"#,##0.00'
FRANC_FORMAT = '"CHF" #,##0.00'
CURRENCY_FORMATS = {  # symbol or ISO code before or after the amount -> Excel number format
    "€": EURO_FORMAT,
    "EUR": EURO_FORMAT,
    "$": DOLLAR_FORMAT,
    "USD": DOLLAR_FORMAT,
    "£": POUND_FORMAT,
    "GBP": POUND_FORMAT,
    "CHF": FRANC_FORMAT,
}

MAX_FILENAME_CHARS = 60
SAFE_LINK_PREFIXES = ("http://", "https://", "mailto:", "tel:")
WORD_LINK_PREFIXES = ("http://", "https://", "mailto:")  # as in Office Documents
MAILTO_FIELDS = {"subject", "body", "cc", "bcc"}  # others, e.g. "attach", are acted on by some mail clients
BROWSER_TIMEOUT_S = (
    30  # Open WebUI waits for the browser without a limit unless WEBSOCKET_EVENT_CALLER_TIMEOUT
)
EMAIL_HTML = (  # format template with a {body} placeholder
    '<!DOCTYPE html><html><head><meta charset="utf-8"></head><body>'
    f'<div style="font-family:{BODY_FONT},Arial,sans-serif;font-size:{BODY_FONT_SIZE.pt:g}pt">{{body}}</div>'
    "</body></html>"
)

# Button labels are language-neutral on purpose: Open WebUI builds the button list without loading
# the valves, so translated labels would only switch after the first click.

TRANSLATIONS: dict[str, dict[str, str]] = {
    "en": {
        "empty_message": "This response is empty, so there is nothing to export.",
        "no_tables": "This response does not contain a table.",
        "eml_downloaded": (
            "This response is too long to hand over to your mail program directly, so an e-mail draft "
            "(.eml) was downloaded. Open the file to edit and send it - Outlook shows it as a new message."
        ),
        "export_failed": "The export failed. Please try again or contact your administrator.",
        "browser_failed": (
            "Your browser did not run the export. Please reload the page and try again. If it keeps "
            "failing, contact your administrator: the page's Content Security Policy has to allow "
            "'unsafe-eval' for Open WebUI's code events."
        ),
        "default_title": "AI response",
        "email_draft_name": "E-mail draft",
        "default_subject": "Draft",
        "sheet_name": "Table {n}",
        "sources": "Sources",
        "image": "[Image]",
        "print_footer": "Created with the AI assistant on {date}. Please check the content before use.",
        "page_number": "Page {page} of {pages}",
        "date_format": "%Y-%m-%d",
        "excel_date_format": "yyyy-mm-dd",
        "decimal_mark": ".",  # how Excel reads "1.250" if the column does not show it
        "paper": "letter",
    },
    "de": {
        "empty_message": "Diese Antwort ist leer, es gibt nichts zu exportieren.",
        "no_tables": "Diese Antwort enthält keine Tabelle.",
        "eml_downloaded": (
            "Die Antwort ist zu lang für die direkte Übergabe an das Mailprogramm. Deshalb wurde ein "
            "E-Mail-Entwurf (.eml) heruntergeladen – Datei öffnen, dann erscheint sie in Outlook als "
            "neue E-Mail."
        ),
        "export_failed": "Der Export ist fehlgeschlagen. Bitte erneut versuchen oder an die IT wenden.",
        "browser_failed": (
            "Der Browser hat den Export nicht ausgeführt. Bitte die Seite neu laden und erneut versuchen. "
            "Klappt es weiterhin nicht, bitte an die IT wenden: Die Content-Security-Policy der Seite muss "
            "'unsafe-eval' für die Code-Events von Open WebUI erlauben."
        ),
        "default_title": "KI-Antwort",
        "email_draft_name": "E-Mail-Entwurf",
        "default_subject": "Entwurf",
        "sheet_name": "Tabelle {n}",
        "sources": "Quellen",
        "image": "[Bild]",
        "print_footer": "Erstellt mit dem KI-Assistenten am {date} – Inhalte bitte prüfen.",
        "page_number": "Seite {page} von {pages}",
        "date_format": "%d.%m.%Y",
        "excel_date_format": "dd.mm.yyyy",
        "decimal_mark": ",",
        "paper": "a4",
    },
}


# Open WebUI shows action icons as <img>, where currentColor is black, and inverts SVG icons by 80% in
# dark mode. This grey ends up close to its own buttons in both modes (gray-600 light, gray-500 dark),
# and the stroke matches their 1.5 in a 28-unit box.
ICON_COLOR = "#5c5c5c"
ICON_STROKE = 1.3


def _svg_icon(paths: str) -> str:
    """Return a stroke-style 24x24 SVG icon as a base64 data URI."""
    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="{ICON_COLOR}" '
        f'stroke-width="{ICON_STROKE}" stroke-linecap="round" stroke-linejoin="round">{paths}</svg>'
    )
    return "data:image/svg+xml;base64," + base64.b64encode(svg.encode()).decode()


# action id -> (button label, icon); each id is handled by Action._export_<id>
ACTIONS = {
    "docx": (
        "Word (.docx)",
        _svg_icon(
            '<path d="M14 3H7a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h10a2 2 0 0 0 2-2V8z"/><path d="M14 3v5h5"/>'
            '<path d="M8.5 12l1.2 5 1.3-3.5 1.3 3.5 1.2-5"/>'
        ),
    ),
    "xlsx": (
        "Excel (.xlsx)",
        _svg_icon('<rect x="4" y="4" width="16" height="16" rx="2"/><path d="M4 9.5h16M4 14.5h16M10 4v16"/>'),
    ),
    "pdf": (
        "PDF",
        _svg_icon(  # a document with the letters PDF below it
            '<path d="M5 11V5a2 2 0 0 1 2-2h7l5 5v3"/><path d="M14 3v5h5"/>'
            '<path d="M4.5 21v-6.5h1.75a1.75 1.75 0 0 1 0 3.5H4.5"/>'
            '<path d="M10.25 14.5v6.5h1a2.25 2.25 0 0 0 2.25-2.25v-2a2.25 2.25 0 0 0-2.25-2.25z"/>'
            '<path d="M19.5 14.5h-3v6.5M16.5 17.75h2.5"/>'
        ),
    ),
    "email": (
        "E-Mail",
        _svg_icon('<rect x="3" y="5" width="18" height="14" rx="2"/><path d="M3.5 6.5l8.5 6.5 8.5-6.5"/>'),
    ),
}

# ---------------------------------------------------------------------------------- message content --

XML_ILLEGAL_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")  # also lone surrogates
# Open WebUI's own blocks in a message: reasoning, tool calls and code interpreter runs (processDetails in
# its utils/index.ts). Other <details> are the model's, e.g. a collapsible solution.
HIDDEN_DETAILS_TYPES = {"reasoning", "tool_calls", "code_interpreter"}
TAG_PATTERN = r"<(?P<close>/)?(?P<tag>details|think|thinking)\b(?P<attrs>[^>]*)>"
TAG_RE = re.compile(TAG_PATTERN, re.IGNORECASE)
LINK_TARGET = r"(?:[^()\s]|\([^()\s]*\))+"  # one level of parentheses, as in "a_(b)"
CONTENT_RE = re.compile(  # what clean_markdown looks at: code (left as it is), an image or a tag
    r"^[ \t]*(?P<fence>(?P<char>[`~])(?P=char){2,}).*?(?:^[ \t]*(?P=fence)(?P=char)*[ \t]*$|\Z)"
    r"|`(?:[^`\n]|\n(?![ \t]*\n))+`"  # inline code, which cannot span a blank line
    rf"|!\[(?P<alt>[^\]]*)\]\((?P<src>{LINK_TARGET})(?:\s+\"[^\"]*\")?\)"
    rf"|{TAG_PATTERN}",
    re.MULTILINE | re.DOTALL | re.IGNORECASE,
)
DETAILS_TYPE_RE = re.compile(r"(?:^|\s)type\s*=\s*[\"']?(?P<type>[\w-]+)", re.IGNORECASE)
SUMMARY_RE = re.compile(r"\s*<summary\b[^>]*>(?P<summary>.*?)</summary\s*>", re.DOTALL | re.IGNORECASE)
# Characters that start Markdown formatting; "]" and "()" only matter after "[", and "\(", "\[" would start
# a formula
MARKDOWN_SPECIAL_RE = re.compile(r"([\\`*_{}\[#+\-.!|~<>$])")
WEB_PREFIXES = ("http://", "https://")


def find_message(body: dict[str, Any]) -> dict[str, Any]:
    """Return the clicked message, falling back to the last assistant message."""
    messages = [m for m in body.get("messages") or [] if isinstance(m, dict)]
    clicked = next((m for m in messages if m.get("id") == body.get("id")), None)
    fallback = next((m for m in reversed(messages) if m.get("role") == "assistant"), None)
    return clicked or fallback or {}


def message_text(message: dict[str, Any]) -> str:
    """Return the text of a message whose content is a string or a list of content parts."""
    content = message.get("content")
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "\n\n".join(
            part.get("text") or ""
            for part in content
            if isinstance(part, dict) and part.get("type") == "text"
        )
    return ""


def message_sources(message: dict[str, Any]) -> list[tuple[str, str | None]]:
    """(name, web address or None) of the sources Open WebUI sends with a message, numbered like its citation
    markers [1], [2] ...: one per distinct document source, as in its Citations.svelte."""
    found: dict[str, tuple[str, str | None]] = {}
    for source in message.get("sources") or []:
        if not isinstance(source, dict):
            continue
        origin = source.get("source") if isinstance(source.get("source"), dict) else {}
        documents, metadata = source.get("document"), source.get("metadata")
        for index in range(len(documents) if isinstance(documents, list) else 0):
            meta = metadata[index] if isinstance(metadata, list) and index < len(metadata) else None
            meta = meta if isinstance(meta, dict) else {}
            key = str(next((v for v in (meta.get("source"), origin.get("id")) if v is not None), "N/A"))
            name, url = meta.get("name") or origin.get("name"), origin.get("url")
            if key.startswith(WEB_PREFIXES):
                name, url = name if meta.get("external") else key, key
            url = url if isinstance(url, str) and url.lower().startswith(WEB_PREFIXES) else None
            found.setdefault(key, (unquote(str(name or key)), url))
    return list(found.values())


def escape_markdown(text: str) -> str:
    return MARKDOWN_SPECIAL_RE.sub(r"\\\1", text)


def _is_hidden_details(tag: re.Match[str]) -> bool:
    kind = DETAILS_TYPE_RE.search(tag["attrs"]) if tag["tag"].lower() == "details" else None
    return bool(kind and kind["type"].lower() in HIDDEN_DETAILS_TYPES)


def _drop_hidden_details(text: str) -> str:
    """Remove Open WebUI's own <details> blocks (also nested ones), also where they look like code: it writes
    them itself, with escaped attributes, so a backtick in the text before one must not keep it. Each ends at
    the </details> that closes it, counting all <details> tags in between; one never closed stays."""
    parts: list[str] = []
    position = 0
    for opening in TAG_RE.finditer(text):
        if opening.start() < position or not _is_hidden_details(opening):  # removed already, or kept
            continue
        depth = 0
        for tag in TAG_RE.finditer(text, opening.start()):
            if tag["tag"].lower() == "details":
                depth += -1 if tag["close"] else 1
                if not depth:
                    parts.append(text[position : opening.start()])
                    position = tag.end()
                    break
    parts.append(text[position:])
    return "".join(parts)


def _tag_pairs(text: str) -> dict[int, tuple[int, int]]:
    """Start of each opening tag outside code -> (start, end) of its closing tag, matched by nesting per tag
    name; tags in code do not count."""
    pairs: dict[int, tuple[int, int]] = {}
    open_tags: dict[str, list[int]] = {}
    for match in CONTENT_RE.finditer(text):
        if match["tag"] is None:  # code or an image
            continue
        starts = open_tags.setdefault(match["tag"].lower(), [])
        if not match["close"]:
            starts.append(match.start())
        elif starts:
            pairs[starts.pop()] = match.span()
    return pairs


def clean_markdown(text: str, image_label: str = "[Image]") -> str:
    """Prepare a message for parsing; code stays as it is:
    - Open WebUI's reasoning and tool-call blocks and <think> blocks (also nested ones) are dropped,
    - other <details> become their summary (bold) and text,
    - images that are not on the web (embedded data, files of the chat) become their description or
      `image_label`; web images stay, they become links,
    - characters invalid in XML (docx/xlsx) and lone surrogates are dropped."""
    text = _drop_hidden_details(XML_ILLEGAL_RE.sub("", text))
    pairs = _tag_pairs(text)
    kept_ends: set[int] = set()  # closing tags of the model's <details>
    parts: list[str] = []
    position = 0
    while match := CONTENT_RE.search(text, position):
        parts.append(text[position : match.start()])
        position = match.end()
        if match["src"] is not None:
            web = match["src"].lower().startswith(WEB_PREFIXES)
            parts.append(match[0] if web else escape_markdown(" ".join(match["alt"].split()) or image_label))
        elif match["tag"] is None:  # code
            parts.append(match[0])
        elif match["close"]:
            parts.append("\n\n" if match.start() in kept_ends else match[0])
        elif match.start() not in pairs:  # never closed: text
            parts.append(match[0])
        elif match["tag"].lower() != "details" or _is_hidden_details(match):
            position = pairs[match.start()][1]
        else:
            end = pairs[match.start()][0]
            kept_ends.add(end)
            summary, label = SUMMARY_RE.match(text, position), ""
            if summary and summary.end() <= end:
                position, label = summary.end(), to_plain(summary["summary"])
            parts.append(f"\n\n**{escape_markdown(label)}**\n\n" if label else "\n\n")
    parts.append(text[position:])
    return "".join(parts).strip()


# ----------------------------------------------------------------------------------- inline markdown --

INLINE_RE = re.compile(
    # LaTeX as Open WebUI shows it: "$5 and $10" is no formula (no space inside the $, no digit after it)
    r"(?P<math>\$\$.+?\$\$|\\\[.+?\\\]|\\\(.+?\\\)|\$(?![\s$])[^$\n]*?(?<!\s)\$(?!\d))"
    r"|\\(?P<escaped>[\\`*_{}\[\]()#+\-.!|~<>$])"
    r"|`(?P<code>[^`]+)`"
    r"|(?P<br><br\s*/?>)"
    rf"|!?\[(?P<label>[^\]]*)\]\((?P<url>{LINK_TARGET})(?:\s+\"[^\"]*\")?\)"
    r"|<(?P<autolink>(?:https?://|mailto:)[^>\s]+)>"
    r"|\*\*\*(?P<bold_italic>.+?)\*\*\*"
    r"|\*\*(?P<bold>(?:[^*]|\*[^*\s][^*]*\*|\*)+?)\*\*"  # may contain *italic*
    r"|(?<!\w)__(?P<bold_u>.+?)__(?!\w)"
    r"|~~(?P<strike>.+?)~~"
    r"|\*(?P<italic>[^*\s](?:[^*]*[^*\s])?)\*"
    r"|(?<!\w)_(?P<italic_u>[^_\s](?:[^_]*[^_\s])?)_(?!\w)",
    re.IGNORECASE | re.DOTALL,  # emphasis may span the lines of a paragraph
)
EMPHASIS_GROUPS: dict[str, dict[str, bool]] = {
    "bold_italic": {"bold": True, "italic": True},
    "bold": {"bold": True},
    "bold_u": {"bold": True},
    "strike": {"strike": True},
    "italic": {"italic": True},
    "italic_u": {"italic": True},
}


@dataclass(frozen=True)
class Run:
    """A piece of inline text with uniform formatting. Line breaks are kept as '\\n' in the text."""

    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False
    url: str | None = None


PLAIN = Run("")


def parse_inline(text: str, style: Run = PLAIN) -> list[Run]:
    """Split inline Markdown into formatted runs (emphasis may be nested); HTML entities such as &amp; in
    the text are decoded, in code and formulas they stay."""
    runs: list[Run] = []
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            runs.append(replace(style, text=html.unescape(text[position : match.start()])))
        runs.extend(_runs_for_match(match, style))
        position = match.end()
    if position < len(text):
        runs.append(replace(style, text=html.unescape(text[position:])))
    return runs


def _runs_for_match(match: re.Match[str], style: Run) -> list[Run]:
    kind = match.lastgroup or ""
    if kind == "math":  # with its delimiters, nothing inside is formatting
        return [replace(style, text=match["math"])]
    if kind == "escaped":
        return [replace(style, text=match["escaped"])]
    if kind == "code":
        return [replace(style, text=match["code"], code=True)]
    if kind == "br":
        return [replace(style, text="\n")]
    if kind == "url":
        label = to_plain(match["label"]) or match["url"]
        return [replace(style, text=label, url=match["url"])]
    if kind == "autolink":
        return [replace(style, text=match["autolink"], url=match["autolink"])]
    return parse_inline(match[kind], replace(style, **EMPHASIS_GROUPS[kind]))


def _run_text(run: Run, with_url: bool) -> str:
    if with_url and run.url and run.url != run.text:
        return f"{run.text} ({run.url})"
    return run.text


def to_plain(text: str, with_urls: bool = False) -> str:
    """Strip inline Markdown; optionally keep link targets as 'label (url)'."""
    return "".join(_run_text(run, with_urls) for run in parse_inline(text)).strip()


def is_safe_url(url: str) -> bool:
    return url.strip().lower().startswith(SAFE_LINK_PREFIXES)


def word_link(run: Run) -> tuple[str, str | None]:
    """The text of a run and its link target if Word and PDF make it clickable, else the target stays in the
    text as 'text (url)'."""
    if run.url and is_word_link(run.url):
        return run.text, run.url.strip()
    return _run_text(run, with_url=True), None


def is_word_link(url: str) -> bool:
    """Links Word gets as real hyperlinks (like Office Documents): http, https and plain mailto links."""
    url = url.strip()
    if not url.lower().startswith(WORD_LINK_PREFIXES):
        return False
    if url.lower().startswith("mailto:") and "?" in url:
        fields = [part.split("=", 1)[0].lower() for part in url.split("?", 1)[1].split("&") if part]
        return all(field in MAILTO_FIELDS for field in fields)
    return True


# ------------------------------------------------------------------------------------ block markdown --

FENCE_RE = re.compile(r"^\s*(?P<fence>`{3,}|~{3,})")
HEADING_RE = re.compile(r"^\s{0,3}(?P<hashes>#{1,6})\s+(?P<text>.*?)(?:\s+#+)?\s*$")
RULE_RE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
LIST_ITEM_RE = re.compile(r"^(?P<indent>[ \t]*)(?:[-*+]|(?P<number>\d{1,9})[.)])\s+(?P<text>.*)$")
QUOTE_RE = re.compile(r"^\s{0,3}>\s?(?P<text>.*)$")
TABLE_DELIMITER_RE = re.compile(r"^:?-+:?$")
UNESCAPED_PIPE_RE = re.compile(r"(?<!\\)\|")
UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]+')
# Open WebUI's citation markers: [1], [1][2], [1, 2], [1#page] (its citation-extension.ts)
CITATION_RE = re.compile(r"\[\d+(?:#[^,\]\s]+)?(?:,\s*\d+(?:#[^,\]\s]+)?)*\]")
URL_SPACE_RE = re.compile(r"[\s<>]")  # would end an autolink


@dataclass
class Heading:
    level: int
    text: str


@dataclass
class Paragraph:
    text: str  # lines are stripped and joined with "\n"


@dataclass
class ListItem:
    text: str
    sublists: list["List"] = field(default_factory=list)  # nested lists; deeper nesting is flattened


@dataclass
class List:
    ordered: bool
    start: int
    items: list[ListItem]


@dataclass
class Quote:
    lines: list[str]


@dataclass
class CodeBlock:
    code: str


@dataclass
class Table:
    rows: list[list[str]]  # first row is the header; all rows have the header's width
    aligns: list[str | None]


@dataclass
class Rule:
    pass


Block = Heading | Paragraph | List | Quote | CodeBlock | Table | Rule


def _split_row(line: str) -> list[str]:
    line = line.strip().removeprefix("|")
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    # GFM unescapes "\|" in cells before inline parsing, so it also works inside `code`
    return [cell.strip().replace("\\|", "|") for cell in UNESCAPED_PIPE_RE.split(line)]


def _table_alignments(lines: list[str], i: int) -> list[str | None] | None:
    """Return column alignments if a GFM table starts at line i, else None."""
    if "|" not in lines[i] or i + 1 >= len(lines) or "|" not in lines[i + 1]:
        return None
    header, delimiter = _split_row(lines[i]), [c.replace(" ", "") for c in _split_row(lines[i + 1])]
    if len(header) != len(delimiter) or not all(TABLE_DELIMITER_RE.match(c) for c in delimiter):
        return None
    return [TABLE_ALIGNMENTS.get((cell.startswith(":"), cell.endswith(":"))) for cell in delimiter]


def _parse_table(lines: list[str], i: int, aligns: list[str | None]) -> tuple[Table, int]:
    width = len(aligns)
    rows = [_split_row(lines[i])]
    i += 2
    while i < len(lines) and lines[i].strip() and "|" in lines[i]:
        rows.append(_split_row(lines[i]))
        i += 1
    # GFM: missing cells are empty, surplus cells are ignored
    return Table([(row + [""] * width)[:width] for row in rows], aligns), i


def _parse_code(lines: list[str], i: int, fence: str) -> tuple[CodeBlock, int]:
    closing = re.compile(rf"^\s*{re.escape(fence[0])}{{{len(fence)},}}\s*$")
    code: list[str] = []
    i += 1
    while i < len(lines) and not closing.match(lines[i]):
        code.append(lines[i])
        i += 1
    return CodeBlock("\n".join(code)), i + 1


def _block_start(lines: list[str], i: int) -> tuple[str, Any]:
    """The kind of block line i opens, in order of precedence ("- - -" is a rule, not a list item), and its
    match: blank, fence, table, heading, rule, list, quote or paragraph."""
    line = lines[i]
    if not line.strip():
        return "blank", None
    if fence := FENCE_RE.match(line):
        return "fence", fence
    if (aligns := _table_alignments(lines, i)) is not None:
        return "table", aligns
    if heading := HEADING_RE.match(line):
        return "heading", heading
    if RULE_RE.match(line):
        return "rule", None
    if item := LIST_ITEM_RE.match(line):
        return "list", item
    if QUOTE_RE.match(line):
        return "quote", None
    return "paragraph", None


def _starts_block(lines: list[str], i: int) -> bool:
    """True if line i is blank or opens a block, i.e. ends a running paragraph."""
    return _block_start(lines, i)[0] != "paragraph"


def _add_list_item(lists: list[List], match: re.Match[str], text: str) -> None:
    ordered = match["number"] is not None
    nested = bool(lists) and len(match["indent"].expandtabs(4)) >= 2
    siblings = lists[-1].items[-1].sublists if nested else lists
    if not siblings or siblings[-1].ordered != ordered:  # a change of list type starts a new list
        siblings.append(List(ordered, int(match["number"] or 1), []))
    siblings[-1].items.append(ListItem(text))


def _parse_lists(lines: list[str], i: int, match: re.Match[str]) -> tuple[list[List], int]:
    """Parse consecutive list items (blank lines between them allowed) into top-level lists."""
    lists: list[List] = []
    while True:
        text = match["text"].strip()
        i += 1
        # indented lines directly below an item continue its text
        while i < len(lines) and lines[i][:1].isspace() and not _starts_block(lines, i):
            text += " " + lines[i].strip()
            i += 1
        _add_list_item(lists, match, text)
        while i < len(lines) and not lines[i].strip():
            i += 1
        if i >= len(lines):
            return lists, i
        kind, match = _block_start(lines, i)
        if kind != "list":
            return lists, i


def parse_blocks(markdown_text: str) -> list[Block]:
    """Parse the subset of Markdown that chat models produce into a flat list of blocks."""
    lines = markdown_text.splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        kind, match = _block_start(lines, i)
        if kind == "blank":
            i += 1
        elif kind == "fence":
            code, i = _parse_code(lines, i, match["fence"])
            blocks.append(code)
        elif kind == "table":
            table, i = _parse_table(lines, i, match)
            blocks.append(table)
        elif kind == "heading":
            if match["text"]:
                blocks.append(Heading(len(match["hashes"]), match["text"]))
            i += 1
        elif kind == "rule":
            blocks.append(Rule())
            i += 1
        elif kind == "list":
            lists, i = _parse_lists(lines, i, match)
            blocks.extend(lists)
        elif kind == "quote":
            quote = Quote([])
            while i < len(lines) and (quote_match := QUOTE_RE.match(lines[i])):
                quote.lines.append(quote_match["text"])
                i += 1
            blocks.append(quote)
        else:
            paragraph = [lines[i].strip()]
            i += 1
            while i < len(lines) and not _starts_block(lines, i):
                paragraph.append(lines[i].strip())
                i += 1
            blocks.append(Paragraph("\n".join(paragraph)))
    return blocks


def document_title(blocks: list[Block]) -> str:
    """Plain text of the first heading, or '' if there is none."""
    return next((to_plain(b.text) for b in blocks if isinstance(b, Heading)), "")


def _inline_texts(blocks: list[Block]) -> Iterator[str]:
    """The inline Markdown of all blocks except code blocks."""
    for block in blocks:
        match block:
            case Heading(_, text) | Paragraph(text):
                yield text
            case List(_, _, items):
                for item in items:
                    yield item.text
                    yield from _inline_texts(item.sublists)
            case Quote(lines):
                yield from lines
            case Table(rows):
                for row in rows:
                    yield from row


def _cites(marker: str, count: int) -> bool:
    """Whether a marker like [1, 2#page] is a citation: Open WebUI only shows it as one if every number is
    one of its `count` sources, so "in [2024]" is text."""
    return all(1 <= int(part.split("#")[0]) <= count for part in marker[1:-1].split(","))


def add_sources(blocks: list[Block], sources: list[tuple[str, str | None]], label: str) -> list[Block]:
    """The blocks and, if the text cites sources ([1], [2, 3], [1#page] outside code), a numbered list of
    them under a bold `label`."""
    runs = (run for text in _inline_texts(blocks) for run in parse_inline(text) if not run.code)
    markers = (marker for run in runs for marker in CITATION_RE.finditer(run.text))
    cited = any(_cites(marker[0], len(sources)) for marker in markers)
    if not (sources and cited):
        return blocks
    items = []
    for name, url in sources:
        link = f"<{URL_SPACE_RE.sub(lambda m: quote(m[0]), url)}>" if url else ""  # an autolink
        text = link if url == name else " – ".join(filter(None, (escape_markdown(name), link)))
        items.append(ListItem(text))
    return [*blocks, Paragraph(f"**{escape_markdown(label)}**"), List(True, 1, items)]


def safe_filename(text: str, fallback: str) -> str:
    """Turn text into a short file name that is valid on Windows, macOS and Linux."""
    name = " ".join(UNSAFE_FILENAME_RE.sub(" ", text).split())
    return name[:MAX_FILENAME_CHARS].strip(" .") or fallback


# ------------------------------------------------------------------------------------------------ Word --
# Same design, valves (accent_color, font, word_template) and link rules as the Word files of Office
# Documents, copied from it, so both plugins give the same result.

DEFAULT_ACCENT = "1F4E79"
WHITE, BLACK = "FFFFFF", "000000"
TEXT_COLOR = "262626"
MUTED_COLOR = "6B6B6B"
RULE_COLOR = "D9D9D9"  # thin lines
FONT_NAME_RE = re.compile(r"[^\W_][\w .-]{0,40}")  # 'Calibri', 'Segoe UI', 'Source Sans 3'
THEME_FONTS_RE = (
    r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")'  # headings and body text of a theme
)


def mix(color: str, other: str, share: float) -> str:
    """`color` with `share` (0 to 1) of `other` mixed in, as RRGGBB."""
    return "".join(
        f"{round(a + (b - a) * share):02X}"
        for a, b in zip(bytes.fromhex(color), bytes.fromhex(other), strict=True)
    )


@dataclass(frozen=True)
class Design:
    """Accent colour (accent_color valve) and font (font valve) of the Word file."""

    accent: str = DEFAULT_ACCENT
    font: str = BODY_FONT

    @property
    def band(self) -> str:  # every second table row
        return mix(self.accent, WHITE, 0.93)

    def row_fill(self, index: int) -> str | None:
        """Header row in the accent colour, every second row after it tinted."""
        return self.accent if index == 0 else self.band if index % 2 == 0 else None


DEFAULT_DESIGN = Design()

# Word wants the children of w:pPr and w:tblPr in a fixed order; new ones go before these.
PPR_AFTER_BORDERS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip
TBLPR_AFTER_BORDERS = (
    "w:shd",
    "w:tblLayout",
    "w:tblCellMar",
    "w:tblLook",
    "w:tblCaption",
    "w:tblDescription",
)
TBLPR_AFTER_CELL_MARGINS = ("w:tblLook", "w:tblCaption", "w:tblDescription")
THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
PAPER_SIZES = {"a4": (Cm(21), Cm(29.7)), "letter": (Cm(21.59), Cm(27.94))}

MAX_TEMPLATE_BYTES = 20_000_000
# With a template, its styles for these are replaced: they get the design. The content uses Body Text, not the
# template's Normal style, which its header and footer build on.
LIST_STYLES = (*BULLET_LIST_STYLES, *ORDERED_LIST_STYLES)
TEMPLATE_STYLES = ("Body Text", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Quote", *LIST_STYLES)
# Settings of a template that would follow every document: a mail merge (Word asks to run its SQL
# command), write protection, and the path of the template it was made from (Word looks for it, e.g. on a
# network share).
TEMPLATE_SETTINGS_REMOVED = ("w:attachedTemplate", "w:mailMerge", "w:documentProtection", "w:writeProtection")
DOTX_MAIN = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"


def read_word_template(path: str) -> bytes | None:
    """The Word file of the word_template valve, a .dotx turned into a .docx; None if there is none or it
    cannot be read (logged: the documents are made without it)."""
    path = path.strip()
    if not path:
        return None
    try:
        with open(path, "rb") as file:
            data = file.read(MAX_TEMPLATE_BYTES + 1)
        if len(data) > MAX_TEMPLATE_BYTES:
            raise ValueError(f"the file is larger than {MAX_TEMPLATE_BYTES // 1_000_000} MB")
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            if sum(item.file_size for item in package.infolist()) > 5 * MAX_TEMPLATE_BYTES:
                raise ValueError("the file is too large unpacked")
            types = package.read("[Content_Types].xml")
            if DOTX_MAIN not in types:
                return data
            buffer = io.BytesIO()  # python-docx opens documents only
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as document:
                for item in package.infolist():
                    content = package.read(item)
                    if item.filename == "[Content_Types].xml":
                        content = content.replace(DOTX_MAIN, DOCX_CT.WML_DOCUMENT_MAIN.encode())
                    document.writestr(item, content)
            return buffer.getvalue()
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as e:
        log.warning("The Word template %r cannot be used, documents are made without it: %s", path, e)
        return None


def _set_style(style: Any, size: float, color: str, bold: bool = False, italic: bool = False,
               before: float | None = None, after: float | None = None,
               font: str = BODY_FONT) -> None:  # fmt: skip
    """Font, colour and spacing of a built-in style; its theme font and colour are removed, they would win."""
    style.font.name = font
    for attr in THEME_FONT_ATTRS:
        style.element.rPr.rFonts.attrib.pop(qn(attr), None)
    style.font.size, style.font.bold, style.font.italic = Pt(size), bold, italic
    style.font.color.rgb = RGBColor.from_string(color)
    for attr in ("w:themeShade", "w:themeTint"):
        style.element.rPr.color.attrib.pop(qn(attr), None)
    if before is not None:
        style.paragraph_format.space_before = Pt(before)
    if after is not None:
        style.paragraph_format.space_after = Pt(after)


def _set_borders(
    parent: Any, tag: str, successors: tuple[str, ...], sides: list[tuple[str, int, str]]
) -> None:
    """Replace the w:pBdr of a paragraph (or style) or the w:tblBorders of a table. Sides in schema order
    (top, left, bottom, right, insideH, insideV), size in eighths of a point, 0 for no line."""
    old = parent.find(qn(tag))
    if old is not None:
        parent.remove(old)
    borders = OxmlElement(tag)
    for side, size, color in sides:
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single" if size else "nil")
        if size:
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:space"), "4" if tag == "w:pBdr" else "0")
            edge.set(qn("w:color"), color)
        borders.append(edge)
    parent.insert_element_before(borders, *successors)


def _set_cell_margins(table: Any, top: int, bottom: int) -> None:
    """Cell margins in twips; left and right stay at Word's 108."""
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", 108), ("bottom", bottom), ("right", 108)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    table._tbl.tblPr.insert_element_before(margins, *TBLPR_AFTER_CELL_MARGINS)


def _set_cell_shading(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


class DocxWriter:
    """Render parsed blocks into a Word document: headings in the accent colour, tables with an accent
    header row, real hyperlinks, optionally on a company template (header, footer, page size, margins)."""

    def __init__(
        self,
        design: Design = DEFAULT_DESIGN,
        template: bytes | None = None,
        paper: str = "letter",
        title: str = "",
    ) -> None:
        self.design, self.paper, self.title = design, paper, title
        self.template = False
        if template:
            try:
                self.doc, self.template = Document(io.BytesIO(template)), True
                self._take_template()
                self._styles()
                return
            except Exception:
                log.warning(
                    "The Word template cannot be used, the document is made without it", exc_info=True
                )
                self.template = False
        self.doc = Document()
        self._styles()

    @property
    def body_style(self) -> str | None:
        return "Body Text" if self.template else None  # None: Normal

    def render(self, blocks: list[Block]) -> bytes:
        self._page_setup()
        for block in blocks:
            match block:
                case Heading(level, text):
                    self.doc.add_heading(to_plain(text), level=min(level, MAX_DOCX_HEADING_LEVEL))
                case Paragraph(text):
                    self._add_runs(self.doc.add_paragraph(style=self.body_style), parse_inline(text))
                case List():
                    self._list(block, level=0)
                case Quote(lines):
                    for line in filter(str.strip, lines):
                        self._add_runs(self.doc.add_paragraph(style="Quote"), parse_inline(line))
                case CodeBlock(code):
                    run = self.doc.add_paragraph(style=self.body_style).add_run(code)
                    run.font.name = CODE_FONT
                    run.font.size = CODE_FONT_SIZE
                case Table():
                    self._table(block)
                case Rule():  # a thin line: an empty paragraph with a bottom border
                    pPr = self.doc.add_paragraph(style=self.body_style)._p.get_or_add_pPr()
                    _set_borders(pPr, "w:pBdr", PPR_AFTER_BORDERS, [("bottom", 6, RULE_COLOR)])
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()

    def _take_template(self) -> None:
        """Keep the template's header, footer, page size and margins; remove its content and give the content
        styles the design (TEMPLATE_STYLES)."""
        body = self.doc.element.body
        last = body.get_or_add_sectPr()  # the last section's settings, with its header and footer
        for tag in (
            "w:headerReference",
            "w:footerReference",
        ):  # a section without its own takes the previous one's
            for earlier in reversed(body.xpath("./w:p/w:pPr/w:sectPr")):
                for reference in earlier.findall(qn(tag)):
                    if not any(
                        own.get(qn("w:type")) == reference.get(qn("w:type")) for own in last.findall(qn(tag))
                    ):
                        last.insert(0, copy.deepcopy(reference))
        for child in list(body):
            if child is not last:
                body.remove(child)
        default = Document()
        styles = self.doc.styles.element
        # The template's default character style, which Word names after its language
        default_font = next(
            (style.styleId for style in styles.findall(qn("w:style"))
             if style.get(qn("w:type")) == "character" and style.get(qn("w:default")) in ("1", "true")),
            None,
        )  # fmt: skip
        renamed = {"Normal": "BodyText", "DefaultParagraphFont": default_font}
        for name in TEMPLATE_STYLES:
            style = default.styles[name].element
            link = style.find(qn("w:link"))
            linked = default.styles.element.get_by_id(link.get(qn("w:val"))) if link is not None else None
            for new in (style, linked) if linked is not None else (style,):
                for old in styles.findall(qn("w:style")):
                    if (
                        old.name_val or ""
                    ).casefold() == new.name_val.casefold() or old.styleId == new.styleId:
                        old_link = old.find(qn("w:link"))  # e.g. "Überschrift 1 Zchn" of "Überschrift 1"
                        old_linked = (
                            styles.get_by_id(old_link.get(qn("w:val"))) if old_link is not None else None
                        )
                        if old_linked is not None:
                            styles.remove(old_linked)
                        if old.getparent() is not None:
                            styles.remove(old)
                new = copy.deepcopy(new)
                for tag in ("w:basedOn", "w:next"):
                    reference = new.find(qn(tag))
                    if reference is None or reference.get(qn("w:val")) not in renamed:
                        continue
                    if new.styleId == "BodyText" or not renamed[reference.get(qn("w:val"))]:
                        new.remove(reference)  # nothing of the template's Normal
                    else:
                        reference.set(qn("w:val"), renamed[reference.get(qn("w:val"))])
                styles.append(new)
        self._take_numbering(default)
        try:
            settings_part = self.doc.part.part_related_by(RELATIONSHIP_TYPE.SETTINGS)
        except KeyError:
            settings_part = None  # a file without settings
        if settings_part is not None:
            for tag in TEMPLATE_SETTINGS_REMOVED:
                for element in settings_part.element.findall(qn(tag)):
                    settings_part.element.remove(element)
                    for node in element.iter():
                        if node.get(qn("r:id")):
                            settings_part.drop_rel(node.get(qn("r:id")))
        core = self.doc.core_properties  # the template's author, title ... are not the document's
        core.author = core.last_modified_by = core.subject = core.keywords = core.comments = core.category = (
            ""
        )
        core.title = self.title

    def _take_numbering(self, default: Any) -> None:
        """The list styles now point to the default template's numbering: copy it in, with new numbers."""
        source = default.part.numbering_part.element
        try:
            numbering = self.doc.part.part_related_by(RELATIONSHIP_TYPE.NUMBERING).element
        except KeyError:
            package, numbering = self.doc.part.package, parse_xml(f"<w:numbering {nsdecls('w')}/>")
            partname = package.next_partname("/word/numbering%d.xml")
            self.doc.part.relate_to(
                NumberingPart(partname, DOCX_CT.WML_NUMBERING, numbering, package),
                RELATIONSHIP_TYPE.NUMBERING,
            )
        abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        for abstract_id, name in enumerate(LIST_STYLES, start=max(abstract_ids, default=-1) + 1):
            num_id = self.doc.styles[name].element.pPr.numPr.numId
            used = source.num_having_numId(num_id.val).abstractNumId.val
            abstract = copy.deepcopy(source.xpath(f"w:abstractNum[@w:abstractNumId='{used}']")[0])
            abstract.set(qn("w:abstractNumId"), str(abstract_id))
            for nsid in abstract.findall(qn("w:nsid")):  # Word's list identity; the template may use the same
                abstract.remove(nsid)
            numbering.insert_element_before(abstract, "w:num", "w:numIdMacAtCleanup")
            num_id.val = numbering.add_num(abstract_id).numId

    def _styles(self) -> None:
        """A calm, modern look: dark grey body text with some air, headings in the accent colour."""
        styles, accent, font = self.doc.styles, self.design.accent, self.design.font
        body = styles[self.body_style or "Normal"]
        _set_style(body, 10.5, TEXT_COLOR, after=6, font=font)
        body.paragraph_format.line_spacing = 1.15
        _set_style(styles["Heading 1"], 16, accent, bold=True, before=20, after=6, font=font)
        _set_style(styles["Heading 2"], 13, TEXT_COLOR, bold=True, before=14, after=4, font=font)
        _set_style(styles["Heading 3"], 11, accent, bold=True, before=10, after=2, font=font)
        _set_style(
            styles["Heading 4"], 10.5, TEXT_COLOR, bold=True, italic=True, before=8, after=2, font=font
        )
        quote = styles["Quote"]
        _set_style(quote, 10.5, MUTED_COLOR, italic=True, before=6, after=10, font=font)
        quote.paragraph_format.left_indent = Cm(0.4)
        _set_borders(quote.element.get_or_add_pPr(), "w:pBdr", PPR_AFTER_BORDERS, [("left", 18, accent)])
        self._for_previews()

    def _for_previews(self) -> None:
        """Word takes fonts and colours from the paragraph styles. Other viewers also read the document
        defaults, the theme fonts and the character styles linked to the paragraph styles, and they may have
        no Symbol font for Word's bullets (boxes instead of dots). All get the design."""
        font, styles = self.design.font, self.doc.styles
        if not self.template:  # the template's header and footer may use its defaults and theme fonts
            for fonts in styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts"):
                for attr in THEME_FONT_ATTRS:
                    fonts.attrib.pop(qn(attr), None)
                for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                    fonts.set(qn(attr), font)
            theme_fonts = re.compile(THEME_FONTS_RE.encode())
            for part in self.doc.part.package.iter_parts():
                if str(part.partname).startswith("/word/theme/"):
                    part._blob = theme_fonts.sub(lambda m: m[1] + font.encode() + m[2], part.blob)
        for style in styles:
            link = style.element.find(qn("w:link"))
            if style.type != WD_STYLE_TYPE.PARAGRAPH or link is None or style.element.rPr is None:
                continue
            linked = styles.element.get_by_id(link.get(qn("w:val")))
            if linked is not None:  # e.g. "Heading 1 Char": same look as its paragraph style
                for old in linked.findall(qn("w:rPr")):
                    linked.remove(old)
                linked.append(copy.deepcopy(style.element.rPr))
        for level in self.doc.part.numbering_part.element.iter(qn("w:lvl")):
            kind = level.find(qn("w:numFmt"))
            if kind is None or kind.get(qn("w:val")) != "bullet":
                continue
            indent = level.find(f"{qn('w:pPr')}/{qn('w:ind')}")
            sub_item = indent is not None and int(indent.get(qn("w:left"), "0")) > 720
            level.find(qn("w:lvlText")).set(qn("w:val"), "–" if sub_item else "•")
            for old in level.findall(qn("w:rPr")):
                level.remove(old)
            props = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            color = OxmlElement("w:color")
            color.set(qn("w:val"), self.design.accent)
            props.append(fonts)
            props.append(color)
            level.append(props)

    def _page_setup(self) -> None:
        """Paper and margins; a template keeps its own."""
        section = self.doc.sections[0]
        width, height = PAPER_SIZES[self.paper]
        page = {"page_width": width, "page_height": height, "left_margin": Cm(2.5), "right_margin": Cm(2.5),
                "top_margin": Cm(2.5), "bottom_margin": Cm(2.2), "header_distance": Cm(1.25),
                "footer_distance": Cm(1.25)}  # fmt: skip
        for name, value in page.items():
            if not self.template or getattr(section, name) is None:  # what a template lacks
                setattr(section, name, value)

    def _list(self, block: List, level: int) -> None:
        style = (ORDERED_LIST_STYLES if block.ordered else BULLET_LIST_STYLES)[level]
        num_id: int | None = None
        for item in block.items:
            paragraph = self.doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(3)
            if block.ordered:
                if num_id is None:  # each list gets its own numbering, so the next one restarts
                    num_id = self._new_numbering(paragraph.style, block.start)
                paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
            self._add_runs(paragraph, parse_inline(item.text))
            for sublist in item.sublists:
                self._list(sublist, level=1)

    def _new_numbering(self, style: "ParagraphStyle", start: int) -> int:
        """Create a numbering instance of the style's list that starts at `start`.

        python-docx has no public API for this; without it every numbered list in the
        document would continue the previous list's numbering.
        """
        numbering = self.doc.part.numbering_part.element
        style_num_id = style.element.pPr.numPr.numId.val
        num = numbering.add_num(numbering.num_having_numId(style_num_id).abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(start)
        return num.numId

    def _table(self, table: Table) -> None:
        """Header row in the accent colour, every second row lightly tinted, thin lines between rows only."""
        doc_table = self.doc.add_table(rows=len(table.rows), cols=len(table.aligns))
        lines = [("top", 0, ""), ("left", 0, ""), ("bottom", 8, self.design.accent), ("right", 0, ""),
                 ("insideH", 4, RULE_COLOR), ("insideV", 0, "")]  # fmt: skip
        _set_borders(doc_table._tbl.tblPr, "w:tblBorders", TBLPR_AFTER_BORDERS, lines)
        _set_cell_margins(doc_table, top=50, bottom=50)
        for row_index, (doc_row, row) in enumerate(zip(doc_table.rows, table.rows, strict=True)):
            if row_index == 0:  # repeat the header row on every page
                header_flag = OxmlElement("w:tblHeader")
                header_flag.set(qn("w:val"), "true")
                doc_row._tr.get_or_add_trPr().append(header_flag)
            fill = self.design.row_fill(row_index)
            for cell, text, align in zip(doc_row.cells, row, table.aligns, strict=True):
                paragraph = cell.paragraphs[0]
                if self.template:
                    paragraph.style = self.body_style
                paragraph.alignment = DOCX_ALIGNMENT.get(align or "")
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                runs = parse_inline(text, Run("", bold=row_index == 0))
                self._add_runs(paragraph, runs, WHITE if row_index == 0 else None)
                if fill:
                    _set_cell_shading(cell, fill)
        self.doc.add_paragraph(style=self.body_style).paragraph_format.space_after = Pt(4)

    def _add_runs(self, paragraph: DocxParagraph, runs: list[Run], color: str | None = None) -> None:
        for run in runs:
            text, link = word_link(run)
            docx_run = paragraph.add_run(text)
            docx_run.bold = run.bold or None
            docx_run.italic = run.italic or None
            if color:
                docx_run.font.color.rgb = RGBColor.from_string(color)
            if run.strike:
                docx_run.font.strike = True
            if run.code:
                docx_run.font.name = CODE_FONT
            if link:
                self._make_link(paragraph, docx_run, link, color or self.design.accent)

    @staticmethod
    def _make_link(paragraph: DocxParagraph, docx_run: Any, url: str, color: str) -> None:
        docx_run.font.underline = True
        docx_run.font.color.rgb = RGBColor.from_string(color)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        docx_run._r.addprevious(link)
        link.append(docx_run._r)


def markdown_to_docx(
    blocks: list[Block], design: Design = DEFAULT_DESIGN, template: bytes | None = None, paper: str = "letter"
) -> bytes:
    return DocxWriter(design, template, paper, document_title(blocks)).render(blocks)


# ----------------------------------------------------------------------------------------------- Excel --

CURRENCY_RE = "|".join(map(re.escape, CURRENCY_FORMATS))
NUMBER_RE = re.compile(  # thousands separators: . , ' ’ space, no-break space, narrow no-break space
    rf"(?P<prefix>{CURRENCY_RE})?\s*"
    r"(?P<sign>[-+]?)"
    r"(?P<int>\d{1,3}(?:(?P<group>[.,'\u2019 \u00a0\u202f])\d{3}(?:(?P=group)\d{3})*)?|\d+)"
    r"(?:(?!(?P=group))(?P<decimal>[.,])(?P<fraction>\d+))?"  # decimal mark differs from the separator
    rf"\s*(?P<suffix>%|{CURRENCY_RE})?"
)


DATE_RE = re.compile(  # 15.03.2024, 1.3.2024 or 2024-03-15
    r"(?P<day>\d{1,2})\.(?P<month>\d{1,2})\.(?P<year>\d{4})"
    r"|(?P<iso_year>\d{4})-(?P<iso_month>\d{2})-(?P<iso_day>\d{2})"
)
MIN_EXCEL_YEAR = 1900
INVALID_SHEET_TITLE_RE = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_TITLE_CHARS = 31
RESERVED_SHEET_TITLES = {"history"}  # Excel keeps it for its change tracking
HYPERLINK_COLOR = "0563C1"  # Office's colour for links


def _ambiguous(match: re.Match[str]) -> bool:
    """A single dot or comma before three digits, e.g. "1.250": a thousands separator or the decimal mark."""
    group = match["group"]
    return group in (".", ",") and not match["fraction"] and match["int"].count(group) == 1


def parse_number(text: str, decimal: str | None = None) -> tuple[str | int | float, str | None]:
    """Convert '1.234,56 €', '1.234,56 EUR', '-3,2 %', '1,234.5' etc. to (number, Excel number format).

    A dot or comma followed by exactly three digits ("2.000", "1,234") is read as the decimal mark
    `decimal` ("." or ","), else as a thousands separator; in "0.125" it is always the decimal mark.
    Values that do not look like a plain number, integers with leading zeros (postcodes, IDs) and
    phone numbers ("+49 711 123", "+1 234") are returned unchanged as (text, None).
    """
    match = NUMBER_RE.fullmatch(text.strip().replace("\u2212", "-"))  # Unicode minus sign
    if not match or (match["prefix"] and match["suffix"]):
        return text, None
    if match["sign"] == "+" and (match["group"] or "").isspace() and not (match["prefix"] or match["suffix"]):
        return text, None  # international phone number
    whole, group, fraction = match["int"], match["group"], match["fraction"] or ""
    if _ambiguous(match) and (whole.startswith("0") or group == decimal):
        (whole, fraction), group = whole.split(group), None
    digits = re.sub(r"\D", "", whole)
    if (len(digits) > 1 and digits.startswith("0")) or len(digits) + len(fraction) > MAX_EXACT_DIGITS:
        return text, None
    number: int | float = (
        float(f"{match['sign']}{digits}.{fraction}") if fraction else int(match["sign"] + digits)
    )
    unit = match["prefix"] or match["suffix"]
    if unit == "%":
        return number / 100, "0%" if float(number).is_integer() else "0.0#%"
    if unit:
        return float(number), CURRENCY_FORMATS[unit]
    if group:  # written with thousands separators: Excel shows them too (but not in years or postcodes)
        return number, "#,##0" + ("." + "0" * len(fraction) if fraction else "")
    return number, None


def decimal_mark(text: str) -> str | None:
    """The decimal mark a number shows: "." in "1.25", "1,234.5", "1,234,567" and "0.125", "," in "1,5" and
    "1.234.567"; None if it shows none ("1.250", "42") and for text."""
    match = NUMBER_RE.fullmatch(text.strip().replace("\u2212", "-"))
    if not match:
        return None
    if match["fraction"]:
        return match["decimal"]
    group = match["group"]
    if group not in (".", ","):
        return None
    if match["int"].count(group) > 1:
        return "," if group == "." else "."
    return group if match["int"].startswith("0") else None


def column_decimal(values: list[str], fallback: str) -> str:
    """The decimal mark of a table column: the one its numbers show, else `fallback` (the language's)."""
    marks = {mark for text in values if (mark := decimal_mark(text))}
    return marks.pop() if len(marks) == 1 else fallback


def parse_date(text: str) -> date | None:
    """A date written as 15.03.2024, 1.3.2024 or 2024-03-15; None for anything else, also for 03/15/2024
    (month or day first?) and dates before 1900, which Excel cannot show."""
    match = DATE_RE.fullmatch(text.strip())
    if not match:
        return None
    year, month, day = (int(match[part] or match[f"iso_{part}"]) for part in ("year", "month", "day"))
    try:
        return date(year, month, day) if year >= MIN_EXCEL_YEAR else None
    except ValueError:  # 31.02.2024
        return None


def cell_value(text: str, decimal: str, date_format: str) -> tuple[str | int | float | date, str | None]:
    """Value and Excel number format of a cell below the header: a date, a number (parse_number) or text."""
    day = parse_date(text)
    return (day, date_format) if day else parse_number(text, decimal)


def tables_with_titles(blocks: list[Block]) -> list[tuple[str, Table]]:
    """The tables, each with the plain text of a heading directly above it ('' if there is none)."""
    return [
        (to_plain(blocks[i - 1].text) if i and isinstance(blocks[i - 1], Heading) else "", block)
        for i, block in enumerate(blocks)
        if isinstance(block, Table)
    ]


def sheet_title(title: str, default: str, used: set[str]) -> str:
    """A valid sheet name not in `used` (compared as Excel does, ignoring case): the title without the
    characters Excel forbids, at most 31 characters, else `default`; a repeated one gets " (2)" etc."""
    name = " ".join(INVALID_SHEET_TITLE_RE.sub(" ", title).split())[:MAX_SHEET_TITLE_CHARS]
    name = name.strip(" '") or default  # Excel does not allow an apostrophe at either end
    candidate, n = name, 1
    while candidate.casefold() in used | RESERVED_SHEET_TITLES:
        n += 1
        suffix = f" ({n})"
        candidate = name[: MAX_SHEET_TITLE_CHARS - len(suffix)].rstrip() + suffix
    used.add(candidate.casefold())
    return candidate


def _cell_link(raw: str) -> str | None:
    """The target if the cell is exactly one link that Word would make clickable too (http, https, mailto)."""
    runs = [run for run in parse_inline(raw) if run.text.strip()]
    if len(runs) == 1 and runs[0].url and is_word_link(runs[0].url):
        return runs[0].url.strip()
    return None


def tables_to_xlsx(tables: list[tuple[str, Table]], sheet_name: str, font: str = BODY_FONT,
                   decimal: str = ".", date_format: str = "yyyy-mm-dd") -> bytes:  # fmt: skip
    """One worksheet per table (see tables_with_titles), named after its title or `sheet_name`, a format
    string with an {n} placeholder. "1.250" is read with the decimal mark its column shows, else with
    `decimal`; dates get `date_format`. A cell that is one link is a hyperlink, other links stay
    visible as 'text (url)'."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    body_font, header_font = Font(name=font), Font(name=font, bold=True)  # a font without a name looks different
    link_font = Font(name=font, underline="single", color=HYPERLINK_COLOR)
    used: set[str] = set()
    for n, (title, table) in enumerate(tables, start=1):
        sheet = workbook.create_sheet(title=sheet_title(title, sheet_name.format(n=n), used))
        widths = [0] * len(table.aligns)
        columns = [[to_plain(row[column]) for row in table.rows[1:]] for column in range(len(widths))]
        marks = [column_decimal(values, decimal) for values in columns]
        for row_index, row in enumerate(table.rows, start=1):
            for column, raw in enumerate(row, start=1):
                link = _cell_link(raw) if row_index > 1 else None
                text = to_plain(raw, with_urls=not link)
                if row_index == 1 or link:
                    value, number_format = text, None
                else:
                    value, number_format = cell_value(text, marks[column - 1], date_format)
                cell = sheet.cell(row=row_index, column=column, value=value if value != "" else None)
                if isinstance(value, str) and value:
                    cell.data_type = "s"  # never let model output become a live formula ("=...")
                if number_format:
                    cell.number_format = number_format
                if link:  # Excel's own style for links, in the font of the sheet
                    cell.hyperlink, cell.style = link, "Hyperlink"
                if "\n" in text:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = header_font if row_index == 1 else link_font if link else body_font
                if row_index == 1:
                    cell.fill = header_fill
                widths[column - 1] = max(widths[column - 1], *(len(line) for line in text.split("\n")))
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max(width + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------------------------ HTML and text --


def _runs_html(runs: list[Run]) -> str:
    parts: list[str] = []
    for run in runs:
        part = html.escape(run.text).replace("\n", "<br>")
        if run.code:
            part = f"<code>{part}</code>"
        if run.strike:
            part = f"<s>{part}</s>"
        if run.italic:
            part = f"<em>{part}</em>"
        if run.bold:
            part = f"<strong>{part}</strong>"
        if run.url and is_safe_url(run.url):
            part = f'<a href="{html.escape(run.url.strip())}">{part}</a>'
        parts.append(part)
    return "".join(parts)


def _inline_html(text: str) -> str:
    return _runs_html(parse_inline(text))


def _list_html(block: List) -> str:
    tag = "ol" if block.ordered else "ul"
    start = f' start="{block.start}"' if block.ordered and block.start != 1 else ""
    items = "".join(
        f"<li>{_inline_html(item.text)}{''.join(map(_list_html, item.sublists))}</li>" for item in block.items
    )
    return f"<{tag}{start}>{items}</{tag}>"


def _table_html(table: Table) -> str:
    def row_html(row: list[str], tag: str) -> str:
        cells = []
        for text, align in zip(row, table.aligns, strict=True):
            style = f' style="text-align:{align}"' if align else ""
            cells.append(f"<{tag}{style}>{_inline_html(text)}</{tag}>")
        return f"<tr>{''.join(cells)}</tr>"

    body = "".join(row_html(row, "td") for row in table.rows[1:])
    return (
        '<table border="1" cellpadding="5" style="border-collapse:collapse">'
        f"<thead>{row_html(table.rows[0], 'th')}</thead><tbody>{body}</tbody></table>"
    )


def _block_html(block: Block) -> str:
    match block:
        case Heading(level, text):
            return f"<h{level}>{_inline_html(text)}</h{level}>"
        case Paragraph(text):
            return f"<p>{_inline_html(text)}</p>"
        case List():
            return _list_html(block)
        case Quote(lines):
            return f"<blockquote>{'<br>'.join(_inline_html(line) for line in lines)}</blockquote>"
        case CodeBlock(code):
            return f"<pre><code>{html.escape(code)}</code></pre>"
        case Table():
            return _table_html(block)
        case Rule():
            return "<hr>"
    raise TypeError(f"Unexpected block: {block!r}")


def blocks_to_html(blocks: list[Block]) -> str:
    """Render blocks as HTML. Every piece of text is escaped; no raw HTML passes through."""
    return "\n".join(map(_block_html, blocks))


def _list_text(block: List, indent: str = "") -> str:
    lines: list[str] = []
    for number, item in enumerate(block.items, start=block.start):
        marker = f"{number}." if block.ordered else "•"
        lines.append(f"{indent}{marker} {to_plain(item.text, with_urls=True)}")
        lines.extend(_list_text(sublist, indent + "    ") for sublist in item.sublists)
    return "\n".join(lines)


def _block_text(block: Block) -> str:
    match block:
        case Heading(_, text):
            return to_plain(text, with_urls=True)
        case Paragraph(text):
            return to_plain(text, with_urls=True)
        case List():
            return _list_text(block)
        case Quote(lines):
            return "\n".join(f"> {to_plain(line, with_urls=True)}" for line in lines)
        case CodeBlock(code):
            return code
        case Table(rows):
            return "\n".join(
                "   ".join(to_plain(cell, with_urls=True) for cell in row).rstrip() for row in rows
            )
    raise TypeError(f"Unexpected block: {block!r}")


def blocks_to_text(blocks: list[Block]) -> str:
    """Render blocks as readable plain text (used for mailto links and the .eml text part)."""
    return "\n\n".join(_block_text(block) for block in blocks if not isinstance(block, Rule))


# ---------------------------------------------------------------------------------------------- E-mail --

SUBJECT_RE = re.compile(r"^(?:subject|betreff)\s*:\s*(?P<subject>.+)$", re.IGNORECASE)


def split_email(blocks: list[Block]) -> tuple[str, list[Block]]:
    """Return (subject, body blocks). A 'Subject:'/'Betreff:' line in a heading or paragraph becomes
    the subject and everything before it (the model's preamble) is dropped; without such a line
    the subject is '' and all blocks are the body."""
    for i, block in enumerate(blocks):
        if not isinstance(block, Heading | Paragraph):
            continue
        lines = block.text.split("\n")
        for j, line in enumerate(lines):
            if match := SUBJECT_RE.match(to_plain(line)):
                rest = "\n".join(lines[j + 1 :])
                return match["subject"].strip(), ([Paragraph(rest)] if rest else []) + blocks[i + 1 :]
    return "", blocks


def build_mailto(subject: str, body_text: str) -> str:
    body = body_text.replace("\r\n", "\n").replace("\n", "\r\n")
    return f"mailto:?subject={quote(subject, safe='')}&body={quote(body, safe='')}"


def build_eml(subject: str, blocks: list[Block], body_text: str) -> bytes:
    """Build an unsent e-mail draft with a plain-text and an HTML alternative."""
    message = EmailMessage()
    message["Subject"] = " ".join(subject.split())
    message["X-Unsent"] = "1"  # Outlook opens the file as a new, editable message
    message.set_content(body_text)
    message.add_alternative(EMAIL_HTML.format(body=blocks_to_html(blocks)), subtype="html")
    return message.as_bytes()


# ------------------------------------------------------------------------------------------------- PDF --
# The plugin writes the PDF itself, so no PDF library is needed: the standard fonts every PDF viewer has
# (Helvetica and Courier, not embedded, WinAnsiEncoding = cp1252), text wrapped by the fonts' real glyph
# widths. Paper, margins, colours, spacing and link rules are those of the Word export.

CM = 72 / 2.54  # points
PDF_PAPER = {"a4": (595.28, 841.89), "letter": (612.0, 792.0)}
PDF_SIDE_MARGIN, PDF_TOP_MARGIN, PDF_BOTTOM_MARGIN = 2.5 * CM, 2.5 * CM, 2.2 * CM  # as in Word
PDF_FOOTER_Y = 1.25 * CM  # baseline of the page footer
PDF_LEADING = 1.35  # line height / font size of text
PDF_BODY_SIZE = 10.5
PDF_HEADINGS = {  # level -> font size, italic, in the accent colour, space before and after (as in Word)
    1: (16, False, True, 20, 6),
    2: (13, False, False, 14, 4),
    3: (11, False, True, 10, 2),
    4: (10.5, True, False, 8, 2),
}
PDF_LIST_INDENT = 18  # per level, as Word's list styles
PDF_QUOTE_INDENT, PDF_QUOTE_BAR = 12, 2.25
PDF_CODE_SIZE, PDF_CODE_LEADING, PDF_CODE_PAD, PDF_CODE_FILL = 9, 1.3, 6, "F4F4F4"
PDF_TABLE_SIZE, PDF_TABLE_SIZE_SMALL, PDF_TABLE_LEADING = 10, 8.5, 1.25  # small: more than 6 columns
PDF_CELL_PAD_X, PDF_CELL_PAD_Y = 5.4, 3
PDF_FOOTER_SIZE = 8
PDF_FONTS = {  # (bold, italic, code) -> standard font
    (False, False, False): "Helvetica",
    (True, False, False): "Helvetica-Bold",
    (False, True, False): "Helvetica-Oblique",
    (True, True, False): "Helvetica-BoldOblique",
    (False, False, True): "Courier",
    (True, False, True): "Courier-Bold",
    (False, True, True): "Courier-Oblique",
    (True, True, True): "Courier-BoldOblique",
}
PDF_FONT_IDS = {name: f"F{n}" for n, name in enumerate(PDF_FONTS.values(), start=1)}

# Advance widths in 1/1000 em of the cp1252 characters 32-255, 16 per line, from Adobe's font metrics
# (AFM). The oblique fonts have the widths of the upright ones; Courier is 600 throughout.
HELVETICA_WIDTHS = """
    278 278 355 556 556 889 667 191 333 333 389 584 278 333 278 278
    556 556 556 556 556 556 556 556 556 556 278 278 584 584 584 556
    1015 667 667 722 722 667 611 778 722 278 500 667 556 833 722 778
    667 778 722 667 611 722 667 944 667 667 611 278 278 278 469 556
    333 556 556 500 556 556 278 556 556 222 222 500 222 833 556 556
    556 556 333 500 278 556 500 722 500 500 500 334 260 334 584 0
    556 0 222 556 333 1000 556 556 333 1000 667 333 1000 0 611 0
    0 222 222 333 333 350 556 1000 333 1000 500 333 944 0 500 667
    278 333 556 556 556 556 260 556 333 737 370 556 584 333 737 333
    400 584 333 333 333 556 537 278 333 333 365 556 834 834 834 611
    667 667 667 667 667 667 1000 722 667 667 667 667 278 278 278 278
    722 722 778 778 778 778 778 584 778 722 722 722 722 667 667 611
    556 556 556 556 556 556 889 500 556 556 556 556 278 278 278 278
    556 556 556 556 556 556 556 584 611 556 556 556 556 500 556 500
"""
HELVETICA_BOLD_WIDTHS = """
    278 333 474 556 556 889 722 238 333 333 389 584 278 333 278 278
    556 556 556 556 556 556 556 556 556 556 333 333 584 584 584 611
    975 722 722 722 722 667 611 778 722 278 556 722 611 833 722 778
    667 778 722 667 611 722 667 944 667 667 611 333 278 333 584 556
    333 556 611 556 611 556 333 611 611 278 278 556 278 889 611 611
    611 611 389 556 333 611 556 778 556 556 500 389 280 389 584 0
    556 0 278 556 500 1000 556 556 333 1000 667 333 1000 0 611 0
    0 278 278 500 500 350 556 1000 333 1000 556 333 944 0 500 667
    278 333 556 556 556 556 280 556 333 737 370 556 584 333 737 333
    400 584 333 333 333 611 556 278 333 333 365 556 834 834 834 611
    722 722 722 722 722 722 1000 722 667 667 667 667 278 278 278 278
    722 722 778 778 778 778 778 584 778 722 722 722 722 667 667 611
    556 556 556 556 556 556 889 556 556 556 556 556 278 278 278 278
    611 611 611 611 611 611 611 584 611 611 611 611 611 556 611 556
"""


def _font_widths(table: str) -> tuple[int, ...]:
    widths = (0,) * 32 + tuple(map(int, table.split()))
    assert len(widths) == 256
    return widths


PDF_WIDTHS = {
    **dict.fromkeys(("Helvetica", "Helvetica-Oblique"), _font_widths(HELVETICA_WIDTHS)),
    **dict.fromkeys(("Helvetica-Bold", "Helvetica-BoldOblique"), _font_widths(HELVETICA_BOLD_WIDTHS)),
    **dict.fromkeys((name for name in PDF_FONTS.values() if name.startswith("Courier")), (600,) * 256),
}

# Characters the standard fonts lack (they cover cp1252 only): common ones get a stand-in, box-drawing
# characters an ASCII one (BOX_DRAWING), emoji are dropped, other characters lose their accents ("č" -> "c")
# or become "?".
PDF_REPLACEMENTS = {
    "\t": " ", "\u00ad": "", "\u2007": "\u00a0", "\u202f": "\u00a0",  # no-break spaces
    "\u2010": "-", "\u2011": "-", "\u2012": "–", "\u2015": "—", "\u2212": "-", "\u2043": "-",  # dashes
    "′": "'", "″": '"', "⁄": "/", "∕": "/", "∗": "*", "∙": "·", "⋅": "·",
    "←": "<-", "→": "->", "↔": "<->", "⇐": "<=", "⇒": "=>", "⇔": "<=>", "↑": "^", "↓": "v",
    "➡": "->", "➔": "->", "➜": "->", "⬅": "<-", "⮕": "->", "⟶": "->", "⟵": "<-", "⟹": "=>", "↳": "->",
    "↪": "->", "⇨": "->", "⇦": "<-", "▶": ">", "►": ">", "▸": ">", "➤": ">", "❯": ">", "◀": "<", "◄": "<",
    "▲": "^", "▼": "v",
    "≤": "<=", "≥": ">=", "≠": "!=", "≈": "~", "≡": "=",
    "✓": "+", "✔": "+", "✅": "+", "✗": "-", "✘": "-", "❌": "-", "✕": "x", "✖": "x",
    "☐": "[ ]", "☑": "[x]", "☒": "[x]", "★": "*", "☆": "*", "⚠": "!",
    "●": "•", "▪": "•", "■": "•", "‣": "•", "◆": "•", "♦": "•", "○": "o", "◦": "o", "□": "[ ]", "✦": "*",
    "⌀": "Ø", "μ": "µ", "Ł": "L", "ł": "l", "Đ": "D", "đ": "d", "ı": "i",
}  # fmt: skip
EMOJI_RANGES = ((0x1F000, 0x1FAFF), (0x2600, 0x27BF), (0x2B00, 0x2BFF), (0x2300, 0x23FF))
BOX_DRAWING = range(0x2500, 0x2580)  # ─ │ ┌ ├ └ ═ ║ ╭ ...
MAX_CACHED_CHARS = 10_000


def _cp1252(char: str) -> bool:
    try:
        char.encode("cp1252")
        return True
    except UnicodeEncodeError:
        return False


def _box_drawing(char: str) -> str:
    """One ASCII character for a box-drawing one, so trees and boxes keep their shape: "├──" -> "+--"."""
    words = set(unicodedata.name(char).split())
    if "DIAGONAL" in words:
        return {"╱": "/", "╲": "\\"}.get(char, "X")
    across, upright = words & {"HORIZONTAL", "LEFT", "RIGHT"}, words & {"VERTICAL", "UP", "DOWN"}
    if across and upright:
        return "+"
    if upright:
        return "|"
    return "=" if "DOUBLE" in words and "DASH" not in words else "-"


def _pdf_char(char: str) -> str:
    """The stand-in of one character in the standard fonts (see PDF_REPLACEMENTS)."""
    if char in PDF_REPLACEMENTS:
        return PDF_REPLACEMENTS[char]
    if ord(char) in BOX_DRAWING:
        return _box_drawing(char)
    category = unicodedata.category(char)
    if char == "\n":
        return char
    if category in ("Cc", "Cf", "Mn", "Me") or any(low <= ord(char) <= high for low, high in EMOJI_RANGES):
        return ""  # control and format characters (zero-width space ...), variation selectors, emoji
    if _cp1252(char):  # also the no-break space, which keeps words together
        return char
    if category.startswith("Z"):
        return " "
    folded = "".join(
        PDF_REPLACEMENTS.get(part, part if _cp1252(part) and not unicodedata.combining(part) else "")
        for part in unicodedata.normalize("NFKD", char)
    )
    return folded or "?"


class _PdfChars(dict[int, str]):
    """str.translate table that works out each character once."""

    def __missing__(self, code: int) -> str:
        value = _pdf_char(chr(code))
        if len(self) < MAX_CACHED_CHARS:
            self[code] = value
        return value


PDF_CHARS = _PdfChars()


def pdf_safe_text(text: str) -> str:
    """Text the standard fonts can show: only cp1252 characters, '\\n' kept, tabs as spaces."""
    return unicodedata.normalize("NFC", text).translate(PDF_CHARS)


def text_width(font: str, size: float, text: str) -> float:
    """Width in points of text (cp1252 characters only) in a standard font."""
    widths = PDF_WIDTHS[font]
    return sum(widths[code] for code in text.encode("cp1252")) * size / 1000


@dataclass(frozen=True)
class PdfStyle:
    font: str
    size: float
    color: str
    url: str | None = None  # a link: underlined, clickable
    strike: bool = False


@dataclass
class PdfWord:
    """Text between two spaces, possibly in several styles; `space` is the space in front of it."""

    parts: list[tuple[PdfStyle, str]] = field(default_factory=list)
    width: float = 0.0
    space: tuple[PdfStyle, str] | None = None
    space_width: float = 0.0


@dataclass
class PdfLine:
    parts: list[tuple[PdfStyle, str]]
    width: float = 0.0


SPACE_RE = re.compile(r"(\n| +)")


def pdf_words(runs: list[Run], size: float, color: str, link_color: str) -> list[PdfWord | None]:
    """Split runs into words for line wrapping; None is a line break. Links as in Word: http, https and
    plain mailto links are clickable, other targets stay text as 'text (url)'."""
    words: list[PdfWord | None] = []
    word: PdfWord | None = None
    space: tuple[PdfStyle, str] | None = None
    for run in runs:
        text, link = word_link(run)
        font = PDF_FONTS[run.bold, run.italic, run.code]
        style = PdfStyle(font, size, link_color if link else color, link, run.strike)
        for token in SPACE_RE.split(pdf_safe_text(text)):
            if token == "\n":
                words.append(None)
                word = space = None
            elif token.startswith(" "):
                word, space = None, space or (style, " ")
            elif token:
                if word is None:
                    space_width = text_width(space[0].font, size, " ") if space else 0.0
                    word = PdfWord(space=space, space_width=space_width)
                    words.append(word)
                    space = None
                word.parts.append((style, token))
                word.width += text_width(font, size, token)
    return words


def _pdf_line(parts: list[tuple[PdfStyle, str]]) -> PdfLine:
    merged: list[tuple[PdfStyle, str]] = []
    for style, text in parts:  # neighbours in the same style become one piece of text
        if merged and merged[-1][0] == style:
            merged[-1] = (style, merged[-1][1] + text)
        else:
            merged.append((style, text))
    return PdfLine(merged, sum(text_width(style.font, style.size, text) for style, text in merged))


def wrap_words(words: list[PdfWord | None], width: float) -> list[PdfLine]:
    """Fill lines up to `width` points; a word longer than a line is broken between characters."""
    lines: list[list[tuple[PdfStyle, str]]] = [[]]
    used, width = 0.0, width + 0.01  # rounding: a cell as wide as its text keeps it on one line
    for word in words:
        if word is None:
            lines.append([])
            used = 0.0
            continue
        space = word.space_width if lines[-1] and word.space else 0.0
        if lines[-1] and used + space + word.width <= width:
            if word.space:
                lines[-1].append(word.space)
            lines[-1].extend(word.parts)
            used += space + word.width
            continue
        if lines[-1]:
            lines.append([])
        if word.width <= width:
            lines[-1].extend(word.parts)
            used = word.width
            continue
        used = 0.0
        for style, text in word.parts:
            widths, start = PDF_WIDTHS[style.font], 0
            for i, code in enumerate(text.encode("cp1252")):
                char_width = widths[code] * style.size / 1000
                if used + char_width > width and used > 0:  # at least one character per line
                    if i > start:
                        lines[-1].append((style, text[start:i]))
                    lines.append([])
                    used, start = 0.0, i
                used += char_width
            lines[-1].append((style, text[start:]))
    return [_pdf_line(parts) for parts in lines]


def _numeric_column(rows: list[list[str]], column: int) -> bool:
    """True if every filled cell below the header is a number, as the Excel export reads them."""
    values = [text for row in rows[1:] if (text := to_plain(row[column]))]
    return bool(values) and all(not isinstance(parse_number(text)[0], str) for text in values)


def _pdf_literal(text: str) -> bytes:
    """A PDF string of cp1252 text."""
    data = text.encode("cp1252")
    return b"(" + data.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def _pdf_unicode(text: str) -> bytes:
    """A PDF text string in UTF-16 (document information such as the title)."""
    text = "".join(char for char in text if unicodedata.category(char) != "Cc")
    return b"<FEFF" + text.encode("utf-16-be", "replace").hex().upper().encode() + b">"


def _pdf_date(moment: datetime) -> bytes:
    offset = moment.utcoffset()
    zone = ""
    if offset is not None:
        minutes = int(offset.total_seconds() // 60)
        sign, hours, minutes = "+" if minutes >= 0 else "-", abs(minutes) // 60, abs(minutes) % 60
        zone = f"{sign}{hours:02d}'{minutes:02d}'" if hours or minutes else "Z"
    return _pdf_literal(moment.strftime("D:%Y%m%d%H%M%S") + zone)


def _rgb(color: str) -> str:
    return " ".join(f"{value / 255:.3f}".rstrip("0").rstrip(".") or "0" for value in bytes.fromhex(color))


class PdfWriter:
    """Lay out parsed blocks on pages and write them as a PDF: headings in the design of the Word export,
    wrapped text, lists, quotes, code blocks, tables whose header row is repeated on every page, clickable
    links, and a footer with the page number and an optional note."""

    def __init__(
        self,
        design: Design = DEFAULT_DESIGN,
        paper: str = "letter",
        footer: str = "",
        page_label: str = "Page {page} of {pages}",
    ) -> None:
        self.design, self.footer, self.page_label = design, footer, page_label
        self.page_width, self.page_height = PDF_PAPER[paper]
        self.left, self.right = PDF_SIDE_MARGIN, self.page_width - PDF_SIDE_MARGIN
        self.top, self.bottom = self.page_height - PDF_TOP_MARGIN, PDF_BOTTOM_MARGIN
        self.pages: list[list[bytes]] = []  # content stream operators per page
        self.links: list[list[list[Any]]] = []  # [x1, y1, x2, y2, url] per page
        self.fonts: set[str] = set()
        self.y = self.after = 0.0  # top of the free space; space still owed after the last block
        self._new_page()

    def render(self, blocks: list[Block], title: str = "", language: str = "en",
               created: datetime | None = None) -> bytes:  # fmt: skip
        for block in blocks:
            match block:
                case Heading(level, text):
                    size, italic, accent, before, after = PDF_HEADINGS[min(level, MAX_DOCX_HEADING_LEVEL)]
                    color = self.design.accent if accent else TEXT_COLOR
                    runs = [Run(to_plain(text), bold=True, italic=italic)]  # as in Word: no links in headings
                    keep = 2 * PDF_BODY_SIZE * PDF_LEADING  # with the first lines of the following text
                    self._paragraph(runs, size=size, color=color, before=before, after=after, keep=keep)
                case Paragraph(text):
                    self._paragraph(parse_inline(text))
                case List():
                    self._list(block, level=0, x=0)
                    self.after = 6
                case Quote(lines):
                    if text := "\n".join(lines).strip():
                        runs = parse_inline(text, Run("", italic=True))
                        self._paragraph(runs, color=MUTED_COLOR, indent=PDF_QUOTE_INDENT, before=6, after=10,
                                        bar=self.design.accent)  # fmt: skip
                case CodeBlock(code):
                    self._code(code)
                case Table():
                    self._table(block)
                case Rule():
                    self._start(8, 1)
                    self._stroke(self.left, self.right, self.y, 0.75, RULE_COLOR)
                    self.after = 8
        for number, (ops, links) in enumerate(zip(self.pages, self.links, strict=True), start=1):
            self.ops, self.page_links = ops, links
            self._footer(number, len(self.pages))
        return self._document(title, language, created or datetime.now())

    # --- layout

    def _new_page(self) -> None:
        self.ops: list[bytes] = []
        self.page_links: list[list[Any]] = []
        self.pages.append(self.ops)
        self.links.append(self.page_links)
        self.y = self.top

    def _start(self, before: float, needed: float = 0.0) -> None:
        """Space before a block (the larger of it and the previous block's space after; none at the top of
        a page), and a new page if the first `needed` points of the block do not fit."""
        gap = 0.0 if self.y >= self.top else max(before, self.after)
        needed = min(needed, self.top - self.bottom)
        if self.y - gap - needed < self.bottom and self.y < self.top:
            self._new_page()
            gap = 0.0
        self.y -= gap
        self.after = 0.0

    def _need(self, height: float) -> None:
        """A new page if `height` points do not fit on this one."""
        if self.y - height < self.bottom - 0.01 and self.y < self.top:
            self._new_page()

    def _paragraph(self, runs: list[Run], *, size: float = PDF_BODY_SIZE, color: str = TEXT_COLOR,
                   indent: float = 0.0, before: float = 0.0, after: float = 6.0, keep: float = 0.0,
                   marker: tuple[float, str, str] | None = None, bar: str | None = None) -> None:  # fmt: skip
        """Wrapped text. `marker` is (x, text, colour) of a list bullet or number in front of the first line,
        `bar` the colour of a quote's line on the left, `keep` the room needed after the first line."""
        height = size * PDF_LEADING
        lines = wrap_words(pdf_words(runs, size, color, self.design.accent), self.right - self.left - indent)
        self._start(before, height + keep)
        room = int((self.y - self.bottom) / height + 1e-6)  # lines that fit on this page
        if 0 < room < len(lines) and self.y < self.top:  # at least two lines on each side of a page break
            room = min(room, len(lines) - 2)
            if room < 2:
                room = 0
                self._new_page()
        for n, line in enumerate(lines):
            if n == room and n:
                self._new_page()
            self._need(height)
            if n == 0 and marker:
                x, text, marker_color = marker
                bullet = PdfLine([(PdfStyle("Helvetica", size, marker_color), text)])
                self._draw(bullet, self.left + x, self.y, height, size)
            if bar:
                self._fill(self.left, self.y - height, PDF_QUOTE_BAR, height, bar)
            self._draw(line, self.left + indent, self.y, height, size)
            self.y -= height
        self.after = after

    def _list(self, block: List, level: int, x: float) -> None:
        """Bullets (• and – in the accent colour) and numbers as in Word; each list counts from its start."""
        count = range(block.start, block.start + len(block.items))
        markers = [f"{n}." for n in count] if block.ordered else ["–" if level else "•"] * len(block.items)
        color = TEXT_COLOR if block.ordered else self.design.accent
        hang = max(PDF_LIST_INDENT, max(text_width("Helvetica", PDF_BODY_SIZE, m) for m in markers) + 6)
        for item, marker in zip(block.items, markers, strict=True):
            self._paragraph(parse_inline(item.text), indent=x + hang, after=3, marker=(x, marker, color))
            for sublist in item.sublists:
                self._list(sublist, level=1, x=x + hang)

    def _code(self, code: str) -> None:
        """Courier on a light background, whitespace kept, long lines broken."""
        size, pad = PDF_CODE_SIZE, PDF_CODE_PAD
        height, style = size * PDF_CODE_LEADING, PdfStyle("Courier", size, TEXT_COLOR)
        per_line = max(1, int((self.right - self.left - 2 * pad) / (0.6 * size)))
        lines = [
            raw[i : i + per_line]
            for raw in pdf_safe_text(code.expandtabs(4)).split("\n")
            for i in range(0, max(len(raw), 1), per_line)
        ]
        self._start(8, min(len(lines), 2) * height + 2 * pad)
        done = 0
        while done < len(lines):
            room = int((self.y - self.bottom - 2 * pad) / height)
            if room < 1 and self.y < self.top:
                self._new_page()
                continue
            chunk = lines[done : done + max(room, 1)]
            box = len(chunk) * height + 2 * pad
            self._fill(self.left, self.y - box, self.right - self.left, box, PDF_CODE_FILL)
            for n, text in enumerate(chunk):
                self._draw(PdfLine([(style, text)]), self.left + pad, self.y - pad - n * height, height, size)
            self.y -= box
            done += len(chunk)
        self.after = 8

    def _table(self, table: Table) -> None:
        """Header row in the accent colour with white bold text, every second row tinted, thin lines between
        rows; number columns right-aligned. A row is only split if it is taller than a page."""
        size = PDF_TABLE_SIZE if len(table.aligns) <= 6 else PDF_TABLE_SIZE_SMALL
        height, accent = size * PDF_TABLE_LEADING, self.design.accent
        aligns = [
            align or ("right" if _numeric_column(table.rows, column) else "left")
            for column, align in enumerate(table.aligns)
        ]
        words = [
            [pdf_words(parse_inline(text, Run("", bold=r == 0)), size, *colors) for text in row]
            for r, row in enumerate(table.rows)
            for colors in [(WHITE, WHITE) if r == 0 else (TEXT_COLOR, accent)]  # text and links
        ]
        widths = self._column_widths(words)
        rows = [
            [wrap_words(cell, width - 2 * PDF_CELL_PAD_X) for cell, width in zip(row, widths, strict=True)]
            for row in words
        ]
        heights = [max(map(len, row)) * height + 2 * PDF_CELL_PAD_Y for row in rows]
        header = rows[0] if heights[0] < (self.top - self.bottom) / 3 else None  # repeated on each page
        first = min(heights[1], 3 * height + 2 * PDF_CELL_PAD_Y) if len(heights) > 1 else 0.0
        self._start(6, heights[0] + first)  # the header row not alone at the bottom of a page
        for r, row in enumerate(rows):
            self._table_row(row, widths, aligns, size, self.design.row_fill(r), header if r else None)
        self._stroke(self.left, self.left + sum(widths), self.y, 1, accent)
        self.after = 10

    def _column_widths(self, cells: list[list[list[PdfWord | None]]]) -> list[float]:
        """Widths from the content within the text width: each column at least its longest word (up to
        30 % of the width) or, if its cells are short, their whole text; the rest is shared by how long the
        cells' lines are."""
        available, pad = self.right - self.left, 2 * PDF_CELL_PAD_X
        low, high = [pad] * len(cells[0]), [pad] * len(cells[0])
        for row in cells:
            for column, words in enumerate(row):
                line = 0.0
                for word in [*words, None]:
                    if word is None:
                        high[column] = max(high[column], line + pad)
                        line = 0.0
                    else:
                        line += (word.space_width if line else 0.0) + word.width
                        low[column] = max(low[column], min(word.width + pad, 0.3 * available))
        low = [b if b <= 0.2 * available else a for a, b in zip(low, high, strict=True)]  # short: one line
        if sum(high) <= available:  # the table spans the text width, like in Word
            return [width * available / sum(high) for width in high]
        if sum(low) >= available:
            return [width * available / sum(low) for width in low]
        share = (available - sum(low)) / (sum(high) - sum(low))
        return [a + (b - a) * share for a, b in zip(low, high, strict=True)]

    def _table_row(self, row: list[list[PdfLine]], widths: list[float], aligns: list[str], size: float,
                   fill: str | None, header: list[list[PdfLine]] | None) -> None:  # fmt: skip
        """One row; on a new page the header row comes first."""
        height, pad = size * PDF_TABLE_LEADING, PDF_CELL_PAD_Y
        header_height = max(map(len, header)) * height + 2 * pad if header else 0.0
        total, start, fresh = max(map(len, row)), 0, False
        while start < total:
            room, rest = int((self.y - self.bottom - 2 * pad) / height + 1e-6), total - start
            if room < rest and not fresh:
                fresh_room = int((self.top - self.bottom - header_height - 2 * pad) / height)
                if rest <= fresh_room or room < min(3, rest):  # split a row only if it is taller than a page
                    self._new_page()
                    if header:
                        lines = max(map(len, header))
                        self._row_part(header, widths, aligns, size, self.design.accent, 0, lines)
                    fresh = True
                    continue
            count = max(1, min(room, rest))
            self._row_part(row, widths, aligns, size, fill, start, count)
            start, fresh = start + count, False

    def _row_part(self, row: list[list[PdfLine]], widths: list[float], aligns: list[str], size: float,
                  fill: str | None, start: int, count: int) -> None:  # fmt: skip
        """Lines start ... start + count of the cells of a row."""
        height = size * PDF_TABLE_LEADING
        box = count * height + 2 * PDF_CELL_PAD_Y
        if fill:
            self._fill(self.left, self.y - box, sum(widths), box, fill)
        x = self.left
        for lines, width, align in zip(row, widths, aligns, strict=True):
            for n, line in enumerate(lines[start : start + count]):
                top = self.y - PDF_CELL_PAD_Y - n * height
                self._draw(line, x + PDF_CELL_PAD_X, top, height, size, width - 2 * PDF_CELL_PAD_X, align)
            x += width
        self.y -= box
        self._stroke(self.left, x, self.y, 0.5, RULE_COLOR)

    def _footer(self, number: int, total: int) -> None:
        style = PdfStyle("Helvetica", PDF_FOOTER_SIZE, MUTED_COLOR)
        label = pdf_safe_text(self.page_label.format(page=number, pages=total))
        width = text_width(style.font, style.size, label)
        self._show(self.right - width, PDF_FOOTER_Y, style, label)
        if self.footer:
            note = wrap_words(pdf_words([Run(self.footer)], style.size, style.color, style.color),
                              self.right - self.left - width - 18)[0]  # fmt: skip
            self._draw(note, self.left, PDF_FOOTER_Y + 0.718 * style.size, 0.925 * style.size, style.size)

    # --- drawing

    def _draw(self, line: PdfLine, x: float, top: float, height: float, size: float, width: float = 0.0,
              align: str = "left") -> None:  # fmt: skip
        """A line of text whose box starts at `top`; centred or right-aligned within `width`."""
        baseline = top - (height - 0.925 * size) / 2 - 0.718 * size  # Helvetica: ascent 0.718, descent 0.207
        if align in ("center", "right"):
            x += (width - line.width) / (2 if align == "center" else 1)
        for style, text in line.parts:
            advance = text_width(style.font, style.size, text)
            if text.strip():
                self._show(x, baseline, style, text)
            if style.url:
                self._fill(x, baseline - 0.12 * style.size, advance, 0.06 * style.size, style.color)
                box = [x, baseline - 0.25 * style.size, x + advance, baseline + 0.85 * style.size, style.url]
                last = self.page_links[-1] if self.page_links else None
                if last and last[4] == style.url and abs(last[2] - x) < 0.01 and abs(last[1] - box[1]) < 0.01:
                    last[2] = box[2]  # the next piece of the same link on this line
                else:
                    self.page_links.append(box)
            if style.strike:
                self._fill(x, baseline + 0.28 * style.size, advance, 0.06 * style.size, style.color)
            x += advance

    def _show(self, x: float, baseline: float, style: PdfStyle, text: str) -> None:
        self.fonts.add(style.font)
        self.ops.append(
            f"BT /{PDF_FONT_IDS[style.font]} {style.size:g} Tf {_rgb(style.color)} rg "
            f"{x:.2f} {baseline:.2f} Td ".encode()
            + _pdf_literal(text)
            + b" Tj ET"
        )

    def _fill(self, x: float, y: float, width: float, height: float, color: str) -> None:
        self.ops.append(f"{_rgb(color)} rg {x:.2f} {y:.2f} {width:.2f} {height:.2f} re f".encode())

    def _stroke(self, x1: float, x2: float, y: float, width: float, color: str) -> None:
        self.ops.append(f"{_rgb(color)} RG {width:g} w {x1:.2f} {y:.2f} m {x2:.2f} {y:.2f} l S".encode())

    # --- file

    def _document(self, title: str, language: str, created: datetime) -> bytes:
        objects: list[bytes] = [b"", b"", b""]  # catalog, page tree and document information come last

        def add(obj: bytes) -> int:
            objects.append(obj)
            return len(objects)

        font = b"<< /Type /Font /Subtype /Type1 /BaseFont /%s /Encoding /WinAnsiEncoding >>"
        fonts = b" ".join(
            b"/%s %d 0 R" % (PDF_FONT_IDS[name].encode(), add(font % name.encode()))
            for name in sorted(self.fonts)
        )
        resources = add(b"<< /Font << %s >> /ProcSet [/PDF /Text] >>" % fonts)
        media_box = f"[0 0 {self.page_width:g} {self.page_height:g}]".encode()
        kids = []
        for ops, links in zip(self.pages, self.links, strict=True):
            stream = zlib.compress(b"\n".join(ops))
            contents = add(
                b"<< /Length %d /Filter /FlateDecode >>\nstream\n%s\nendstream" % (len(stream), stream)
            )
            annotations = [
                add(
                    f"<< /Type /Annot /Subtype /Link /Rect [{x1:.2f} {y1:.2f} {x2:.2f} {y2:.2f}] "
                    f"/Border [0 0 0] /A << /Type /Action /S /URI /URI ".encode()
                    + _pdf_literal(quote(url, safe="!#$%&'()*+,-./:;=?@[]~_"))
                    + b" >> >>"
                )
                for x1, y1, x2, y2, url in links
            ]
            annots = b" /Annots [%s]" % b" ".join(b"%d 0 R" % a for a in annotations) if annotations else b""
            kids.append(add(
                b"<< /Type /Page /Parent 2 0 R /MediaBox %s /Resources %d 0 R /Contents %d 0 R%s >>"
                % (media_box, resources, contents, annots)
            ))  # fmt: skip
        objects[0] = (
            b"<< /Type /Catalog /Pages 2 0 R /Lang %s /ViewerPreferences << /DisplayDocTitle true >> >>"
            % _pdf_literal(language)
        )
        objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
            b" ".join(b"%d 0 R" % kid for kid in kids),
            len(kids),
        )
        objects[2] = b"<< /Title %s /Producer (Response Export for Open WebUI) /CreationDate %s >>" % (
            _pdf_unicode(title),
            _pdf_date(created),
        )
        out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = []
        for number, obj in enumerate(objects, start=1):
            offsets.append(len(out))
            out += b"%d 0 obj\n%s\nendobj\n" % (number, obj)
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
        out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
        out += b"trailer\n<< /Size %d /Root 1 0 R /Info 3 0 R >>\n" % (len(objects) + 1)
        out += b"startxref\n%d\n%%%%EOF\n" % xref
        return bytes(out)


def markdown_to_pdf(blocks: list[Block], design: Design = DEFAULT_DESIGN, paper: str = "letter",
                    footer: str = "", page_label: str = "Page {page} of {pages}", title: str = "",
                    language: str = "en", created: datetime | None = None) -> bytes:  # fmt: skip
    return PdfWriter(design, paper, footer, page_label).render(blocks, title, language, created)


# ------------------------------------------------------------------------------------- browser snippets --
# The snippets run inside an async function in the user's browser. They must return quickly: the server
# waits for the result with a timeout.


def _click_link_js(href: str, filename: str | None = None) -> str:
    """Click a temporary link; `href` is a JavaScript expression, `filename` turns it into a download."""
    download = f"\nlink.download = {json.dumps(filename)};" if filename is not None else ""
    return f"""
const link = document.createElement('a');
link.href = {href};{download}
document.body.appendChild(link);
link.click();
link.remove();
"""


def download_js(data: bytes, filename: str, mime: str) -> str:
    return f"""
const text = atob({json.dumps(base64.b64encode(data).decode())});
const bytes = new Uint8Array(text.length);
for (let i = 0; i < text.length; i++) bytes[i] = text.charCodeAt(i);
const url = URL.createObjectURL(new Blob([bytes], {{ type: {json.dumps(mime)} }}));
{_click_link_js("url", filename)}
setTimeout(() => URL.revokeObjectURL(url), 10000);
return true;
"""


def open_url_js(url: str) -> str:
    return f"""{_click_link_js(json.dumps(url))}
return true;
"""


# The interface language (Open WebUI keeps it in localStorage.locale) and the time zone of the browser.
BROWSER_SETTINGS_JS = """
let locale = '';
let zone = '';
try { locale = localStorage.getItem('locale') || ''; } catch {}
try { zone = Intl.DateTimeFormat().resolvedOptions().timeZone || ''; } catch {}
return [String(locale || navigator.language || ''), String(zone)];
"""


# ---------------------------------------------------------------------------------------------- Action --


class BrowserError(Exception):
    """The browser did not run a snippet: blocked by the page's CSP, tab closed or disconnected, no answer."""


def zone(name: str) -> tzinfo | None:
    """The IANA time zone, None if the name is empty or unknown."""
    try:
        return ZoneInfo(name.strip()) if name.strip() else None
    except (ZoneInfoNotFoundError, ValueError, OSError):  # OSError: on some systems for a name too long for a file
        return None


@dataclass
class Browser:
    """Thin wrapper around Open WebUI's event callbacks."""

    emitter: EventCallback | None
    caller: EventCallback | None

    async def notify(self, level: Literal["info", "success", "warning", "error"], text: str) -> None:
        if self.emitter is not None:
            await self.emitter({"type": "notification", "data": {"type": level, "content": text}})

    async def call_js(self, code: str) -> Any:
        """Run a snippet in the browser and return its result. A snippet that throws never calls back, a
        disconnected session answers {'error': ...}: both raise BrowserError."""
        if self.caller is None:
            raise BrowserError("__event_call__ is not available, cannot run code in the browser")
        try:
            result = await asyncio.wait_for(
                self.caller({"type": "execute", "data": {"code": code}}), BROWSER_TIMEOUT_S
            )
        except TimeoutError as e:
            raise BrowserError(f"the browser did not answer within {BROWSER_TIMEOUT_S} s") from e
        if isinstance(result, dict) and "error" in result:
            raise BrowserError(str(result["error"])[:200])
        return result

    async def run_js(self, code: str) -> None:
        """Run one of the export snippets, which all return true."""
        result = await self.call_js(code)
        if result is not True:
            raise BrowserError(f"unexpected result from the browser: {str(result)[:200]}")

    async def download(self, data: bytes, filename: str, mime: str) -> None:
        await self.run_js(download_js(data, filename, mime))

    async def settings(self) -> tuple[str, str]:
        """The interface language (e.g. 'de-DE') and time zone of the user's browser."""
        result = await self.call_js(BROWSER_SETTINGS_JS)
        if not (isinstance(result, list) and len(result) == 2):
            raise BrowserError(f"unexpected result from the browser: {str(result)[:200]}")
        return str(result[0] or ""), str(result[1] or "")


@dataclass
class Export:
    """What an export needs besides the content: texts in the user's language, time zone, browser and the
    sources of the message (message_sources)."""

    texts: dict[str, str]
    language: str
    timezone: tzinfo | None
    browser: Browser
    sources: list[tuple[str, str | None]] = field(default_factory=list)

    def now(self) -> datetime:
        return datetime.now(self.timezone)

    def with_sources(self, blocks: list[Block]) -> list[Block]:
        return add_sources(blocks, self.sources, self.texts["sources"])


class Action:
    class Valves(BaseModel):
        language: Literal["auto", "en", "de"] = Field(
            default="auto",
            description="Language of the messages, default file names, the PDF footer and the paper size of "
            "Word and PDF files (en: Letter, de: A4). auto = the user's interface language (English if there "
            "is no translation).",
        )
        timezone: str = Field(
            default="",
            description="IANA time zone for file name timestamps and the date in the PDF footer, "
            "e.g. 'Europe/Berlin'. Empty or unknown = the time zone of the user's browser.",
        )
        max_mailto_length: int = Field(
            default=1900,
            ge=0,
            description="Longest mailto link to hand to the mail program (Outlook on Windows handles about "
            "2000 characters). Longer answers are downloaded as an .eml draft instead.",
        )
        always_eml: bool = Field(
            default=False,
            description="Always download a formatted .eml draft instead of opening a plain-text mailto link.",
        )
        print_footer: bool = Field(
            default=True,
            description="Add a small 'created with the AI assistant' note with the date to the footer of "
            "PDF pages (the page number is always there).",
        )
        accent_color: str = Field(
            default=DEFAULT_ACCENT,
            description="Accent colour of Word and PDF files (headings, table header rows, links) as a hex "
            "value without '#', e.g. 1F4E79. Same as the valve of Office Documents.",
        )
        font: str = Field(
            default=BODY_FONT,
            description="Font of Word documents and Excel sheets, e.g. Calibri, Arial or Aptos. It has to be "
            "installed where the files are opened; otherwise Office shows a similar one. PDF files always use "
            "the standard PDF fonts Helvetica and Courier.",
        )
        word_template: str = Field(
            default="",
            description="Optional Word file (.docx or .dotx) on the Open WebUI server, e.g. a letterhead: "
            "Word documents get its header, footer, page size and margins, e.g. "
            "/app/backend/data/templates/letterhead.dotx. Not used for PDF files.",
        )

    def __init__(self) -> None:
        self.valves = self.Valves()

    @property
    def actions(self) -> list[dict[str, str]]:
        """Sub-actions shown as buttons (language-neutral labels, see ACTIONS)."""
        return [
            {"id": action_id, "name": label, "icon_url": icon} for action_id, (label, icon) in ACTIONS.items()
        ]

    async def action(
        self,
        body: dict[str, Any],
        __id__: str | None = None,
        __event_emitter__: EventCallback | None = None,
        __event_call__: EventCallback | None = None,
    ) -> None:
        browser = Browser(__event_emitter__, __event_call__)
        # Without a sub-action (e.g. an API call) Open WebUI passes the function's own ID: export to Word.
        action_id = __id__ if __id__ in ACTIONS else "docx"
        try:
            export = await self._export_context(browser)
        except BrowserError as e:
            log.warning("Response export: the browser cannot run the export: %s", e)
            language = self.valves.language if self.valves.language != "auto" else "en"
            await browser.notify("error", TRANSLATIONS[language]["browser_failed"])
            return
        message = find_message(body)
        content = clean_markdown(message_text(message), export.texts["image"])
        export.sources = message_sources(message)
        if not content:
            await browser.notify("warning", export.texts["empty_message"])
            return
        try:
            await getattr(self, f"_export_{action_id}")(content, export)
        except BrowserError as e:
            log.warning("Response export %r: the browser did not run it: %s", action_id, e)
            await browser.notify("error", export.texts["browser_failed"])
        except Exception:
            log.exception("Response export %r failed", action_id)
            await browser.notify("error", export.texts["export_failed"])

    async def _export_context(self, browser: Browser) -> Export:
        """Language and time zone from the valves; 'auto' and an empty time zone come from the browser."""
        language, timezone = self.valves.language, zone(self.valves.timezone)
        if language == "auto" or timezone is None:
            locale, browser_zone = await browser.settings()
            if language == "auto":
                code = locale.replace("_", "-").split("-")[0].strip().lower()
                language = code if code in TRANSLATIONS else "en"
            timezone = timezone or zone(browser_zone)
        return Export(TRANSLATIONS[language], language, timezone, browser)

    def _design(self) -> Design:
        accent = self.valves.accent_color.strip().lstrip("#")
        font = " ".join(self.valves.font.split())
        return Design(
            accent.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", accent) else DEFAULT_ACCENT,
            font if FONT_NAME_RE.fullmatch(font) else BODY_FONT,
        )

    async def _export_docx(self, content: str, export: Export) -> None:
        blocks = parse_blocks(content)

        def build() -> bytes:
            template = read_word_template(self.valves.word_template)
            content = export.with_sources(blocks)
            return markdown_to_docx(content, self._design(), template, export.texts["paper"])

        data = await asyncio.to_thread(build)
        await export.browser.download(data, f"{self._file_stem(blocks, export)}.docx", DOCX_MIME)

    async def _export_xlsx(self, content: str, export: Export) -> None:
        blocks, texts = parse_blocks(content), export.texts
        tables = tables_with_titles(blocks)
        if not tables:
            await export.browser.notify("warning", texts["no_tables"])
            return
        data = await asyncio.to_thread(tables_to_xlsx, tables, texts["sheet_name"], self._design().font,
                                       texts["decimal_mark"], texts["excel_date_format"])  # fmt: skip
        await export.browser.download(data, f"{self._file_stem(blocks, export)}.xlsx", XLSX_MIME)

    async def _export_pdf(self, content: str, export: Export) -> None:
        blocks = parse_blocks(content)
        footer = ""
        if self.valves.print_footer:
            footer = export.texts["print_footer"].format(
                date=export.now().strftime(export.texts["date_format"])
            )
        texts, design = export.texts, self._design()
        title = document_title(blocks) or texts["default_title"]

        def build() -> bytes:
            return markdown_to_pdf(export.with_sources(blocks), design, texts["paper"], footer,
                                   texts["page_number"], title, export.language, export.now())  # fmt: skip

        data = await asyncio.to_thread(build)
        await export.browser.download(data, f"{self._file_stem(blocks, export)}.pdf", PDF_MIME)

    async def _export_email(self, content: str, export: Export) -> None:
        texts, browser = export.texts, export.browser
        subject, blocks = split_email(parse_blocks(content))
        subject = subject or document_title(blocks)
        blocks = export.with_sources(blocks)
        body_text = blocks_to_text(blocks)
        if not self.valves.always_eml:
            url = build_mailto(subject, body_text)
            if len(url) <= self.valves.max_mailto_length:
                await browser.run_js(open_url_js(url))
                return
        data = await asyncio.to_thread(build_eml, subject or texts["default_subject"], blocks, body_text)
        await browser.download(data, f"{safe_filename(subject, texts['email_draft_name'])}.eml", EML_MIME)
        if not self.valves.always_eml:
            await browser.notify("info", texts["eml_downloaded"])

    @staticmethod
    def _file_stem(blocks: list[Block], export: Export) -> str:
        """File name without extension: first heading (or a default) plus a timestamp."""
        title = safe_filename(document_title(blocks), export.texts["default_title"])
        return f"{title} {export.now():%Y-%m-%d %H%M}"
