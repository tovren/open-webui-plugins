"""Tests for the Open WebUI plugins in the parent folder.

Run from anywhere (Python 3.11+):
    pip install python-docx openpyxl pydantic pytest pypdf
    python -I -m pytest -q path/to/tests

pypdf reads the PDF files of Response Export; without it those tests are skipped. Optional checks
run automatically when the tools are installed: Node.js (executes the generated browser snippets and
office-file-ui/loader.js against small DOM stubs), LibreOffice (converts the .docx to PDF; needs
pdftotext from poppler), qpdf (checks the PDF files) and pdftoppm from poppler (renders them).
"""

import ast
import asyncio
import base64
import email
import email.policy
import importlib.util
import io
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from types import ModuleType, SimpleNamespace
from typing import Any
from urllib.parse import parse_qs, urlsplit
from zoneinfo import ZoneInfo

import pytest
from docx import Document
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
STUB_JS = Path(__file__).resolve().parent / "browser_stub.js"
PLUGIN_FILES = ("response_export.py", "context_guard.py", "company_context.py", "reasoning_toggle.py")
ALLOWED_THIRD_PARTY = {"docx", "openpyxl", "pydantic"}
LAZY_IMPORTS = {"open_webui"}  # only imported inside functions, when running in Open WebUI
FRONTMATTER_KEYS = {
    "title",
    "author",
    "author_url",
    "version",
    "license",
    "required_open_webui_version",
    "description",
}


def plugin_path(file_name: str) -> Path:
    """Each plugin lives in its own folder, e.g. response-export/response_export.py."""
    return REPO / Path(file_name).stem.replace("_", "-") / file_name


def load_plugin(file_name: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(f"plugin_{Path(file_name).stem}", plugin_path(file_name))
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


rx = load_plugin("response_export.py")
cg = load_plugin("context_guard.py")
cc = load_plugin("company_context.py")
rt = load_plugin("reasoning_toggle.py")

SAMPLE = """<details type="reasoning" done="true"><summary>Thought for 3 seconds</summary>
SECRET REASONING
</details>
# Quarterly report C#

Intro line one
continues **bold *nested*** and `code` and [docs](https://example.org/a_(b)).
Steps:
1. First
   continued
2. Second
   - sub a
   - sub b
3. Third

Between lists.

1. Restart
2. Again

| Item | Amount | Share |
|:-----|-------:|:-----:|
| A \\| B | 1.234,56 € | -3,2 % |
| ragged |
| C | 2.000 | 1,5 | surplus |

> A quote
> second line

```bash
# not a heading
echo "<b>hi</b>"
```

<script>alert(1)</script> [bad](javascript:alert(1))
"""


class Recorder:
    """Collects the events an action/filter sends to Open WebUI; answers like the browser would."""

    def __init__(self, locale: str = "en-US", timezone: str = "Europe/Berlin", result: Any = True) -> None:
        self.events: list[dict[str, Any]] = []
        self.calls: list[dict[str, Any]] = []
        self.probes = 0  # requests for the browser's language and time zone
        self.locale, self.timezone, self.result = locale, timezone, result

    async def emit(self, event: dict[str, Any]) -> None:
        self.events.append(event)

    async def call(self, event: dict[str, Any]) -> Any:
        if event["data"]["code"] == rx.BROWSER_SETTINGS_JS:
            self.probes += 1
            return [self.locale, self.timezone]
        self.calls.append(event)
        return self.result

    def notifications(self) -> list[tuple[str, str]]:
        return [(e["data"]["type"], e["data"]["content"]) for e in self.events if e["type"] == "notification"]

    def code(self) -> str:
        assert len(self.calls) == 1, self.calls
        assert self.calls[0]["type"] == "execute"
        return self.calls[0]["data"]["code"]


def js_value(code: str, prefix: str) -> Any:
    """Decode the JSON literal that follows `prefix` in a generated JS snippet."""
    start = code.index(prefix) + len(prefix)
    return json.JSONDecoder().raw_decode(code, start)[0]


def make_filter(module: ModuleType, **valves: Any) -> Any:
    flt = module.Filter()
    flt.valves = flt.Valves(**valves)
    return flt


def run_action(
    action_id: str | None,
    content: Any = SAMPLE,
    with_call: bool = True,
    body: dict[str, Any] | None = None,
    recorder: Recorder | None = None,
    **valves: Any,
) -> Recorder:
    action = rx.Action()
    action.valves = action.Valves(**valves)
    recorder = recorder or Recorder()
    body = body or {
        "id": "m2",
        "messages": [
            {"id": "m1", "role": "user", "content": "Please write a report"},
            {"id": "m2", "role": "assistant", "content": content},
        ],
    }
    asyncio.run(
        action.action(
            body,
            __id__=action_id,
            __event_emitter__=recorder.emit,
            __event_call__=recorder.call if with_call else None,
        )
    )
    return recorder


def downloaded(recorder: Recorder) -> tuple[bytes, str, str]:
    code = recorder.code()
    data = base64.b64decode(js_value(code, "atob("))
    return data, js_value(code, "link.download = "), js_value(code, "{ type: ")


# ------------------------------------------------------------------------------------- all plugins --


@pytest.mark.parametrize("file_name", PLUGIN_FILES)
def test_frontmatter(file_name: str) -> None:
    """Replicates Open WebUI's frontmatter parser: no stray 'key:' lines in the docstring."""
    lines = (plugin_path(file_name)).read_text(encoding="utf-8").splitlines()
    assert lines[0] == '"""'
    pattern = re.compile(r"^\s*([a-z_]+):\s*(.*)\s*$", re.IGNORECASE)
    meta = {}
    for line in lines[1:]:
        if '"""' in line:
            break
        if match := pattern.match(line):
            meta[match[1]] = match[2]
    assert set(meta) == FRONTMATTER_KEYS
    assert re.fullmatch(r"\d+\.\d+\.\d+", meta["version"])
    assert meta["license"] == "MIT"
    assert meta["required_open_webui_version"] == "0.11.4"
    assert "Security & privacy" in "\n".join(lines[: lines.index('"""', 1)])


@pytest.mark.parametrize("file_name", PLUGIN_FILES)
def test_only_bundled_libraries(file_name: str) -> None:
    tree = ast.parse((plugin_path(file_name)).read_text(encoding="utf-8"))
    modules = {alias.name for node in ast.walk(tree) if isinstance(node, ast.Import) for alias in node.names}
    modules |= {node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom) and node.module}
    top_level = {name.split(".")[0] for name in modules}
    assert top_level <= set(sys.stdlib_module_names) | ALLOWED_THIRD_PARTY | LAZY_IMPORTS


# -------------------------------------------------------------------------------- response export --


def test_action_buttons_are_language_neutral() -> None:
    action = rx.Action()
    expected = ["Word (.docx)", "Excel (.xlsx)", "PDF", "E-Mail"]
    assert [a["id"] for a in action.actions] == ["docx", "xlsx", "pdf", "email"]
    assert [a["name"] for a in action.actions] == expected
    action.valves = action.Valves(language="de")
    assert [a["name"] for a in action.actions] == expected
    for item in action.actions:
        svg = base64.b64decode(item["icon_url"].split(",", 1)[1]).decode()
        assert item["icon_url"].startswith("data:image/svg+xml;base64,")
        assert 'fill="none"' in svg and 'stroke="currentColor"' in svg and 'stroke-width="1.5"' in svg


def test_docx_structure() -> None:
    data, filename, mime = downloaded(run_action("docx"))
    assert mime == rx.DOCX_MIME
    assert re.fullmatch(r"Quarterly report C# \d{4}-\d{2}-\d{2} \d{4}\.docx", filename)
    doc = Document(io.BytesIO(data))
    paragraphs = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    all_text = "\n".join(text for _, text in paragraphs)
    assert ("Heading 1", "Quarterly report C#") in paragraphs
    assert "SECRET REASONING" not in all_text and "Thought for" not in all_text
    assert ("Normal", '# not a heading\necho "<b>hi</b>"') in paragraphs
    assert ("List Number", "First continued") in paragraphs
    assert ("List Bullet 2", "sub a") in paragraphs
    assert ("Quote", "A quote") in paragraphs and ("Quote", "second line") in paragraphs
    intro = next(p for p in doc.paragraphs if p.text.startswith("Intro"))
    assert intro.text == "Intro line one\ncontinues bold nested and code and docs.\nSteps:"
    assert [r.text for r in intro.runs if r.bold] == ["bold ", "nested"]
    assert [r.text for r in intro.runs if r.italic] == ["nested"]
    assert [r.text for r in intro.runs if r.font.name == "Consolas"] == ["code"]
    (link,) = intro.hyperlinks  # a real hyperlink, not "docs (url)"
    assert link.text == "docs" and link.address == "https://example.org/a_(b)"
    assert link.runs[0].font.underline and str(link.runs[0].font.color.rgb) == rx.DEFAULT_ACCENT
    code_run = next(p for p in doc.paragraphs if p.text.startswith("# not")).runs[0]
    assert code_run.font.name == "Consolas"
    assert "<script>alert(1)</script> bad (javascript:alert(1))" in all_text  # plain text, harmless
    assert not any(p.hyperlinks for p in doc.paragraphs if p.text.startswith("<script>"))

    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Item", "Amount", "Share"]
    assert all(run.bold for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)
    header_fill = table.rows[0].cells[0]._tc.xpath("./w:tcPr/w:shd/@w:fill")
    assert header_fill == [rx.DEFAULT_ACCENT]
    assert table._tbl.xpath("./w:tblPr/w:tblBorders/w:bottom/@w:color") == [rx.DEFAULT_ACCENT]
    assert [c.text for c in table.rows[1].cells] == ["A | B", "1.234,56 €", "-3,2 %"]
    assert [c.text for c in table.rows[2].cells] == ["ragged", "", ""]
    assert [c.text for c in table.rows[3].cells] == ["C", "2.000", "1,5"]
    assert not any(run.bold for run in table.rows[1].cells[0].paragraphs[0].runs)


def test_docx_numbered_lists_restart() -> None:
    data, _, _ = downloaded(run_action("docx"))
    doc = Document(io.BytesIO(data))
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    num_ids = [p._p.pPr.numPr.numId.val for p in numbered]
    assert [p.text for p in numbered] == ["First continued", "Second", "Third", "Restart", "Again"]
    assert num_ids[0] == num_ids[1] == num_ids[2] != num_ids[3] == num_ids[4]
    numbering = doc.part.numbering_part.element
    for num_id in (num_ids[0], num_ids[3]):
        num = numbering.num_having_numId(num_id)
        assert num.xpath("./w:lvlOverride/w:startOverride/@w:val") == ["1"]


def test_docx_list_starting_number_is_kept() -> None:
    data, _, _ = downloaded(run_action("docx", "Intro\n\n3. three\n4. four"))
    doc = Document(io.BytesIO(data))
    item = next(p for p in doc.paragraphs if p.text == "three")
    num = doc.part.numbering_part.element.num_having_numId(item._p.pPr.numPr.numId.val)
    assert num.xpath("./w:lvlOverride/w:startOverride/@w:val") == ["3"]


@pytest.mark.parametrize(
    ("text", "value", "number_format"),
    [
        ("42", 42, None),
        ("2026", 2026, None),  # a year or postcode keeps its digits as written
        ("70173", 70173, None),
        ("-17", -17, None),
        ("1,5", 1.5, None),
        ("0,5", 0.5, None),
        ("3.14159", 3.14159, None),
        ("2.000", 2000, "#,##0"),
        ("1.234.567", 1234567, "#,##0"),
        ("1.234,56", 1234.56, "#,##0.00"),
        ("1,234.56", 1234.56, "#,##0.00"),
        ("1 234,5", 1234.5, "#,##0.0"),
        ("1 000", 1000, "#,##0"),
        ("1'000'000", 1000000, "#,##0"),
        ("−5", -5, None),
        ("12%", 0.12, "0%"),
        ("-3,2 %", -0.032, "0.0#%"),
        ("1.234,56 €", 1234.56, '#,##0.00 "€"'),
        ("€ -12,50", -12.5, '#,##0.00 "€"'),
        ("$1,234.50", 1234.5, '"$"#,##0.00'),
        ("1.234,56 EUR", 1234.56, '#,##0.00 "€"'),
        ("USD 12", 12.0, '"$"#,##0.00'),
        ("£5 %", "£5 %", None),
        ("01067", "01067", None),
        ("1 2", "1 2", None),
        ("1.2.3", "1.2.3", None),
        ("12.10.2026", "12.10.2026", None),
        ("1.234,567.8", "1.234,567.8", None),
        ("+49 30 123456", "+49 30 123456", None),
        ("+49 711 123", "+49 711 123", None),  # phone numbers, not 49711123
        ("+1 234", "+1 234", None),
        ("+41\u00a079 123 456", "+41\u00a079 123 456", None),
        ("+5", 5, None),
        ("+1.234,56", 1234.56, "#,##0.00"),
        ("+1 000 €", 1000.0, '#,##0.00 "€"'),
        ("1 234", 1234, "#,##0"),
        ("CHF 1'234.50", 1234.5, '"CHF" #,##0.00'),
        ("1’234.50 CHF", 1234.5, '"CHF" #,##0.00'),  # Swiss thousands separator U+2019
        ("CHF -12", -12.0, '"CHF" #,##0.00'),
        ("1234567890123456", "1234567890123456", None),
        ("=SUM(A1:A2)", "=SUM(A1:A2)", None),
        ("%", "%", None),
        ("abc", "abc", None),
    ],
)
def test_parse_number(text: str, value: Any, number_format: str | None) -> None:
    result, fmt = rx.parse_number(text)
    assert fmt == number_format
    assert result == pytest.approx(value) if isinstance(value, float) else result == value
    assert type(result) is type(value)


@pytest.mark.parametrize(
    ("text", "decimal", "value", "number_format"),
    [
        ("1.085", ".", 1.085, None),
        ("1.085", ",", 1085, "#,##0"),
        ("1,234", ",", 1.234, None),
        ("1,234", ".", 1234, "#,##0"),
        ("-3.125 %", ".", -0.03125, "0.0#%"),
        ("2.000 €", ",", 2000.0, '#,##0.00 "€"'),
        ("0.125", ",", 0.125, None),  # never a thousands separator after a single 0
        ("1.234.567", ".", 1234567, "#,##0"),  # unambiguous: the decimal mark does not matter
        ("1.234,5", ".", 1234.5, "#,##0.0"),
    ],
)
def test_parse_number_with_a_decimal_mark(
    text: str, decimal: str, value: Any, number_format: str | None
) -> None:
    result, fmt = rx.parse_number(text, decimal)
    assert fmt == number_format
    assert result == pytest.approx(value) if isinstance(value, float) else result == value
    assert type(result) is type(value)


def xlsx_sheets(content: str, **valves: Any) -> Any:
    data, _, _ = downloaded(run_action("xlsx", content, **valves))
    return load_workbook(io.BytesIO(data))


def column_values(sheet: Any, column: str) -> list[Any]:
    return [cell.value for cell in sheet[column][1:]]


def test_xlsx_ambiguous_numbers_follow_their_column() -> None:
    content = (
        "| Rate | Amount | Mixed | Text |\n|---|---|---|---|\n| 1.085 | 2.000 | 1.250 | a |\n"
        "| 0.125 | 15.000 | 1,5 | 1.085 |\n| 1.25 | 1.250.000 | 2.5 | b |\n| 3.125 | 800 | | |"
    )
    for language in ("en", "de"):  # the column decides, not the language
        sheet = xlsx_sheets(content, language=language).active
        assert column_values(sheet, "A") == pytest.approx([1.085, 0.125, 1.25, 3.125])
        assert column_values(sheet, "B") == [2000, 15000, 1250000, 800]
        assert sheet["B2"].number_format == "#,##0"
    # columns that show no or both decimal marks: the one of the language
    assert xlsx_sheets(content, language="en").active["C2"].value == pytest.approx(1.25)
    assert xlsx_sheets(content, language="de").active["C2"].value == 1250
    assert xlsx_sheets(content, language="en").active["D3"].value == pytest.approx(1.085)
    assert xlsx_sheets(content, language="de").active["D3"].value == 1085
    assert xlsx_sheets("| x |\n|---|\n| 1,234 |", language="de").active["A2"].value == pytest.approx(1.234)
    assert xlsx_sheets("| x |\n|---|\n| 1,234 |", language="en").active["A2"].value == 1234


@pytest.mark.parametrize(("language", "number_format"), [("en", "yyyy-mm-dd"), ("de", "dd.mm.yyyy")])
def test_xlsx_dates(language: str, number_format: str) -> None:
    content = (
        "| Date |\n|---|\n| 15.03.2024 |\n| 1.3.2024 |\n| 2024-03-12 |\n| 03/12/2024 |\n"
        "| 31.02.2024 |\n| 01.01.1899 |\n| 2024-3-12 |"
    )
    sheet = xlsx_sheets(content, language=language).active
    dates = [datetime(2024, 3, 15), datetime(2024, 3, 1), datetime(2024, 3, 12)]
    assert column_values(sheet, "A") == [*dates, "03/12/2024", "31.02.2024", "01.01.1899", "2024-3-12"]
    assert all(cell.is_date and cell.number_format == number_format for cell in sheet["A"][1:4])
    assert all(cell.data_type == "s" for cell in sheet["A"][4:])  # month or day first, impossible, too early
    assert rx.parse_date("12.10.2026") == rx.date(2026, 10, 12) and rx.parse_date("2026") is None


def test_xlsx_links() -> None:
    content = (
        "| Name | Link |\n|---|---|\n| [Docs](https://example.org/d) | **[bold](http://example.org)** |\n"
        "| see [a](https://a.example) here | <mailto:x@example.org> |\n"
        "| [js](javascript:alert(1)) | [attach](mailto:x@example.org?attach=/etc/passwd) |\n"
        "| [Header link](https://example.org) | ![logo](https://example.org/l.png) |"
    )
    sheet = xlsx_sheets(content).active
    cells = [cell for row in sheet for cell in row if cell.hyperlink]
    links = {cell.coordinate: (cell.value, cell.hyperlink.target) for cell in cells}
    assert links == {
        "A2": ("Docs", "https://example.org/d"),
        "B2": ("bold", "http://example.org"),
        "B3": ("mailto:x@example.org", "mailto:x@example.org"),
        "A5": ("Header link", "https://example.org"),
        "B5": ("logo", "https://example.org/l.png"),
    }
    assert sheet["A2"].font.u == "single" and sheet["A2"].font.color.rgb.endswith(rx.HYPERLINK_COLOR)
    assert sheet["A2"].font.name == "Calibri" and sheet["A2"].style == "Hyperlink"
    assert sheet["A3"].value == "see a (https://a.example) here"
    assert sheet["A4"].value == "js (javascript:alert(1))"
    assert sheet["B4"].value == "attach (mailto:x@example.org?attach=/etc/passwd)"
    header = xlsx_sheets("| [Name](https://example.org) |\n|---|\n| x |").active["A1"]
    assert header.value == "Name (https://example.org)" and not header.hyperlink and header.font.bold


def test_xlsx_sheet_names_from_headings() -> None:
    content = (
        "# Report\n\n## Prices: 2024/25 [EUR]\n\n| a |\n|---|\n| 1 |\n\nText\n\n| b |\n|---|\n| 2 |\n\n"
        "## Prices: 2024/25 [EUR]\n| c |\n|---|\n| 3 |\n\n"
        "### Quarterly figures of all regions\n\n| d |\n|---|\n| 4 |\n\n"
        "## 'History'\n\n| e |\n|---|\n| 5 |\n\n## Table 2\n\n| f |\n|---|\n| 6 |\n\n"
        "## *?\n\n| g |\n|---|\n| 7 |"
    )
    names = xlsx_sheets(content).sheetnames
    assert names == [
        "Prices 2024 25 EUR",
        "Table 2",
        "Prices 2024 25 EUR (2)",
        "Quarterly figures of all region",  # at most 31 characters
        "History (2)",  # Excel keeps "History" for itself
        "Table 2 (2)",
        "Table 7",  # nothing left of the heading
    ]
    assert all(len(name) <= 31 for name in names)
    assert xlsx_sheets(content, language="de").sheetnames[1] == "Tabelle 2"


def test_xlsx_values_and_formats() -> None:
    content = SAMPLE + (
        "\n\n| Code | Formula | Total |\n|---|---|---|\n| 01067 | =1+1 | 12% |\n"
        "\n```\n| in | code |\n|---|---|\n| x | y |\n```\n"
    )
    data, filename, mime = downloaded(run_action("xlsx", content))
    assert mime == rx.XLSX_MIME and filename.endswith(".xlsx")
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Table 1", "Table 2"]
    sheet = workbook["Table 1"]
    assert [c.value for c in sheet[1]] == ["Item", "Amount", "Share"]
    assert all(c.font.bold and c.fill.fgColor.rgb.endswith("DDE4EE") for c in sheet[1])
    assert {c.font.name for row in sheet.iter_rows() for c in row} == {"Calibri"}
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C4"
    assert sheet["A2"].value == "A | B"
    assert sheet["B2"].value == pytest.approx(1234.56) and sheet["B2"].number_format == '#,##0.00 "€"'
    assert sheet["C2"].value == pytest.approx(-0.032) and sheet["C2"].number_format == "0.0#%"
    assert sheet["B3"].value is None and sheet["B4"].value == 2000 and sheet["C4"].value == 1.5
    assert sheet.max_column == 3  # surplus cell of a ragged row is dropped like in GFM
    assert 8 <= sheet.column_dimensions["B"].width <= 60
    second = workbook["Table 2"]
    assert second["A2"].value == "01067" and second["A2"].data_type == "s"
    assert second["B2"].value == "=1+1" and second["B2"].data_type == "s"  # text, not a formula
    assert second["C2"].value == pytest.approx(0.12) and second["C2"].number_format == "0%"


def test_xlsx_german_sheet_names() -> None:
    data, _, _ = downloaded(run_action("xlsx", language="de"))
    assert load_workbook(io.BytesIO(data)).sheetnames == ["Tabelle 1"]


@pytest.mark.parametrize(("language", "text"), [("en", "does not contain a table"), ("de", "keine Tabelle")])
def test_xlsx_without_table_warns(language: str, text: str) -> None:
    recorder = run_action("xlsx", "# Just text", language=language)
    assert not recorder.calls
    assert recorder.notifications()[0][0] == "warning" and text in recorder.notifications()[0][1]


EMAIL = """Sure, here is a draft for you:

**Subject:** Meeting on Friday

Dear Ms. Smith,

thank you for your message. Agenda:
- Budget
- Timeline

Best regards"""


def test_email_mailto() -> None:
    recorder = run_action("email", EMAIL)
    url = js_value(recorder.code(), "link.href = ")
    assert url.startswith("mailto:?")
    query = parse_qs(urlsplit(url).query)
    assert query["subject"] == ["Meeting on Friday"]
    body = query["body"][0]
    assert body.startswith("Dear Ms. Smith,\r\n\r\nthank you")
    assert "\r\n• Budget\r\n• Timeline\r\n" in body
    assert "here is a draft" not in body
    assert not recorder.notifications()


@pytest.mark.parametrize(("language", "text"), [("en", "too long"), ("de", "zu lang")])
def test_email_eml_fallback(language: str, text: str) -> None:
    content = EMAIL + "\n\n| Item | Cost |\n|---|---|\n| <b>A</b> | 5 € |\n" + "More text. " * 300
    recorder = run_action("email", content, language=language)
    data, filename, mime = downloaded(recorder)
    assert filename == "Meeting on Friday.eml" and mime == "message/rfc822"
    message = email.message_from_bytes(data, policy=email.policy.default)
    assert message["Subject"] == "Meeting on Friday"
    assert message["X-Unsent"] == "1"
    assert message.get_content_type() == "multipart/alternative"
    parts = {part.get_content_type(): part.get_content() for part in message.iter_parts()}
    assert set(parts) == {"text/plain", "text/html"}
    assert "here is a draft" not in parts["text/plain"] and "• Budget" in parts["text/plain"]
    assert '<table border="1"' in parts["text/html"] and "<li>Budget</li>" in parts["text/html"]
    assert "&lt;b&gt;A&lt;/b&gt;" in parts["text/html"] and "<b>A</b>" not in parts["text/html"]
    assert [kind for kind, _ in recorder.notifications()] == ["info"]
    assert text in recorder.notifications()[0][1]


def test_email_always_eml_without_subject_uses_defaults() -> None:
    recorder = run_action("email", "Just a body", always_eml=True, language="de")
    data, filename, _ = downloaded(recorder)
    assert filename == "E-Mail-Entwurf.eml"
    assert email.message_from_bytes(data, policy=email.policy.default)["Subject"] == "Entwurf"
    assert not recorder.notifications()


def test_email_subject_falls_back_to_heading() -> None:
    content = "```\nSubject: in code\n```\n# Offer *2026*\nText"
    blocks = rx.parse_blocks(content)
    assert rx.split_email(blocks) == ("", blocks)
    url = js_value(run_action("email", content).code(), "link.href = ")
    assert parse_qs(urlsplit(url).query)["subject"] == ["Offer 2026"]


def test_email_subject_line() -> None:
    blocks = rx.parse_blocks("Here you go:\n## Subject: Hi\nDear Bob,\nthanks")
    assert rx.split_email(blocks) == ("Hi", [rx.Paragraph("Dear Bob,\nthanks")])
    blocks = rx.parse_blocks("Draft:\nBetreff: Angebot\nHallo Bob")
    assert rx.split_email(blocks) == ("Angebot", [rx.Paragraph("Hallo Bob")])
    blocks = rx.parse_blocks("#Subject: not a heading")
    assert rx.split_email(blocks)[0] == ""


# As Open WebUI 0.11.4 to 0.12.3 send them with the clicked message (chatActionHandler in Chat.svelte)
SOURCES = [
    {
        "source": {"id": "file-1", "name": "Handbook.pdf", "type": "file"},
        "document": ["chunk 1", "chunk 2", "chunk 3"],
        "metadata": [
            {"source": "file-1", "name": "Handbook%20v2.pdf"},
            {"source": "file-1"},
            {"source": "file-2"},
        ],
        "distances": [0.8, 0.7, 0.6],
    },
    {
        "source": {"type": "web_search", "name": "web search", "urls": ["https://example.org/a b"]},
        "document": ["result"],
        "metadata": [{"source": "https://example.org/a b", "title": "A"}],
    },
    {"source": {"name": "Weather tool"}, "document": ["sunny"], "metadata": [{}]},
    {"source": {"id": "wiki", "name": "Wiki", "url": "https://wiki.example/x"}, "document": ["text"]},
    {
        "source": {"name": "News"},
        "document": ["n"],
        "metadata": [{"source": "https://news.example/1", "external": True}],
    },
    {},
    "broken",
    {"source": {"name": "no documents"}},
]


def sources_body(content: str, sources: Any = SOURCES) -> dict[str, Any]:
    message = {"id": "m2", "role": "assistant", "content": content, "sources": sources}
    return {"id": "m2", "messages": [message]}


def test_message_sources_are_numbered_like_open_webui() -> None:
    assert rx.message_sources(sources_body("")["messages"][0]) == [
        ("Handbook v2.pdf", None),  # two chunks of one file are one source
        ("Handbook.pdf", None),  # another file of the same source
        ("https://example.org/a b", "https://example.org/a b"),
        ("Weather tool", None),
        ("Wiki", "https://wiki.example/x"),
        ("News", "https://news.example/1"),  # an external source keeps its name (0.12.3)
    ]
    assert rx.message_sources({}) == [] and rx.message_sources({"sources": "x"}) == []


@pytest.mark.parametrize(("language", "label"), [("en", "Sources"), ("de", "Quellen")])
def test_sources_are_listed(language: str, label: str) -> None:
    content = "# Report\n\nSales rose [1][2, 3] and fell [4#page=2].\n\n| a | b |\n|---|---|\n| [5] | 1 |"
    body = sources_body(content)
    doc = Document(io.BytesIO(downloaded(run_action("docx", body=body, language=language))[0]))
    texts = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert texts[2:] == [
        ("Normal", label),
        *[("List Number", text) for text in ("Handbook v2.pdf", "Handbook.pdf", "https://example.org/a%20b",
                                             "Weather tool", "Wiki – https://wiki.example/x",
                                             "News – https://news.example/1")],
    ]  # fmt: skip
    assert all(run.bold for run in next(p for p in doc.paragraphs if p.text == label).runs)
    assert [link.address for p in doc.paragraphs for link in p.hyperlinks] == [
        "https://example.org/a%20b", "https://wiki.example/x", "https://news.example/1"
    ]  # fmt: skip
    assert downloaded(run_action("docx", body=body))[1].startswith("Report ")  # still the title
    if pypdf is not None:
        data, _, _ = downloaded(run_action("pdf", body=body, language=language))
        text = " ".join(page_texts(data)[0].split())
        assert f"{label} 1. Handbook v2.pdf 2. Handbook.pdf" in text
        assert "6. News – https://news.example/1" in text
        assert link_targets(data)[-1] == "https://news.example/1"
    url = js_value(run_action("email", body=body, language=language).code(), "link.href = ")
    mail_body = parse_qs(urlsplit(url).query)["body"][0]
    assert f"\r\n\r\n{label}\r\n\r\n1. Handbook v2.pdf\r\n2. Handbook.pdf\r\n" in mail_body


@pytest.mark.parametrize(
    ("content", "sources"),
    [
        ("No markers here.", SOURCES),
        ("Only in code: `a[1]`\n\n```\nb[2]\n```", SOURCES),
        ("A footnote[^1] and a [link](https://example.org) and [x]", SOURCES),
        ("Cited [1] without sources", None),
        ("Cited [1] with empty sources", [{}]),
    ],
)
def test_no_sources_list_without_citations_or_sources(content: str, sources: Any) -> None:
    doc = Document(io.BytesIO(downloaded(run_action("docx", body=sources_body(content, sources)))[0]))
    assert not any(p.style.name == "List Number" for p in doc.paragraphs)
    assert "Sources" not in [p.text for p in doc.paragraphs]


@pytest.mark.parametrize(
    ("content", "listed"),
    [("Im Jahr [2024] wuchs der Umsatz.", False), ("[0] and [1, 3]", False), ("Er wuchs [2].", True)],
)
def test_only_numbers_of_sources_are_citations(content: str, listed: bool) -> None:
    body = sources_body(content, SOURCES[:1])  # two sources
    doc = Document(io.BytesIO(downloaded(run_action("docx", body=body))[0]))
    assert any(p.style.name == "List Number" for p in doc.paragraphs) == listed


def test_emphasis_across_lines() -> None:
    blocks = rx.parse_blocks("**bold\ncontinued** x")
    assert rx.blocks_to_html(blocks) == "<p><strong>bold<br>continued</strong> x</p>"
    assert rx.blocks_to_text(blocks) == "bold\ncontinued x"


def test_list_of_parts_content() -> None:
    content = [
        {"type": "text", "text": "# Parts title\n\n| a | b |\n|---|---|\n| 1 | 2 |"},
        {"type": "image_url", "image_url": {"url": "data:image/png;base64,AAAA"}},
        {"type": "text", "text": "Closing line"},
    ]
    data, filename, _ = downloaded(run_action("docx", content))
    assert filename.startswith("Parts title ")
    doc = Document(io.BytesIO(data))
    assert "Closing line" in [p.text for p in doc.paragraphs]
    assert doc.tables[0].rows[1].cells[1].text == "2"


def test_clicked_message_is_exported() -> None:
    body = {
        "id": "a1",
        "messages": [
            {"id": "a1", "role": "assistant", "content": "# First"},
            {"id": "a2", "role": "assistant", "content": "# Second"},
        ],
    }
    assert downloaded(run_action("docx", body=body))[1].startswith("First ")


@pytest.mark.parametrize("content", ["", "   ", None, [], '<details type="reasoning">x</details>'])
def test_empty_message_warns(content: Any) -> None:
    recorder = run_action("docx", content)
    assert not recorder.calls
    assert recorder.notifications() == [("warning", rx.TRANSLATIONS["en"]["empty_message"])]


def test_export_error_shows_toast(monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture) -> None:
    def broken(*_args: Any) -> bytes:
        raise ValueError("boom")

    monkeypatch.setattr(rx, "markdown_to_docx", broken)
    recorder = run_action("docx", language="de")
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["de"]["export_failed"])]
    assert "boom" in caplog.text


def test_missing_event_call_shows_toast() -> None:
    recorder = run_action("pdf", with_call=False)
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["en"]["browser_failed"])]
    recorder = run_action("pdf", with_call=False, language="de", timezone="Europe/Berlin")
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["de"]["browser_failed"])]


def test_table_edge_cases() -> None:
    blocks = rx.parse_blocks("Text with | a pipe\n---\n\n| `a\\|b` | c |\n| - | - |\n| 1 |\n")
    assert isinstance(blocks[0], rx.Paragraph) and isinstance(blocks[1], rx.Rule)  # not a table
    table = blocks[2]
    assert table.rows == [["`a|b`", "c"], ["1", ""]]
    assert rx.to_plain(table.rows[0][0]) == "a|b"


def test_filename_helpers() -> None:
    assert rx.safe_filename('a/b\\c:d*e?f"g<h>i|j\nk', "x") == "a b c d e f g h i j k"
    assert rx.safe_filename("  ...  ", "fallback") == "fallback"
    assert rx.safe_filename("x" * 100, "f") == "x" * rx.MAX_FILENAME_CHARS
    assert (
        rx.document_title(rx.parse_blocks("```\n# shell comment\n```\n## Real **title** ##")) == "Real title"
    )


def export_context(recorder: Recorder | None = None, **valves: Any) -> Any:
    action = rx.Action()
    action.valves = action.Valves(**valves)
    recorder = recorder or Recorder()
    return asyncio.run(action._export_context(rx.Browser(recorder.emit, recorder.call))), recorder


def test_timezone_valve() -> None:
    export, recorder = export_context(timezone=" Asia/Tokyo ", language="de")
    assert export.now().utcoffset().total_seconds() == 9 * 3600 and recorder.probes == 0
    export, recorder = export_context(timezone="Not/AZone", language="de")  # unknown: the browser's zone
    assert export.now().utcoffset() == datetime.now(ZoneInfo("Europe/Berlin")).utcoffset() and recorder.probes == 1
    # empty: the browser's time zone, asked together with the language
    export, recorder = export_context(Recorder(timezone="America/New_York"), language="de")
    assert str(export.timezone) == "America/New_York" and recorder.probes == 1
    export, _ = export_context(Recorder(timezone="Not/AZone"), language="de")
    assert export.timezone is None and isinstance(export.now(), datetime)  # server time


@pytest.mark.parametrize(
    ("locale", "language"),
    [("de-DE", "de"), ("de", "de"), ("de_CH", "de"), ("en-GB", "en"), ("fr-FR", "en"), ("", "en")],
)
def test_language_auto_follows_the_browser(locale: str, language: str) -> None:
    export, recorder = export_context(Recorder(locale=locale), timezone="UTC")
    assert export.language == language and export.texts is rx.TRANSLATIONS[language] and recorder.probes == 1
    assert rx.Action.Valves().language == "auto"


def test_language_auto_in_exports() -> None:
    data, filename, _ = downloaded(run_action("xlsx", recorder=Recorder(locale="de-DE")))
    assert load_workbook(io.BytesIO(data)).sheetnames == ["Tabelle 1"]
    recorder = run_action("xlsx", "# Just text", recorder=Recorder(locale="de-DE"))
    assert recorder.notifications() == [("warning", rx.TRANSLATIONS["de"]["no_tables"])]


@pytest.mark.parametrize(
    "answer",
    [
        {"error": "Client session disconnected."},  # Open WebUI's answer for a closed tab
        {"error": "Event call timed out. The browser tab may be inactive or closed."},
        None,
        False,
    ],
)
def test_browser_errors_show_a_message(answer: Any, caplog: pytest.LogCaptureFixture) -> None:
    recorder = run_action("docx", recorder=Recorder(result=answer), language="de", timezone="UTC")
    assert len(recorder.calls) == 1
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["de"]["browser_failed"])]
    assert "Content-Security-Policy" in rx.TRANSLATIONS["de"]["browser_failed"]
    assert "Content Security Policy" in rx.TRANSLATIONS["en"]["browser_failed"]
    assert "did not run it" in caplog.text


def test_browser_that_never_answers_times_out(monkeypatch: pytest.MonkeyPatch) -> None:
    """A snippet that throws in the browser (e.g. CSP blocks new Function) never calls back."""
    monkeypatch.setattr(rx, "BROWSER_TIMEOUT_S", 0.05)
    recorder = Recorder()

    async def hanging(event: dict[str, Any]) -> Any:
        await asyncio.sleep(10)

    action = rx.Action()
    asyncio.run(
        action.action({"messages": [{"role": "assistant", "content": "x"}]}, "docx", recorder.emit, hanging)
    )
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["en"]["browser_failed"])]


def test_browser_settings_error_stops_early() -> None:
    recorder = Recorder()

    async def disconnected(event: dict[str, Any]) -> Any:
        recorder.calls.append(event)
        return {"error": "Client session disconnected."}

    asyncio.run(rx.Action().action({"messages": []}, "pdf", recorder.emit, disconnected))
    assert len(recorder.calls) == 1 and recorder.calls[0]["data"]["code"] == rx.BROWSER_SETTINGS_JS
    assert recorder.notifications() == [("error", rx.TRANSLATIONS["en"]["browser_failed"])]


@pytest.mark.parametrize("action_id", ["response_export", "", None])
def test_without_sub_action_exports_word(action_id: str | None) -> None:
    """Open WebUI passes the function ID as __id__ when there is no sub-action."""
    _, filename, mime = downloaded(run_action(action_id))
    assert mime == rx.DOCX_MIME and filename.endswith(".docx")


def test_nested_details_are_removed() -> None:
    text = (
        '<details type="tool_calls"><summary>Tool</summary>outer <details type="reasoning">inner'
        "</details> tail</details>Answer <think>a <think>b</think> c</think>end"
    )
    assert rx.clean_markdown(text) == "Answer end"
    assert rx.clean_markdown("<details>open without end\nText") == "<details>open without end\nText"
    text = '<thinking>a ` b</thinking>A <details type="code_interpreter">\n```python\nx\n```\n</details>B'
    assert rx.clean_markdown(text) == "A B"  # a backtick in the reasoning does not keep it


def test_open_webui_blocks_are_removed_after_a_lone_backtick() -> None:
    text = (
        'Ich prüfe das kurz, it`s quick:\n<details type="tool_calls" id="x" done="true" arguments="{}" '
        'result="&quot;use `ls`&quot;">\n<summary>Tool Executed</summary>\n</details>\nErgebnis'
    )
    assert rx.clean_markdown(text) == "Ich prüfe das kurz, it`s quick:\n\nErgebnis"


def test_closing_tag_in_code_does_not_close_details() -> None:
    text = "<details>\n<summary>Hint</summary>\n\nClose it with `</details>`.\n</details>\n\nAfter"
    assert rx.parse_blocks(rx.clean_markdown(text)) == [
        rx.Paragraph("**Hint**"),
        rx.Paragraph("Close it with `</details>`."),
        rx.Paragraph("After"),
    ]


def test_code_keeps_tags_and_images() -> None:
    text = (
        "HTML:\n\n```html\n<details><summary>More</summary>\nHidden text\n</details>\n![x](data:a)\n```\n\n"
        "~~~\n<think>kept</think>\n~~~\n\nUse `<think>` and `</details>` like <think>this</think> here."
    )
    assert rx.clean_markdown(text) == text.replace("<think>this</think>", "")
    paragraphs = [p.text for p in Document(io.BytesIO(downloaded(run_action("docx", text))[0])).paragraphs]
    assert "<details><summary>More</summary>\nHidden text\n</details>\n![x](data:a)" in paragraphs
    assert "<think>kept</think>" in paragraphs and "Use <think> and </details> like  here." in paragraphs


def test_details_of_the_model_are_kept() -> None:
    text = (
        "Task: 2x = 4\n\n<details>\n<summary>**Solution** 1.</summary>\n\nx = **2**\n"
        '<details type="reasoning">hidden</details>\n</details>\n\n<details open><p>No summary</p></details>'
    )
    assert rx.parse_blocks(rx.clean_markdown(text)) == [
        rx.Paragraph("Task: 2x = 4"),
        rx.Paragraph("**Solution 1\\.**"),
        rx.Paragraph("x = **2**"),
        rx.Paragraph("<p>No summary</p>"),
    ]
    doc = Document(io.BytesIO(downloaded(run_action("docx", text))[0]))
    summary = next(p for p in doc.paragraphs if p.text == "Solution 1.")
    assert all(run.bold for run in summary.runs) and "hidden" not in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize(("language", "placeholder"), [("en", "[Image]"), ("de", "[Bild]")])
def test_images_without_a_web_address(language: str, placeholder: str) -> None:
    picture = base64.b64encode(bytes(range(256)) * 1600).decode()  # 400 KB
    content = (
        f"# Chart\n\nBefore ![Sales 2025](data:image/png;base64,{picture}) after\n\n"
        "![](/api/v1/files/1234-abcd/content)\n\n"
        "![a *b*](./chart.png \"title\") ![web](https://example.org/p.png)"
    )
    doc = Document(io.BytesIO(downloaded(run_action("docx", content, language=language))[0]))
    texts = [p.text for p in doc.paragraphs]
    assert texts[1:] == ["Before Sales 2025 after", placeholder, "a *b* web"]
    assert [link.address for p in doc.paragraphs for link in p.hyperlinks] == ["https://example.org/p.png"]
    if pypdf is not None:
        data, _, _ = downloaded(run_action("pdf", content, language=language))
        assert len(read_pdf(data).pages) == 1 and placeholder in page_texts(data)[0]


def test_latex_stays_as_written() -> None:
    text = r"Energy \[ E = mc^2 \] and \(a_1\), $a*b*c$ and $$\frac{1}{2}$$ cost $5 and $10, \$3, **$x$**"
    runs = rx.parse_inline(text)
    assert "".join(run.text for run in runs) == (
        r"Energy \[ E = mc^2 \] and \(a_1\), $a*b*c$ and $$\frac{1}{2}$$ cost $5 and $10, $3, $x$"
    )
    assert [run.text for run in runs if run.italic or run.bold] == ["$x$"]
    assert rx.to_plain("$5-$10 and $ 5 $") == "$5-$10 and $ 5 $"  # amounts, no formulas
    assert rx.to_plain("$$\nx_1 = *y*\n$$") == "$$\nx_1 = *y*\n$$"


def test_html_entities_and_lone_surrogates() -> None:
    content = "A &amp; B&nbsp;C &lt;tag&gt; &copy; `&amp;` \ud83d end\n\n```\n&lt;\n```"
    doc = Document(io.BytesIO(downloaded(run_action("docx", content))[0]))
    assert [p.text for p in doc.paragraphs] == ["A & B C <tag> © &amp;  end", "&lt;"]
    data, _, _ = downloaded(run_action("email", content, always_eml=True))
    message = email.message_from_bytes(data, policy=email.policy.default)
    parts = {part.get_content_type(): part.get_content() for part in message.iter_parts()}
    assert "A & B C <tag> ©" in parts["text/plain"]
    assert "A &amp; B C &lt;tag&gt; ©" in parts["text/html"]  # escaped once, shown as written


def test_docx_rule_is_a_line() -> None:
    data, _, _ = downloaded(run_action("docx", "Above\n\n---\n\nBelow"))
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    rule = doc.paragraphs[texts.index("Above") + 1]
    assert rule.text == "" and texts.index("Below") == texts.index("Above") + 2
    assert rule._p.xpath("./w:pPr/w:pBdr/w:bottom/@w:val") == ["single"]


def test_docx_links_only_for_safe_targets() -> None:
    content = (
        "[web](http://a.example) [mail](mailto:x@example.org?subject=Hi) "
        "[attach](mailto:x@example.org?attach=/etc/passwd) [phone](tel:+49711) [file](file:///etc/passwd) <https://b.example>"
    )
    doc = Document(io.BytesIO(downloaded(run_action("docx", content))[0]))
    paragraph = doc.paragraphs[0]
    assert [(link.text, link.address) for link in paragraph.hyperlinks] == [
        ("web", "http://a.example"),
        ("mail", "mailto:x@example.org?subject=Hi"),
        ("https://b.example", "https://b.example"),
    ]
    text = paragraph.text
    assert "attach (mailto:x@example.org?attach=/etc/passwd)" in text and "phone (tel:+49711)" in text
    assert "file (file:///etc/passwd)" in text


def test_docx_design_valves() -> None:
    data, _, _ = downloaded(run_action("docx", accent_color="#c00000", font="Arial", language="de"))
    doc = Document(io.BytesIO(data))
    assert doc.styles["Normal"].font.name == "Arial"
    assert (
        str(doc.styles["Heading 1"].font.color.rgb) == "C00000"
        and doc.styles["Heading 1"].font.name == "Arial"
    )
    assert doc.tables[0].rows[0].cells[0]._tc.xpath("./w:tcPr/w:shd/@w:fill") == ["C00000"]
    section = doc.sections[0]
    assert (
        round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7
    )  # A4 for German
    # invalid values fall back to the defaults, like in Office Documents
    data, _, _ = downloaded(run_action("docx", accent_color="red", font="<b>"))
    doc = Document(io.BytesIO(data))
    assert doc.styles["Normal"].font.name == "Calibri"
    assert str(doc.styles["Heading 1"].font.color.rgb) == rx.DEFAULT_ACCENT
    assert round(doc.sections[0].page_width.cm, 2) == 21.59  # Letter for English
    for name in ("accent_color", "font", "word_template"):
        assert rx.Action.Valves.model_fields[name].default == od_valve_default(name)


def od_valve_default(name: str) -> Any:
    """Default of a valve of Office Documents, read from its source (it needs python-pptx to import)."""
    tree = ast.parse(plugin_path("office_documents.py").read_text(encoding="utf-8"))
    constants = {
        target.id: node.value
        for node in tree.body
        if isinstance(node, ast.Assign)
        for target in node.targets
        if isinstance(target, ast.Name)
    }
    for node in ast.walk(tree):
        if (
            isinstance(node, ast.AnnAssign)
            and getattr(node.target, "id", None) == name
            and isinstance(node.value, ast.Call)  # name: str = Field(default=...)
        ):
            default = next(k.value for k in node.value.keywords if k.arg == "default")
            return ast.literal_eval(constants.get(getattr(default, "id", ""), default))
    raise KeyError(name)


def make_template(tmp: Path, suffix: str = ".docx") -> Path:
    template = Document()
    template.sections[0].header.paragraphs[0].text = "Example Ltd letterhead"
    template.sections[0].footer.paragraphs[0].text = "Example Ltd, Main Street 1"
    template.sections[0].left_margin = template.sections[0].right_margin = rx.Cm(4)
    template.add_paragraph("Template body text that must not appear")
    template.core_properties.author = "Template author"
    path = tmp / f"letterhead{suffix}"
    template.save(path)
    if suffix == ".dotx":  # same package with the template content type
        data = path.read_bytes()
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(path, "w") as target:
            for item in source.infolist():
                content = source.read(item)
                if item.filename == "[Content_Types].xml":
                    content = content.replace(rx.DOCX_CT.WML_DOCUMENT_MAIN.encode(), rx.DOTX_MAIN)
                target.writestr(item, content)
    return path


@pytest.mark.parametrize("suffix", [".docx", ".dotx"])
def test_docx_word_template(tmp_path: Path, suffix: str) -> None:
    path = make_template(tmp_path, suffix)
    data, _, _ = downloaded(run_action("docx", word_template=str(path)))
    doc = Document(io.BytesIO(data))
    section = doc.sections[0]
    assert section.header.paragraphs[0].text == "Example Ltd letterhead"
    assert section.footer.paragraphs[0].text == "Example Ltd, Main Street 1"
    assert round(section.left_margin.cm, 1) == 4.0
    texts = [p.text for p in doc.paragraphs]
    assert "Template body text that must not appear" not in texts
    assert (
        "Quarterly report C#" in texts
        and doc.paragraphs[
            texts.index("Intro line one\ncontinues bold nested and code and docs.\nSteps:")
        ].style.name
        == "Body Text"
    )
    assert doc.core_properties.author == "" and doc.core_properties.title == "Quarterly report C#"
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in numbered] == ["First continued", "Second", "Third", "Restart", "Again"]
    assert doc.tables[0].rows[0].cells[0]._tc.xpath("./w:tcPr/w:shd/@w:fill") == [rx.DEFAULT_ACCENT]


def test_docx_unreadable_template_is_skipped(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    (tmp_path / "broken.docx").write_bytes(b"not a zip file")
    for path in (tmp_path / "missing.dotx", tmp_path / "broken.docx"):
        data, _, _ = downloaded(run_action("docx", word_template=str(path)))
        assert "Quarterly report C#" in [p.text for p in Document(io.BytesIO(data)).paragraphs]
        assert str(path) in caplog.text


# ------------------------------------------------------------------------------------------- PDF --

try:
    import pypdf
    from pypdf.generic import ContentStream
except ImportError:  # the PDF tests that read the file are skipped
    pypdf = None

needs_pypdf = pytest.mark.skipif(pypdf is None, reason="pypdf not installed")

GERMAN = "ä ö ü Ä Ö Ü ß € „Anführung“ ‚einfach‘ – — … • 20 °C § 5 m² m³ ½ 3 × 4"
LONG_TABLE = "# Materialliste\n\n| Nr. | Artikel | Betrag |\n|---|---|---|\n" + "".join(
    f"| {n} | Artikel {n:03d} Kabelkanal weiß | {n},50 € |\n" for n in range(1, 61)
)


def make_pdf(markdown: str, **options: Any) -> bytes:
    return rx.markdown_to_pdf(rx.parse_blocks(markdown), **options)


def read_pdf(data: bytes) -> Any:
    return pypdf.PdfReader(io.BytesIO(data))


def page_texts(data: bytes) -> list[str]:
    return [page.extract_text() for page in read_pdf(data).pages]


def link_targets(data: bytes) -> list[str]:
    return [
        annotation.get_object()["/A"]["/URI"]
        for page in read_pdf(data).pages
        for annotation in page.get("/Annots", [])
    ]


def text_pieces(data: bytes) -> list[tuple[int, str, float, float, float, str]]:
    """(page, font, size, x, y, text) of every piece of text, read from the content streams."""
    reader, pieces = read_pdf(data), []
    for number, page in enumerate(reader.pages):
        fonts = {name: ref.get_object()["/BaseFont"][1:] for name, ref in page["/Resources"]["/Font"].items()}
        font, size, x, y = "", 0.0, 0.0, 0.0
        for operands, operator in ContentStream(page.get_contents(), reader).operations:
            if operator == b"Tf":
                font, size = fonts[operands[0]], float(operands[1])
            elif operator == b"Td":
                x, y = float(operands[0]), float(operands[1])
            elif operator == b"Tj":
                text = operands[0].original_bytes.decode("cp1252")
                pieces.append((number, font, size, x, y, text))
    return pieces


def body_text(data: bytes) -> str:
    """The text of all pages without the footers, without spaces and line breaks."""
    return "".join(text for *_, y, text in text_pieces(data) if y > rx.PDF_FOOTER_Y + 1).replace(" ", "")


def assert_inside_margins(data: bytes, paper: str = "letter") -> None:
    width, height = rx.PDF_PAPER[paper]
    for _, font, size, x, y, text in text_pieces(data):
        assert x >= rx.PDF_SIDE_MARGIN - 0.01, text
        assert x + rx.text_width(font, size, text) <= width - rx.PDF_SIDE_MARGIN + 0.01, text
        assert rx.PDF_FOOTER_Y - 0.01 <= y <= height - rx.PDF_TOP_MARGIN, text


def test_pdf_is_downloaded() -> None:
    data, filename, mime = downloaded(run_action("pdf"))
    assert mime == rx.PDF_MIME == "application/pdf"
    assert re.fullmatch(r"Quarterly report C# \d{4}-\d{2}-\d{2} \d{4}\.pdf", filename)
    assert data.startswith(b"%PDF-1.4\n") and data.endswith(b"%%EOF\n")


@needs_pypdf
@pytest.mark.parametrize(
    ("language", "size", "footer", "page"),
    [
        ("en", (612, 792), r"Created with the AI assistant on \d{4}-\d{2}-\d{2}\. Please",
         "Page 1 of 1"),
        ("de", (595.28, 841.89), r"Erstellt mit dem KI-Assistenten am \d{2}\.\d{2}\.\d{4} – Inhalte",
         "Seite 1 von 1"),
    ],
)  # fmt: skip
def test_pdf_content(language: str, size: tuple[float, float], footer: str, page: str) -> None:
    data, _, _ = downloaded(run_action("pdf", language=language))
    reader = read_pdf(data)
    assert reader.metadata.title == "Quarterly report C#"
    assert len(reader.pages) == 1
    assert [float(v) for v in reader.pages[0].mediabox[2:]] == pytest.approx(size, abs=0.01)
    assert reader.trailer["/Root"]["/Lang"] == language
    text = reader.pages[0].extract_text()
    assert re.search(footer, text) and page in text
    assert "SECRET REASONING" not in text
    assert "<script>alert(1)</script>" in text  # shown as text
    assert "bad (javascript:alert(1))" in text  # as in Word: not a link, the target stays visible
    lines = [" ".join(line.split()) for line in text.splitlines()]
    for line in ("1. First continued", "2. Second", "– sub a", "3. Third", "1. Restart", "2. Again"):
        assert line in lines
    assert "Intro line one" in lines and "A quote" in lines and "echo \"<b>hi</b>\"" in lines
    assert link_targets(data) == ["https://example.org/a_(b)"]
    assert_inside_margins(data, "a4" if language == "de" else "letter")


@needs_pypdf
def test_pdf_without_footer_note() -> None:
    data, _, _ = downloaded(run_action("pdf", print_footer=False, language="de"))
    text = page_texts(data)[0]
    assert "Erstellt mit" not in text and "Seite 1 von 1" in text


@needs_pypdf
def test_pdf_german_characters() -> None:
    data = make_pdf(f"# Übersicht für Größen\n\n{GERMAN}\n\n| Größe | Maß |\n|---|---|\n| {GERMAN} | ö |",
                    paper="a4", title="Übersicht für Größen")  # fmt: skip
    assert read_pdf(data).metadata.title == "Übersicht für Größen"
    text = page_texts(data)[0]
    assert "Übersicht für Größen" in text
    for word in GERMAN.split():
        assert word in text, word
    assert {font for _, font, *_ in text_pieces(data)} == {"Helvetica", "Helvetica-Bold"}
    fonts = read_pdf(data).pages[0]["/Resources"]["/Font"]
    assert all(font.get_object()["/Encoding"] == "/WinAnsiEncoding" for font in fonts.values())


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("a → b ← c ⇒ d", "a -> b <- c => d"),
        ("≤ ≥ ≠ ≈ − ✓ ✗ ✅ ❌ ☐ ☑", "<= >= != ~ - + - + - [ ] [x]"),
        ("Start 🚀 👍🏽 👨‍👩‍👧 🇩🇪 ❤️ 1️⃣ Ende", "Start      1 Ende"),
        ("zero​width soft­hyphen⁠", "zerowidth softhyphen"),
        ("10 km 5 %", "10 km 5 %"),
        ("Łódź Čeština Ș ﬁ ⅓ x⁴", "Lódz Ceština S fi 1/3 x4"),
        ("漢字 αβ Жж", "?? ?? ??"),
        ("tab\there\r\x07", "tab here"),
        (GERMAN, GERMAN),
        ("├── src/\n│   └── main.py\n└── README.md", "+-- src/\n|   +-- main.py\n+-- README.md"),
        ("┌─┬─┐ ╔═╦═╗ ╭─╮ ┃━┃ ┄┆ ╱╲╳", "+-+-+ +=+=+ +-+ |-| -| /\\X"),
        ("↑ ↓ ⟶ ⟹ ↳ ▶ ► ◀ ▲ ▼ ➤ ❯ ✕ ✖ ◆ ✦", "^ v -> => -> > > < ^ v > > x x • *"),
    ],
)
def test_pdf_characters_outside_the_standard_fonts(text: str, expected: str) -> None:
    assert rx.pdf_safe_text(text) == expected
    rx.pdf_safe_text(expected).encode("cp1252")


def test_pdf_box_drawing_covers_the_block() -> None:
    stand_ins = {rx.pdf_safe_text(chr(code)) for code in rx.BOX_DRAWING}
    assert stand_ins == {"-", "=", "|", "+", "/", "\\", "X"}  # one character each: trees keep their shape


@needs_pypdf
def test_pdf_directory_tree() -> None:
    data = make_pdf("```\nproject/\n├── src/\n│   └── main.py\n└── README.md\n```")
    code = [text for _, font, _, _, _, text in text_pieces(data) if font == "Courier"]
    assert code == ["project/", "+-- src/", "|   +-- main.py", "+-- README.md"]


@needs_pypdf
def test_pdf_unknown_characters_do_not_break_the_export() -> None:
    data, _, _ = downloaded(run_action("pdf", "# 🚀 Plan → Ziel\n\n✅ fertig, 漢字 und \U0010ffff."))
    assert read_pdf(data).metadata.title == "🚀 Plan → Ziel"  # the title is Unicode
    text = page_texts(data)[0]
    assert "Plan -> Ziel" in text and "+ fertig, ?? und ?." in text


@needs_pypdf
def test_pdf_empty_document() -> None:
    reader = read_pdf(rx.markdown_to_pdf([]))
    assert len(reader.pages) == 1 and reader.pages[0].extract_text().strip() == "Page 1 of 1"
    recorder = run_action("pdf", "<think>only reasoning</think>")
    assert not recorder.calls and recorder.notifications()[0][0] == "warning"


@needs_pypdf
def test_pdf_long_words_are_broken() -> None:
    word, url = "x" * 5000, "https://example.org/" + "a" * 400
    data = make_pdf(f"{word}\n\n[{url}]({url})\n\n`{'y' * 300}`\n\n| a | b |\n|---|---|\n| {'z' * 500} | 1 |")
    assert_inside_margins(data)
    text = body_text(data)
    assert word in text and url in text and "y" * 300 in text and "z" * 500 in text
    links = link_targets(data)
    assert len(links) == 6 and set(links) == {url}  # one clickable area per line


def test_wrap_words_keeps_to_the_width() -> None:
    runs = rx.parse_inline("Ein **langer** Satz mit `Code` und " + "Wort" * 100 + " und Ende")
    words = rx.pdf_words(runs, 10.5, rx.TEXT_COLOR, rx.DEFAULT_ACCENT)
    for width in (5, 60, 200, 450):
        lines = rx.wrap_words(words, width)
        assert all(line.width <= width + 0.01 for line in lines if len(line.parts[0][1]) > 1)
        assert "".join(text for line in lines for _, text in line.parts).replace(" ", "") == (
            "EinlangerSatzmitCodeund" + "Wort" * 100 + "undEnde"
        )


@needs_pypdf
def test_pdf_list_numbering() -> None:
    data = make_pdf(
        "Intro\n\n98. Achtundneunzig\n99. Neunundneunzig\n100. Hundert\n\nText\n\n"
        "1. Eins\n   1. Eins-a\n   2. Eins-b\n2. Zwei\n   - Punkt\n\n- Oben\n  - Unten"
    )
    lines = [" ".join(line.split()) for line in page_texts(data)[0].splitlines()]
    expected = ("98. Achtundneunzig", "100. Hundert", "1. Eins", "1. Eins-a", "2. Eins-b", "2. Zwei")
    for line in (*expected, "– Punkt", "• Oben", "– Unten"):
        assert line in lines, line
    pieces = {text: x for _, _, _, x, _, text in reversed(text_pieces(data))}  # the first of each
    assert pieces["Achtundneunzig"] == pieces["Hundert"] > pieces["100."]  # the same indent for all items
    assert pieces["Eins-a"] > pieces["Eins"] and pieces["Unten"] > pieces["Oben"]
    assert pieces["1."] == pieces["98."] == pieces["•"] == pytest.approx(rx.PDF_SIDE_MARGIN, abs=0.01)


@needs_pypdf
def test_pdf_code_block() -> None:
    data = make_pdf("Before\n\n```python\ndef f():\n\treturn  1  # <b>\n```\n\nAfter")
    code = [(font, text) for _, font, _, _, _, text in text_pieces(data) if font.startswith("Courier")]
    assert code == [("Courier", "def f():"), ("Courier", "    return  1  # <b>")]
    stream = read_pdf(data).pages[0].get_contents().get_data()
    assert b"0.957 0.957 0.957 rg" in stream  # the light background


@needs_pypdf
def test_pdf_long_table_repeats_its_header() -> None:
    data = make_pdf(LONG_TABLE, paper="a4")
    texts = page_texts(data)
    assert len(texts) == 2
    assert all("Nr. Artikel Betrag" in " ".join(text.split()) for text in texts)
    for n in range(1, 61):  # each row is on one page, complete
        assert sum(f"{n} Artikel {n:03d} Kabelkanal weiß {n},50 €" in " ".join(t.split()) for t in texts) == 1
    assert_inside_margins(data, "a4")


@needs_pypdf
def test_pdf_table_layout() -> None:
    data = make_pdf(
        "| Item | Text | Amount | Left |\n|---|---|---|:--|\n"
        "| a | x | 1.234,56 € | 1 |\n| b | y | 2 € | 22 |\n"
        "| **Total** | [z](https://example.org/z) | **1.236,56 €** | 3 |"
    )
    pieces = [(font, size, x, text) for _, font, size, x, _, text in text_pieces(data)]
    right = {round(x + rx.text_width(font, size, text), 2) for font, size, x, text in pieces if "€" in text}
    assert len(right) == 1  # number columns are right-aligned, also the bold total
    left = {x for _, _, x, text in pieces if text in ("1", "22", "3")}
    assert len(left) == 1  # an explicit alignment wins
    header = [(font, text) for font, _, _, text in pieces[:4]]
    assert header == [("Helvetica-Bold", name) for name in ("Item", "Text", "Amount", "Left")]
    stream = read_pdf(data).pages[0].get_contents().get_data()
    assert b"0.122 0.306 0.475 rg" in stream  # header row in the accent colour 1F4E79
    assert f"{rx._rgb(rx.DEFAULT_DESIGN.band)} rg".encode() in stream  # banded rows
    assert link_targets(data) == ["https://example.org/z"]
    assert rx._numeric_column([["h"], ["1,5 %"], [""], ["**2**"]], 0)
    assert not rx._numeric_column([["h"], ["1"], ["+49 711 123456"]], 0)


@needs_pypdf
def test_pdf_row_taller_than_a_page_is_split() -> None:
    tall = "<br>".join(f"line {n}" for n in range(150))
    data = make_pdf(f"| Key | Value |\n|---|---|\n| a | {tall} |\n| b | after |")
    texts = page_texts(data)
    assert len(texts) >= 3 and all("Key Value" in text for text in texts)
    assert body_text(data).startswith("KeyValuealine0line1line2")
    lines = [line.strip() for text in texts for line in text.splitlines()]
    assert all(f"line {n}" in lines for n in range(1, 150)) and "b after" in lines
    assert_inside_margins(data)


@needs_pypdf
def test_pdf_links_only_for_safe_targets() -> None:
    data = make_pdf(
        "[web](https://example.org/ä?q=1) [plain](http://example.org) [mail](mailto:a@example.org) "
        "[mail with cc](mailto:a@example.org?cc=b@example.org) [attach](mailto:a@example.org?attach=x) "
        "[js](javascript:alert(1)) [tel](tel:+49711) [file](file:///etc/passwd) <https://example.org/auto> "
        "![picture](https://example.org/p.png)"
    )
    assert link_targets(data) == [
        "https://example.org/%C3%A4?q=1",
        "http://example.org",
        "mailto:a@example.org",
        "mailto:a@example.org?cc=b@example.org",
        "https://example.org/auto",
        "https://example.org/p.png",  # images are not loaded, they become links as in Word
    ]
    text = " ".join(page_texts(data)[0].split())
    for shown in ("attach (mailto:a@example.org?attach=x)", "js (javascript:alert(1))", "tel (tel:+49711)"):
        assert shown in text


@needs_pypdf
def test_pdf_accent_colour_valve() -> None:
    content = "# Title\n\n[link](https://example.org)"
    data, _, _ = downloaded(run_action("pdf", content, accent_color="#c00000"))
    stream = read_pdf(data).pages[0].get_contents().get_data()
    assert b"0.753 0 0 rg" in stream


def test_pdf_icon_is_a_document() -> None:
    svg = base64.b64decode(rx.ACTIONS["pdf"][1].split(",", 1)[1]).decode()
    assert "M14 3v5h5" in svg  # the folded corner of the Word icon
    assert "<rect" not in svg  # not the printer icon of 1.1.0


# ------------------------------------------------------------------- browser snippets (Node.js) --

node = pytest.mark.skipif(shutil.which("node") is None, reason="Node.js not installed")


def run_in_node(code: str) -> dict[str, Any]:
    result = subprocess.run(
        ["node", str(STUB_JS)], input=code, capture_output=True, text=True, timeout=30, check=True
    )
    return json.loads(result.stdout)


@node
def test_js_download_runs() -> None:
    report = run_in_node(run_action("docx").code())
    assert report["returned"] is True
    assert report["blobs"][0]["type"] == rx.DOCX_MIME and report["blobs"][0]["size"] > 1000
    assert report["clicks"][0]["download"].endswith(".docx")


@node
def test_js_pdf_download_runs() -> None:
    report = run_in_node(run_action("pdf").code())
    assert report["returned"] is True
    assert report["blobs"][0]["type"] == rx.PDF_MIME and report["blobs"][0]["size"] > 1000
    assert report["clicks"][0]["download"].endswith(".pdf")


@node
def test_js_mailto_runs() -> None:
    report = run_in_node(run_action("email", EMAIL).code())
    assert report["clicks"][0]["href"].startswith("mailto:?subject=Meeting%20on%20Friday")


@node
def test_js_browser_settings_runs() -> None:
    report = run_in_node(rx.BROWSER_SETTINGS_JS)
    locale, timezone = report["returned"]
    assert locale == "de-CH" and timezone and rx.zone(timezone) is not None


LOADER_JS = REPO / "office-file-ui" / "loader.js"
LOADER_STUB_JS = Path(__file__).resolve().parent / "loader_stub.js"


def run_loader(scenario: str) -> dict[str, Any]:
    """Runs loader.js against a small DOM stub with a virtual clock (see loader_stub.js)."""
    result = subprocess.run(
        ["node", str(LOADER_STUB_JS), scenario], capture_output=True, text=True, timeout=60, check=True
    )
    return json.loads(result.stdout)


@node
def test_loader_js_syntax() -> None:
    subprocess.run(
        ["node", "--check", str(LOADER_JS)], capture_output=True, text=True, timeout=30, check=True
    )


@node
def test_loader_reloads_the_chat_only_a_few_times() -> None:
    """A chip without a matching file in the saved chat, while an answer streams for two minutes."""
    report = run_loader("missing-file")
    assert report["fetches"] == 3 and report["icon"] is None
    assert report["scans"] <= 120 / 0.15 + 5  # page changes every 100 ms, handled at most every 150 ms
    report = run_loader("found-file")
    assert report["fetches"] == 1 and report["icon"] == "f1"
    report = run_loader("flaky-server")  # failed requests do not use up the three tries
    assert report["fetches"] == 4 and report["icon"] == "f1"


@node
@pytest.mark.parametrize(
    ("scenario", "clicks", "display", "selected"),
    [
        ("tabs", 1, "none", 1),  # Preview selected: the bar is hidden
        ("tabs-no-switch", 1, "", 0),  # the click did not switch: the bar stays
        ("tabs-reordered", 0, "", 1),  # an unexpected first state: nothing is clicked or hidden
    ],
)
def test_loader_tab_switch(scenario: str, clicks: int, display: str, selected: int) -> None:
    report = run_loader(scenario)
    assert (report["clicks"], report["barDisplay"], report["selected"]) == (clicks, display, selected)
    if scenario == "tabs":
        assert report["barDisplayAfterReset"] == ""  # back on "Content": the bar is shown again


# ------------------------------------------------------------------------ optional rendering checks --


HAS_POPPLER = shutil.which("pdftotext") is not None and shutil.which("pdfinfo") is not None


def pdf_text(pdf: Path, layout: bool = False) -> str:
    options = ["-layout"] if layout else []
    return subprocess.run(
        ["pdftotext", *options, str(pdf), "-"], capture_output=True, text=True, check=True
    ).stdout


@pytest.mark.skipif(not (shutil.which("soffice") and HAS_POPPLER), reason="LibreOffice or poppler missing")
def test_docx_renders_in_libreoffice() -> None:
    data, _, _ = downloaded(run_action("docx"))
    with tempfile.TemporaryDirectory() as tmp:
        Path(tmp, "report.docx").write_bytes(data)
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{tmp}/profile",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(Path(tmp, "report.docx")),
            ],
            capture_output=True,
            timeout=120,
            check=True,
        )
        text = pdf_text(Path(tmp, "report.pdf"), layout=True)
    lines = [" ".join(line.split()) for line in text.splitlines()]
    assert "3. Third" in lines
    assert "1. Restart" in lines and "2. Again" in lines  # second list restarts at 1


@pytest.mark.skipif(not shutil.which("qpdf"), reason="qpdf not installed")
@pytest.mark.parametrize("markdown", [SAMPLE, LONG_TABLE, "x"])
def test_pdf_passes_qpdf_check(markdown: str) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp, "export.pdf")
        pdf.write_bytes(make_pdf(markdown))
        result = subprocess.run(["qpdf", "--check", str(pdf)], capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "No syntax or stream encoding errors found" in result.stdout


@pytest.mark.skipif(not (shutil.which("pdftoppm") and HAS_POPPLER), reason="poppler not installed")
def test_pdf_renders_with_poppler() -> None:
    data, _, _ = downloaded(run_action("pdf", language="de"))
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp, "export.pdf")
        pdf.write_bytes(data)
        command = ["pdftoppm", "-png", "-r", "50", "-f", "1", "-l", "1", str(pdf), f"{tmp}/page"]
        result = subprocess.run(command, capture_output=True, text=True, timeout=60)
        assert result.returncode == 0 and not result.stderr, result.stderr  # poppler reports damaged files
        assert Path(tmp, "page-1.png").read_bytes().startswith(b"\x89PNG")
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        text = pdf_text(pdf)
    assert re.search(r"Page size:.*\(A4\)", info) and re.search(r"Title:\s+Quarterly report C#", info)
    assert "Quarterly report C#" in text and "Erstellt mit dem KI-Assistenten" in text



# --------------------------------------------------------------------------------- context guard --


def body_with_chars(chars: int, model: str = "m") -> dict[str, Any]:
    return {"model": model, "messages": [{"role": "user", "content": "x" * chars}]}


def run_guard(
    flt: Any, body: dict[str, Any], metadata: dict[str, Any] | None = None
) -> tuple[dict[str, Any], Recorder]:
    recorder = Recorder()
    result = asyncio.run(flt.request(body, __event_emitter__=recorder.emit, __metadata__=metadata))
    return result, recorder


def test_guard_ok() -> None:
    body = body_with_chars(1000)
    result, recorder = run_guard(make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000), body)
    assert result is body and result == body_with_chars(1000)
    assert recorder.events == []


@pytest.mark.parametrize(
    ("language", "text"), [("en", "very long (about 80%"), ("de", "sehr lang (ca. 80 %")]
)
def test_guard_warns(language: str, text: str) -> None:
    # usable = 8000 tokens; (22384 + 16) chars / 3.5 = 6400 tokens = 80 %
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000, language=language)
    _, recorder = run_guard(flt, body_with_chars(22384))
    assert len(recorder.events) == 1 and recorder.events[0]["type"] == "status"
    assert text in recorder.events[0]["data"]["description"] and recorder.events[0]["data"]["done"] is True


@pytest.mark.parametrize(
    ("language", "text"), [("en", "too large for the AI model (about 101%"), ("de", "ca. 101 %")]
)
def test_guard_blocks(language: str, text: str) -> None:
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000, language=language)
    with pytest.raises(cg.ContextLimitError) as error:
        run_guard(flt, body_with_chars(8001 * 35 // 10))
    assert text in str(error.value)


def test_guard_counts_parts_images_and_tools() -> None:
    body = {
        "model": "m",
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": "x" * 349},
                    {"type": "image_url", "image_url": {"url": "data:..."}},
                    "not-a-dict",
                ],
            },
            {"role": "assistant", "content": None, "tool_calls": [{"id": "1", "function": {"name": "f"}}]},
        ],
        "tools": [{"type": "function", "function": {"name": "f"}}],
    }
    tool_chars = len(json.dumps(body["tools"])) + len(json.dumps(body["messages"][1]["tool_calls"][0]))
    expected = -(-(349 + 2 * cg.MESSAGE_OVERHEAD_CHARS + tool_chars) * 10 // 35) + 300
    assert cg.estimate_tokens(body, 3.5, 300) == expected


def test_guard_counts_reasoning_sent_back() -> None:
    reasoning = "r" * 35000
    for field in ("reasoning_content", "reasoning", "thinking"):  # llama.cpp, other providers, Ollama
        body = {"model": "m", "messages": [{"role": "assistant", "content": "ok", field: reasoning}]}
        assert cg.estimate_tokens(body, 3.5, 300) == -(-(2 + 35000 + cg.MESSAGE_OVERHEAD_CHARS) * 10 // 35)
    body = {"model": "m", "messages": [{"role": "assistant", "content": "ok", "reasoning_content": None}]}
    assert cg.estimate_tokens(body, 3.5, 300) == 6
    with pytest.raises(cg.ContextLimitError):
        run_guard(
            make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000),
            {
                "model": "m",
                "messages": [{"role": "assistant", "content": "", "reasoning_content": reasoning}],
            },
        )


MODEL_PROMPT = "You are the e-mail assistant of Example Ltd.\nAlways sign with the company name." * 100


@pytest.mark.parametrize("key", ["system_prompt", "base_system_prompt"])  # Open WebUI 0.11.4, 0.12.x
def test_guard_counts_the_model_system_prompt(key: str) -> None:
    """Open WebUI adds the model's system prompt only after the request hook (in the router)."""
    chat_prompt = "Answer briefly.\nUse bullet points."
    user = {"role": "user", "content": "Hi"}
    # metadata = "model prompt\nsystem message", set before the hook
    cases = [
        ([user], MODEL_PROMPT, len(MODEL_PROMPT) + 1),  # no system message yet
        (
            [{"role": "system", "content": chat_prompt}, user],
            f"{MODEL_PROMPT}\n{chat_prompt}",
            len(MODEL_PROMPT) + 1,
        ),
        # RAG context appended to the system message (RAG_SYSTEM_CONTEXT)
        (
            [{"role": "system", "content": f"{chat_prompt}\n<context>doc</context>"}, user],
            f"{MODEL_PROMPT}\n{chat_prompt}",
            len(MODEL_PROMPT) + 1,
        ),
        # tool-call rounds: the body already contains the model prompt
        (
            [{"role": "system", "content": f"{MODEL_PROMPT}\n{chat_prompt}"}, user],
            f"{MODEL_PROMPT}\n{chat_prompt}",
            0,
        ),
        # no model prompt: metadata holds only the system message
        ([{"role": "system", "content": chat_prompt}, user], chat_prompt, 0),
        (
            [{"role": "system", "content": [{"type": "text", "text": chat_prompt}]}, user],
            f"{MODEL_PROMPT}\n{chat_prompt}",
            len(MODEL_PROMPT) + 1,
        ),
        ([user], None, 0),
    ]
    for messages, prompt, expected in cases:
        assert cg.missing_system_prompt_chars(messages, {key: prompt}) == expected, (messages, prompt)
    assert cg.missing_system_prompt_chars([user], None) == 0
    body = {"model": "m", "messages": [user]}
    assert cg.estimate_tokens(body, 3.5, 300, {key: MODEL_PROMPT}) - cg.estimate_tokens(
        body, 3.5, 300
    ) == pytest.approx((len(MODEL_PROMPT) + 1) / 3.5, abs=1)
    # blocked only because of the model's prompt
    flt = make_filter(cg, context_tokens=4000, answer_reserve_tokens=2000)  # usable 2000 tokens
    run_guard(flt, body_with_chars(1000))
    with pytest.raises(cg.ContextLimitError):
        run_guard(flt, body_with_chars(1000), {key: MODEL_PROMPT})


def test_guard_workspace_model_uses_base_model_size() -> None:
    flt = make_filter(
        cg,
        context_tokens=100000,
        answer_reserve_tokens=100,
        model_context_tokens="qwen3-8b=1000, emails=50000",
    )
    assistant = {"id": "support", "info": {"base_model_id": "qwen3-8b"}}
    body = body_with_chars(10000, model="support")  # in the hook, body["model"] is the workspace model's ID
    with pytest.raises(cg.ContextLimitError):
        asyncio.run(flt.request(body, __model__=assistant))
    # an own entry of the workspace model wins over its base model
    assert asyncio.run(
        flt.request(
            body_with_chars(10000, model="emails"),
            __model__={"id": "emails", "info": {"base_model_id": "qwen3-8b"}},
        )
    )
    assert asyncio.run(flt.request(body_with_chars(10000, model="other"), __model__={"info": None}))
    assert cg.model_ids({"model": "a"}, {"id": "a", "info": {"base_model_id": "b"}}) == ["a", "b"]


@pytest.mark.parametrize(
    ("locale", "text"),
    [("de-DE", "ca. 101 %"), ("en-GB", "about 101%"), ("fr", "about 101%"), (None, "about 101%")],
)
def test_guard_language_auto(locale: str | None, text: str) -> None:
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    assert flt.valves.language == "auto"
    metadata = {"variables": {"{{USER_LANGUAGE}}": locale}} if locale else {}
    with pytest.raises(cg.ContextLimitError) as error:
        run_guard(flt, body_with_chars(8001 * 35 // 10), metadata)
    assert text in str(error.value)


def test_guard_per_model_override_and_misconfiguration(caplog: pytest.LogCaptureFixture) -> None:
    flt = make_filter(
        cg,
        context_tokens=100000,
        answer_reserve_tokens=100,
        model_context_tokens="small:latest=1000, broken, x=abc",
    )
    with pytest.raises(cg.ContextLimitError):
        run_guard(flt, body_with_chars(10000, model="small:latest"))
    run_guard(flt, body_with_chars(10000, model="other"))
    assert "broken" in caplog.text and "x=abc" in caplog.text
    # a reserve as large as the context (a small Ollama num_ctx) keeps half of it free, not all of it
    small = make_filter(cg, context_tokens=32768)
    assert run_guard(small, {**body_with_chars(1500 * 35 // 10), "options": {"num_ctx": 4096}})[1].events == []
    with pytest.raises(cg.ContextLimitError):
        run_guard(small, {**body_with_chars(2100 * 35 // 10), "options": {"num_ctx": 4096}})


def stub_compaction_config(monkeypatch: pytest.MonkeyPatch, values: dict[str, Any]) -> None:
    """Stands in for open_webui.models.config.Config, which holds the Context Compaction settings."""

    async def get_many(*keys: str) -> dict[str, Any]:
        return {key: values[key] for key in keys if key in values}

    modules = {
        "open_webui": ModuleType("open_webui"),
        "open_webui.models": ModuleType("open_webui.models"),
        "open_webui.models.config": SimpleNamespace(Config=SimpleNamespace(get_many=get_many)),
    }
    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)


def test_guard_compaction_hides_hint_for_normal_growth(monkeypatch: pytest.MonkeyPatch) -> None:
    # usable = 8000 tokens, compaction from 6500 (+ 1500 for the summary request): at 80 % the chat
    # is summarized before the limit
    stub_compaction_config(
        monkeypatch,
        {"chat.context_compaction.enable": True, "chat.context_compaction.token_threshold": 6500},
    )
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    assert run_guard(flt, body_with_chars(22384))[1].events == []
    # still above the compaction threshold after compaction (e.g. large documents): hint is shown
    _, recorder = run_guard(flt, body_with_chars(7200 * 35 // 10))
    assert len(recorder.events) == 1 and "very long" in recorder.events[0]["data"]["description"]
    with pytest.raises(cg.ContextLimitError):
        run_guard(flt, body_with_chars(8001 * 35 // 10))


def test_guard_compaction_threshold_too_high(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    # default threshold 80000 is above the usable 8000 tokens: compaction never runs in time
    stub_compaction_config(monkeypatch, {"chat.context_compaction.enable": True})
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    run_guard(flt, body_with_chars(1000))
    assert "Context Compaction threshold" not in caplog.text  # only checked once a chat gets long
    assert len(run_guard(flt, body_with_chars(22384))[1].events) == 1
    with pytest.raises(cg.ContextLimitError):
        run_guard(flt, body_with_chars(8001 * 35 // 10))
    assert caplog.text.count("Context Compaction threshold for model 'm' (80000 tokens)") == 1


def test_guard_compaction_threshold_leaves_no_room_for_the_summary(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    # 7000 is below the usable 8000 tokens, but the summary request (whole chat + about 1500 tokens)
    # would not fit: vLLM rejects it and Open WebUI only cuts the old messages short
    stub_compaction_config(
        monkeypatch,
        {"chat.context_compaction.enable": True, "chat.context_compaction.token_threshold": 7000},
    )
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    assert len(run_guard(flt, body_with_chars(22384))[1].events) == 1
    assert "Context Compaction threshold for model 'm' (7000 tokens) leaves too little room" in caplog.text


def summary_body(summary: str, chars: int) -> dict[str, Any]:
    system = f"You are helpful.\n[CONVERSATION SUMMARY]\n{summary}"
    return {
        "model": "m",
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": "x" * chars}],
    }


def test_guard_detects_the_fallback_summary(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    stub_compaction_config(
        monkeypatch,
        {"chat.context_compaction.enable": True, "chat.context_compaction.token_threshold": 6500},
    )
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    real = summary_body("The user asked about A; decided B.\n- open: C", 22000)
    assert run_guard(flt, real, {"chat_id": "c1"})[1].events == []
    assert "could not get a summary" not in caplog.text
    # failed summary request: the first 500 characters per message instead of a summary
    fallback = summary_body("Earlier summary\n- user: Question one\n- assistant: Answer one", 22000)
    _, recorder = run_guard(flt, fallback, {"chat_id": "c1"})
    assert len(recorder.events) == 1  # compaction does not really help: the hint is shown
    run_guard(flt, summary_body("- user: tiny", 100), {"chat_id": "c1"})
    run_guard(flt, summary_body("- user: tiny", 100), {"chat_id": "c2"})
    assert caplog.text.count("could not get a summary from the model for chat c1") == 1
    assert caplog.text.count("for chat c2") == 1
    assert "Question one" not in caplog.text and "tiny" not in caplog.text  # no message content


@pytest.mark.parametrize(
    "metadata",
    [{"base_user_prompt": "x" * 26000}, {"user_message": {"role": "user", "content": "x" * 26000}}],
)
def test_guard_names_a_long_pasted_message(metadata: dict[str, Any]) -> None:
    # 0.12 keeps the typed text in base_user_prompt, 0.11 in user_message
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000, language="de")
    with pytest.raises(cg.ContextLimitError) as error:
        run_guard(flt, body_with_chars(30000), metadata)
    assert str(error.value).startswith("Ihre Nachricht ist zu lang für das KI-Modell (zusammen mit dem Chat ca. 108 %")


def test_guard_long_document_in_the_message_is_not_the_users_text() -> None:
    # document excerpts are added to the user message; the typed question itself is short
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000)
    with pytest.raises(cg.ContextLimitError) as error:
        run_guard(flt, body_with_chars(30000), {"base_user_prompt": "Summarize the file"})
    assert str(error.value).startswith("This request is too large")


@pytest.mark.parametrize("language", ["en", "de"])
def test_guard_messages_need_no_line_breaks(language: str) -> None:
    # Open WebUI's error box shows plain text: line breaks would be lost
    assert not any("\n" in text for text in cg.TRANSLATIONS[language].values())


def test_guard_compaction_threshold_resolution(monkeypatch: pytest.MonkeyPatch) -> None:
    per_model = {"params": {"compact_token_threshold": 5000}}
    assert asyncio.run(cg.compaction_settings(per_model)) is None  # outside Open WebUI
    stub_compaction_config(monkeypatch, {"chat.context_compaction.enable": False})
    assert asyncio.run(cg.compaction_settings(per_model)) is None
    settings = {"chat.context_compaction.enable": True, "chat.context_compaction.token_threshold": "20000"}
    stub_compaction_config(monkeypatch, settings)
    assert asyncio.run(cg.compaction_settings(None)).threshold == 20000
    assert asyncio.run(cg.compaction_settings(per_model)).threshold == 5000
    assert asyncio.run(cg.compaction_settings({"params": {"compact_token_threshold": 30000}})).threshold == 20000
    stub_compaction_config(monkeypatch, {**settings, "chat.context_compaction.token_cap": 40000})
    assert asyncio.run(cg.compaction_settings({"params": {"compact_token_threshold": 30000}})).threshold == 30000


def test_guard_keeps_room_for_the_requested_max_tokens() -> None:
    # 32k context, 25,005 tokens prompt: fits with the reserve of 6000, but not with max_tokens 8192,
    # which vLLM adds to the prompt and rejects with HTTP 400
    flt = make_filter(cg, context_tokens=32768)
    prompt = 25005 * 35 // 10
    assert run_guard(flt, body_with_chars(prompt))[1].events
    for key in ("max_tokens", "max_completion_tokens"):
        with pytest.raises(cg.ContextLimitError):
            run_guard(flt, {**body_with_chars(prompt), key: 8192})
    assert run_guard(flt, {**body_with_chars(prompt), "max_tokens": 2000})[1].events  # reserve is larger
    # a max_tokens at or above the context size is left to the server (llama-server just stops there)
    assert run_guard(flt, {**body_with_chars(prompt), "max_tokens": 32768})[1].events


def app_request(models: Any) -> SimpleNamespace:
    return SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(MODELS=models)))


class RedisModels:
    """Stands in for Open WebUI's RedisDict, the model list with Redis: a mapping, but no dict."""

    def __init__(self, models: dict[str, Any]) -> None:
        self.models = models

    def get(self, key: str, default: Any = None) -> Any:
        return self.models.get(key, default)

    def __contains__(self, key: str) -> bool:
        return key in self.models

    def __len__(self) -> int:
        return len(self.models)


def test_guard_reads_the_context_size_from_ollama_and_vllm() -> None:
    flt = make_filter(cg, context_tokens=32768, answer_reserve_tokens=1000)
    body = body_with_chars(20000 * 35 // 10)
    assert asyncio.run(flt.request(body))  # fits into 32k
    with pytest.raises(cg.ContextLimitError):  # Ollama: the model's num_ctx comes with the request
        asyncio.run(flt.request({**body, "options": {"num_ctx": 8192}}))
    vllm = {"id": "m", "max_model_len": 8192, "openai": {"id": "m", "max_model_len": 8192}}
    with pytest.raises(cg.ContextLimitError):
        asyncio.run(flt.request(body, __model__=vllm))
    with pytest.raises(cg.ContextLimitError):
        asyncio.run(flt.request(body, __model__={"id": "m", "openai": {"max_model_len": 8192}}))
    # a workspace model has no max_model_len of its own: the base model's entry in the model list counts
    assistant = {"id": "support", "info": {"base_model_id": "m"}}
    support = body_with_chars(20000 * 35 // 10, model="support")
    assert asyncio.run(flt.request(support, __model__=assistant))
    for models in ({"m": vllm}, RedisModels({"m": vllm})):
        with pytest.raises(cg.ContextLimitError):
            asyncio.run(flt.request(support, __model__=assistant, __request__=app_request(models)))
    # an entry in model_context_tokens wins
    flt.valves.model_context_tokens = "m=40960"
    assert asyncio.run(flt.request({**body, "options": {"num_ctx": 8192}}))
    assert asyncio.run(flt.request(body, __model__=vllm))


def test_guard_compaction_checks_the_model_that_writes_the_summary(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    # the chat model has room for threshold 6500 + 1500, but the summary goes to a small task model
    settings = {
        "chat.context_compaction.enable": True,
        "chat.context_compaction.token_threshold": 6500,
        "chat.context_compaction.model": "small",
    }
    stub_compaction_config(monkeypatch, settings)
    flt = make_filter(cg, context_tokens=10000, answer_reserve_tokens=2000, model_context_tokens="small=4096")
    request = app_request({"m": {"id": "m"}, "small": {"id": "small"}})
    recorder = Recorder()
    asyncio.run(flt.request(body_with_chars(22384), __event_emitter__=recorder.emit, __request__=request))
    assert len(recorder.events) == 1
    assert "usable context of 'small' (2096 tokens)" in caplog.text
    # a compaction model that is not in the model list is not used: Open WebUI takes the chat model
    recorder = Recorder()
    request = app_request({"m": {"id": "m"}})
    asyncio.run(flt.request(body_with_chars(22384), __event_emitter__=recorder.emit, __request__=request))
    assert recorder.events == []
    # a longer summary (max_tokens of the task model parameters) needs more room
    stub_compaction_config(
        monkeypatch, {**settings, "chat.context_compaction.model": "", "task.model.params": {"max_tokens": 2000}}
    )
    assert len(run_guard(flt, body_with_chars(22384))[1].events) == 1
    assert "plus about 2500 tokens and must fit into the usable context of 'm' (8000 tokens)" in caplog.text


# ------------------------------------------------------------------------------- company context --

MONDAY = datetime(2026, 10, 12, 14, 5, tzinfo=ZoneInfo("Europe/Berlin"))


def test_note_english() -> None:
    flt = make_filter(cc, company_name="Example Ltd", timezone="Europe/Berlin")
    note = flt.build_note({"name": "Ada\nLovelace"}, MONDAY)
    assert note.splitlines() == [
        "[Company context]",
        "You are the internal AI assistant of Example Ltd.",
        (
            "Always answer in English unless the user writes in another language "
            "or explicitly asks for a different one."
        ),
        cc.TRANSLATIONS["en"]["formal"],
        "Today is Monday, 12 October 2026 (calendar week 42).",
        "The user's name is Ada Lovelace.",
    ]


def test_note_german_informal_with_rules() -> None:
    flt = make_filter(
        cc,
        language="de",
        company_name="Musterfirma GmbH",
        form_of_address="informal",
        include_user_name=False,
        extra_rules="Wir schreiben Kunden gegenüber in der Wir-Form.",
    )
    note = flt.build_note({"name": "Max"}, MONDAY)
    lines = note.splitlines()
    assert lines[:2] == ["[Firmen-Kontext]", "Du bist der interne KI-Assistent der Firma Musterfirma GmbH."]
    assert lines[2].startswith("Antworte immer auf Deutsch, außer")
    assert "„du“" in lines[3] and "Du-Form" in lines[3]
    assert lines[4:] == [
        "Wir schreiben Kunden gegenüber in der Wir-Form.",
        "Heute ist Montag, 12. Oktober 2026 (Kalenderwoche 42).",
    ]
    assert "Max" not in note


def test_note_without_company_and_custom_response_language() -> None:
    note = make_filter(cc, response_language="French").build_note(None, MONDAY)
    assert "internal AI assistant" not in note
    assert "Always answer in French unless" in note


def test_note_company_profile_comes_before_date_and_name() -> None:
    flt = make_filter(
        cc, language="de", company_name="Musterfirma GmbH", company_profile="  Standorte: Stuttgart, Ulm\n"
    )
    note = flt.build_note({"name": "Max"}, MONDAY)
    paragraphs = note.split("\n\n")
    assert len(paragraphs) == 4
    assert paragraphs[0].startswith("[Firmen-Kontext]\n") and "Heute" not in paragraphs[0]
    assert paragraphs[1] == "Informationen zur Firma:\nStandorte: Stuttgart, Ulm"
    assert paragraphs[2] == cc.TRANSLATIONS["de"]["no_invention"]
    assert paragraphs[3] == "Heute ist Montag, 12. Oktober 2026 (Kalenderwoche 42).\nDer Nutzer heißt Max."
    # the start up to the date is the same for every user, so the model server can share its cache
    other = flt.build_note({"name": "Erika"}, datetime(2026, 10, 13))
    assert other.split("\n\n")[:3] == paragraphs[:3]


def test_parse_group_profiles() -> None:
    text = (
        "Text without group\r\n## Vertrieb\r\nPreise: siehe Preisliste\n### Ansprechpartner\nFrau Muster\n"
        "##   Personal  ##\n\n## vertrieb\nZweiter Teil\n#Kein Abschnitt\n## Leer"
    )
    sections, text_before = cc.parse_group_profiles(text)
    assert text_before
    assert sections == [
        ("Vertrieb", "Preise: siehe Preisliste\n### Ansprechpartner\nFrau Muster\nZweiter Teil\n#Kein Abschnitt"),
        ("Personal", ""),
        ("Leer", ""),
    ]
    assert cc.parse_group_profiles("") == ([], False)
    assert cc.parse_group_profiles("## A\nx")[1] is False


def test_note_group_sections() -> None:
    flt = make_filter(cc, include_user_name=False)
    note = flt.build_note(None, MONDAY, "en", [("Sales", "Price list: P:/Sales"), ("Staff", "")])
    paragraphs = note.split("\n\n")
    assert paragraphs[1:4] == [
        "Information for Sales:\nPrice list: P:/Sales",
        "The user belongs to: Sales, Staff.",
        cc.TRANSLATIONS["en"]["no_invention"],
    ]
    # a heading without text only names the group; without names nothing is added
    assert flt.build_note(None, MONDAY, "en", [("Staff", "")]).split("\n\n")[1] == "The user belongs to: Staff."
    flt = make_filter(cc, include_user_name=False, include_group_names=False)
    assert "\n\n" not in flt.build_note(None, MONDAY, "en", [("Staff", "")])
    assert "belongs to" not in flt.build_note(None, MONDAY, "en", [("Sales", "x")])


def stub_groups(
    monkeypatch: pytest.MonkeyPatch, members: dict[str, list[str]], nested: bool = True, fail: bool = False
) -> list[str]:
    """Stands in for open_webui.models.groups.Groups; nested=False behaves like 0.11 (no
    include_inherited). Returns the list of calls."""
    calls: list[str] = []

    async def get_groups_by_member_id(user_id: str, **kwargs: Any) -> list[Any]:
        if kwargs and not nested:
            raise TypeError("unexpected keyword argument 'include_inherited'")
        if fail:
            raise RuntimeError("database gone")
        calls.append(f"member {user_id} {kwargs}")
        return [SimpleNamespace(name=name) for name in members.get(user_id, [])]

    async def get_all_groups() -> list[Any]:
        calls.append("all")
        return [SimpleNamespace(name=name) for names in members.values() for name in names]

    modules = {
        "open_webui": ModuleType("open_webui"),
        "open_webui.models": ModuleType("open_webui.models"),
        "open_webui.models.groups": SimpleNamespace(
            Groups=SimpleNamespace(get_groups_by_member_id=get_groups_by_member_id, get_all_groups=get_all_groups)
        ),
    }
    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)
    return calls


PROFILES = "## Vertrieb\nPreisliste: P:/Vertrieb\n## Personal\nUrlaub über das Portal\n## Lager\n## Einkauf alt\nx"


def test_inlet_adds_the_sections_of_the_users_groups(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    calls = stub_groups(monkeypatch, {"u1": ["personal", "Vertrieb", "Admins"], "u2": ["Lager"]})
    flt = make_filter(cc, language="de", group_profiles=PROFILES)
    with caplog.at_level(logging.WARNING):
        note = asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1", "name": "Ada"}))["messages"][0]["content"]
        asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1", "name": "Ada"}))
    # in the order of the sections, not of the database; groups without a section are not named
    assert "Informationen für Vertrieb:\nPreisliste: P:/Vertrieb\n\nInformationen für Personal:" in note
    assert "Der Nutzer gehört zu: Vertrieb, Personal." in note and "Admins" not in note and "Lager" not in note
    assert calls.count("all") == 1 and calls[-1] == "member u1 {'include_inherited': True}"
    assert caplog.text.count("names groups that do not exist: Einkauf alt") == 1

    # checked again after GROUP_CHECK_SECONDS, reported again only if the unknown names changed
    flt._group_check = (PROFILES, flt._group_check[1] - cc.GROUP_CHECK_SECONDS, ["Einkauf alt"])
    asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1"}))
    assert calls.count("all") == 2 and caplog.text.count("do not exist") == 1
    flt.valves.group_profiles = PROFILES + "\n## Lagr"
    with caplog.at_level(logging.WARNING):
        asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1"}))
    assert "names groups that do not exist: Einkauf alt, Lagr" in caplog.text

    note = asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u2"}))["messages"][0]["content"]
    assert "Der Nutzer gehört zu: Lager." in note and "Informationen für" not in note
    assert cc.TRANSLATIONS["de"]["no_invention"] not in note
    note = asyncio.run(flt.inlet({"messages": []}, __user__={"name": "API"}))["messages"][0]["content"]
    assert "gehört zu" not in note


def test_inlet_groups_on_open_webui_0_11_and_without_groups(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    calls = stub_groups(monkeypatch, {"u1": ["Vertrieb"]}, nested=False)
    flt = make_filter(cc, language="de", group_profiles=PROFILES)
    note = asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1"}))["messages"][0]["content"]
    assert "Informationen für Vertrieb:" in note and calls[-1] == "member u1 {}"

    # the groups cannot be read: the chat goes on without them, one warning
    stub_groups(monkeypatch, {}, fail=True)
    flt = make_filter(cc, group_profiles=PROFILES)
    with caplog.at_level(logging.WARNING):
        for _ in range(2):
            note = asyncio.run(flt.inlet({"messages": []}, __user__={"id": "u1"}))["messages"][0]["content"]
    assert "Information for" not in note and note.startswith("[Company context]")
    assert caplog.text.count("cannot read the user's groups") == 1

    # without sections Open WebUI is not asked at all
    calls = stub_groups(monkeypatch, {"u1": ["Vertrieb"]})
    asyncio.run(make_filter(cc, group_profiles="only text").inlet({"messages": []}, __user__={"id": "u1"}))
    assert calls == []


def test_inlet_warns_once_about_a_long_note(caplog: pytest.LogCaptureFixture) -> None:
    flt = make_filter(cc, company_profile="Fakt. " * 1500)
    with caplog.at_level(logging.WARNING):
        for _ in range(3):
            run_inlet(flt, {"messages": []})
    assert caplog.text.count("every request carries it") == 1


def run_inlet(flt: Any, body: dict[str, Any], metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    return asyncio.run(flt.inlet(body, __user__={"name": "Ada"}, __metadata__=metadata))


def browser_variables(language: str = "en-US", timezone: str = "Europe/Berlin") -> dict[str, Any]:
    """__metadata__ of a chat request: Open WebUI's frontend sends these variables (0.11.4 and 0.12.1)."""
    return {"variables": {"{{USER_LANGUAGE}}": language, "{{CURRENT_TIMEZONE}}": timezone}}


def test_inlet_adds_system_message_and_is_idempotent() -> None:
    flt = make_filter(cc, timezone="Europe/Berlin")
    body = run_inlet(flt, {"messages": [{"role": "user", "content": "Hi"}]})
    body = run_inlet(flt, body)
    assert [m["role"] for m in body["messages"]] == ["system", "user"]
    assert body["messages"][0]["content"].count("[Company context]") == 1
    assert re.search(
        r"^Today is \w+, \d{1,2} \w+ \d{4} \(calendar week \d{1,2}\)\.$",
        body["messages"][0]["content"],
        re.MULTILINE,
    )


def test_inlet_prepends_to_existing_system_prompt() -> None:
    flt = make_filter(cc, language="de")
    body = {
        "messages": [{"role": "system", "content": "You are a pirate."}, {"role": "user", "content": "Hi"}]
    }
    body = run_inlet(flt, run_inlet(flt, body))
    system = body["messages"][0]["content"]
    assert system.startswith("[Firmen-Kontext]\n") and system.endswith("\n\nYou are a pirate.")
    assert system.count("[Firmen-Kontext]") == 1 and len(body["messages"]) == 2


def test_inlet_handles_list_system_content_and_missing_messages() -> None:
    flt = make_filter(cc)
    body = run_inlet(flt, {"messages": [{"role": "system", "content": [{"type": "text", "text": "Base"}]}]})
    body = run_inlet(flt, body)
    parts = body["messages"][0]["content"]
    assert len(parts) == 2 and parts[0]["text"].startswith("[Company context]") and parts[1]["text"] == "Base"
    assert run_inlet(flt, {})["messages"][0]["role"] == "system"


def test_unknown_timezone_falls_back_to_the_browser() -> None:
    flt = make_filter(cc, timezone="Mars/Olympus")  # loading a stored typo must not fail every chat
    now = flt._now({"variables": {"{{CURRENT_TIMEZONE}}": "Asia/Tokyo"}})
    assert now.utcoffset().total_seconds() == 9 * 3600
    flt = make_filter(cc, timezone="Europe/" + "x" * 300)  # depending on the system, ZoneInfo raises OSError here
    assert flt._now({"variables": {"{{CURRENT_TIMEZONE}}": "y" * 300}}).tzinfo is None
    assert rx.zone("Europe/" + "x" * 300) is None
    assert cc.Filter.Valves().timezone == "" and cc.Filter.Valves().language == "auto"


def test_date_in_the_browser_time_zone() -> None:
    flt = make_filter(cc)
    tokyo, honolulu = (
        flt._now(browser_variables(timezone="Asia/Tokyo")),
        flt._now(browser_variables(timezone="Pacific/Honolulu")),
    )
    assert (
        tokyo.utcoffset().total_seconds() == 9 * 3600 and honolulu.utcoffset().total_seconds() == -10 * 3600
    )
    # the valve wins over the browser; an unknown or missing browser zone means server time
    assert make_filter(cc, timezone="Asia/Tokyo")._now(
        browser_variables(timezone="Pacific/Honolulu")
    ).tzinfo == ZoneInfo("Asia/Tokyo")
    for metadata in (
        browser_variables(timezone="Not/AZone"),
        browser_variables(timezone="../../etc/passwd"),
        {},
        None,
        {"variables": None},
    ):
        assert flt._now(metadata).tzinfo is None


def test_note_date_follows_the_browser_time_zone(monkeypatch: pytest.MonkeyPatch) -> None:
    # 23:30 UTC on Sunday 11 October: already Monday in Berlin, still Sunday in New York
    class FixedDatetime(datetime):
        @classmethod
        def now(cls, tz: Any = None) -> datetime:
            moment = datetime(2026, 10, 11, 23, 30, tzinfo=ZoneInfo("UTC"))
            return moment.astimezone(tz) if tz else moment.replace(tzinfo=None)

    monkeypatch.setattr(cc, "datetime", FixedDatetime)
    flt = make_filter(cc)
    berlin = run_inlet(flt, {"messages": []}, browser_variables(timezone="Europe/Berlin"))["messages"][0][
        "content"
    ]
    new_york = run_inlet(flt, {"messages": []}, browser_variables(timezone="America/New_York"))["messages"][
        0
    ]["content"]
    assert "Today is Monday, 12 October 2026 (calendar week 42)." in berlin
    assert "Today is Sunday, 11 October 2026 (calendar week 41)." in new_york


@pytest.mark.parametrize(
    ("day", "week"), [(datetime(2027, 1, 1), 53), (datetime(2025, 12, 29), 1), (datetime(2026, 10, 21), 43)]
)
def test_note_calendar_week_is_iso(day: datetime, week: int) -> None:
    assert f"(calendar week {week})." in make_filter(cc, language="en").build_note(None, day)
    assert f"(Kalenderwoche {week})." in make_filter(cc, language="de").build_note(None, day)


@pytest.mark.parametrize(
    ("locale", "marker"),
    [
        ("de-DE", "[Firmen-Kontext]"),
        ("de", "[Firmen-Kontext]"),
        ("en-US", "[Company context]"),
        ("fr-FR", "[Company context]"),
    ],
)
def test_note_language_auto(locale: str, marker: str) -> None:
    body = run_inlet(make_filter(cc), {"messages": []}, browser_variables(language=locale))
    assert body["messages"][0]["content"].startswith(marker + "\n")
    # explicit values still win; without the variable (API request) the note is English
    body = run_inlet(make_filter(cc, language="de"), {"messages": []}, browser_variables(language="en-US"))
    assert body["messages"][0]["content"].startswith("[Firmen-Kontext]")
    assert run_inlet(make_filter(cc), {"messages": []})["messages"][0]["content"].startswith(
        "[Company context]"
    )


# ------------------------------------------------------------------------------ reasoning toggle --


def run_toggle(flt: Any, body: dict[str, Any], model: dict[str, Any] | None = None) -> dict[str, Any]:
    return asyncio.run(flt.inlet(body, __model__=model))


def test_toggle_button() -> None:
    flt = rt.Filter()
    assert flt.toggle is True
    prefix = "data:image/svg+xml;base64,"
    assert flt.icon.startswith(prefix)
    assert base64.b64decode(flt.icon[len(prefix) :]).decode() == rt.ICON_SVG


def test_toggle_enables_thinking_and_keeps_other_kwargs() -> None:
    body = {"model": "m", "messages": [], "chat_template_kwargs": {"enable_thinking": False, "custom": 1}}
    result = run_toggle(make_filter(rt), body)
    assert result == {
        "model": "m",
        "messages": [],
        "chat_template_kwargs": {"enable_thinking": True, "custom": 1},
    }
    assert run_toggle(make_filter(rt), {"model": "m", "chat_template_kwargs": "broken"})[
        "chat_template_kwargs"
    ] == {"enable_thinking": True}


@pytest.mark.parametrize(("default", "expected"), [("none", None), ("low", "low")])
def test_toggle_clears_only_a_none_default(default: str, expected: str | None) -> None:
    body = run_toggle(make_filter(rt), {"model": "m", "reasoning_effort": default})
    assert "reasoning_effort" in body and body["reasoning_effort"] == expected


def test_toggle_per_model_effort_and_budget(caplog: pytest.LogCaptureFixture) -> None:
    flt = make_filter(
        rt,
        model_reasoning_effort="qwen=XHigh, broken, gemma=none, x=ultra",
        thinking_budget=2000,
        model_thinking_budget="qwen=8000, gemma=0, y=abc",
    )
    body = run_toggle(flt, {"model": "qwen", "reasoning_effort": "none"})
    assert body["reasoning_effort"] == "xhigh"
    assert body["thinking_budget_tokens"] == body["thinking_token_budget"] == 8000
    # workspace model: settings of its base model apply
    workspace = {"id": "emails", "info": {"base_model_id": "qwen"}}
    assert run_toggle(flt, {"model": "emails"}, workspace)["reasoning_effort"] == "xhigh"
    # per-model budget 0 = no limit, despite the global budget
    body = run_toggle(flt, {"model": "gemma"})
    assert not {"reasoning_effort", "thinking_budget_tokens", "thinking_token_budget"} & body.keys()
    body = run_toggle(flt, {"model": "other"})
    assert body["thinking_budget_tokens"] == 2000 and "reasoning_effort" not in body
    for entry in ("broken", "gemma=none", "x=ultra", "y=abc"):
        assert repr(entry) in caplog.text


def test_toggle_model_ids() -> None:
    assert rt.model_ids({"model": "a"}, {"id": "a", "info": {"base_model_id": "b"}}) == ["a", "b"]
    assert rt.model_ids({}, None) == []
    assert rt.model_ids({"model": "a"}, {"info": None}) == ["a"]
